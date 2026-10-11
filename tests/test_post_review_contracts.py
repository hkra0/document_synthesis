# -*- coding: utf-8 -*-
"""N0–N7 regression contracts extracted from the post-review failures."""

import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Mm, Pt
from docx.oxml.ns import nsdecls, qn

from lib.build_plan import prepare_project_build
from lib.composition import assemble_document, build_cover, mark_end, mark_start
from lib.config import ConfigError, ProjectConfig
from lib.content_integrity import (
    ContentIntegrityError,
    build_expected_inventory,
    verify_content_integrity,
    verify_delivery_format,
)
from lib.delivery import validate_delivery
from lib.docx_inspector import BlockInspection, NodeRef
from lib.field_updater import build_field_index, required_pageref_targets, verify_final_field_caches
from lib.manifest_migration import migrate_manifest_file
from lib.format_analysis import analyze_format_sample
from lib.verification_contracts import (
    build_expected_delivery,
    build_verification_context,
    occurrence_bookmark_name,
    stamp_output_occurrences,
)
from lib.content_integrity import verify_generated_content


def can_symlink() -> bool:
    try:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "target.txt"
            src.write_text("ok", encoding="utf-8")
            link = Path(tmp) / "link.txt"
            link.symlink_to(src)
            return True
    except (OSError, NotImplementedError):
        return False


def _config(source: Path, *, mode: str = "restyle", page_policy: str = "target", cover=None, parts=None):
    data = {
        "schema_version": 3,
        "project_name": "契约测试项目",
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": {"mode": mode, "page_policy": page_policy},
        "source": {"strategy": "docx_document", "file": str(source)},
        "output": {"documents": [{"id": "main", "filename": "main.docx", "parts": parts or ["body"]}]},
    }
    if cover is not None:
        data["cover"] = cover
    return ProjectConfig(data, source.parent)


def _bounded_doc(texts, path: Path):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    mark_start(doc, "body")
    mark_end(doc, "body")
    doc.save(path)


class PostReviewContractsTest(unittest.TestCase):
    """Every test has a stable defect ID in its method name."""

    def test_N1_C01_pageref_cache_is_checked_independently_of_toc(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "pageref.docx"
            doc = Document()
            target = doc.add_paragraph("第四页目标")
            start = target._p.get_or_add_pPr()
            bookmark_start = target._p.xpath("./w:bookmarkStart")
            if not bookmark_start:
                from lib.composition import add_bookmark
                add_bookmark(target._p, "ordinary_target", 1)
            ref = doc.add_paragraph("引用")
            ref._p.append(__import__("docx.oxml", fromlist=["parse_xml"]).parse_xml(
                f'<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'w:instr="PAGEREF ordinary_target"><w:r><w:t>1</w:t></w:r></w:fldSimple>'
            ))
            doc.save(path)

            index = build_field_index(path)
            self.assertEqual(index.pageref_targets, ("ordinary_target",))
            measured_index = index.with_measurements(
                {"ordinary_target": {"physical_page": 4, "printed_page": 4, "expected_label": "4"}}
            )
            self.assertEqual(measured_index.observations[0].status, "failed")
            self.assertEqual(measured_index.observations[0].measurement["physical_page"], 4)
            report = verify_final_field_caches(
                path,
                {"ordinary_target": {"physical_page": 4, "printed_page": 4, "expected_label": "4"}},
            )
            self.assertFalse(report["passed"])
            self.assertEqual(report["failures"][0]["actual"], "1")

            doc = Document(path)
            field = next(doc.element.body.iter(qn("w:fldSimple")))
            next(field.iter(qn("w:t"))).text = "4"
            doc.save(path)
            self.assertTrue(verify_final_field_caches(
                path,
                {"ordinary_target": {"physical_page": 4, "printed_page": 4, "expected_label": "4"}},
            )["passed"])

    def test_N1_smoke_measures_pageref_and_checks_final_disk_cache(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            delivered = root / "delivered.docx"
            doc = Document()
            doc.add_paragraph("封面")
            mark_start(doc, "cover")
            mark_end(doc, "cover")
            doc.save(delivered)
            spec = {"id": "main", "filename": "delivered.docx", "parts": ["cover"]}
            measured = {
                "_Synth_cover": {"physical_page": 1, "printed_page": 1, "expected_label": "1"},
                "ordinary_target": {"physical_page": 4, "printed_page": 4, "expected_label": "4"},
            }
            with patch("lib.delivery.required_pageref_targets", side_effect=[["ordinary_target"], ["ordinary_target"]]) as targets, \
                    patch("lib.delivery.inspect_document", return_value=measured) as inspect, \
                    patch("lib.delivery.validate_measured_delivery", return_value=measured), \
                    patch("lib.delivery.verify_final_field_caches", return_value={"passed": True, "failures": []}) as caches:
                self.assertEqual(
                    validate_delivery(delivered, spec, [], root / "not-created.pdf"),
                    measured,
                )
            self.assertEqual(targets.call_count, 2)
            self.assertEqual(inspect.call_args.args[2], ["_Synth_cover", "ordinary_target"])
            caches.assert_called_once()

    def test_N1_pageref_index_keeps_duplicate_instances_across_body_and_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "duplicate-pageref.docx"
            doc = Document()
            body_paragraph = doc.add_paragraph("正文引用 ")
            body_paragraph._p.append(parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr="PAGEREF ordinary_target">'
                '<w:r><w:t>4</w:t></w:r></w:fldSimple>'
            ))
            cell_paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
            cell_paragraph._p.append(parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr="PAGEREF ordinary_target">'
                '<w:r><w:t>4</w:t></w:r></w:fldSimple>'
            ))
            doc.save(path)

            index = build_field_index(path)
            pagerefs = [item for item in index.observations if item.command == "PAGEREF"]
            self.assertEqual(len(pagerefs), 2)
            self.assertEqual(len({item.field_id for item in pagerefs}), 2)
            self.assertEqual(required_pageref_targets(path), ["ordinary_target"])
            report = verify_final_field_caches(
                path,
                {"ordinary_target": {"physical_page": 4, "printed_page": 4, "expected_label": "4"}},
            )
            self.assertTrue(report["passed"], report)
            self.assertEqual(report["checked_count"], 2)

    def test_N2_C02_v3_cover_does_not_discover_legacy_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            legacy = root / "封面+目录.docx"
            legacy_doc = Document()
            legacy_doc.add_paragraph("旧业务封面")
            legacy_doc.save(legacy)
            source = root / "source.docx"
            Document().save(source)
            config = _config(source, cover={
                "main_title": "新标题",
                "author": "作者",
                "date": "2026",
            })
            self.assertEqual(config.get_template_path(), None)
            cover = build_cover(config)
            text = "".join(node.text or "" for node in cover.element.body.iter(qn("w:t")))
            self.assertIn("新标题", text)
            self.assertNotIn("旧业务封面", text)

    def test_N2_C02_explicit_template_requires_all_nonempty_placeholders(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "cover.docx"
            template_doc = Document()
            template_doc.add_paragraph("{{MAIN_TITLE}}")
            template_doc.save(template)
            source = root / "source.docx"
            Document().save(source)
            config = _config(source, cover={
                "mode": "template",
                "template": str(template),
                "main_title": "标题",
                "author": "作者",
            })
            with self.assertRaisesRegex(ValueError, "COVER_TEMPLATE_MISSING_FIELDS"):
                build_cover(config)

    def test_N2_generated_cover_rejects_explicit_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "cover.docx"
            Document().save(template)
            source = root / "source.docx"
            Document().save(source)
            with self.assertRaisesRegex(ConfigError, "cover.mode=generated 不能同时指定 template"):
                _config(source, cover={
                    "mode": "generated",
                    "template": str(template),
                    "main_title": "标题",
                })

    def test_N2_static_template_is_not_rewritten_or_treated_as_generated(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "fixed-cover.docx"
            template_doc = Document()
            template_doc.add_paragraph("固定封面内容")
            template_doc.save(template)
            source = root / "source.docx"
            Document().save(source)
            config = _config(source, cover={
                "mode": "static_template",
                "template": str(template),
            })
            cover = build_cover(config)
            text = "".join(node.text or "" for node in cover.element.body.iter(qn("w:t")))
            self.assertEqual(text, "固定封面内容")
            self.assertEqual(config.cover_spec.dynamic_fields, {})

    def test_N2_cover_modes_reject_missing_template_and_static_dynamic_conflicts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            Document().save(source)
            template = root / "fixed-cover.docx"
            Document().save(template)
            with self.subTest(case="template-without-path"):
                with self.assertRaisesRegex(ConfigError, "cover.mode=template 必须提供有效 template 路径"):
                    _config(source, cover={"mode": "template"})
            with self.subTest(case="static-template-with-dynamic-field"):
                with self.assertRaisesRegex(ConfigError, "cover.mode=static_template 不能同时设置动态封面字段"):
                    _config(source, cover={
                        "mode": "static_template",
                        "template": str(template),
                        "author": "不应被忽略",
                    })

    def test_N2_cover_template_is_frozen_by_the_build_plan(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "cover.docx"
            template_doc = Document()
            template_doc.add_paragraph("{{MAIN_TITLE}}")
            template_doc.save(template)
            source = root / "source.docx"
            _bounded_doc(["正文"], source)
            config = _config(source, cover={"template": str(template), "main_title": "标题"}, parts=["cover", "body"])
            prepared = prepare_project_build(source, config)
            template.write_bytes(template.read_bytes() + b"changed")
            with self.assertRaisesRegex(ValueError, "封面模板自准备后已发生变更"):
                prepared.verify_inputs_unchanged()

    def test_N3_prepared_config_contract_rejects_mutation_after_prepare(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            _bounded_doc(["正文"], source)
            config = _config(source, parts=["body"])
            prepared = prepare_project_build(source, config)

            self.assertEqual(len(prepared.config_hash), 64)
            self.assertEqual(
                prepared.to_public_dict()["config_contract_sha256"],
                prepared.config_hash,
            )
            # ProjectConfig is intentionally still the runtime object consumed
            # by the renderer, so the prepared contract must detect a direct
            # mutation instead of relying on PreparedBuild's shallow freeze.
            config.documents[0]["filename"] = "tampered.docx"
            with self.assertRaisesRegex(ValueError, "配置自准备后已发生变更"):
                prepared.verify_inputs_unchanged()

    def test_N2_generated_content_contract_counts_repeated_template_placeholders(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "cover.docx"
            template_doc = Document()
            template_doc.add_paragraph("{{MAIN_TITLE}} / {{MAIN_TITLE}}")
            template_doc.save(template)
            source = root / "source.docx"
            _bounded_doc(["正文"], source)
            config = _config(
                source,
                cover={"mode": "template", "template": str(template), "main_title": "标题"},
                parts=["cover", "body"],
            )
            prepared = prepare_project_build(source, config)
            expected = build_expected_delivery(prepared, config.documents[0])
            self.assertEqual(len(expected.generated_content), 1)
            self.assertEqual(expected.generated_content[0]["expected_count"], 2)

    def test_N3_N4_C03_preserve_uses_source_geometry_and_inline_emphasis(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            section = doc.sections[0]
            section.page_width = Mm(180)
            section.page_height = Mm(240)
            p = doc.add_paragraph()
            first = p.add_run("保留")
            first.bold = True
            first.font.size = Pt(12)
            second = p.add_run("强调")
            second.font.size = Pt(12)
            doc.save(source)
            config = _config(source, mode="preserve", page_policy="source")
            prepared = prepare_project_build(source, config)
            body = root / "body.docx"
            body.write_bytes(source.read_bytes())
            delivered = root / "delivered.docx"
            assemble_document(config, config.documents[0], body, [], {}, delivered)
            context = build_verification_context(prepared, None, config.documents[0], delivered)
            report = verify_delivery_format(delivered, config.resolved_format, verification_context=context)
            self.assertTrue(report.passed, report.violations)
            self.assertEqual(report.details["page_geometry"][0]["expected"]["width_mm"], 180.0)

    def test_N4_inline_emphasis_is_checked_by_logical_range_after_run_merge(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source_doc = Document()
            source_doc.sections[0].page_width = Mm(180)
            source_doc.sections[0].page_height = Mm(240)
            paragraph = source_doc.add_paragraph()
            paragraph.add_run("保留").bold = True
            paragraph.add_run("强调").italic = True
            source_doc.save(source)
            config = _config(source, mode="preserve", page_policy="source")
            prepared = prepare_project_build(source, config)

            # Simulate a save/import that merges two source runs into one run
            # carrying only the first run's style.  Text-only or run-index
            # verification would incorrectly accept this output.
            delivered = Document()
            delivered.sections[0].page_width = Mm(180)
            delivered.sections[0].page_height = Mm(240)
            merged = delivered.add_paragraph().add_run("保留强调")
            merged.bold = True
            mark_start(delivered, "body")
            mark_end(delivered, "body")
            delivered_path = root / "delivered.docx"
            delivered.save(delivered_path)

            context = build_verification_context(
                prepared, None, config.documents[0], delivered_path
            )
            report = verify_delivery_format(
                delivered_path,
                config.resolved_format,
                verification_context=context,
            )
            self.assertFalse(report.passed)
            self.assertTrue(any("斜体状态" in item for item in report.violations))

    def test_N3_N4_C04_expected_denominator_rejects_one_missing_of_five(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            for index in range(5):
                doc.add_paragraph(f"来源段落 {index}")
            doc.save(source)
            config = _config(source)
            prepared = prepare_project_build(source, config)
            run_dir = root / "run"
            rendered = __import__("lib.engine", fromlist=["UnifiedSynthesizer"]).UnifiedSynthesizer.render_body(prepared, run_dir)
            delivered = root / "delivered.docx"
            assemble_document(config, config.documents[0], rendered.path, list(rendered.nodes), {}, delivered)
            # Corrupt one managed source node after the output map has been created.
            corrupted = Document(delivered)
            corrupted.paragraphs[2].runs[0].font.size = Pt(70)
            corrupted.save(delivered)
            context = build_verification_context(prepared, rendered, config.documents[0], delivered)
            self.assertEqual(context.config_hash, prepared.config_hash)
            self.assertEqual(context.to_dict()["config_hash"], prepared.config_hash)
            self.assertEqual(len(context.expected_delivery.ordered_nodes), 5)
            self.assertEqual(len(context.output_node_map.locations), 5)
            report = verify_delivery_format(delivered, config.resolved_format, verification_context=context)
            self.assertFalse(report.passed)
            self.assertTrue(any("字号" in item for item in report.violations))

    def test_N4_source_page_expectations_include_all_source_sections(self):
        from types import SimpleNamespace
        from lib.format_schema import PageSpec

        first = PageSpec(width_mm=180.0, height_mm=240.0)
        second = PageSpec(width_mm=210.0, height_mm=297.0)
        prepared = SimpleNamespace(
            parts={},
            inspections=(
                SimpleNamespace(blocks=(), sections=(SimpleNamespace(page_spec=first),)),
                SimpleNamespace(blocks=(), sections=(SimpleNamespace(page_spec=second),)),
            ),
            config=SimpleNamespace(
                regions={},
                formatting=SimpleNamespace(mode="preserve", page_policy="source", inline_emphasis="preserve"),
            ),
            source_order=(),
            assignments=(),
        )
        from lib.verification_contracts import build_expected_delivery

        expected = build_expected_delivery(prepared, {"id": "main", "parts": ["body"]})
        self.assertEqual([item["width_mm"] for item in expected.page_expectations], [180.0, 210.0])
        self.assertEqual([item["height_mm"] for item in expected.page_expectations], [240.0, 297.0])

    def test_N3_output_mapping_rejects_same_shape_wrong_text_and_tracks_reuse(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            _bounded_doc(["A source paragraph", "B source paragraph"], source)
            config = _config(source)
            prepared = prepare_project_build(source, config)

            wrong = root / "wrong.docx"
            _bounded_doc(["Unrelated replacement", "B source paragraph"], wrong)
            context = build_verification_context(prepared, None, config.documents[0], wrong)
            self.assertEqual(len(context.output_node_map.locations), 1)
            self.assertEqual(len(context.output_node_map.missing_occurrences), 1)
            self.assertTrue(all(
                item.get("text_hash_matches") is True
                for locations in context.output_node_map.locations.values()
                for item in locations
            ))
            first_occurrence = next(iter(context.output_node_map.locations))
            self.assertEqual(
                context.output_node_map.locations[first_occurrence][0]["text_range"],
                {"start": 0, "end": len("A source paragraph")},
            )

            duplicated = root / "duplicated.docx"
            _bounded_doc(["A source paragraph", "B source paragraph", "A source paragraph"], duplicated)
            duplicate_context = build_verification_context(prepared, None, config.documents[0], duplicated)
            self.assertEqual(len(duplicate_context.output_node_map.locations), 2)
            self.assertEqual(len(duplicate_context.output_node_map.duplicate_occurrences), 1)

            reused_source = root / "reused-source.docx"
            _bounded_doc(["A source paragraph", "A source paragraph"], reused_source)
            reused_config = _config(reused_source)
            reused_prepared = prepare_project_build(reused_source, reused_config)
            reused_context = build_verification_context(
                reused_prepared, None, reused_config.documents[0], reused_source
            )
            self.assertEqual(len(reused_context.expected_delivery.ordered_nodes), 2)
            self.assertEqual(len(reused_context.output_node_map.locations), 2)
            self.assertEqual(reused_context.output_node_map.duplicate_occurrences, ())

    def test_N3_source_occurrences_distinguish_same_named_source_nodes(self):
        from types import SimpleNamespace

        def block(source_sha256):
            return BlockInspection(
                node=NodeRef(
                    source_sha256=source_sha256,
                    part_uri="word/document.xml",
                    element_path="/w:document/w:body/w:p[1]",
                    text_hash="same-title-hash",
                ),
                structure_type="paragraph",
                story_type="body",
                visible_text="同名标题",
                p_style_id=None,
                p_style_name=None,
                outline_level=1,
                effective_run=None,
                effective_paragraph=None,
                protected_objects=[],
                runs_count=1,
            )

        prepared = SimpleNamespace(
            inspections=(
                SimpleNamespace(blocks=[block("sha-first")]),
                SimpleNamespace(blocks=[block("sha-second")]),
            ),
            assignments=(),
            parts={},
            source_order=("/sources/one/manuscript.docx", "/sources/two/manuscript.docx"),
            config=SimpleNamespace(
                regions={},
                formatting=SimpleNamespace(mode="preserve", page_policy="target", inline_emphasis="preserve"),
            ),
        )
        expected = build_expected_delivery(prepared, {"id": "main", "parts": ["body"]})
        occurrences = [item.occurrence for item in expected.ordered_nodes]
        self.assertEqual(len(occurrences), 2)
        self.assertEqual([item.source_order_index for item in occurrences], [0, 1])
        self.assertNotEqual(occurrences[0].occurrence_id, occurrences[1].occurrence_id)
        self.assertNotEqual(occurrences[0].source_sha256, occurrences[1].source_sha256)

    def test_N3_provenance_bookmarks_survive_save_and_reject_tampered_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            _bounded_doc(["A source paragraph", "B source paragraph"], source)
            config = _config(source)
            prepared = prepare_project_build(source, config)
            delivered = root / "delivered.docx"
            _bounded_doc(["A source paragraph", "B source paragraph"], delivered)
            expected = build_expected_delivery(prepared, config.documents[0])

            stamped = stamp_output_occurrences(delivered, expected)
            self.assertEqual(len(stamped.transformations), 2)
            names = {
                marker.get(qn("w:name"))
                for marker in Document(delivered).element.body.iter(qn("w:bookmarkStart"))
            }
            self.assertTrue(all(
                occurrence_bookmark_name(node.occurrence.occurrence_id) in names
                for node in expected.ordered_nodes
            ))

            tampered = Document(delivered)
            tampered.paragraphs[0].runs[0].text = "A different paragraph"
            tampered.save(delivered)
            context = build_verification_context(prepared, None, config.documents[0], delivered)
            self.assertEqual(len(context.output_node_map.locations), 2)
            self.assertFalse(all(
                location.get("text_hash_matches")
                for locations in context.output_node_map.locations.values()
                for location in locations
            ))

    def test_N5_C05_order_and_extra_instances_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.add_paragraph("A source paragraph")
            doc.add_paragraph("B source paragraph")
            doc.save(source)
            prepared = type("Prepared", (), {
                "source_order": (str(source),),
                "source_docx_path": source,
                "nodes": (),
                "parts": {},
                "config": type("Config", (), {"regions": {}})(),
            })()
            expected = build_expected_inventory(prepared, {"id": "main", "parts": ["body"]})
            reordered = root / "reordered.docx"
            _bounded_doc(["B source paragraph", "A source paragraph"], reordered)
            with self.assertRaises(ContentIntegrityError):
                verify_content_integrity(expected, reordered)
            duplicated = root / "duplicated.docx"
            _bounded_doc(["A source paragraph", "B source paragraph", "A source paragraph"], duplicated)
            with self.assertRaises(ContentIntegrityError):
                verify_content_integrity(expected, duplicated)
            valid = root / "valid.docx"
            _bounded_doc(["A source paragraph", "B source paragraph"], valid)
            self.assertTrue(verify_content_integrity(expected, valid))

    def test_N5_declared_reuse_of_one_source_region_preserves_instance_counts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source_doc = Document()
            source_doc.add_paragraph("可复用选区")
            source_doc.add_paragraph("未选内容")
            source_doc.save(source)
            config = ProjectConfig({
                "schema_version": 3,
                "project_name": "重复选区契约",
                "format": {"ref": "preset:academic-basic@1.0.0"},
                "formatting": {"mode": "preserve"},
                "source": {
                    "strategy": "docx_document",
                    "file": str(source),
                    "regions": {
                        "excerpt": {
                            "start": "/w:document/w:body/w:p[1]",
                            "end": "/w:document/w:body/w:p[2]",
                        },
                        "excluded": {
                            "start": "/w:document/w:body/w:p[2]",
                            "end": "end_of_document",
                            "exclude": True,
                        },
                    },
                },
                "layout": {"parts": {
                    "excerpt_a": {"kind": "content", "source_region": "excerpt"},
                    "excerpt_b": {"kind": "content", "source_region": "excerpt"},
                }},
                "output": {"documents": [{
                    "id": "main",
                    "filename": "reused-region.docx",
                    "parts": ["excerpt_a", "excerpt_b"],
                }]},
            }, root)
            prepared = prepare_project_build(source, config)
            expected = build_expected_inventory(prepared, config.documents[0])
            self.assertEqual(expected.semantic.paragraphs, ["可复用选区", "可复用选区"])

            delivered = root / "reused-region.docx"
            assemble_document(config, config.documents[0], source, [], {}, delivered)
            self.assertTrue(verify_content_integrity(expected, delivered))

    def test_N5_same_text_from_two_sources_is_a_valid_two_instance_sequence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = root / "first.docx"
            second = root / "second.docx"
            _bounded_doc(["相同来源段落"], first)
            _bounded_doc(["相同来源段落"], second)
            prepared = type("Prepared", (), {
                "source_order": (str(first), str(second)),
                "source_docx_path": None,
                "nodes": (),
                "parts": {},
                "config": type("Config", (), {"regions": {}})(),
            })()
            expected = build_expected_inventory(prepared, {"id": "main", "parts": ["body"]})
            delivered = root / "two-sources.docx"
            _bounded_doc(["相同来源段落", "相同来源段落"], delivered)
            self.assertTrue(verify_content_integrity(expected, delivered))

    def test_N5_generated_content_is_scoped_and_counted_by_part(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            expected = type("Expected", (), {
                "generated_content": ({
                    "part_id": "cover",
                    "kind": "cover_field",
                    "field": "main_title",
                    "expected_text": "论文标题",
                    "expected_text_hash": "fixture",
                    "expected_count": 1,
                    "position": {"part_id": "cover", "field": "main_title"},
                },)
            })()

            valid = Document()
            valid.add_paragraph("论文标题")
            valid.add_paragraph("作者")
            mark_start(valid, "cover")
            mark_end(valid, "cover")
            valid.add_paragraph("论文标题")
            valid_path = root / "valid.docx"
            valid.save(valid_path)
            report = verify_generated_content(valid_path, expected)
            self.assertTrue(report["passed"], report)
            self.assertEqual(report["checks"][0]["actual_count"], 1)

            missing_from_cover = Document()
            missing_from_cover.add_paragraph("作者")
            mark_start(missing_from_cover, "cover")
            mark_end(missing_from_cover, "cover")
            missing_from_cover.add_paragraph("论文标题")
            missing_path = root / "missing-from-cover.docx"
            missing_from_cover.save(missing_path)
            missing_report = verify_generated_content(missing_path, expected)
            self.assertFalse(missing_report["passed"])
            self.assertEqual(missing_report["checks"][0]["actual_count"], 0)

            duplicate = Document()
            duplicate.add_paragraph("论文标题")
            duplicate.add_paragraph("作者")
            duplicate.add_paragraph("论文标题")
            mark_start(duplicate, "cover")
            mark_end(duplicate, "cover")
            duplicate.add_paragraph("正文")
            duplicate_path = root / "duplicate.docx"
            duplicate.save(duplicate_path)
            duplicate_report = verify_generated_content(duplicate_path, expected)
            self.assertFalse(duplicate_report["passed"])
            self.assertEqual(duplicate_report["checks"][0]["actual_count"], 2)

    def test_N5_generated_content_markers_bind_the_field_instance(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source_doc = Document()
            source_doc.add_paragraph("正文")
            source_doc.save(source)
            config = _config(
                source,
                cover={"main_title": "论文标题"},
                parts=["cover", "body"],
            )
            prepared = prepare_project_build(source, config)
            expected = build_expected_delivery(prepared, config.documents[0])
            cover = build_cover(config)
            # An unmarked duplicate in the same part must not satisfy the
            # generated field contract.
            cover.add_paragraph("论文标题")
            path = root / "cover.docx"
            cover.save(path)
            self.assertTrue(verify_generated_content(path, expected)["passed"])

            # The same marker must survive front-part import and bookmark ID
            # renumbering in a complete assembled delivery.
            assembled = root / "assembled.docx"
            assemble_document(config, config.documents[0], source, [], {}, assembled)
            assembled_report = verify_generated_content(assembled, expected)
            self.assertTrue(assembled_report["passed"], assembled_report)

            tampered = Document(path)
            marker = next(
                item for item in tampered.element.body.iter(qn("w:bookmarkStart"))
                if (item.get(qn("w:name")) or "").startswith("_SynthGen_")
            )
            marker_id = marker.get(qn("w:id"))
            elements = list(tampered.element.body.iter())
            start = elements.index(marker)
            end = next(
                index for index in range(start + 1, len(elements))
                if elements[index].tag == qn("w:bookmarkEnd")
                and elements[index].get(qn("w:id")) == marker_id
            )
            text_node = next(
                item for item in elements[start + 1:end]
                if item.tag == qn("w:t")
            )
            text_node.text = "错误字段"
            tampered.save(path)
            report = verify_generated_content(path, expected)
            self.assertFalse(report["passed"], report)
            self.assertEqual(report["checks"][0]["marker_values"], ["错误字段"])

    def test_N5_template_markers_cover_split_runs_and_table_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "cover.docx"
            template_doc = Document()
            paragraph = template_doc.add_paragraph()
            paragraph.add_run("{{MAIN_")
            paragraph.add_run("TITLE}}")
            table = template_doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "{{AUTHOR}}"
            template_doc.save(template)
            source = root / "source.docx"
            _bounded_doc(["正文"], source)
            config = _config(
                source,
                cover={
                    "mode": "template",
                    "template": str(template),
                    "main_title": "论文标题",
                    "author": "作者",
                },
                parts=["cover"],
            )
            prepared = prepare_project_build(source, config)
            expected = build_expected_delivery(prepared, config.documents[0])
            cover = build_cover(config)
            path = root / "split-and-table-cover.docx"
            cover.save(path)
            report = verify_generated_content(path, expected)
            self.assertTrue(report["passed"], report)
            self.assertEqual(
                {check["field"] for check in report["checks"]},
                {"main_title", "author"},
            )

    def test_N6_C06_migration_rejects_path_alias_and_hardlink(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "manifest.json"
            source.write_text(json.dumps({
                "schema_version": 2,
                "project_name": "migration-contract",
            }), encoding="utf-8")
            before = source.read_bytes()
            aliases = [root / "nested" / ".." / "manifest.json"]
            hardlink = root / "manifest-hardlink.json"
            try:
                hardlink.hardlink_to(source)
                aliases.append(hardlink)
            except OSError:
                pass
            for alias in aliases:
                with self.subTest(alias=str(alias)):
                    with self.assertRaisesRegex(ConfigError, "MIGRATION_SOURCE_EQUALS_OUTPUT"):
                        migrate_manifest_file(source, alias, replace=True)
                    self.assertEqual(source.read_bytes(), before)

            destination = root / "destination.json"
            destination.write_text("old", encoding="utf-8")
            with self.assertRaises(ConfigError):
                migrate_manifest_file(source, destination)
            migrate_manifest_file(source, destination, replace=True)
            self.assertEqual(source.read_bytes(), before)
            self.assertNotEqual(destination.read_text(encoding="utf-8"), "old")

    @unittest.skipUnless(can_symlink(), "系统不支持符号链接或无权限（Windows 需开启开发者模式）")
    def test_N6_C06_migration_rejects_symlink_alias(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "manifest.json"
            source.write_text(json.dumps({
                "schema_version": 2,
                "project_name": "migration-contract",
            }), encoding="utf-8")
            before = source.read_bytes()
            symlink = root / "manifest-link.json"
            symlink.symlink_to(source)
            with self.assertRaisesRegex(ConfigError, "MIGRATION_SOURCE_EQUALS_OUTPUT"):
                migrate_manifest_file(source, symlink, replace=True)
            self.assertEqual(source.read_bytes(), before)

    def test_N7_C07_unstructured_analysis_returns_ranked_evidence_not_a_single_signal(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "unstructured.docx"
            doc = Document()
            heading = doc.add_paragraph()
            heading.alignment = 1
            heading.paragraph_format.space_before = Pt(12)
            heading.add_run("一、无大纲标题").bold = True
            body = doc.add_paragraph(
                "这是一段足够长的正文，用来证明字号或首段位置不能单独决定语义角色。"
                "正文还包含多个句子和完整的说明内容。"
            )
            body.runs[0].font.size = Pt(12)
            heading.runs[0].font.size = Pt(16)
            doc.save(path)
            report = analyze_format_sample(path)
            evidence = report["unstructured_candidate_evidence"]
            self.assertTrue(evidence)
            first = next(iter(evidence.values()))
            self.assertLessEqual(len(first["candidates"]), 3)
            self.assertTrue(any(item["signals"] for item in first["candidates"]))
            self.assertIn("candidate_evidence", report["candidate_roles"].get("body", {}))
            body_evidence = evidence["/w:document/w:body/w:p[2]"]
            self.assertEqual(body_evidence["selected_role"], "body")
            self.assertFalse(
                any(
                    item["role"] == "title" or item["role"].startswith("heading.")
                    for item in body_evidence["candidates"]
                )
            )


if __name__ == "__main__":
    unittest.main()
