# -*- coding: utf-8 -*-
"""
构建计划与源文件哈希校验单元测试 (tests/test_build_plan.py)
验证 SHA-256 哈希计算、构建前源文件防篡改校验与内容块统计。
"""

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from lib.engine import BuildError, UnifiedSynthesizer, compute_file_sha256


class BuildPlanTest(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_dir = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_compute_file_sha256(self):
        """测试 SHA-256 哈希计算一致性"""
        test_file = self.test_dir / "sample.txt"
        test_file.write_text("Hello World", encoding="utf-8")
        
        h1 = compute_file_sha256(test_file)
        self.assertEqual(len(h1), 64)
        
        # 相同内容哈希相同
        h2 = compute_file_sha256(test_file)
        self.assertEqual(h1, h2)
        
        # 修改后哈希改变
        test_file.write_text("Modified World", encoding="utf-8")
        h3 = compute_file_sha256(test_file)
        self.assertNotEqual(h1, h3)

    def test_plan_records_source_hashes_and_block_counts(self):
        """测试 plan() 会记录 source_hashes、content_block_count 和 heading_count"""
        doc_path = self.test_dir / "doc.docx"
        doc = Document()
        p = doc.add_paragraph("标题一")
        p._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
        doc.add_paragraph("正文")
        doc.save(str(doc_path))

        manifest = {
            "project_name": "Plan Hash Test",
            "schema_version": 2,
            "source": {
                "strategy": "docx_document",
                "file": "doc.docx"
            },
            "output": {
                "documents": [
                    {"id": "main", "filename": "main.docx", "parts": ["body"]}
                ]
            }
        }
        (self.test_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

        plan = UnifiedSynthesizer.plan(self.test_dir)
        self.assertIn("source_hashes", plan)
        resolved_doc = str(doc_path.resolve())
        self.assertIn(resolved_doc, plan["source_hashes"])
        self.assertEqual(plan["source_hashes"][resolved_doc], compute_file_sha256(doc_path))
        self.assertEqual(plan["content_block_count"], 2)
        self.assertEqual(plan["heading_count"], 1)

    def test_verify_build_plan_hashes_detects_tampered_or_deleted_file(self):
        """测试 verify_build_plan_hashes 校验成功与篡改/删除拦截"""
        f1 = self.test_dir / "source1.docx"
        doc = Document()
        doc.add_paragraph("内容 1")
        doc.save(str(f1))

        plan = {
            "source_hashes": {
                str(f1): compute_file_sha256(f1)
            }
        }

        # 1. 正常情况无异常
        UnifiedSynthesizer.verify_build_plan_hashes(plan)

        # 2. 文件被修改
        doc.add_paragraph("篡改内容")
        doc.save(str(f1))
        with self.assertRaises(BuildError) as cm:
            UnifiedSynthesizer.verify_build_plan_hashes(plan)
        self.assertIn("哈希不匹配", str(cm.exception))

        # 3. 文件被删除
        f1.unlink()
        with self.assertRaises(BuildError) as cm:
            UnifiedSynthesizer.verify_build_plan_hashes(plan)
        self.assertIn("已丢失", str(cm.exception))


if __name__ == "__main__":
    unittest.main()
