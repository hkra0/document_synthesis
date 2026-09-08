"""交付回归检查按声明清单与实际来源运行，不依赖 Office。"""

from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

from smoke_test import (
    declared_delivery_paths,
    _source_docx_paths,
    _check_body_text,
    _metadata_contract_violations,
    test_project,
)


class SmokeOutputsTest(unittest.TestCase):
    def test_metadata_contract_requires_current_source_config_and_delivery_hashes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            delivery = root / "main.docx"
            Document().save(source)
            Document().save(delivery)
            source_hash = __import__("hashlib").sha256(source.read_bytes()).hexdigest()
            delivery_hash = __import__("hashlib").sha256(delivery.read_bytes()).hexdigest()
            prepared = type("Prepared", (), {
                "config_hash": "config-hash",
                "format_hash": "format-hash",
                "source_hashes": {str(source): source_hash},
            })()
            plan = {
                "source_hashes": {str(source): source_hash},
                "format_hash": "format-hash",
            }
            metadata = {
                "source_hashes": {str(source): source_hash},
                "configuration": {"config_contract_sha256": "config-hash"},
                "format": {"source_sha256": "format-hash"},
                "deliveries": {
                    "main": {
                        "filename": "main.docx",
                        "parts": ["body"],
                        "sha256": delivery_hash,
                    }
                },
            }
            paths = [({"id": "main", "filename": "main.docx", "parts": ["body"]}, delivery)]
            self.assertEqual(_metadata_contract_violations(metadata, plan, prepared, paths), [])
            metadata["deliveries"]["main"]["sha256"] = "stale"
            self.assertIn("交付物 main SHA-256 与元数据不一致", _metadata_contract_violations(metadata, plan, prepared, paths))

    def test_old_filename_does_not_replace_missing_declared_delivery(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            Document().save(output / "example_目录+正文.docx")
            plan = {"documents": [{"id": "main", "filename": "complete.docx", "parts": ["cover", "toc", "body"]}]}
            with self.assertRaisesRegex(ValueError, "complete.docx"):
                declared_delivery_paths(plan, output)

    def test_every_declared_file_is_required(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            Document().save(output / "main.docx")
            plan = {"documents": [
                {"id": "main", "filename": "main.docx", "parts": ["toc", "body"]},
                {"id": "front", "filename": "front.docx", "parts": ["cover", "toc"]},
            ]}
            with self.assertRaisesRegex(ValueError, "front.docx"):
                declared_delivery_paths(plan, output)
            Document().save(output / "front.docx")
            self.assertEqual(len(declared_delivery_paths(plan, output)), 2)

    def test_file_source_is_checked_once(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.docx"
            Document().save(source)
            self.assertEqual(_source_docx_paths(source, [{"file": source.name}] * 3), [source])

    def test_only_plan_files_are_checked(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp)
            Document().save(source / "used.docx")
            self.assertEqual(_source_docx_paths(source, [
                {"type": "folder"}, {"file": "used.docx"}, {"file": "used.docx"},
            ]), [(source / "used.docx").resolve()])

    def test_text_check_covers_tables_and_split_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, target = Path(tmp) / "source.docx", Path(tmp) / "target.docx"
            doc = Document()
            doc.add_paragraph("Important paragraph")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
            doc.save(source)
            dest = Document()
            paragraph = dest.add_paragraph()
            paragraph.add_run("Important ")
            paragraph.add_run("paragraph")
            dest.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
            dest.save(target)
            _check_body_text(target, [source], [])
            dest.tables[0].cell(0, 0).text = "Missing"
            dest.save(target)
            with self.assertRaisesRegex(ValueError, "Tableevidence"):
                _check_body_text(target, [source], [])

    def test_configured_heading_replacement_is_permitted(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, target = Path(tmp) / "source.docx", Path(tmp) / "target.docx"
            doc = Document()
            doc.add_paragraph("Original heading")
            doc.save(source)
            dest = Document()
            dest.add_paragraph("Updated heading")
            dest.save(target)
            _check_body_text(target, [source], [{"raw_text": "Original heading", "title": "Updated heading"}])

    def test_full_check_validates_body_before_referenced_front(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp)
            for name in ("main.docx", "front.docx"):
                Document().save(source / name)
            plan = {"project": "test", "nodes": [], "documents": [
                {"id": "front", "filename": "front.docx", "parts": ["cover", "toc"], "toc": {"reference": "main"}},
                {"id": "main", "filename": "main.docx", "parts": ["toc", "body"], "toc": {"reference": "main"}},
            ]}
            page_map = {"heading": {"physical_page": 3, "printed_page": 1}}
            with patch("smoke_test.UnifiedSynthesizer.plan", return_value=plan), patch(
                "lib.delivery.validate_delivery", return_value=page_map
            ) as validate:
                self.assertTrue(test_project(source, source))
            self.assertEqual([call.args[1]["id"] for call in validate.call_args_list], ["main", "front"])
            self.assertIsNone(validate.call_args_list[0].kwargs["reference_map"])
            self.assertEqual(validate.call_args_list[1].kwargs["reference_map"], page_map)

    def test_structure_only_does_not_call_word_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp)
            Document().save(source / "main.docx")
            plan = {"project": "test", "nodes": [], "documents": [
                {"id": "main", "filename": "main.docx", "parts": ["body"]},
            ]}
            with patch("smoke_test.UnifiedSynthesizer.plan", return_value=plan), patch(
                "lib.delivery.validate_delivery"
            ) as word, patch("lib.delivery.validate_delivery_structure") as structure:
                self.assertTrue(test_project(source, source, structure_only=True))
                word.assert_not_called()
                structure.assert_called_once()

    def test_default_output_uses_configured_project_name_not_source_stem(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source-name"
            source.mkdir()
            output_root = Path(tmp) / "output"
            actual = output_root / "Configured Project"
            actual.mkdir(parents=True)
            Document().save(actual / "main.docx")
            plan = {"project": "Configured Project", "nodes": [], "documents": [
                {"id": "main", "filename": "main.docx", "parts": ["body"]},
            ]}
            with patch("smoke_test.DEFAULT_OUTPUT_DIR", output_root), patch(
                "smoke_test.UnifiedSynthesizer.plan", return_value=plan
            ), patch("lib.delivery.validate_delivery_structure"):
                self.assertTrue(test_project(source, structure_only=True))


if __name__ == "__main__":
    unittest.main()
