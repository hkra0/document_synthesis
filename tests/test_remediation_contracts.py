# -*- coding: utf-8 -*-
"""S0 formal contracts for the N0–N10 remediation defects.

The D01–D04 assertions are written against the desired contract.  Positive
and negative cases remain explicit so a passing test cannot be obtained by
silently weakening the gate; no case is skipped or marked expected-failure.
"""

import copy
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.shared import RGBColor

from lib.composition import assemble_document
from lib.content_integrity import (
    ExpectedInventory,
    extract_semantic_inventory,
    verify_content_integrity,
    verify_delivery_format,
)
from lib.engine import UnifiedSynthesizer
from lib.format_analysis import analyze_format_sample
from lib.verification_contracts import build_verification_context


ROOT = Path(__file__).resolve().parents[1]
CONTRACTS = ROOT / "docs" / "acceptance" / "remediation" / "scenario-contracts.json"
IMAGE = ROOT / "docs" / "acceptance" / "r0-r11-review" / "thesis-fixture" / "figure.png"


def _write_manifest(path: Path, source: Path, *, formatting=None) -> Path:
    formatting = formatting or {"mode": "preserve", "page_policy": "source"}
    manifest = path / "manifest.json"
    manifest.write_text(json.dumps({
        "schema_version": 3,
        "project_name": "S0 remediation contract",
        "source": {"strategy": "docx_document"},
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": formatting,
        "output": {"documents": [{"id": "main", "filename": "result.docx", "parts": ["body"]}]},
    }), encoding="utf-8")
    return manifest


def _assemble_preserve_case(
    root: Path,
    *,
    sizes=(12, 24),
    with_story=False,
    second_paragraph=False,
    formatting=None,
) -> tuple[Path, object, object]:
    source = root / "source.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    for index, size in enumerate(sizes):
        run = paragraph.add_run(f"S0 protected source span {index}. ")
        run.font.name = "Arial"
        run.font.size = Pt(size)
    if second_paragraph:
        managed = doc.add_paragraph("S1 managed source paragraph")
        managed.runs[0].font.name = "Times New Roman"
        managed.runs[0].font.size = Pt(9)
    if with_story:
        section = doc.sections[0]
        section.header.paragraphs[0].text = "S0 DEFAULT HEADER"
        section.footer.paragraphs[0].text = "S0 DEFAULT FOOTER"
        section.header_distance = Mm(12)
        section.footer_distance = Mm(18)
    doc.save(source)
    manifest = _write_manifest(root, source, formatting=formatting)
    plan = UnifiedSynthesizer.prepare_build(source, manifest)
    rendered = UnifiedSynthesizer.render_body(plan, root / "rendered")
    output = root / "assembled.docx"
    assemble_document(plan.config, plan.config.documents[0], rendered.path, rendered.nodes, {}, output)
    return output, plan, rendered


def _format_report(path: Path, plan, rendered):
    context = build_verification_context(plan, rendered, plan.config.documents[0], path)
    return verify_delivery_format(
        path,
        plan.config.resolved_format,
        verification_context=context,
    )


class RemediationContractTest(unittest.TestCase):
    def _assemble_linked_story_case(self, root: Path):
        source = root / "linked-stories.docx"
        doc = Document()
        first = doc.sections[0]
        first.header.paragraphs[0].text = "S2 SHARED HEADER"
        first.footer.paragraphs[0].text = "S2 SHARED FOOTER"
        doc.add_paragraph("S2 first section body")
        linked = doc.add_section(1)
        doc.add_paragraph("S2 linked second section body")
        separate = doc.add_section(1)
        separate.header.is_linked_to_previous = False
        separate.footer.is_linked_to_previous = False
        separate.header.paragraphs[0].text = "S2 SEPARATE HEADER"
        separate.footer.paragraphs[0].text = "S2 SEPARATE FOOTER"
        doc.add_paragraph("S2 separate third section body")
        doc.save(source)
        manifest = _write_manifest(root, source)
        plan = UnifiedSynthesizer.prepare_build(source, manifest)
        rendered = UnifiedSynthesizer.render_body(plan, root / "rendered")
        output = root / "linked-stories-assembled.docx"
        assemble_document(plan.config, plan.config.documents[0], rendered.path, rendered.nodes, {}, output)
        return output, plan, rendered

    def test_S0_scenario_contracts_are_explicit_and_evidence_backed(self):
        payload = json.loads(CONTRACTS.read_text(encoding="utf-8"))
        self.assertEqual(payload["schema_version"], 1)
        defects = {item["defect"] for item in payload["scenarios"]}
        self.assertEqual(defects, {"D01", "D02", "D03", "D04", "D05", "D06"})
        for item in payload["scenarios"]:
            for key in ("id", "defect", "preconditions", "fixture_assertions", "command", "artifact_assertions", "required", "evidence"):
                self.assertIn(key, item)
            self.assertTrue(item["required"])
            self.assertEqual(item["status"], "verified")
            self.assertTrue(item["evidence"])
            self.assertTrue(
                "historical_probe" in item["evidence"]
                or "current_tests" in item["evidence"]
                or "current_test" in item["evidence"]
            )

    def test_S0_D01_mixed_inline_sizes_are_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp))
            report = _format_report(output, plan, rendered)
            self.assertTrue(report.passed, report.to_dict())
            self.assertGreaterEqual(
                sum(item["attribute"].endswith("size_pt") for item in report.details["items"]),
                2,
            )

    def test_S0_D01_corrupted_second_inline_size_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp))
            changed = Document(output)
            changed.paragraphs[0].runs[1].font.size = Pt(12)
            changed.save(Path(tmp) / "corrupted.docx")
            report = _format_report(Path(tmp) / "corrupted.docx", plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S1_T02_corrupted_second_inline_font_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp))
            changed = Document(output)
            changed.paragraphs[0].runs[1].font.name = "Courier New"
            changed.save(Path(tmp) / "corrupted-font.docx")
            report = _format_report(Path(tmp) / "corrupted-font.docx", plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S1_T03_same_format_run_split_is_accepted(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp), sizes=(12,))
            changed = Document(output)
            paragraph = changed.paragraphs[0]
            text = paragraph.runs[0].text
            paragraph.runs[0].text = text[:10]
            split = paragraph.add_run(text[10:])
            split.font.name = "Arial"
            split.font.size = Pt(12)
            changed.save(Path(tmp) / "split.docx")
            report = _format_report(Path(tmp) / "split.docx", plan, rendered)
            self.assertTrue(report.passed, report.to_dict())

    def test_S1_T04_same_format_adjacent_runs_can_merge(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp), sizes=(12, 12))
            changed = Document(output)
            paragraph = changed.paragraphs[0]
            merged_text = "".join(run.text or "" for run in paragraph.runs)
            for run in list(paragraph.runs)[1:]:
                run._element.getparent().remove(run._element)
            paragraph.runs[0].text = merged_text
            changed.save(Path(tmp) / "merged.docx")
            report = _format_report(Path(tmp) / "merged.docx", plan, rendered)
            self.assertTrue(report.passed, report.to_dict())

    def test_S1_T05_whitespace_tab_and_break_are_part_of_logical_coordinates(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("left ")
            first.font.name = "Arial"
            first.font.size = Pt(12)
            first.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
            first.add_tab()
            second = paragraph.add_run("middle")
            second.font.name = "Arial"
            second.font.size = Pt(18)
            second.add_break()
            third = paragraph.add_run("right")
            third.font.name = "Arial"
            third.font.size = Pt(12)
            doc.save(source)
            manifest = _write_manifest(root, source)
            plan = UnifiedSynthesizer.prepare_build(source, manifest)
            rendered = UnifiedSynthesizer.render_body(plan, root / "rendered")
            output = root / "assembled.docx"
            assemble_document(plan.config, plan.config.documents[0], rendered.path, rendered.nodes, {}, output)
            report = _format_report(output, plan, rendered)
            self.assertTrue(report.passed, report.to_dict())
            self.assertGreaterEqual(
                sum(item["attribute"].startswith("run[") for item in report.details["items"]),
                6,
            )

    def test_S1_T06_corrupted_subinterval_after_break_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("before")
            first.font.name = "Arial"
            first.font.size = Pt(12)
            first.add_break()
            second = paragraph.add_run("after")
            second.font.name = "Arial"
            second.font.size = Pt(18)
            doc.save(source)
            manifest = _write_manifest(root, source)
            plan = UnifiedSynthesizer.prepare_build(source, manifest)
            rendered = UnifiedSynthesizer.render_body(plan, root / "rendered")
            output = root / "assembled.docx"
            assemble_document(plan.config, plan.config.documents[0], rendered.path, rendered.nodes, {}, output)
            changed = Document(output)
            changed.paragraphs[0].runs[-1].font.size = Pt(12)
            changed.save(root / "corrupted-break.docx")
            report = _format_report(root / "corrupted-break.docx", plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S1_T07_mixed_protected_scope_rejects_protected_override(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            formatting = {
                "mode": "mixed",
                "page_policy": "target",
                "scopes": [{
                    "file": "source.docx",
                    "node_range": "/w:document/w:body/w:p[1]",
                    "mode": "preserve",
                }],
            }
            output, plan, rendered = _assemble_preserve_case(
                root,
                second_paragraph=True,
                formatting=formatting,
            )
            report = _format_report(output, plan, rendered)
            self.assertTrue(report.passed, report.to_dict())
            changed = Document(output)
            changed.paragraphs[0].runs[1].font.size = Pt(12)
            changed.save(root / "mixed-corrupted.docx")
            corrupted_report = _format_report(root / "mixed-corrupted.docx", plan, rendered)
            self.assertFalse(corrupted_report.passed, corrupted_report.to_dict())

    def test_S0_D02_header_content_is_a_protected_story(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp), sizes=(12, 12), with_story=True)
            changed = Document(output)
            changed.sections[0].header.paragraphs[0].text = "S0 WRONG HEADER"
            changed.save(Path(tmp) / "wrong-header.docx")
            report = _format_report(Path(tmp) / "wrong-header.docx", plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S0_D02_header_distance_is_a_protected_geometry(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp), sizes=(12, 12), with_story=True)
            changed = Document(output)
            changed.sections[0].header_distance = Mm(40)
            changed.save(Path(tmp) / "wrong-distance.docx")
            report = _format_report(Path(tmp) / "wrong-distance.docx", plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S2_header_effective_font_is_a_protected_story_property(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(
                Path(tmp), sizes=(12, 12), with_story=True
            )
            changed = Document(output)
            changed.sections[0].header.paragraphs[0].runs[0].font.size = Pt(24)
            changed.save(Path(tmp) / "wrong-header-font.docx")
            report = _format_report(Path(tmp) / "wrong-header-font.docx", plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S2_shared_and_unlinked_story_bindings_are_distinct(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output, plan, rendered = self._assemble_linked_story_case(root)
            report = _format_report(output, plan, rendered)
            self.assertTrue(report.passed, report.to_dict())
            context = build_verification_context(plan, rendered, plan.config.documents[0], output)
            expected_bindings = {
                (item["section_index"], item["story_type"], item["variant"]): item
                for item in context.expected_delivery.story_expectations
            }
            self.assertTrue(expected_bindings[(1, "header", "default")]["linked_to_previous"])
            self.assertFalse(expected_bindings[(2, "header", "default")]["linked_to_previous"])
            self.assertGreaterEqual(len(report.details["story_checks"]), 6)
            changed = Document(output)
            changed.sections[2].header.paragraphs[0].text = "S2 WRONG SEPARATE HEADER"
            changed.save(root / "wrong-separate-story.docx")
            corrupted = _format_report(root / "wrong-separate-story.docx", plan, rendered)
            self.assertFalse(corrupted.passed, corrupted.to_dict())

    def test_S2_footer_content_and_distance_are_protected(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, plan, rendered = _assemble_preserve_case(Path(tmp), sizes=(12, 12), with_story=True)
            changed = Document(output)
            changed.sections[0].footer.paragraphs[0].text = "S0 WRONG FOOTER"
            changed.sections[0].footer_distance = Mm(40)
            changed.save(Path(tmp) / "wrong-footer.docx")
            report = _format_report(Path(tmp) / "wrong-footer.docx", plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S2_missing_story_relationship_target_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output, plan, rendered = _assemble_preserve_case(
                root, sizes=(12, 12), with_story=True
            )
            corrupted_path = root / "wrong-story-target.docx"
            with zipfile.ZipFile(output) as source_zip, zipfile.ZipFile(corrupted_path, "w") as target_zip:
                for info in source_zip.infolist():
                    payload = source_zip.read(info.filename)
                    if info.filename == "word/_rels/document.xml.rels":
                        text = payload.decode("utf-8")
                        self.assertIn("header1.xml", text)
                        self.assertIn("footer1.xml", text)
                        payload = text.replace("header1.xml", "footer1.xml", 1).encode("utf-8")
                    target_zip.writestr(info, payload)
            report = _format_report(corrupted_path, plan, rendered)
            self.assertFalse(report.passed, report.to_dict())

    def test_S0_D03_extra_image_instance_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            paragraph = doc.add_paragraph("S0 image source")
            paragraph.add_run().add_picture(str(IMAGE), width=Mm(15))
            doc.save(source)
            expected = ExpectedInventory(extract_semantic_inventory(source))
            clone = copy.deepcopy(next(doc.element.body.iter(qn("w:drawing"))))
            clone.find(".//" + qn("wp:docPr")).set("id", "999")
            doc.paragraphs[0]._p.append(clone)
            corrupted = root / "extra-image.docx"
            doc.save(corrupted)
            try:
                accepted = verify_content_integrity(expected, corrupted)
            except Exception:
                accepted = False
            self.assertFalse(accepted)

    def test_S3_legal_image_instance_is_accepted(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            paragraph = doc.add_paragraph("S0 image source")
            paragraph.add_run().add_picture(str(IMAGE), width=Mm(15))
            doc.save(source)
            expected = ExpectedInventory(extract_semantic_inventory(source))
            self.assertTrue(verify_content_integrity(expected, source))

    def test_S0_D04_missing_intermediate_level_is_not_compressed(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "missing-level.docx"
            doc = Document()
            for heading, size in (("1. Introduction", 18), ("1.1.1. Detailed protocol", 14)):
                run = doc.add_paragraph().add_run(heading)
                run.font.name = "Arial"
                run.font.size = Pt(size)
                run.bold = True
                for _ in range(3):
                    body = doc.add_paragraph().add_run(
                        "S0 body paragraph with sufficient ordinary text for the body baseline. " * 2
                    )
                    body.font.name = "Arial"
                    body.font.size = Pt(12)
            doc.save(source)
            report = analyze_format_sample(source)
            evidence = report["unstructured_candidate_evidence"]
            second = evidence["/w:document/w:body/w:p[5]"]
            roles = {item["role"] for item in second["candidates"]}
            self.assertIn("heading.3", roles, evidence)
            self.assertNotEqual(second["selected_role"], "heading.2", evidence)
            self.assertEqual(second["selected_role"], "heading.3", evidence)


if __name__ == "__main__":
    unittest.main()
