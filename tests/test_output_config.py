"""Delivery contracts and explicit overrides do not require Office."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from contextlib import ExitStack

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from lib.config import ConfigError, ProjectConfig, load_project_config
from lib.engine import UnifiedSynthesizer


def document(doc_id="main", filename="main.docx", parts=None, **options):
    return {"id": doc_id, "filename": filename, "parts": parts or ["cover", "toc", "body"], **options}


class OutputConfigTests(unittest.TestCase):
    def config(self, documents=None, **extra):
        raw = {"schema_version": 2, "project_name": "示例", **extra}
        if documents is not None:
            raw["output"] = {"documents": documents}
        return ProjectConfig(raw)

    def test_v2_default_is_one_complete_document(self):
        config = self.config()
        self.assertEqual(config.documents, [document(filename="示例_完整文档.docx", toc={"reference": "main", "links": "internal"})])
        self.assertEqual(config.schema_version, 2)
        self.assertEqual(config.warnings, [])

    def test_v1_preserves_all_legacy_deliveries(self):
        config = ProjectConfig({"schema_version": 1, "project_name": "示例"})
        self.assertEqual([doc["id"] for doc in config.documents], ["compiled_document", "cover_toc", "toc_body"])
        self.assertEqual(config.documents[1]["toc"], {"reference": "toc_body", "links": "none"})
        self.assertEqual(config.output["compiled_document"], "示例_合成材料.docx")
        self.assertTrue(config.warnings)

    def test_split_documents_replace_default(self):
        config = self.config([
            document("front", "front.docx", ["cover", "toc"], toc={"reference": "main"}),
            document(parts=["toc", "body"]),
        ])
        self.assertEqual(len(config.documents), 2)
        self.assertEqual(config.documents[0]["toc"]["links"], "none")

    def test_parts_allow_all_nonempty_ordered_subsets(self):
        for parts in (["cover"], ["body"], ["cover", "body"], ["toc", "body"], ["cover", "toc", "body"]):
            with self.subTest(parts=parts):
                self.assertEqual(self.config([document(parts=parts)]).documents[0]["parts"], parts)
        for parts in ([], ["body", "cover"], ["body", "body"], ["other"], "body", None):
            with self.subTest(parts=parts), self.assertRaises(ConfigError):
                self.config([{**document(), "parts": parts}])

    def test_empty_null_and_wrong_shape_rejected(self):
        for value in ([], None, "main", {}):
            with self.subTest(value=value), self.assertRaises(ConfigError):
                self.config(output={"documents": value})
        for value in (None, [], "none"):
            with self.subTest(value=value), self.assertRaises(ConfigError):
                self.config(output=value)

    def test_invalid_references_and_links_rejected(self):
        cases = [
            [document(parts=["toc"])],
            [document(parts=["toc"], toc={"reference": "missing"})],
            [document("front", "front.docx", ["toc"], toc={"reference": "main", "links": "internal"}), document()],
            [document(toc={"reference": "elsewhere"})],
            [document(toc={"links": "external"})],
            [document(parts=["body"], toc={})],
            [document("front", "front.docx", ["toc"], toc={"reference": "cover"}), document("cover", "cover.docx", ["cover"])],
        ]
        for docs in cases:
            with self.subTest(documents=docs), self.assertRaises(ConfigError):
                self.config(docs)

    def test_filenames_and_ids_are_safe_and_unique(self):
        bad_filenames = (
            "../bad.docx", "/bad.docx", "dir/bad.docx", "dir\\bad.docx", "C:bad.docx", "bad.pdf",
            "~$bad.docx", "bad\x00.docx",
            # Windows reserved device names
            "CON.docx", "con.docx", "PRN.docx", "aux.docx", "NUL.docx", "com1.docx", "COM9.docx", "lpt1.docx",
            # Windows forbidden characters
            "bad<name.docx", "bad>name.docx", "bad:name.docx", "bad\"name.docx",
            "bad|name.docx", "bad?name.docx", "bad*name.docx",
            # Trailing dots and spaces
            "bad .docx", "bad..docx", "bad.docx.", "bad.docx "
        )
        for filename in bad_filenames:
            with self.subTest(filename=filename), self.assertRaises(ConfigError):
                self.config([document(filename=filename)])

        # Valid filenames must pass
        for filename in ("main.docx", "report-v1.docx", "2026_公文.docx", "Document (Final).docx"):
            with self.subTest(filename=filename):
                self.assertEqual(self.config([document(filename=filename)]).documents[0]["filename"], filename)

        for docs in ([document(), document("second", "MAIN.DOCX")], [document(), document("MAIN", "second.docx")]):
            with self.subTest(documents=docs), self.assertRaises(ConfigError):
                self.config(docs)
        with self.assertRaises(ConfigError):
            ProjectConfig({"schema_version": 1, "project_name": "示例", "output": {"compiled_document": "示例_目录+正文.docx"}})

    def test_unknown_fields_and_old_output_keys_rejected(self):
        for docs in ([document(unknown=True)], [document(toc={"unknown": True})]):
            with self.subTest(documents=docs), self.assertRaises(ConfigError):
                self.config(docs)
        with self.assertRaises(ConfigError):
            self.config(output={"toc_body": "body.docx"})

    def test_project_name_is_safe_even_when_output_filenames_are_explicit(self):
        bad_names = (
            None, "", ".", "..", "../escape", "/absolute", "nested/project", "nested\\project", "C:escape",
            ".hidden", "~$lock", " leading", "trailing ", "line\nfeed", "tab\tname", "nul\x00name", "del\x7fname",
            # Windows reserved device names
            "CON", "con", "PRN", "AUX", "aux", "NUL", "COM1", "com9", "LPT1", "lpt8", "con.nested",
            # Windows forbidden characters
            "proj<name", "proj>name", "proj:name", "proj\"name", "proj|name", "proj?name", "proj*name",
            # Trailing dot
            "project.", "project..", "project. "
        )
        for name in bad_names:
            with self.subTest(name=name), self.assertRaisesRegex(ConfigError, "project_name"):
                self.config([document(filename="safe.docx")], project_name=name)
        for name in ("普通项目", "My Project", "Project.v2", "2026-09_材料"):
            with self.subTest(name=name):
                self.assertEqual(self.config(project_name=name).project_name, name)

    def test_profile_lookup_name_is_validated_before_reading_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ConfigError, "project_name"):
                load_project_config(Path(directory), project_name="../escape")

    def test_override_deep_merge_arrays_replace_and_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manifest = root / "manifest.json"
            first, second = root / "first.json", root / "second.json"
            manifest.write_text(json.dumps({"schema_version": 2, "project_name": "例", "cover": {"author": "保留", "date": "旧"}, "tree_order": ["a", "b"], "output": {"documents": [document("old", "old.docx"), document("second", "second.docx")]}}), encoding="utf-8")
            first.write_text(json.dumps({"cover": {"date": "新"}, "output": {"documents": [document(parts=["body"])]}, "tree_order": ["c"]}), encoding="utf-8")
            second.write_text(json.dumps({"cover": {"template": False}}), encoding="utf-8")
            config = load_project_config(root, override_paths=[first, second])
            self.assertEqual(config.cover["author"], "保留")
            self.assertEqual(config.cover["date"], "新")
            self.assertFalse(config.cover["template"])
            self.assertEqual(config.tree_order, ["c"])
            self.assertEqual(config.documents, [document(parts=["body"])])
            self.assertEqual(config.manifest_path, manifest.resolve())
            self.assertEqual(config.override_paths, [first.resolve(), second.resolve()])

    def test_override_null_is_not_deletion_and_missing_file_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            patch = root / "patch.json"
            patch.write_text('{"output":null}', encoding="utf-8")
            with self.assertRaises(ConfigError):
                load_project_config(root, override_paths=[patch])
            with self.assertRaises(ConfigError):
                load_project_config(root, override_paths=[root / "missing.json"])

    def test_schema_upgrade_rejects_retained_legacy_output_keys(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "manifest.json").write_text(json.dumps({"schema_version": 1, "project_name": "例", "output": {"toc_body": "old.docx"}}), encoding="utf-8")
            patch = root / "patch.json"
            patch.write_text('{"schema_version":2}', encoding="utf-8")
            with self.assertRaisesRegex(ConfigError, "迁移"):
                load_project_config(root, override_paths=[patch])

    def test_no_configuration_uses_v2(self):
        with tempfile.TemporaryDirectory() as directory:
            config = load_project_config(Path(directory), project_name="fictional-test-no-profile")
            self.assertEqual(config.schema_version, 2)
            self.assertIsNone(config.manifest_path)
            self.assertEqual(len(config.documents), 1)

    def test_plan_is_read_only_for_all_source_strategies_and_overrides(self):
        for strategy in ("directory_tree", "explicit_tree", "highlighted_docx"):
            with self.subTest(strategy=strategy), tempfile.TemporaryDirectory() as directory:
                root = Path(directory).resolve()
                source = root / "source"
                source.mkdir()
                fixture = source / "1_fixture.docx"
                doc = Document()
                run = doc.add_paragraph().add_run("一、Fictional heading")
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                doc.add_paragraph("Fictional body")
                doc.save(fixture)
                raw = {"schema_version": 2, "project_name": "ReadOnlyPlan", "cover": {"template": False}, "source": {"strategy": strategy}}
                if strategy == "explicit_tree":
                    raw["source"]["tree"] = [{"type": "docx", "file": fixture.name, "title": "Explicit heading", "level": 1}]
                manifest, override = root / "manifest.json", root / "override.json"
                manifest.write_text(json.dumps(raw), encoding="utf-8")
                override.write_text(json.dumps({"output": {"documents": [document(parts=["body"])]}}), encoding="utf-8")
                before = {p.relative_to(root): p.read_bytes() for p in root.rglob("*") if p.is_file()}
                with ExitStack() as stack:
                    for target in ("lib.engine._run_applescript", "lib.engine.prepare_conversions", "lib.pagination.inspect_document", "subprocess.run", "docx.document.Document.save", "pathlib.Path.mkdir"):
                        stack.enter_context(patch(target, side_effect=AssertionError(f"plan must not call {target}")))
                    plan = UnifiedSynthesizer.plan(fixture if strategy == "highlighted_docx" else source, manifest, override_paths=[override])
                after = {p.relative_to(root): p.read_bytes() for p in root.rglob("*") if p.is_file()}
                self.assertEqual(before, after)
                self.assertGreater(plan["node_count"], 0)
                self.assertEqual(plan["documents"], [document(parts=["body"])])
                self.assertEqual(plan["manifest_path"], str(manifest.resolve()))
                self.assertEqual(plan["override_paths"], [str(override.resolve())])

    def test_cover_free_plan_does_not_resolve_template(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            doc = Document()
            doc.add_paragraph("Fictional body")
            doc.save(root / "fixture.docx")
            manifest = root / "manifest.json"
            manifest.write_text(json.dumps({"schema_version": 2, "project_name": "NoCover", "output": {"documents": [document(parts=["body"])]}}), encoding="utf-8")
            with patch.object(ProjectConfig, "get_template_path", side_effect=AssertionError("cover-free plan must not resolve template")):
                plan = UnifiedSynthesizer.plan(root, manifest)
            self.assertIsNone(plan["template"])


if __name__ == "__main__":
    unittest.main()
