"""Tests for Page Sequences (layout.page_sequences), Roman Numerals, and Word QA."""

from pathlib import Path
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib.composition import assemble_document, add_bookmark
from lib.config import ProjectConfig, ConfigError
from lib.delivery import (
    _check_page_map,
    _enrich_page_map_labels,
    build_deliveries,
    validate_delivery_structure,
    validate_measured_delivery,
)
from lib.document_parts import DocumentPart
from lib.pagination import extract_pdf_page_labels, page_records_from_map
from lib.pagination_types import (
    PageSequence,
    PageRecord,
    format_page_number,
    int_to_roman,
    roman_to_int,
    int_to_letters,
)
from lib.qa import OfficeExportError, run_qa_assertions


def make_v3_manifest(**kwargs):
    base = {
        "schema_version": 3,
        "project_name": "Test Project",
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": {"mode": "preserve"},
        "cover": {"template": False},
    }
    base.update(kwargs)
    return base


class NumeralFormattingTest(unittest.TestCase):
    def test_page_record_tracks_unverified_and_verified_labels(self):
        pending = PageRecord.from_dict(
            {"physical_page": 2, "printed_page": 1, "expected_label": "a"},
            section_id=2,
            sequence_id="front",
        )
        self.assertEqual(pending.verification_status, "unverified")
        verified = PageRecord.from_dict(
            {"physical_page": 2, "printed_page": 1, "expected_label": "a", "observed_label": "a"},
            section_id=2,
            sequence_id="front",
        )
        self.assertEqual(verified.verification_status, "verified")
        self.assertEqual(verified.to_dict()["number_value"], 1)

    def test_page_record_conversion_preserves_part_sequence(self):
        parts = {
            "preface": type("Part", (), {"page_sequence": "front"})(),
            "body": type("Part", (), {"page_sequence": "main"})(),
        }
        sequences = {
            "front": PageSequence("front", "lowerLetter", 1),
            "main": PageSequence("main", "decimal", 1),
        }
        records = page_records_from_map(
            {
                "_Synth_part_preface_start": {"physical_page": 1, "printed_page": 1, "expected_label": "a"},
                "_Synth_body": {"physical_page": 3, "printed_page": 1, "expected_label": "1"},
            },
            part_boundaries={"preface": "_Synth_part_preface_start", "body": "_Synth_body"},
            part_specs=parts,
            sequences=sequences,
        )
        self.assertEqual(records["_Synth_part_preface_start"].sequence_id, "front")
        self.assertEqual(records["_Synth_body"].expected_label, "1")

    def test_unnumbered_cover_does_not_require_visible_label(self):
        spec = {"parts": ["cover", "toc", "body"]}
        parts = {
            "cover": SimpleNamespace(page_sequence=None, kind="cover"),
            "toc": SimpleNamespace(page_sequence="front", kind="generated_toc"),
            "body": SimpleNamespace(page_sequence="main", kind="content"),
        }
        sequences = {
            "front": PageSequence("front", "lowerRoman", 1),
            "main": PageSequence("main", "decimal", 1),
        }
        page_map = {
            "_Synth_cover": {"physical_page": 1, "printed_page": 1},
            "_Synth_toc": {"physical_page": 2, "printed_page": 1, "observed_label": "i"},
            "_Synth_body": {"physical_page": 4, "printed_page": 1, "observed_label": "1"},
        }

        _enrich_page_map_labels(page_map, spec, [], parts, sequences)
        self.assertFalse(page_map["_Synth_cover"]["label_required"])
        self.assertEqual(page_map["_Synth_cover"]["expected_label"], "")
        _check_page_map(
            page_map,
            ["_Synth_cover", "_Synth_toc", "_Synth_body"],
            require_observed_labels=True,
        )

    def test_pdf_page_labels_support_letter_sequences(self):
        with tempfile.TemporaryDirectory() as td:
            pdf_path = Path(td) / "letters.pdf"
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 750), "- A -", fontsize=10)
                pdf.save(pdf_path)
            self.assertEqual(extract_pdf_page_labels(pdf_path), {1: "A"})

    def test_roman_numerals(self):
        self.assertEqual(int_to_roman(1), "I")
        self.assertEqual(int_to_roman(4), "IV")
        self.assertEqual(int_to_roman(9), "IX")
        self.assertEqual(int_to_roman(14), "XIV")
        self.assertEqual(int_to_roman(42), "XLII")
        self.assertEqual(int_to_roman(99), "XCIX")
        self.assertEqual(int_to_roman(2026), "MMXXVI")

        self.assertEqual(roman_to_int("i"), 1)
        self.assertEqual(roman_to_int("iv"), 4)
        self.assertEqual(roman_to_int("XIV"), 14)
        self.assertEqual(roman_to_int("xlii"), 42)
        self.assertIsNone(roman_to_int("not_a_roman"))
        self.assertIsNone(roman_to_int(""))

    def test_letter_numerals(self):
        self.assertEqual(int_to_letters(1, lowercase=True), "a")
        self.assertEqual(int_to_letters(26, lowercase=True), "z")
        self.assertEqual(int_to_letters(27, lowercase=True), "aa")
        self.assertEqual(int_to_letters(1, lowercase=False), "A")

    def test_format_page_number(self):
        self.assertEqual(format_page_number(5, "decimal"), "5")
        self.assertEqual(format_page_number(5, "lowerRoman"), "v")
        self.assertEqual(format_page_number(5, "upperRoman"), "V")
        self.assertEqual(format_page_number(5, "lowerLetter"), "e")
        self.assertEqual(format_page_number(5, "upperLetter"), "E")

    def test_page_sequence_validation(self):
        seq = PageSequence(id="front", format="lowerRoman", start=1)
        self.assertEqual(seq.format, "lowerRoman")
        self.assertEqual(seq.start, 1)

        with self.assertRaises(ValueError):
            PageSequence(id="invalid", format="invalid_format")
        with self.assertRaises(ValueError):
            PageSequence(id="invalid", format="decimal", start=0)


class PageSequenceConfigTest(unittest.TestCase):
    def test_valid_page_sequences(self):
        cfg = ProjectConfig(make_v3_manifest(
            layout={
                "page_sequences": {
                    "front": {"format": "lowerRoman", "start": 1},
                    "main": {"format": "decimal", "start": 1},
                },
                "parts": {
                    "preface": {"kind": "content", "page_sequence": "front"},
                    "body": {"kind": "content", "page_sequence": "main"},
                }
            },
            output={
                "documents": [{"id": "main", "filename": "main.docx", "parts": ["preface", "body"]}]
            },
        ))
        self.assertIn("front", cfg.page_sequences)
        self.assertIn("main", cfg.page_sequences)
        self.assertEqual(cfg.page_sequences["front"].format, "lowerRoman")
        self.assertEqual(cfg.page_sequences["main"].format, "decimal")

    def test_invalid_sequence_format(self):
        with self.assertRaises(ConfigError) as ctx:
            ProjectConfig(make_v3_manifest(
                layout={
                    "page_sequences": {
                        "front": {"format": "hexadecimal"}
                    }
                }
            ))
        self.assertIn("format 无效", str(ctx.exception))

    def test_part_references_undeclared_sequence(self):
        with self.assertRaises(ConfigError) as ctx:
            ProjectConfig(make_v3_manifest(
                layout={
                    "page_sequences": {
                        "front": {"format": "lowerRoman"}
                    },
                    "parts": {
                        "preface": {"kind": "content", "page_sequence": "nonexistent"}
                    }
                },
                output={
                    "documents": [{"id": "main", "filename": "main.docx", "parts": ["preface"]}]
                }
            ))
        self.assertIn("未在 layout.page_sequences 中定义", str(ctx.exception))


class PageSequenceAssemblyAndQATest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)

        # Create a body docx with bookmarks
        self.body_path = self.root / "body.docx"
        doc = Document()
        p1 = doc.add_paragraph("Abstract Heading")
        add_bookmark(p1._p, "_Toc_auto_001", 1)
        doc.add_paragraph("Abstract content goes here.")
        p2 = doc.add_paragraph("Chapter 1")
        add_bookmark(p2._p, "_Toc_auto_002", 2)
        doc.add_paragraph("Chapter 1 body content.")
        doc.save(self.body_path)

        self.nodes = [
            {"title": "Abstract Heading", "level": 1, "bookmark_name": "_Toc_auto_001", "bm_id": 1},
            {"title": "Chapter 1", "level": 1, "bookmark_name": "_Toc_auto_002", "bm_id": 2},
        ]

    def test_assemble_with_roman_and_decimal_sequences(self):
        cfg = ProjectConfig(make_v3_manifest(
            source={
                "regions": {
                    "abstract": {
                        "start": "/w:document/w:body/w:p[1]",
                        "end": "/w:document/w:body/w:p[3]",
                    },
                    "main": {
                        "start": "/w:document/w:body/w:p[3]",
                        "end": "end_of_document",
                    },
                }
            },
            layout={
                "page_sequences": {
                    "front": {"format": "lowerRoman", "start": 1},
                    "main": {"format": "decimal", "start": 1},
                },
                "parts": {
                    "preface": {"kind": "content", "source_region": "abstract", "page_sequence": "front"},
                    "toc": {"kind": "generated_toc", "page_sequence": "front"},
                    "body": {"kind": "content", "source_region": "main", "page_sequence": "main"},
                }
            },
            output={
                "documents": [{"id": "main", "filename": "main.docx", "parts": ["preface", "toc", "body"]}]
            },
        ))

        out_path = self.root / "assembled.docx"
        spec = cfg.documents[0]
        # pages mapping with expected roman and decimal labels
        pages = {
            "_Toc_auto_001": {"physical_page": 1, "printed_page": 1, "expected_label": "i"},
            "_Toc_auto_002": {"physical_page": 3, "printed_page": 1, "expected_label": "1"},
        }
        assemble_document(
            config=cfg,
            spec=spec,
            body_path=self.body_path,
            nodes=self.nodes,
            pages=pages,
            out_path=out_path,
        )
        self.assertTrue(out_path.exists())

        # Inspect resulting document
        doc = Document(str(out_path))
        # Find all pgNumType elements
        pg_num_types = doc.element.body.xpath(".//w:pgNumType")
        self.assertTrue(len(pg_num_types) >= 2)
        fmts = [p.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fmt") for p in pg_num_types]
        self.assertIn("lowerRoman", fmts)
        self.assertIn("decimal", fmts)

    def test_single_page_short_body_is_not_an_orphan_tail(self):
        pdf_path = self.root / "short-body.pdf"
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 72), "短文档正文", fontsize=12)
            page.insert_text((72, 750), "- 1 -", fontsize=10)
            pdf.save(str(pdf_path))

        self.assertTrue(run_qa_assertions(str(pdf_path), ignore_front_pages=0))

    def test_odd_page_break_whitelists_blank_page(self):
        cfg = ProjectConfig(make_v3_manifest(
            source={
                "regions": {
                    "abstract": {
                        "start": "/w:document/w:body/w:p[1]",
                        "end": "/w:document/w:body/w:p[3]",
                    },
                    "main": {
                        "start": "/w:document/w:body/w:p[3]",
                        "end": "end_of_document",
                    },
                }
            },
            layout={
                "page_sequences": {
                    "front": {"format": "lowerRoman", "start": 1},
                    "main": {"format": "decimal", "start": 1},
                },
                "parts": {
                    "preface": {"kind": "content", "source_region": "abstract", "page_sequence": "front"},
                    "body": {"kind": "content", "source_region": "main", "page_sequence": "main", "section_type": "oddPage"},
                }
            },
            output={
                "documents": [{"id": "main", "filename": "main.docx", "parts": ["preface", "body"]}]
            },
        ))

        pdf_path = self.root / "test_odd.pdf"
        # Create a 3-page PDF where page 2 is blank due to oddPage break
        with pymupdf.open() as pdf:
            # Page 1: preface
            p1 = pdf.new_page()
            p1.insert_text((72, 72), "Preface Line 1\nPreface Line 2\nPreface Line 3\nPreface Line 4", fontsize=12)
            p1.insert_text((72, 750), "- i -", fontsize=10)
            # Page 2: intentionally blank page inserted by oddPage break!
            pdf.new_page()
            # Page 3: body
            p3 = pdf.new_page()
            p3.insert_text((72, 72), "Chapter Line 1\nChapter Line 2\nChapter Line 3\nChapter Line 4", fontsize=12)
            p3.insert_text((72, 750), "- 1 -", fontsize=10)
            pdf.save(str(pdf_path))

        # Without allowed_blank_pages, QA detects page 2 as blank
        self.assertFalse(run_qa_assertions(str(pdf_path), ignore_front_pages=0))

        # With allowed_blank_pages={2}, QA passes!
        self.assertTrue(run_qa_assertions(str(pdf_path), ignore_front_pages=0, allowed_blank_pages={2}))

    def test_validate_measured_delivery_with_page_sequences(self):
        cfg = ProjectConfig(make_v3_manifest(
            source={
                "regions": {
                    "abstract": {
                        "start": "/w:document/w:body/w:p[1]",
                        "end": "/w:document/w:body/w:p[3]",
                    },
                    "main": {
                        "start": "/w:document/w:body/w:p[3]",
                        "end": "end_of_document",
                    },
                }
            },
            layout={
                "page_sequences": {
                    "front": {"format": "lowerRoman", "start": 1},
                    "main": {"format": "decimal", "start": 1},
                },
                "parts": {
                    "preface": {"kind": "content", "source_region": "abstract", "page_sequence": "front"},
                    "body": {"kind": "content", "source_region": "main", "page_sequence": "main", "section_type": "oddPage"},
                }
            },
            output={
                "documents": [{"id": "main", "filename": "main.docx", "parts": ["preface", "body"]}]
            },
        ))

        out_path = self.root / "assembled.docx"
        spec = cfg.documents[0]
        pages = {
            "_Toc_auto_001": {"physical_page": 1, "printed_page": 1, "expected_label": "i"},
            "_Toc_auto_002": {"physical_page": 3, "printed_page": 1, "expected_label": "1"},
        }
        assemble_document(
            config=cfg,
            spec=spec,
            body_path=self.body_path,
            nodes=self.nodes,
            pages=pages,
            out_path=out_path,
        )

        pdf_path = self.root / "test_seq.pdf"
        with pymupdf.open() as pdf:
            p1 = pdf.new_page()
            p1.insert_text((72, 72), "Preface Line 1\nPreface Line 2\nPreface Line 3\nPreface Line 4", fontsize=12)
            p1.insert_text((72, 750), "- i -", fontsize=10)
            # Page 2: empty oddPage break
            pdf.new_page()
            p3 = pdf.new_page()
            p3.insert_text((72, 72), "Chapter Line 1\nChapter Line 2\nChapter Line 3\nChapter Line 4", fontsize=12)
            p3.insert_text((72, 750), "- 1 -", fontsize=10)
            pdf.save(str(pdf_path))

        page_map = {
            "_Synth_part_preface_start": {"physical_page": 1, "printed_page": 1, "expected_label": "i"},
            "_Synth_part_preface_end": {"physical_page": 1, "printed_page": 1, "expected_label": "i"},
            "_Synth_body": {"physical_page": 3, "printed_page": 1, "expected_label": "1"},
            "_Toc_auto_001": {"physical_page": 1, "printed_page": 1, "expected_label": "i"},
            "_Toc_auto_002": {"physical_page": 3, "printed_page": 1, "expected_label": "1"},
        }

        # validate_measured_delivery should recognize page 2 as allowed oddPage blank and pass!
        res = validate_measured_delivery(
            out_path, spec, self.nodes, pdf_path, page_map,
            parts_registry=cfg.parts, page_sequences=cfg.page_sequences
        )
        self.assertIn("_Toc_auto_001", res)
        self.assertEqual(res["_Toc_auto_001"]["expected_label"], "i")
        self.assertEqual(res["_Toc_auto_002"]["expected_label"], "1")


if __name__ == "__main__":
    unittest.main()
