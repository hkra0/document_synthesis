"""Bookmark pagination checks; real Word coverage is explicitly opt-in."""

import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib.pagination import _parse_page_map, inspect_document, validate_document_structure
from lib.office import BackendStatus, set_backend
from lib.office.mac_applescript import MacAppleScriptBackend
from lib.qa import OfficeExportError, _word_access_directory


def bookmark(paragraph, name, number):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


class PaginationTest(unittest.TestCase):
    def test_parse_distinguishes_physical_and_printed_pages(self):
        self.assertEqual(_parse_page_map("body\t3\t1\nother\t4\t2\n", ["body", "other"], 5), {
            "body": {"physical_page": 3, "printed_page": 1},
            "other": {"physical_page": 4, "printed_page": 2},
        })

    def test_missing_ambiguous_or_invalid_pages_never_fall_back(self):
        for report in ("", "body\t1\t1\nbody\t2\t2", "body\t6\t1", "body\t1\t0", "body\t1\t?", "other\t1\t1"):
            with self.subTest(report=report), self.assertRaises(OfficeExportError):
                _parse_page_map(report, ["body"], 5)

    def test_structure_rejects_missing_and_duplicate_targets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.docx"
            doc = Document()
            bookmark(doc.add_paragraph("Repeated title"), "first", 1)
            doc.save(path)
            validate_document_structure(path, ["first"])
            with self.assertRaisesRegex(OfficeExportError, "缺少书签"):
                validate_document_structure(path, ["absent"])
            bookmark(doc.add_paragraph("Repeated title"), "first", 2)
            doc.save(path)
            with self.assertRaisesRegex(OfficeExportError, "重复书签"):
                validate_document_structure(path)

    def test_structure_rejects_dangling_internal_link_but_not_external_anchor(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.docx"
            doc = Document()
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), "missing")
            doc.add_paragraph()._p.append(link)
            doc.save(path)
            with self.assertRaisesRegex(OfficeExportError, "超链接"):
                validate_document_structure(path)
            link.set(qn("r:id"), "rIdExternal")
            doc.save(path)
            validate_document_structure(path)

    def test_failed_export_preserves_existing_pdf_and_input(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "fixture.docx", Path(directory) / "existing.pdf"
            Document().save(source)
            original = source.read_bytes()
            output.write_bytes(b"previous verified export")
            completed = type("Completed", (), {"returncode": 1, "stdout": "", "stderr": "failure"})()
            with patch("lib.pagination.word_export_status", return_value=(True, "ok")), patch("lib.pagination.subprocess.run", return_value=completed):
                with self.assertRaises(OfficeExportError):
                    inspect_document(source, output, [])
            self.assertEqual(output.read_bytes(), b"previous verified export")
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(sorted(p.name for p in Path(directory).iterdir()), ["existing.pdf", "fixture.docx"])

    def test_success_checks_actual_export_and_preserves_input(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "fixture.docx", Path(directory) / "output.pdf"
            doc = Document()
            bookmark(doc.add_paragraph("Repeated title"), "body", 1)
            doc.save(source)
            original = source.read_bytes()

            def export(*args, **kwargs):
                private = _word_access_directory()
                inspection = next(private.glob("pagination-*.docx")).with_suffix(".pdf")
                writer = PdfWriter()
                writer.add_blank_page(width=612, height=792)
                writer.write(inspection)
                self.assertIn("start bookmarkStart end bookmarkStart", args[0][-1])
                return type("Completed", (), {"returncode": 0, "stdout": "body\t1\t1\n", "stderr": ""})()

            # Drive the macOS backend on any host: Word is reported installed and
            # osascript is replaced, so this checks the backend wiring offline.
            set_backend(MacAppleScriptBackend())
            try:
                with patch.object(
                    MacAppleScriptBackend, "static_status",
                    return_value=BackendStatus(True, "ok", "exact"),
                ), patch("lib.office.mac_applescript.subprocess.run", side_effect=export):
                    self.assertEqual(inspect_document(source, output, ["body"]), {"body": {"physical_page": 1, "printed_page": 1}})
            finally:
                set_backend(None)
            self.assertTrue(output.is_file())
            self.assertEqual(source.read_bytes(), original)

    @unittest.skipUnless(os.environ.get("DOCUMENT_SYNTHESIS_WORD_TEST") == "1", "opt-in local Word test")
    def test_real_word_repeated_titles_and_section_page_restart(self):
        with tempfile.TemporaryDirectory(prefix="pagination-fictional-", dir=os.environ.get("DOCUMENT_SYNTHESIS_TEST_DIR")) as directory:
            source, output = Path(directory) / "fixture.docx", Path(directory) / "output.pdf"
            doc = Document()
            bookmark(doc.add_paragraph("Fictional cover"), "_Synth_cover", 1)
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            numbering = OxmlElement("w:pgNumType")
            numbering.set(qn("w:start"), "1")
            section._sectPr.append(numbering)
            bookmark(doc.add_paragraph("Repeated title"), "_Toc_auto_001", 2)
            doc.add_page_break()
            bookmark(doc.add_paragraph("Repeated title"), "_Toc_auto_002", 3)
            doc.save(source)
            original = source.read_bytes()
            self.assertEqual(inspect_document(source, output, ["_Synth_cover", "_Toc_auto_001", "_Toc_auto_002"]), {
                "_Synth_cover": {"physical_page": 1, "printed_page": 1},
                "_Toc_auto_001": {"physical_page": 2, "printed_page": 1},
                "_Toc_auto_002": {"physical_page": 3, "printed_page": 2},
            })
            self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
