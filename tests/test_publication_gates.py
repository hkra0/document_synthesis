# -*- coding: utf-8 -*-
"""R0 发布门禁与事务契约。

这里的 Word 边界均被替换为确定的 staging 函数，测试只验证生产引擎是否
在发布前执行完整性/格式门禁，以及失败时是否保留已有成果。
"""

import builtins
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Mm, Pt

from lib.content_integrity import verify_delivery_format
from lib.delivery import CheckResult, DeliveryGateError, DeliveryReport, assert_publishable
from lib.engine import BuildError, UnifiedSynthesizer
from lib.format_schema import (
    HeaderFooterSpec,
    PageSpec,
    ParagraphStyle,
    ResolvedFormat,
    RoleSpec,
    RunStyle,
    StyleDefinition,
    TocSpec,
)
from lib.style_applier import apply_roles


def _manifest(source_name, *, strategy="docx_document", parts=None):
    return {
        "schema_version": 3,
        "project_name": "R0PublicationFixture",
        "source": {"strategy": strategy, "file": source_name},
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": {"mode": "restyle"},
        "output": {
            "documents": [{
                "id": "main",
                "filename": "r0-published.docx",
                "parts": parts or ["body"],
            }]
        },
    }


def _sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _custom_format():
    return ResolvedFormat(
        id="r0-page-format",
        version="1.0.0",
        page=PageSpec(width_mm=210.0, height_mm=297.0),
        styles={
            "body": StyleDefinition(
                name="R0 Body",
                run=RunStyle(east_asia="宋体", latin="Times New Roman", size_pt=12.0),
                paragraph=ParagraphStyle(alignment="justify"),
            )
        },
        roles={"body": RoleSpec(style="body")},
        toc=TocSpec(),
        header=HeaderFooterSpec(),
        footer=HeaderFooterSpec(),
        required_capabilities=["styles.fonts.basic"],
        provenance={},
        content_hash="r0-publication-contract",
    )


class PublicationGateContractTest(unittest.TestCase):
    @unittest.skipUnless(
        os.environ.get("DOCUMENT_SYNTHESIS_WORD_TEST") == "1",
        "BLOCKED: real staging gate injection requires Word Automation",
    )
    def test_real_word_gate_injection_after_staging_blocks_publication(self):
        """S7: corrupt only the real post-Word staging artifact before QA."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source_doc = Document()
            for index in range(12):
                source_doc.add_paragraph(
                    f"S7 staging gate paragraph {index}. "
                    "This fictional paragraph provides enough material for the real Word path. " * 3
                )
            source_doc.save(source)
            manifest = root / "manifest.json"
            manifest.write_text(json.dumps(_manifest(source.name)), encoding="utf-8")
            output = root / "output"
            injection = {}

            from lib.delivery import build_deliveries as real_build_deliveries

            def corrupt_after_real_build(config, body, nodes, run_dir, **kwargs):
                staged = real_build_deliveries(config, body, nodes, run_dir, **kwargs)
                staged_path = Path(staged["main"])
                changed = Document(staged_path)
                changed.paragraphs[0].runs[0].font.size = Pt(70)
                changed.save(staged_path)
                injection.update({"stage": "after_real_build", "path": str(staged_path)})
                return staged

            with patch("lib.delivery.build_deliveries", side_effect=corrupt_after_real_build):
                with self.assertRaises(BuildError):
                    UnifiedSynthesizer.synthesize(source, output, manifest)

            self.assertEqual(injection.get("stage"), "after_real_build")
            self.assertFalse((output / "r0-published.docx").exists())
            self.assertFalse((output / "build-metadata.json").exists())
            diagnostics = json.loads((output / "build-diagnostics.json").read_text(encoding="utf-8"))
            self.assertEqual(diagnostics["qa"]["format_verification"], "failed")
            self.assertIn("70", json.dumps(diagnostics, ensure_ascii=False))

    @unittest.skipUnless(
        os.environ.get("DOCUMENT_SYNTHESIS_WORD_TEST") == "1",
        "BLOCKED: real CLI transaction case requires Word Automation",
    )
    def test_real_cli_second_failure_keeps_previous_delivery_and_metadata(self):
        """S7: a real second CLI invocation fails before replacing prior output."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source_doc = Document()
            source_doc.add_paragraph("S7 fictional publication body")
            source_doc.save(source)
            manifest = root / "manifest.json"
            manifest.write_text(json.dumps(_manifest(source.name)), encoding="utf-8")
            output = root / "output"
            command = [
                sys.executable,
                "synthesize.py",
                "--source",
                str(source),
                "--manifest",
                str(manifest),
                "--output-dir",
                str(output),
            ]
            first = subprocess.run(command, cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True)
            self.assertEqual(first.returncode, 0, first.stdout + first.stderr)
            delivery = output / "r0-published.docx"
            metadata = output / "build-metadata.json"
            before = (_sha256(delivery), _sha256(metadata))

            invalid = _manifest(source.name)
            invalid["formatting"] = {
                "mode": "restyle",
                "role_map": "expired-role-map.json",
                "on_unmapped": "error",
            }
            manifest.write_text(json.dumps(invalid), encoding="utf-8")
            second = subprocess.run(command, cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True)
            self.assertNotEqual(second.returncode, 0, second.stdout + second.stderr)
            self.assertIn("RoleMap 文件不存在", second.stdout + second.stderr)
            self.assertEqual(before, (_sha256(delivery), _sha256(metadata)))

    def test_required_not_run_and_unsupported_checks_block_publication(self):
        for status in ("not_run", "unsupported"):
            with self.subTest(status=status):
                report = DeliveryReport(
                    delivery_id="main",
                    filename="main.docx",
                    checks=(CheckResult("format_verification", status, True, ("fixture",)),),
                )
                with self.assertRaises(DeliveryGateError):
                    assert_publishable([report])

    def test_format_violation_is_a_hard_publication_failure(self):
        """F02 integration without Word：格式报告失败不得被降级为 warnings。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source_doc = Document()
            source_doc.add_paragraph("Publication gate body")
            source_doc.save(source)
            manifest = root / "manifest.json"
            manifest.write_text(json.dumps(_manifest(source.name)), encoding="utf-8")
            output = root / "output"
            output.mkdir()
            old_delivery = output / "r0-published.docx"
            old_meta = output / "build-metadata.json"
            Document().save(old_delivery)
            old_meta.write_text('{"status":"old-success"}', encoding="utf-8")
            before = (_sha256(old_delivery), _sha256(old_meta))

            def corrupt_stage(config, body, nodes, run_dir, **kwargs):
                target = Path(run_dir) / config.documents[0]["filename"]
                staged = Document(body)
                staged.paragraphs[0].runs[0].font.size = Pt(70)
                staged.save(target)
                return {"main": target}

            with patch("lib.delivery.build_deliveries", side_effect=corrupt_stage):
                with self.assertRaises(BuildError):
                    UnifiedSynthesizer.synthesize(source, output, manifest)

            self.assertEqual(before, (_sha256(old_delivery), _sha256(old_meta)))
            diagnostics = json.loads((output / "build-diagnostics.json").read_text(encoding="utf-8"))
            self.assertEqual(diagnostics["qa"]["format_verification"], "failed")
            report = diagnostics["qa"]["delivery_reports"]["main"]
            self.assertFalse(report["publishable"])
            self.assertTrue(report["source_nodes"])

    def test_page_geometry_violation_is_reported_before_publication(self):
        """F02/W04 integration without Word：Letter 交付不得满足 A4 契约。"""
        resolved = _custom_format()
        doc = Document()
        doc.sections[0].page_width = Mm(215.9)
        doc.sections[0].page_height = Mm(279.4)
        doc.add_paragraph("Wrong page geometry")
        apply_roles(doc, {"/w:document/w:body/w:p[1]": "body"}, resolved)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "letter.docx"
            doc.save(path)
            report = verify_delivery_format(path, resolved)
            self.assertFalse(report.passed)
            self.assertTrue(any("页面" in message or "尺寸" in message for message in report.violations))

    def test_custom_content_part_cannot_publish_without_content_check(self):
        """F02 integration without Word：自定义 content 部件不能绕过完整性检查。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_dir = root / "sources"
            source_dir.mkdir()
            source = source_dir / "source.docx"
            from PIL import Image
            image = root / "source.png"
            Image.new("RGB", (16, 16), "green").save(image)
            doc = Document()
            doc.add_paragraph("Required custom-part text")
            doc.add_picture(str(image))
            doc.save(source)

            manifest_data = _manifest(source.name, strategy="directory_tree", parts=["main_text"])
            manifest_data["source"] = {"strategy": "directory_tree"}
            manifest_data["layout"] = {
                "parts": {"main_text": {"kind": "content"}}
            }
            manifest = root / "manifest.json"
            manifest.write_text(json.dumps(manifest_data), encoding="utf-8")

            def missing_object_stage(config, body, nodes, run_dir, **kwargs):
                target = Path(run_dir) / config.documents[0]["filename"]
                incomplete = Document()
                incomplete.add_paragraph("Required custom-part text")
                incomplete.save(target)
                return {"main": target}

            with patch("lib.delivery.build_deliveries", side_effect=missing_object_stage):
                with self.assertRaises(BuildError):
                    UnifiedSynthesizer.synthesize(source_dir, root / "output", manifest)

    def test_metadata_write_failure_keeps_previous_delivery_and_metadata(self):
        """F02 transaction contract：元数据写入失败不能损坏上一份成果。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.add_paragraph("Metadata transaction body")
            doc.save(source)
            manifest = root / "manifest.json"
            manifest.write_text(json.dumps(_manifest(source.name)), encoding="utf-8")
            output = root / "output"
            output.mkdir()
            old_delivery = output / "r0-published.docx"
            old_meta = output / "build-metadata.json"
            Document().save(old_delivery)
            old_meta.write_text('{"status":"previous"}', encoding="utf-8")
            before = (old_delivery.read_bytes(), old_meta.read_bytes())

            def stage(config, body, nodes, run_dir, **kwargs):
                target = Path(run_dir) / config.documents[0]["filename"]
                shutil.copy2(body, target)
                return {"main": target}

            real_open = builtins.open

            def fail_metadata(path, *args, **kwargs):
                if str(path).endswith("build-metadata.json"):
                    raise OSError("simulated metadata write failure")
                return real_open(path, *args, **kwargs)

            with patch("lib.delivery.build_deliveries", side_effect=stage), \
                    patch("builtins.open", side_effect=fail_metadata):
                with self.assertRaises(BuildError):
                    UnifiedSynthesizer.synthesize(source, output, manifest)

            self.assertEqual(before, (old_delivery.read_bytes(), old_meta.read_bytes()))


if __name__ == "__main__":
    unittest.main()
