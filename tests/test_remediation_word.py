# -*- coding: utf-8 -*-
"""S6 real Word fixtures for the post-N0–N10 remediation plan.

The suite is opt-in and fail-closed: a missing Word Automation capability is
reported as BLOCKED rather than converted into a passing offline result.
"""

from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import copy
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.enum.section import WD_ORIENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.oxml.ns import nsdecls
from lxml import etree
from docx.shared import Inches, Mm, Pt, RGBColor

from lib.content_integrity import (
    ContentIntegrityError,
    ExpectedInventory,
    build_expected_inventory,
    extract_semantic_inventory,
    verify_delivery_format,
    verify_content_integrity,
)
from lib.docx_inspector import inspect_docx
from lib.engine import UnifiedSynthesizer
from lib.qa import export_docx_to_pdf, word_automation_status
from lib.verification_contracts import build_verification_context


REPO = Path(__file__).resolve().parents[1]


def _manifest(root: Path, filename: str, project: str, *, formatting=None, source_file=None, cover=None, parts=None) -> Path:
    path = root / f"{project}.manifest.json"
    source = {"strategy": "docx_document"}
    if source_file:
        source["file"] = str(source_file)
    payload = {
        "schema_version": 3,
        "project_name": project,
        "source": source,
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": formatting or {"mode": "preserve", "page_policy": "source"},
        "output": {"documents": [{"id": "main", "filename": filename, "parts": parts or ["body"]}]},
    }
    if cover is not None:
        payload["cover"] = cover
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _cli(source: Path, manifest: Path, output: Path) -> Path:
    common = [sys.executable, "synthesize.py", "--source", str(source), "--manifest", str(manifest)]
    plan = subprocess.run(
        common + ["--plan"], cwd=REPO, capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=120,
    )
    if plan.returncode:
        raise AssertionError(plan.stdout + plan.stderr)
    doctor = subprocess.run(
        common + ["--doctor"], cwd=REPO, capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=120,
    )
    if doctor.returncode:
        raise AssertionError(doctor.stdout + doctor.stderr)
    build = subprocess.run(
        common + ["--output-dir", str(output)], cwd=REPO, capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=900,
    )
    if build.returncode:
        raise AssertionError(build.stdout + build.stderr)
    delivered = output / json.loads(manifest.read_text(encoding="utf-8"))["output"]["documents"][0]["filename"]
    if not delivered.is_file():
        raise AssertionError(build.stdout + build.stderr)
    return delivered


def _enable_even_headers(document: Document) -> None:
    settings = document.settings.element
    if settings.find(qn("w:evenAndOddHeaders")) is None:
        settings.append(OxmlElement("w:evenAndOddHeaders"))


def _part_text(path: Path, part_id: str) -> str:
    doc = Document(str(path))
    starts = {"cover": "_Synth_cover", "toc": "_Synth_toc", "body": "_Synth_body"}
    ends = {
        "cover": "_Synth_part_cover_end",
        "toc": "_Synth_part_toc_end",
        "body": "_Synth_part_body_end",
    }
    children = list(doc.element.body)
    start = next(i for i, child in enumerate(children) if any(
        marker.get(qn("w:name")) == starts[part_id]
        for marker in child.iter(qn("w:bookmarkStart"))
    ))
    end = next(i for i in range(start, len(children)) if any(
        marker.get(qn("w:name")) == ends[part_id]
        for marker in children[i].iter(qn("w:bookmarkStart"))
    ))
    return "".join(
        node.text or ""
        for child in children[start:end + 1]
        for node in child.iter(qn("w:t"))
    )


def _write_notes_source(path: Path, label: str, url: str, image_path: Path) -> None:
    """Create a Word-valid source with one footnote and one endnote.

    The note bodies deliberately contain both an external hyperlink and an
    embedded image.  The drawing is borrowed from python-docx's valid inline
    drawing shape, then rebound to a relationship owned by the note part.
    """
    document = Document()
    paragraph = document.add_paragraph(f"{label} body")
    footnote_ref = paragraph.add_run()
    footnote_ref._r.append(parse_xml(
        f'<w:footnoteReference {nsdecls("w")} w:id="1"/>'
    ))
    endnote_ref = paragraph.add_run()
    endnote_ref._r.append(parse_xml(
        f'<w:endnoteReference {nsdecls("w")} w:id="1"/>'
    ))

    image_probe = document.add_paragraph().add_run()
    image_probe.add_picture(str(image_path), width=Mm(12))
    drawing = copy.deepcopy(image_probe._r.find(qn("w:drawing")))
    image_paragraph = image_probe._r.getparent()
    image_paragraph.getparent().remove(image_paragraph)
    image_part = next(
        rel.target_part for rel in document.part.rels.values()
        if rel.reltype == RT.IMAGE
    )

    def note_part(tag: str, ref_tag: str, content_type: str, rel_type: str, filename: str):
        part = Part(
            PackURI(f"/word/{filename}"),
            content_type,
            b"",
            document.part.package,
        )
        link_rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
        image_rid = part.relate_to(image_part, RT.IMAGE)
        note = OxmlElement(f"w:{tag}")
        note.set(qn("w:id"), "1")
        p = OxmlElement("w:p")
        run = OxmlElement("w:r")
        ref = OxmlElement(f"w:{ref_tag}")
        run.append(ref)
        text = OxmlElement("w:t")
        text.text = f" {label} {tag} note"
        run.append(text)
        p.append(run)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), link_rid)
        link_run = OxmlElement("w:r")
        link_text = OxmlElement("w:t")
        link_text.text = f" {label} link"
        link_run.append(link_text)
        hyperlink.append(link_run)
        p.append(hyperlink)
        drawing_copy = copy.deepcopy(drawing)
        for blip in drawing_copy.iter(qn("a:blip")):
            blip.set(qn("r:embed"), image_rid)
        image_run = OxmlElement("w:r")
        image_run.append(drawing_copy)
        p.append(image_run)
        note.append(p)
        root = parse_xml(f'<w:{tag}s {nsdecls("w")}/>')
        root.append(note)
        part._blob = etree.tostring(
            root, encoding="utf-8", xml_declaration=True, standalone="yes"
        )
        document.part.relate_to(part, rel_type)

    note_part(
        "footnote", "footnoteRef",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        RT.FOOTNOTES, "footnotes.xml",
    )
    note_part(
        "endnote", "endnoteRef",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
        RT.ENDNOTES, "endnotes.xml",
    )
    document.save(path)


class RemediationWordTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if os.environ.get("DOCUMENT_SYNTHESIS_WORD_TEST") != "1":
            raise unittest.SkipTest("BLOCKED: set DOCUMENT_SYNTHESIS_WORD_TEST=1 to run SW01–SW07")
        available, reason = word_automation_status()
        if not available:
            raise unittest.SkipTest(f"BLOCKED: Word Automation unavailable: {reason}")

    def test_SW01_real_word_preserves_logical_inline_spans_and_pdf_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("SW01 protected 12pt ")
            first.font.name = "Arial"
            first.font.size = Pt(12)
            first.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
            second = paragraph.add_run("and 24pt interval.")
            second.font.name = "Times New Roman"
            second.font.size = Pt(24)
            second.font.color.rgb = RGBColor(0x80, 0x00, 0x00)
            doc.save(source)
            delivered = _cli(source, _manifest(root, "sw01.docx", "SW01"), root / "out")
            inspection = inspect_docx(delivered)
            spans = inspection.blocks[0].effective_run_spans
            self.assertEqual([round(span["style"].size_pt) for span in spans], [12, 24])
            self.assertEqual([span["style"].latin for span in spans], ["Arial", "Times New Roman"])
            pdf = root / "sw01.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            import pymupdf
            with pymupdf.open(pdf) as rendered:
                text = "".join(page.get_text() for page in rendered)
            self.assertIn("SW01 protected 12pt", text)
            self.assertIn("and 24pt interval", text)

    def test_SW01_real_word_corrupted_logical_span_is_rejected_by_final_gate(self):
        """A real CLI output must fail when one protected span is damaged."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("SW01 damage protected 12pt ")
            first.font.name = "Arial"
            first.font.size = Pt(12)
            first.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
            second = paragraph.add_run("and damaged 24pt interval.")
            second.font.name = "Times New Roman"
            second.font.size = Pt(24)
            second.font.color.rgb = RGBColor(0x80, 0x00, 0x00)
            doc.save(source)

            manifest = _manifest(root, "sw01-damage.docx", "SW01-damage")
            delivered = _cli(source, manifest, root / "out")

            corrupted_doc = Document(str(delivered))
            damaged_run = next(
                run for run in corrupted_doc.paragraphs[0].runs
                if "damaged 24pt" in (run.text or "")
            )
            damaged_run.font.size = Pt(12)
            corrupted = root / "sw01-damaged-output.docx"
            corrupted_doc.save(corrupted)

            prepared = UnifiedSynthesizer.prepare_build(source, manifest)
            context = build_verification_context(
                prepared,
                None,
                prepared.config.documents[0],
                corrupted,
            )
            report = verify_delivery_format(
                corrupted,
                prepared.config.resolved_format,
                verification_context=context,
            )
            self.assertFalse(report.passed, report.to_dict())
            self.assertTrue(
                any("字号" in violation for violation in report.violations),
                report.to_dict(),
            )

    def test_SW02_real_word_preserves_default_first_even_story_bindings(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            section = doc.sections[0]
            section.different_first_page_header_footer = True
            section.header.paragraphs[0].text = "SW02 DEFAULT HEADER"
            header_table = section.header.add_table(rows=1, cols=2, width=section.page_width)
            header_table.cell(0, 0).text = "SW02 HEADER TABLE A"
            header_table.cell(0, 1).text = "SW02 HEADER TABLE B"
            section.first_page_header.paragraphs[0].text = "SW02 FIRST HEADER"
            _enable_even_headers(doc)
            section.even_page_header.paragraphs[0].text = "SW02 EVEN HEADER"
            section.footer.paragraphs[0].text = "SW02 DEFAULT FOOTER"
            for index in range(18):
                doc.add_paragraph(f"SW02 body page material {index}. " * 20)
            doc.save(source)
            delivered = _cli(source, _manifest(root, "sw02.docx", "SW02"), root / "out")
            inspection = inspect_docx(delivered)
            stories = {(story.story_type, story.variant): story for story in inspection.stories}
            self.assertIn(("header", "default"), stories)
            self.assertIn(("header", "first"), stories)
            self.assertIn(("header", "even"), stories)
            self.assertIn("SW02 FIRST HEADER", " ".join(stories[("header", "first")].blocks[0].visible_text for _ in [0]))
            default_story = stories[("header", "default")]
            self.assertTrue(any(
                block.structure_type == "table"
                and "SW02 HEADER TABLE A" in block.visible_text
                and "SW02 HEADER TABLE B" in block.visible_text
                for block in default_story.blocks
            ))
            self.assertTrue(any(item.different_first_page for item in inspection.sections))
            pdf = root / "sw02.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            import pymupdf
            with pymupdf.open(pdf) as rendered:
                text = "\n".join(page.get_text() for page in rendered)
            self.assertIn("SW02 DEFAULT HEADER", text)
            self.assertIn("SW02 FIRST HEADER", text)
            self.assertIn("SW02 EVEN HEADER", text)
            self.assertIn("SW02 HEADER TABLE A", text)
            self.assertIn("SW02 HEADER TABLE B", text)

    def test_SW03_real_word_complete_thesis_parts_have_fixture_and_pdf_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = root / "thesis-fixture"
            created = subprocess.run(
                [sys.executable, str(REPO / "docs" / "acceptance" / "create_complete_thesis_fixture.py"), str(fixture)],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
            self.assertEqual(created.returncode, 0, created.stdout + created.stderr)
            manifest = fixture / "manifest.json"
            source = fixture / "source" / "thesis-source.docx"
            source_xml = __import__("zipfile").ZipFile(source).read("word/document.xml")
            for token in ("声明", "中文摘要", "Abstract", "参考文献", "附录 A"):
                self.assertIn(token, source_xml.decode("utf-8"))
            manifest_data = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual(
                manifest_data["output"]["documents"][0]["parts"],
                ["cover", "toc", "body"],
            )
            delivered = _cli(fixture, manifest, root / "out")
            self.assertIn("匿名学位论文样例", _part_text(delivered, "cover"))
            body_text = _part_text(delivered, "body")
            for token in ("声明", "中文摘要", "Abstract", "参考文献", "附录 A"):
                self.assertIn(token, body_text)
            package = __import__("zipfile").ZipFile(delivered)
            self.assertIn("word/footnotes.xml", package.namelist())
            self.assertIn("word/endnotes.xml", package.namelist())
            pdf = root / "sw03.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            import pymupdf
            with pymupdf.open(pdf) as rendered:
                text = "\n".join(page.get_text() for page in rendered)
            for token in ("声明", "中文摘要", "Abstract", "参考文献", "附录 A"):
                self.assertIn(token, text)

    def test_SW04_real_word_landscape_section_and_story_geometry_are_bound(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            first = doc.sections[0]
            first.header.paragraphs[0].text = "SW04 PORTRAIT HEADER"
            doc.add_paragraph("SW04 portrait material. " * 60)
            landscape = doc.add_section(WD_SECTION_START.NEW_PAGE)
            landscape.orientation = WD_ORIENT.LANDSCAPE
            landscape.page_width = Mm(297)
            landscape.page_height = Mm(210)
            landscape.header.is_linked_to_previous = False
            landscape.header.paragraphs[0].text = "SW04 LANDSCAPE HEADER"
            landscape.header_distance = Mm(11)
            landscape.footer_distance = Mm(19)
            doc.add_paragraph("SW04 landscape table section material. " * 70)
            doc.save(source)
            self.assertEqual(len(Document(source).sections), 2)
            self.assertAlmostEqual(Document(source).sections[1].page_width.mm, 297, delta=0.1)
            delivered = _cli(source, _manifest(root, "sw04.docx", "SW04"), root / "out")
            inspection = inspect_docx(delivered)
            self.assertGreaterEqual(len(inspection.sections), 2)
            self.assertAlmostEqual(inspection.sections[1].page_spec.width_mm, 297, delta=0.2)
            self.assertEqual(inspection.sections[1].page_spec.orientation, "landscape")
            story_text = " ".join(
                block.visible_text
                for story in inspection.stories
                for block in story.blocks
            )
            self.assertIn("SW04 PORTRAIT HEADER", story_text)
            self.assertIn("SW04 LANDSCAPE HEADER", story_text)
            pdf = root / "sw04.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            import pymupdf
            with pymupdf.open(pdf) as rendered:
                text = "\n".join(page.get_text() for page in rendered)
            self.assertIn("SW04 PORTRAIT HEADER", text)
            self.assertIn("SW04 LANDSCAPE HEADER", text)

    def test_SW02_real_word_two_section_shared_and_unlinked_bindings(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            document = Document()
            first = document.sections[0]
            first.header.paragraphs[0].text = "SW02 SHARED HEADER"
            first.footer.paragraphs[0].text = "SW02 SHARED FOOTER"
            document.add_paragraph("SW02 first section material. " * 45)

            linked = document.add_section(WD_SECTION_START.NEW_PAGE)
            # python-docx defaults a new section to linked stories.
            document.add_paragraph("SW02 linked section material. " * 45)

            separate = document.add_section(WD_SECTION_START.NEW_PAGE)
            separate.header.is_linked_to_previous = False
            separate.footer.is_linked_to_previous = False
            separate.header.paragraphs[0].text = "SW02 SEPARATE HEADER"
            separate.footer.paragraphs[0].text = "SW02 SEPARATE FOOTER"
            document.add_paragraph("SW02 separate section material. " * 45)
            document.save(source)

            source_doc = Document(source)
            self.assertEqual(len(source_doc.sections), 3)
            self.assertTrue(source_doc.sections[1].header.is_linked_to_previous)
            self.assertFalse(source_doc.sections[2].header.is_linked_to_previous)

            manifest = _manifest(root, "sw02-shared-unlinked.docx", "SW02-shared-unlinked")
            delivered = _cli(source, manifest, root / "out")
            inspection = inspect_docx(delivered)
            bindings = {
                (item.section_index, item.story_type, item.variant): item
                for item in inspection.section_bindings
            }
            self.assertTrue(bindings[(1, "header", "default")].linked_to_previous)
            self.assertFalse(bindings[(2, "header", "default")].linked_to_previous)
            self.assertTrue(bindings[(1, "footer", "default")].linked_to_previous)
            self.assertFalse(bindings[(2, "footer", "default")].linked_to_previous)
            self.assertEqual(
                bindings[(0, "header", "default")].part_uri,
                bindings[(1, "header", "default")].part_uri,
            )
            self.assertNotEqual(
                bindings[(1, "header", "default")].part_uri,
                bindings[(2, "header", "default")].part_uri,
            )

            pdf = root / "sw02-shared-unlinked.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            import pymupdf
            with pymupdf.open(pdf) as rendered:
                text = "\n".join(page.get_text() for page in rendered)
            self.assertIn("SW02 SHARED HEADER", text)
            self.assertIn("SW02 SEPARATE HEADER", text)

    def test_SW02_real_word_wrong_story_relationship_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            document = Document()
            document.sections[0].header.paragraphs[0].text = "SW02 RELATION HEADER"
            document.sections[0].footer.paragraphs[0].text = "SW02 RELATION FOOTER"
            document.add_paragraph("SW02 relationship target material. " * 80)
            document.save(source)
            manifest = _manifest(root, "sw02-relation.docx", "SW02-relation")
            delivered = _cli(source, manifest, root / "out")

            corrupted = root / "sw02-wrong-story-target.docx"
            with zipfile.ZipFile(delivered) as source_zip, zipfile.ZipFile(corrupted, "w") as target_zip:
                for info in source_zip.infolist():
                    payload = source_zip.read(info.filename)
                    if info.filename == "word/_rels/document.xml.rels":
                        text = payload.decode("utf-8")
                        self.assertIn("header1.xml", text)
                        self.assertIn("footer1.xml", text)
                        payload = text.replace("header1.xml", "footer1.xml", 1).encode("utf-8")
                    target_zip.writestr(info, payload)

            prepared = UnifiedSynthesizer.prepare_build(source, manifest)
            context = build_verification_context(
                prepared,
                None,
                prepared.config.documents[0],
                corrupted,
            )
            report = verify_delivery_format(
                corrupted,
                prepared.config.resolved_format,
                verification_context=context,
            )
            self.assertFalse(report.passed, report.to_dict())
            self.assertTrue(any("story" in item or "绑定" in item for item in report.violations))

    def test_SW05_real_word_multi_source_notes_relationship_matrix_is_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_a = root / "note-a.png"
            image_b = root / "note-b.png"
            from PIL import Image
            Image.new("RGB", (48, 30), color="darkgreen").save(image_a)
            Image.new("RGB", (48, 30), color="darkorange").save(image_b)
            source_a = root / "01-source-a.docx"
            source_b = root / "02-source-b.docx"
            _write_notes_source(source_a, "SW05-A", "https://example.com/sw05-a", image_a)
            _write_notes_source(source_b, "SW05-B", "https://example.com/sw05-b", image_b)
            manifest = root / "SW05.manifest.json"
            manifest.write_text(json.dumps({
                "schema_version": 3,
                "project_name": "SW05",
                "source": {
                    "strategy": "explicit_tree",
                    "tree": [
                        {"level": 1, "title": "Source A", "type": "docx", "file": source_a.name, "include_in_toc": False},
                        {"level": 1, "title": "Source B", "type": "docx", "file": source_b.name, "include_in_toc": False},
                    ],
                },
                "format": {"ref": "preset:academic-basic@1.0.0"},
                "formatting": {"mode": "preserve", "page_policy": "source"},
                "output": {"documents": [{"id": "main", "filename": "sw05.docx", "parts": ["body"]}]},
            }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

            delivered = _cli(root, manifest, root.parent / f"sw05-out-{root.name}")
            prepared = UnifiedSynthesizer.prepare_build(root, manifest)
            expected = build_expected_inventory(prepared, spec=prepared.config.documents[0])
            mapping_evidence = {}
            self.assertTrue(
                verify_content_integrity(expected, delivered, evidence_out=mapping_evidence),
                mapping_evidence,
            )
            source_mappings = [
                item for item in mapping_evidence["object_mapping"]
                if item["kind"] == "source_object"
            ]
            self.assertEqual(len(source_mappings), 4)
            self.assertEqual(
                {item["source_sha256"] for item in source_mappings},
                {
                    extract_semantic_inventory(source_a).source_sha256,
                    extract_semantic_inventory(source_b).source_sha256,
                },
            )
            self.assertTrue(all(item["output_part_ids"] == ["body"] for item in source_mappings))
            with zipfile.ZipFile(delivered) as package:
                names = set(package.namelist())
                self.assertIn("word/footnotes.xml", names)
                self.assertIn("word/endnotes.xml", names)
                footnotes = parse_xml(package.read("word/footnotes.xml"))
                endnotes = parse_xml(package.read("word/endnotes.xml"))
                self.assertEqual(
                    {item.get(qn("w:id")) for item in footnotes.findall(qn("w:footnote"))},
                    {"-1", "0", "1", "2"},
                )
                self.assertEqual(
                    {item.get(qn("w:id")) for item in endnotes.findall(qn("w:endnote"))},
                    {"-1", "0", "1", "2"},
                )
                footnote_text = "".join(footnotes.itertext())
                endnote_text = "".join(endnotes.itertext())
                for token in ("SW05-A", "SW05-B", "SW05-A link", "SW05-B link"):
                    self.assertIn(token, footnote_text)
                    self.assertIn(token, endnote_text)
                for rel_name, url_a, url_b in (
                    ("word/_rels/footnotes.xml.rels", "sw05-a", "sw05-b"),
                    ("word/_rels/endnotes.xml.rels", "sw05-a", "sw05-b"),
                ):
                    rels = package.read(rel_name).decode("utf-8")
                    self.assertIn(url_a, rels)
                    self.assertIn(url_b, rels)
                    self.assertGreaterEqual(rels.count("relationships/image"), 2)

            pdf = root / "sw05.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            import pymupdf
            with pymupdf.open(pdf) as rendered:
                text = "\n".join(page.get_text() for page in rendered)
            for token in ("SW05-A body", "SW05-B body", "SW05-A footnote note", "SW05-B endnote note"):
                self.assertIn(token, text)

    def test_SW06_real_word_object_instances_and_extra_image_are_checked(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            image = root / "figure.png"
            from PIL import Image
            Image.new("RGB", (32, 32), color="navy").save(image)
            doc = Document()
            paragraph = doc.add_paragraph("SW06 source image")
            paragraph.add_run().add_picture(str(image), width=Mm(18))
            paragraph.add_run().add_picture(str(image), width=Mm(18))
            doc.save(source)
            delivered = _cli(source, _manifest(root, "sw06.docx", "SW06"), root / "out")
            expected = ExpectedInventory(extract_semantic_inventory(source))
            self.assertTrue(verify_content_integrity(expected, delivered))
            self.assertEqual(len(extract_semantic_inventory(delivered).media_occurrences), 2)
            changed = Document(delivered)
            clone = changed.paragraphs[0].add_run()
            clone.add_picture(str(image), width=Mm(18))
            corrupted = root / "sw06-extra.docx"
            changed.save(corrupted)
            with self.assertRaises(ContentIntegrityError):
                verify_content_integrity(expected, corrupted)
            pdf = root / "sw06.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))

    def test_SW06_real_word_generated_cover_image_has_finite_license(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            image = root / "cover-logo.png"
            from PIL import Image
            Image.new("RGB", (42, 28), color="forestgreen").save(image)
            source_doc = Document()
            source_doc.add_paragraph("SW06 body source")
            source_doc.add_picture(str(image), width=Mm(16))
            source_doc.save(source)

            template = root / "cover-template.docx"
            template_doc = Document()
            template_doc.add_paragraph("{{MAIN_TITLE}}")
            template_doc.add_picture(str(image), width=Mm(16))
            template_doc.save(template)
            manifest = _manifest(
                root,
                "sw06-cover-image.docx",
                "SW06-cover-image",
                formatting={"mode": "restyle", "page_policy": "target"},
                source_file=source.name,
                cover={"mode": "template", "template": template.name, "main_title": "SW06 COVER"},
                parts=["cover", "body"],
            )

            delivered = _cli(source, manifest, root / "out")
            prepared = UnifiedSynthesizer.prepare_build(source, manifest)
            expected = build_expected_inventory(prepared, spec=prepared.config.documents[0])
            self.assertEqual(
                sum(item.get("part_id") == "cover" for item in expected.allowed_generated_objects),
                1,
            )
            evidence = {}
            self.assertTrue(verify_content_integrity(expected, delivered, evidence_out=evidence))
            generated_mappings = [
                item for item in evidence["object_mapping"]
                if item["kind"] == "generated_object"
            ]
            self.assertEqual([item["part_id"] for item in generated_mappings], ["cover"])
            self.assertEqual([item["output_part_ids"] for item in generated_mappings], [["cover"]])

            corrupted = Document(delivered)
            extra = corrupted.add_picture(str(image), width=Mm(16))
            corrupted_path = root / "sw06-cover-extra.docx"
            corrupted.save(corrupted_path)
            with self.assertRaises(ContentIntegrityError):
                verify_content_integrity(expected, corrupted_path)

            pdf = root / "sw06-cover-image.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))

    def test_SW07_real_word_pending_missing_level_requires_manual_target_confirmation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "target.docx"
            document = Document()
            first = document.add_paragraph("1. Target introduction")
            first.runs[0].bold = True
            first.runs[0].font.size = Pt(18)
            body_a = document.add_paragraph("Target body material A. " * 28)
            body_a.runs[0].font.size = Pt(12)
            missing_level = document.add_paragraph("1.1.1. Target detailed protocol")
            missing_level.runs[0].bold = True
            missing_level.runs[0].font.size = Pt(14)
            body_b = document.add_paragraph("Target body material B. " * 28)
            body_b.runs[0].font.size = Pt(12)
            document.save(source)

            review_dir = root / "review"
            analysis = subprocess.run(
                [
                    sys.executable, "synthesize.py", "--analyze-format", str(source),
                    "--analysis-dir", str(review_dir), "--replace-output",
                ], cwd=REPO, capture_output=True, text=True,
                encoding="utf-8", errors="replace", timeout=120,
            )
            self.assertEqual(analysis.returncode, 0, analysis.stdout + analysis.stderr)
            analysis_data = json.loads((review_dir / "analysis.json").read_text(encoding="utf-8"))
            candidate = analysis_data["unstructured_candidate_evidence"]["/w:document/w:body/w:p[3]"]
            self.assertEqual(candidate["selected_role"], "heading.3")
            self.assertEqual(candidate["review_status"], "pending_review")
            self.assertTrue(any(item["kind"] == "numbering_vs_size_level" for item in candidate["conflicts"]))

            decisions = {
                "decisions_schema_version": 1,
                "report_id": analysis_data["report_id"],
                "source_sha256": analysis_data["source_sha256"],
                "role_styles": {
                    role: item["cluster_id"]
                    for role, item in analysis_data.get("candidate_roles", {}).items()
                },
                "style_overrides": {},
                "missing_roles": {
                    item["role"]: "inherit"
                    for item in analysis_data.get("missing_roles", [])
                },
                "node_roles": {"/w:document/w:body/w:p[3]": "heading.3"},
                "on_unmapped": "error",
            }
            decisions_path = review_dir / "decisions.final.json"
            decisions_path.write_text(json.dumps(decisions, ensure_ascii=False, indent=2), encoding="utf-8")
            mapping_path = review_dir / "target-roles.json"
            compiled = subprocess.run(
                [
                    sys.executable, "synthesize.py", "--compile-mapping",
                    str(review_dir / "analysis.json"), "--decisions", str(decisions_path),
                    "--mapping-out", str(mapping_path), "--replace-output",
                ], cwd=REPO, capture_output=True, text=True,
                encoding="utf-8", errors="replace", timeout=120,
            )
            self.assertEqual(compiled.returncode, 0, compiled.stdout + compiled.stderr)
            mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
            manual_assignment = next(
                item for item in mapping["assignments"]
                if item["node_ref"]["element_path"] == "/w:document/w:body/w:p[3]"
            )
            self.assertEqual(manual_assignment["role"], "heading.3")
            self.assertTrue(manual_assignment["confirmed"])
            self.assertEqual(manual_assignment["provenance"], "explicit")

            manifest = _manifest(
                root,
                "sw07.docx",
                "SW07",
                formatting={
                    "mode": "restyle",
                    "role_map": "review/target-roles.json",
                    "on_unmapped": "error",
                },
            )
            delivered = _cli(root, manifest, root.parent / f"sw07-out-{root.name}")
            delivered_doc = Document(str(delivered))
            self.assertEqual(delivered_doc.paragraphs[2].style.style_id, "SynthHeading3")
            pdf = root / "sw07.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            import pymupdf
            with pymupdf.open(pdf) as rendered:
                text = "\n".join(page.get_text() for page in rendered)
            self.assertIn("1.1.1. Target detailed protocol", text)


if __name__ == "__main__":
    unittest.main()
