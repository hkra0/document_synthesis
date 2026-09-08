# -*- coding: utf-8 -*-
"""
内容语义完整性与最终有效格式核验单元测试 (tests/test_content_integrity.py)
"""

import io
from pathlib import Path
import tempfile
import unittest
from types import SimpleNamespace

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import lxml.etree
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from lib.content_integrity import (
    ContentIntegrityError,
    ExpectedInventory,
    FormatVerificationReport,
    SemanticInventory,
    TableTopology,
    extract_semantic_inventory,
    build_expected_inventory,
    merge_semantic_inventories,
    normalize_text,
    verify_content_integrity,
    verify_delivery_format,
    verify_generated_object_allowances,
)
from lib.format_schema import (
    LengthValue,
    LineSpacing,
    PageSpec,
    ParagraphStyle,
    ResolvedFormat,
    RoleSpec,
    RunStyle,
    StyleDefinition,
    TocSpec,
    HeaderFooterSpec,
)
from lib.style_applier import apply_roles, install_styles
from lib.composition import add_bookmark


class ContentIntegrityTest(unittest.TestCase):
    def test_expected_inventory_merges_prepared_sources_in_declared_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = root / "01.docx"
            second = root / "02.docx"
            first_doc = Document()
            first_doc.add_paragraph("first source")
            first_doc.save(first)
            second_doc = Document()
            second_doc.add_paragraph("second source")
            second_doc.add_paragraph("second source")
            second_doc.save(second)

            prepared = SimpleNamespace(
                source_order=(str(first), str(second)),
                source_docx_path=None,
                nodes=(),
                parts={},
                config=SimpleNamespace(regions={}),
            )
            expected = build_expected_inventory(prepared, spec={"id": "main", "parts": ["body"]})
            self.assertEqual(expected.source_order, (str(first), str(second)))
            self.assertEqual(expected.semantic.paragraph_counts["secondsource"], 2)
            self.assertEqual(expected.coverage["source_files"], 2)
            self.assertEqual(expected.allowed_id_remaps["relationship_id"], "allowed_if_target_and_type_preserved")

    def test_expected_inventory_uses_declared_source_region(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.add_paragraph("selected")
            doc.add_paragraph("excluded from this part")
            doc.save(source)
            from lib.document_parts import DocumentPart, SourceRegion

            prepared = SimpleNamespace(
                source_order=(str(source),),
                source_docx_path=source,
                nodes=(),
                parts={"excerpt": DocumentPart("excerpt", "content", source_region="excerpt")},
                config=SimpleNamespace(regions={
                    "excerpt": SourceRegion(
                        "excerpt",
                        "/w:document/w:body/w:p[1]",
                        "/w:document/w:body/w:p[2]",
                    ),
                    "excluded": SourceRegion(
                        "excluded",
                        "/w:document/w:body/w:p[2]",
                        "end_of_document",
                        exclude=True,
                    ),
                }),
            )
            expected = build_expected_inventory(prepared, spec={"id": "excerpt", "parts": ["excerpt"]})
            self.assertIn("selected", expected.semantic.paragraph_counts)
            self.assertNotIn("excludedfromthispart", expected.semantic.paragraph_counts)
            self.assertEqual(expected.source_regions, {"excerpt": "excerpt"})
    def test_extract_semantic_inventory_captures_all_elements(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "sample.docx"
            doc = Document()
            doc.add_paragraph("第一段正文。")
            doc.add_paragraph("第二段重复内容。")
            doc.add_paragraph("第二段重复内容。")

            # 表格
            tbl = doc.add_table(rows=2, cols=2)
            tbl.cell(0, 0).text = "A1"
            tbl.cell(0, 1).text = "B1"
            tbl.cell(1, 0).text = "A2"
            tbl.cell(1, 1).text = "B2"

            # 图片
            from PIL import Image as PILImage
            img_stream = io.BytesIO()
            PILImage.new("RGB", (10, 10), color="red").save(img_stream, format="PNG")
            img_stream.seek(0)
            doc.add_picture(img_stream, width=Inches(1))

            # OMML 公式
            p_math = doc.add_paragraph()
            oMath = lxml.etree.SubElement(
                p_math._p,
                "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath",
            )
            r_math = lxml.etree.SubElement(
                oMath,
                "{http://schemas.openxmlformats.org/officeDocument/2006/math}r",
            )
            t_math = lxml.etree.SubElement(
                r_math,
                "{http://schemas.openxmlformats.org/officeDocument/2006/math}t",
            )
            t_math.text = "x+y=z"

            doc.save(str(doc_path))

            inv = extract_semantic_inventory(doc_path)
            self.assertEqual(inv.paragraph_counts["第一段正文。"], 1)
            self.assertEqual(inv.paragraph_counts["第二段重复内容。"], 2)
            self.assertEqual(len(inv.tables), 1)
            self.assertEqual(inv.tables[0].row_count, 2)
            self.assertEqual(inv.tables[0].col_count, 2)
            self.assertEqual(inv.tables[0].cell_texts, ["A1", "B1", "A2", "B2"])
            self.assertGreater(len(inv.media_hashes), 0)
            self.assertEqual(len(inv.media_occurrences), 1)
            self.assertEqual(inv.media_occurrences[0].kind, "image")
            self.assertEqual(inv.media_occurrences[0].host_element, "/w:document/w:body/w:p[5]")
            self.assertTrue(inv.media_occurrences[0].relationship_id)
            self.assertTrue(inv.media_occurrences[0].target_part.endswith("image1.png"))
            self.assertEqual(len(inv.math_formulas), 1)

    def test_formula_fingerprint_ignores_word_namespace_rewrites(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"

            src = Document()
            src_math = parse_xml(
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                '<m:r><m:t>x+y=z</m:t></m:r></m:oMath>'
            )
            src.add_paragraph()._p.append(src_math)
            src.save(src_path)

            dst = Document()
            dst_math = parse_xml(
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
                'xmlns:unused="urn:word-rewrite-example">'
                '<m:r><m:t>x+y=z</m:t></m:r></m:oMath>'
            )
            dst.add_paragraph()._p.append(dst_math)
            dst.save(dst_path)

            self.assertTrue(
                verify_content_integrity(extract_semantic_inventory(src_path), dst_path)
            )

    def test_verify_content_integrity_success(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"
            doc = Document()
            doc.add_paragraph("关键段落一")
            doc.add_paragraph("关键段落二")
            doc.save(str(src_path))
            doc.save(str(dst_path))

            inv = extract_semantic_inventory(src_path)
            self.assertTrue(verify_content_integrity(inv, dst_path))

    def test_verify_content_integrity_detects_omitted_duplicate_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"

            # 源文档中同一段落出现 2 次
            src_doc = Document()
            src_doc.add_paragraph("重复通知：下周开会。")
            src_doc.add_paragraph("普通正文。")
            src_doc.add_paragraph("重复通知：下周开会。")
            src_doc.save(str(src_path))

            # 目标文档中意外只保留了 1 次（传统子串查找会误判通过）
            dst_doc = Document()
            dst_doc.add_paragraph("重复通知：下周开会。")
            dst_doc.add_paragraph("普通正文。")
            dst_doc.save(str(dst_path))

            inv = extract_semantic_inventory(src_path)
            with self.assertRaises(ContentIntegrityError) as ctx:
                verify_content_integrity(inv, dst_path)
            self.assertIn("重复通知：下周开会。", str(ctx.exception))
            self.assertIn("2 != 1", str(ctx.exception))

    def test_verify_content_integrity_accepts_declared_replacements(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"

            src_doc = Document()
            src_doc.add_paragraph("原始粗糙标题")
            src_doc.add_paragraph("正文第一段。")
            src_doc.save(str(src_path))

            dst_doc = Document()
            dst_doc.add_paragraph("第一章 规范化标题")
            dst_doc.add_paragraph("正文第一段。")
            dst_doc.save(str(dst_path))

            inv = extract_semantic_inventory(src_path)
            replacements = {"原始粗糙标题": "第一章 规范化标题"}
            self.assertTrue(verify_content_integrity(inv, dst_path, replacements=replacements))

    def test_verify_content_integrity_detects_tampered_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"

            src_doc = Document()
            t1 = src_doc.add_table(rows=1, cols=2)
            t1.cell(0, 0).text = "指标名称"
            t1.cell(0, 1).text = "99.8%"
            src_doc.save(str(src_path))

            # 交付物中缺失表格
            dst_doc = Document()
            dst_doc.add_paragraph("指标名称 99.8%")
            dst_doc.save(str(dst_path))

            inv = extract_semantic_inventory(src_path)
            with self.assertRaises(ContentIntegrityError) as ctx:
                verify_content_integrity(inv, dst_path)
            self.assertIn("表格", str(ctx.exception))

    def test_verify_content_integrity_checks_media_instance_multiplicity(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"
            from PIL import Image as PILImage

            image = Path(tmp) / "same.png"
            PILImage.new("RGB", (10, 10), color="red").save(image)
            src = Document()
            src.add_picture(str(image), width=Inches(1))
            src.add_picture(str(image), width=Inches(1))
            src.save(src_path)
            dst = Document()
            dst.add_picture(str(image), width=Inches(1))
            dst.save(dst_path)

            source_inventory = extract_semantic_inventory(src_path)
            self.assertEqual(len(source_inventory.media_instances), 2)
            self.assertEqual(len(source_inventory.media_occurrences), 2)
            self.assertEqual(
                [item.local_index for item in source_inventory.media_occurrences],
                [1, 1],
            )
            self.assertEqual(
                len({item.host_element for item in source_inventory.media_occurrences}),
                2,
            )
            with self.assertRaises(ContentIntegrityError):
                verify_content_integrity(source_inventory, dst_path)

    def test_verify_content_integrity_rejects_moved_media_occurrence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "same.png"
            from PIL import Image as PILImage
            PILImage.new("RGB", (10, 10), color="red").save(image)

            src = Document()
            src.add_paragraph().add_run().add_picture(str(image), width=Inches(1))
            src.save(root / "src.docx")

            dst = Document()
            dst.add_paragraph("new host")
            dst.add_paragraph().add_run().add_picture(str(image), width=Inches(1))
            dst.save(root / "dst.docx")

            source_inventory = extract_semantic_inventory(root / "src.docx")
            with self.assertRaises(ContentIntegrityError) as ctx:
                verify_content_integrity(source_inventory, root / "dst.docx")
            self.assertIn("出现位置", str(ctx.exception))

    def test_expected_inventory_rejects_moved_anchored_media_occurrence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "same.png"
            from PIL import Image as PILImage
            PILImage.new("RGB", (10, 10), color="red").save(image)
            src = Document()
            src.add_paragraph("source host").add_run().add_picture(str(image), width=Inches(1))
            src.add_paragraph("different host")
            src.save(root / "src.docx")
            dst = Document()
            dst.add_paragraph("source host")
            dst.add_paragraph("different host").add_run().add_picture(str(image), width=Inches(1))
            dst.save(root / "dst.docx")
            expected = ExpectedInventory(extract_semantic_inventory(root / "src.docx"))
            with self.assertRaises(ContentIntegrityError) as ctx:
                verify_content_integrity(expected, root / "dst.docx")
            self.assertIn("宿主锚点", str(ctx.exception))

    def test_expected_inventory_rejects_media_moved_between_identical_hosts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            from PIL import Image as PILImage

            first_image = root / "first.png"
            second_image = root / "second.png"
            PILImage.new("RGB", (10, 10), color="red").save(first_image)
            PILImage.new("RGB", (10, 10), color="blue").save(second_image)
            source = Document()
            source.add_paragraph("same host").add_run().add_picture(str(first_image), width=Inches(1))
            source.add_paragraph("same host").add_run().add_picture(str(second_image), width=Inches(1))
            source.save(root / "source.docx")

            moved = Document()
            moved.add_paragraph("same host").add_run().add_picture(str(second_image), width=Inches(1))
            moved.add_paragraph("same host").add_run().add_picture(str(first_image), width=Inches(1))
            moved.save(root / "moved.docx")

            expected = ExpectedInventory(extract_semantic_inventory(root / "source.docx"))
            with self.assertRaises(ContentIntegrityError) as ctx:
                verify_content_integrity(expected, root / "moved.docx")
            self.assertIn("宿主锚点", str(ctx.exception))

    def test_verify_content_integrity_checks_formula_instance_multiplicity(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"

            def add_formula(document):
                paragraph = document.add_paragraph()
                omath = lxml.etree.SubElement(
                    paragraph._p,
                    "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath",
                )
                run = lxml.etree.SubElement(
                    omath,
                    "{http://schemas.openxmlformats.org/officeDocument/2006/math}r",
                )
                text = lxml.etree.SubElement(
                    run,
                    "{http://schemas.openxmlformats.org/officeDocument/2006/math}t",
                )
                text.text = "x+y=z"

            src = Document()
            add_formula(src)
            add_formula(src)
            src.save(src_path)
            dst = Document()
            add_formula(dst)
            dst.save(dst_path)

            source_inventory = extract_semantic_inventory(src_path)
            with self.assertRaises(ContentIntegrityError):
                verify_content_integrity(source_inventory, dst_path)

    def test_verify_content_integrity_checks_hyperlink_targets(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"

            def add_link(document, target):
                paragraph = document.add_paragraph()
                rid = document.part.relate_to(target, RT.HYPERLINK, is_external=True)
                hyperlink = parse_xml(
                    f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}">'
                    f'<w:r><w:t>link</w:t></w:r></w:hyperlink>'
                )
                paragraph._p.append(hyperlink)

            src = Document()
            add_link(src, "https://example.invalid/source")
            src.save(src_path)
            dst = Document()
            add_link(dst, "https://example.invalid/rewritten")
            dst.save(dst_path)

            source_inventory = extract_semantic_inventory(src_path)
            with self.assertRaises(ContentIntegrityError):
                verify_content_integrity(source_inventory, dst_path)

    def test_verify_content_integrity_allows_controlled_field_cache_update(self):
        with tempfile.TemporaryDirectory() as tmp:
            src_path = Path(tmp) / "src.docx"
            dst_path = Path(tmp) / "dst.docx"

            def make_doc(cache):
                document = Document()
                paragraph = document.add_paragraph("Figure ")
                paragraph._p.append(parse_xml(
                    f'<w:fldSimple {nsdecls("w")} w:instr="SEQ Figure">'
                    f'<w:r><w:t>{cache}</w:t></w:r></w:fldSimple>'
                ))
                paragraph.add_run(" caption")
                return document

            make_doc("1").save(src_path)
            make_doc("2").save(dst_path)
            source_inventory = extract_semantic_inventory(src_path)
            self.assertTrue(verify_content_integrity(source_inventory, dst_path))

    def test_expected_inventory_verifies_only_declared_content_part(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_path = root / "source.docx"
            target_path = root / "target.docx"
            source = Document()
            source.add_paragraph("managed body")
            source.save(source_path)

            target = Document()
            target.add_paragraph("managed body")
            from lib.composition import mark_start, mark_end
            mark_start(target, "body")
            mark_end(target, "body")
            cover = target.add_paragraph("generated cover text")
            target.element.body.insert(0, cover._p)
            target.save(target_path)

            expected = build_expected_inventory(
                SimpleNamespace(
                    source_order=(str(source_path),),
                    source_docx_path=source_path,
                    nodes=(),
                    parts={"body": SimpleNamespace(source_region="entire_document")},
                    config=SimpleNamespace(regions={}),
                ),
                spec={"id": "main", "parts": ["body"]},
            )
            self.assertEqual(expected.content_parts, ("body",))
            self.assertTrue(verify_content_integrity(expected, target_path))

    def test_expected_inventory_maps_multi_source_objects_to_output_anchors(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            from PIL import Image as PILImage

            first_image = root / "first.png"
            second_image = root / "second.png"
            PILImage.new("RGB", (10, 10), color="red").save(first_image)
            PILImage.new("RGB", (10, 10), color="blue").save(second_image)

            first = Document()
            first.add_paragraph("first source host").add_run().add_picture(str(first_image), width=Inches(1))
            first.save(root / "first.docx")
            second = Document()
            second.add_paragraph("second source host").add_run().add_picture(str(second_image), width=Inches(1))
            second.save(root / "second.docx")

            expected = ExpectedInventory(
                merge_semantic_inventories([
                    extract_semantic_inventory(root / "first.docx"),
                    extract_semantic_inventory(root / "second.docx"),
                ]),
                source_order=(str(root / "first.docx"), str(root / "second.docx")),
            )
            delivery = Document()
            delivery.add_paragraph("first source host").add_run().add_picture(str(first_image), width=Inches(1))
            delivery.add_paragraph("second source host").add_run().add_picture(str(second_image), width=Inches(1))
            delivery_path = root / "delivery.docx"
            delivery.save(delivery_path)

            evidence = {}
            self.assertTrue(verify_content_integrity(expected, delivery_path, evidence_out=evidence))
            mappings = [item for item in evidence["object_mapping"] if item["kind"] == "source_object"]
            self.assertEqual(len(mappings), 2)
            self.assertEqual(
                {item["source_sha256"] for item in mappings},
                {
                    extract_semantic_inventory(root / "first.docx").source_sha256,
                    extract_semantic_inventory(root / "second.docx").source_sha256,
                },
            )
            self.assertTrue(all(item["output_host_element"] for item in mappings))

    def test_generated_cover_image_requires_one_declared_finite_allowance(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            from PIL import Image as PILImage

            image = root / "cover-logo.png"
            PILImage.new("RGB", (10, 10), color="green").save(image)
            source = root / "source.docx"
            source_doc = Document()
            source_doc.add_paragraph("managed source")
            source_doc.save(source)

            template = root / "cover-template.docx"
            template_doc = Document()
            template_doc.add_paragraph("{{MAIN_TITLE}}")
            template_doc.add_picture(str(image), width=Inches(1))
            template_doc.save(template)

            prepared = SimpleNamespace(
                source_order=(str(source),),
                source_docx_path=source,
                nodes=(),
                parts={"body": SimpleNamespace(source_region="entire_document")},
                config=SimpleNamespace(regions={}),
                cover_template_path=template,
            )
            expected = build_expected_inventory(
                prepared,
                spec={"id": "main", "parts": ["cover", "body"]},
            )
            self.assertEqual(
                [item["rule"] for item in expected.allowed_generated_objects],
                ["cover_template_occurrence"],
            )

            delivery = Document()
            cover_start = delivery.add_paragraph("cover")
            add_bookmark(cover_start._p, "_Synth_cover", 1)
            cover_image = delivery.add_paragraph()
            cover_image.add_run().add_picture(str(image), width=Inches(1))
            add_bookmark(cover_image._p, "_Synth_part_cover_end", 2)
            body_start = delivery.add_paragraph()
            add_bookmark(body_start._p, "_Synth_body", 3)
            delivery.add_paragraph("managed source")
            body_end = delivery.add_paragraph()
            add_bookmark(body_end._p, "_Synth_part_body_end", 4)
            delivery_path = root / "delivery.docx"
            delivery.save(delivery_path)

            evidence = {}
            self.assertTrue(verify_content_integrity(expected, delivery_path, evidence_out=evidence))
            generated_mappings = [
                item for item in evidence["object_mapping"]
                if item["kind"] == "generated_object"
            ]
            self.assertEqual([item["part_id"] for item in generated_mappings], ["cover"])
            self.assertEqual([item["output_part_ids"] for item in generated_mappings], [["cover"]])

            duplicate = Document(delivery_path)
            duplicate.add_picture(str(image), width=Inches(1))
            duplicate_path = root / "duplicate.docx"
            duplicate.save(duplicate_path)
            with self.assertRaisesRegex(ContentIntegrityError, "生成许可"):
                verify_content_integrity(expected, duplicate_path)

            report = verify_generated_object_allowances(delivery_path, expected)
            self.assertTrue(report["passed"], report)

    def test_verify_delivery_format_checks_effective_styles(self):
        resolved_format = ResolvedFormat(
            id="test-format",
            version="1.0.0",
            page=PageSpec(),
            styles={
                "body": StyleDefinition(
                    name="Synth Body",
                    run=RunStyle(east_asia="宋体", latin="Times New Roman", size_pt=12.0),
                    paragraph=ParagraphStyle(alignment="justify"),
                ),
                "heading.1": StyleDefinition(
                    name="Synth Heading 1",
                    run=RunStyle(east_asia="黑体", latin="Arial", size_pt=16.0, bold=True),
                    paragraph=ParagraphStyle(alignment="left"),
                ),
            },
            roles={
                "body": RoleSpec(style="body"),
                "heading.1": RoleSpec(style="heading.1", outline_level=1, include_in_toc=True),
            },
            toc=TocSpec(),
            header=HeaderFooterSpec(),
            footer=HeaderFooterSpec(),
            required_capabilities=["styles.fonts.basic"],
            provenance={},
            content_hash="mock-hash",
        )

        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "delivered.docx"
            doc = Document()
            install_styles(doc, resolved_format)
            section = doc.sections[0]
            section.page_width = Mm(210.0)
            section.page_height = Mm(297.0)
            section.top_margin = Mm(25.4)
            section.bottom_margin = Mm(25.4)
            section.left_margin = Mm(31.8)
            section.right_margin = Mm(31.8)

            p1 = doc.add_paragraph("第一章 综述")
            p2 = doc.add_paragraph("正文测试内容。")

            role_map = {
                "/w:document/w:body/w:p[1]": "heading.1",
                "/w:document/w:body/w:p[2]": "body",
            }
            apply_roles(doc, role_map, resolved_format)
            doc.save(str(doc_path))

            report = verify_delivery_format(doc_path, resolved_format)
            self.assertTrue(report.passed)
            self.assertEqual(report.violation_count, 0)
            self.assertIn("body", report.verified_roles)
            self.assertIn("heading.1", report.verified_roles)
            self.assertTrue(report.details["items"])
            self.assertIn("expected", report.details["items"][0])
            self.assertIn("actual", report.details["items"][0])
            self.assertIn("provenance", report.details["items"][0])
            self.assertIn("coverage", report.details["items"][0])

    def test_verify_delivery_format_checks_every_run_not_first_run_only(self):
        resolved_format = ResolvedFormat(
            id="test-format",
            version="1.0.0",
            page=PageSpec(),
            styles={
                "body": StyleDefinition(
                    name="Synth Body",
                    run=RunStyle(east_asia="宋体", latin="Times New Roman", size_pt=12.0),
                    paragraph=ParagraphStyle(alignment="justify"),
                )
            },
            roles={"body": RoleSpec(style="body")},
            toc=TocSpec(),
            header=HeaderFooterSpec(),
            footer=HeaderFooterSpec(),
            required_capabilities=["styles.fonts.basic"],
            provenance={},
            content_hash="mixed-run",
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "mixed-run.docx"
            doc = Document()
            install_styles(doc, resolved_format)
            section = doc.sections[0]
            section.page_width = Mm(resolved_format.page.width_mm)
            section.page_height = Mm(resolved_format.page.height_mm)
            section.top_margin = Mm(resolved_format.page.margin_top_mm)
            section.bottom_margin = Mm(resolved_format.page.margin_bottom_mm)
            section.left_margin = Mm(resolved_format.page.margin_left_mm)
            section.right_margin = Mm(resolved_format.page.margin_right_mm)
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("first")
            second = paragraph.add_run("second")
            apply_roles(doc, {"/w:document/w:body/w:p[1]": "body"}, resolved_format)
            # 在目标样式已应用后篡改第二个 run，回归“不能只看首个 run”。
            first.font.size = Pt(12)
            second.font.size = Pt(18)
            doc.save(path)

            report = verify_delivery_format(path, resolved_format)
            self.assertFalse(report.passed)
            self.assertTrue(any("字号" in item for item in report.violations))


if __name__ == "__main__":
    unittest.main()
