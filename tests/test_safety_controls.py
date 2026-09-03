"""第一阶段安全门槛的无 Office 回归测试。"""

from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.config import ConfigError, ProjectConfig
from lib.engine import UnifiedSynthesizer
from lib.qa import OfficeExportError, get_exact_printed_heading_pages, word_automation_status
from lib.scanner import flatten_tree_nodes
from lib.source_strategies import build_outline
from lib.renderers import render_docx_file, render_pdf_file, sanitize_ole_and_external_links, set_pdf_render_cache_dir
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx import Document


class SafetyControlsTest(unittest.TestCase):
    def test_word_export_failure_never_returns_fake_page_one_map(self):
        item = {"title": "一、测试章节"}
        with patch("lib.qa.export_docx_to_pdf", return_value=False):
            with self.assertRaises(OfficeExportError):
                get_exact_printed_heading_pages("missing.docx", [item])

    def test_word_automation_probe_distinguishes_tcc_permission_failure(self):
        completed = type("Completed", (), {
            "returncode": 1,
            "stdout": "",
            "stderr": "Not authorized to send Apple events to Microsoft Word. (-1743)",
        })()
        with patch("lib.qa.word_export_status", return_value=(True, "Word 已安装")), patch(
            "lib.qa.subprocess.run", return_value=completed
        ):
            passed, message = word_automation_status()
        self.assertFalse(passed)
        self.assertIn("自动化", message)
        self.assertIn("Full Disk Access", message)

    def test_word_automation_probe_distinguishes_restricted_gui_context(self):
        completed = type("Completed", (), {
            "returncode": 1,
            "stdout": "",
            "stderr": 'Can’t get application id "com.microsoft.Word". (-1728)',
        })()
        with patch("lib.qa.word_export_status", return_value=(True, "Word 已安装")), patch(
            "lib.qa.subprocess.run", return_value=completed
        ):
            passed, message = word_automation_status()
        self.assertFalse(passed)
        self.assertIn("受限上下文", message)

    def test_word_automation_probe_reports_success_without_touching_documents(self):
        completed = type("Completed", (), {
            "returncode": 0,
            "stdout": "Microsoft Word\n",
            "stderr": "",
        })()
        with patch("lib.qa.word_export_status", return_value=(True, "Word 已安装")), patch(
            "lib.qa.subprocess.run", return_value=completed
        ) as run:
            passed, message = word_automation_status()
        self.assertTrue(passed)
        self.assertIn("Automation 探针成功", message)
        self.assertIn("get name", run.call_args.args[0][-1])

    def test_explicit_false_cover_template_disables_auto_discovery(self):
        config = ProjectConfig({
            "schema_version": 1,
            "project_name": "测试项目",
            "cover": {"template": False},
        }, ROOT / "input")
        self.assertIsNone(config.get_template_path())

    def test_anonymous_demo_has_a_read_only_plan_without_a_template(self):
        demo_root = ROOT / "examples" / "minimal-demo"
        plan = UnifiedSynthesizer.plan(demo_root / "source", demo_root / "manifest.json")
        self.assertEqual(plan["project"], "Document Synthesis Demo")
        self.assertIsNone(plan["template"])
        self.assertEqual(plan["node_count"], 2)

    def test_no_automation_script_closes_every_office_document(self):
        source_files = [
            ROOT / "lib" / "qa.py",
            ROOT / "lib" / "engine.py",
            ROOT / "lib" / "highlighted_docx.py",
        ]
        combined = "\n".join(path.read_text(encoding="utf-8") for path in source_files)
        self.assertNotIn("close every document", combined)
        self.assertNotIn("close every presentation", combined)

    def test_explicit_tree_is_validated_and_gets_sequential_bookmarks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_dir = Path(temp_dir)
            (source_dir / "evidence.docx").touch()
            config = ProjectConfig({
                "schema_version": 1,
                "project_name": "测试项目",
                "source": {
                    "strategy": "explicit_tree",
                    "tree": [{
                        "level": 1,
                        "title": "一、材料",
                        "type": "folder",
                        "children": [{
                            "level": 2,
                            "title": "1. 佐证",
                            "type": "docx",
                            "file": "evidence.docx",
                        }],
                    }],
                },
            }, source_dir)
            nodes = flatten_tree_nodes(build_outline(source_dir, config))
            self.assertEqual([node["bm_id"] for node in nodes], [1, 2])
            self.assertEqual([node["bookmark_name"] for node in nodes], ["_Toc_auto_001", "_Toc_auto_002"])

    def test_invalid_versioned_config_is_rejected_instead_of_silent_fallback(self):
        with self.assertRaises(ConfigError):
            ProjectConfig({"project_name": "测试项目"})
        with self.assertRaises(ConfigError):
            ProjectConfig({
                "schema_version": 1,
                "project_name": "测试项目",
                "output": {"compiled_document": "../unsafe.docx"},
            })

    def test_docx_outline_writes_bookmarks_into_original_heading_paragraphs(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "evidence.docx"
            source = Document()
            source.add_paragraph("一、原始一级标题")
            source.add_paragraph("1. 原始二级标题")
            source.save(source_path)
            config = ProjectConfig({
                "schema_version": 1,
                "project_name": "测试项目",
                "source": {"strategy": "explicit_tree", "tree": [{
                    "level": 1,
                    "title": "材料",
                    "type": "docx_outline",
                    "file": "evidence.docx",
                    "outline_rules": [
                        {"level": 1, "pattern": "^一、"},
                        {"level": 2, "pattern": "^1\\."},
                    ],
                }]},
            }, temp_path)
            outline = build_outline(temp_path, config)[0]
            output = Document()
            render_docx_file(output, source_path, embedded_outline=outline["children"])
            bookmark_names = [
                element.get(qn("w:name"))
                for element in output._element.body.iter(qn("w:bookmarkStart"))
            ]
            self.assertEqual(
                bookmark_names,
                [child["bookmark_name"] for child in outline["children"]],
            )

    def test_ole_sanitizer_handles_cloned_elements_without_xpath_namespaces(self):
        element = parse_xml(
            f'<w:p {nsdecls("w")}><w:object><w:fldSimple w:instr=" LINK Excel "/></w:object></w:p>'
        )
        sanitize_ole_and_external_links(element)
        self.assertEqual(len(list(element.iter(qn("w:object")))), 0)
        self.assertEqual(len(list(element.iter(qn("w:fldSimple")))), 0)
        self.assertEqual(len(list(element.iter(qn("w:pict")))), 1)

    def test_qa_accepts_a_scanned_page_that_has_visible_pixels_but_no_text(self):
        import pymupdf
        from lib.qa import run_qa_assertions

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "scan.pdf"
            pdf = pymupdf.open()
            page = pdf.new_page()
            page.draw_rect(pymupdf.Rect(80, 100, 400, 500), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
            pdf.save(pdf_path)
            pdf.close()
            self.assertTrue(run_qa_assertions(str(pdf_path), ignore_front_pages=0))

    def test_pdf_renderer_uses_page_break_before_instead_of_empty_page_break_paragraphs(self):
        import pymupdf

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            pdf_path = temp_path / "two-pages.pdf"
            pdf = pymupdf.open()
            for text in ("第一页", "第二页"):
                page = pdf.new_page()
                page.insert_text((72, 72), text)
            pdf.save(pdf_path)
            pdf.close()

            set_pdf_render_cache_dir(temp_path / "cache")
            output = Document()
            render_pdf_file(output, pdf_path)
            xml = output._element.body.xml
            self.assertEqual(xml.count("<w:pageBreakBefore"), 1)
            self.assertNotIn('<w:br w:type="page"', xml)


if __name__ == "__main__":
    unittest.main()
