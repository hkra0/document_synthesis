# -*- coding: utf-8 -*-
"""
HTML 格式校正报告生成与格式包编译器单元测试 (tests/test_format_review.py)
验证 HTML 报告自包含与转义、决策合并与 JSON 校验、零样本正文泄露保护。
"""

import json
import copy
import tempfile
import unittest
from pathlib import Path

import jsonschema
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt

from lib.config import ConfigError
from lib.contracts import ContractError, validate_decisions_data
from lib.format_analysis import analyze_format_sample
from lib.format_review import (
    build_decisions_example,
    generate_html_review,
    compile_format_package,
    compile_role_mapping,
    load_format_schema,
)


class FormatReviewTest(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_dir = Path(self.temp_dir.name)

        # 构造包含特殊字符的样本文档
        self.sample_path = self.test_dir / "sample_special.docx"
        doc = Document()
        p1 = doc.add_paragraph("安全测试 <script>alert('xss')</script> & 标题 \"Special\"")
        p1._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
        p1.runs[0].font.name = "黑体"
        p1.runs[0].font.size = Pt(18)

        p2 = doc.add_paragraph("秘密机密正文片段：敏感数据 123456789，严禁泄漏到格式包定义中。")
        p2.runs[0].font.name = "仿宋"
        p2.runs[0].font.size = Pt(12)
        doc.save(str(self.sample_path))

        self.analysis_data = analyze_format_sample(self.sample_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_generate_html_review_is_self_contained_and_escaped(self):
        """测试生成的 HTML 校正报告严格转义特殊字符且自包含零外部资源依赖"""
        out_html = self.test_dir / "review.html"
        generated_path = generate_html_review(self.analysis_data, out_html)

        self.assertTrue(generated_path.is_file())
        html_content = generated_path.read_text(encoding="utf-8")

        # 检查特殊字符被转义，不存在裸的 <script>alert
        self.assertNotIn("<script>alert('xss')</script>", html_content)
        self.assertIn("&lt;script&gt;alert(&#x27;xss&#x27;)&lt;/script&gt;", html_content)
        self.assertIn("&amp;", html_content)

        # 检查无外部网络 CDN 链接
        self.assertNotIn("http://", html_content)
        self.assertNotIn("https://", html_content)

        # 检查核心功能板块存在
        self.assertIn("样本文档格式分析与校正报告", html_content)
        self.assertIn("视觉样式聚类与语义角色候选", html_content)
        self.assertIn("虚构内容排版预览", html_content)
        self.assertIn("decisionsPre", html_content)
        self.assertIn("node-role-select", html_content)
        self.assertIn("data-initial-role", html_content)
        self.assertIn("currentDecisions.node_roles", html_content)
        self.assertIn('value="heading.9"', html_content)
        self.assertIn("if (!roleStyles[role])", html_content)
        self.assertIn("review_audit", html_content)
        self.assertIn("reviewerAttested", html_content)
        self.assertIn("recordReviewOperation", html_content)

    def test_review_audit_is_bound_to_report_and_never_fakes_attestation(self):
        decisions = build_decisions_example(self.analysis_data)
        decisions["review_audit"] = {
            "audit_schema_version": 1,
            "report_id": self.analysis_data["report_id"],
            "source_sha256": self.analysis_data["source_sha256"],
            "origin": "test_fixture",
            "automation": True,
            "reviewer_attested": False,
            "operation_count": 1,
            "operations": [{
                "sequence": 1,
                "operation": "set_node_role",
                "target": "/w:document/w:body/w:p[1]",
                "from": "heading.2",
                "to": "heading.3",
            }],
        }
        validate_decisions_data(decisions, self.analysis_data)

        browser_attested = copy.deepcopy(decisions)
        browser_attested["review_audit"].update({
            "origin": "browser_ui",
            "automation": False,
            "reviewer_attested": True,
            "started_at": "2026-09-08T00:00:00Z",
            "completed_at": "2026-09-08T00:01:00Z",
        })
        # The contract accepts a real browser attestation shape; this test
        # does not claim that a human performed the action.
        validate_decisions_data(browser_attested, self.analysis_data)

        automated_attestation = copy.deepcopy(decisions)
        automated_attestation["review_audit"]["reviewer_attested"] = True
        with self.assertRaisesRegex(ContractError, "自动化 review_audit"):
            validate_decisions_data(automated_attestation, self.analysis_data)

        imported_attestation = copy.deepcopy(decisions)
        imported_attestation["review_audit"]["origin"] = "imported"
        imported_attestation["review_audit"]["automation"] = False
        imported_attestation["review_audit"]["reviewer_attested"] = True
        with self.assertRaisesRegex(ContractError, "只有 browser_ui"):
            validate_decisions_data(imported_attestation, self.analysis_data)

        mismatched_audit = copy.deepcopy(decisions)
        mismatched_audit["review_audit"]["source_sha256"] = "0" * 64
        with self.assertRaisesRegex(ContractError, "review_audit 的源哈希"):
            validate_decisions_data(mismatched_audit, self.analysis_data)

    def test_compile_format_package_matches_schema_and_zero_sample_leak(self):
        """测试格式包编译通过 Schema 校验，且零样本正文泄露"""
        decisions_data = {
            "report_id": self.analysis_data["report_id"],
            "source_sha256": self.analysis_data["source_sha256"],
            "role_styles": {
                "body": self.analysis_data["candidate_roles"]["body"]["cluster_id"],
                "heading.1": self.analysis_data["candidate_roles"]["heading.1"]["cluster_id"],
            },
            "style_overrides": {},
            "missing_roles": {},
        }

        out_format_json = self.test_dir / "compiled_format.json"
        pkg = compile_format_package(
            analysis_data=self.analysis_data,
            decisions_data=decisions_data,
            format_id="test-compiled-format",
            format_version="1.0.0",
            output_path=out_format_json,
        )

        self.assertTrue(out_format_json.is_file())
        self.assertEqual(pkg["id"], "test-compiled-format")
        self.assertEqual(pkg["version"], "1.0.0")

        # 验证符合 schemas/format-v1.schema.json
        schema = load_format_schema()
        jsonschema.validate(instance=pkg, schema=schema)

        # 严格断言：格式包 JSON 字符串中绝对没有原始样本文档的涉密正文字符串
        pkg_json_str = out_format_json.read_text(encoding="utf-8")
        self.assertNotIn("秘密机密正文片段", pkg_json_str)
        self.assertNotIn("123456789", pkg_json_str)
        self.assertNotIn("严禁泄漏", pkg_json_str)

    def test_compile_accepts_missing_role_resolved_by_manual_cluster_assignment(self):
        decisions_data = {
            "decisions_schema_version": 1,
            "report_id": self.analysis_data["report_id"],
            "source_sha256": self.analysis_data["source_sha256"],
            "role_styles": {
                "body": "cluster-02",
                "heading.1": "cluster-01",
                # Resolve a role that the first-pass report marked missing.
                "heading.2": "cluster-01",
            },
            "style_overrides": {},
            "missing_roles": {
                "title": "reject",
                "heading.3": "reject",
                "table.body": "reject",
            },
        }
        package = compile_format_package(
            self.analysis_data,
            decisions_data,
            format_id="manual-cluster-resolution",
        )
        self.assertIn("heading.2", package["roles"])

    def test_compile_format_package_rejects_mismatched_report_or_hash(self):
        """测试决策文件与分析报告不匹配时拒绝编译"""
        mismatched_decisions = {
            "report_id": "wrong-id",
            "source_sha256": self.analysis_data["source_sha256"],
            "role_styles": {},
        }
        with self.assertRaises(ConfigError):
            compile_format_package(self.analysis_data, mismatched_decisions)

        mismatched_hash = {
            "report_id": self.analysis_data["report_id"],
            "source_sha256": "0" * 64,
            "role_styles": {},
        }
        with self.assertRaises(ConfigError):
            compile_format_package(self.analysis_data, mismatched_hash)

    def test_compile_role_mapping(self):
        """测试编译目标文档角色映射文件 roles.json"""
        decisions_data = {
            "report_id": self.analysis_data["report_id"],
            "source_sha256": self.analysis_data["source_sha256"],
            "node_roles": {
                "/w:document/w:body/w:p[1]": "heading.1",
                "/w:document/w:body/w:p[2]": "body",
            },
        }
        out_roles_json = self.test_dir / "roles.json"
        mapping = compile_role_mapping(self.analysis_data, decisions_data, output_path=out_roles_json)

        self.assertTrue(out_roles_json.is_file())
        self.assertEqual(mapping["source_sha256"], self.analysis_data["source_sha256"])
        self.assertEqual(len(mapping["assignments"]), len(self.analysis_data["node_roles"]))


if __name__ == "__main__":
    unittest.main()
