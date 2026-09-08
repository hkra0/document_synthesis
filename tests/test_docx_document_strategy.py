# -*- coding: utf-8 -*-
"""
单 DOCX 来源策略测试 (tests/test_docx_document_strategy.py)
验证 docx_document 策略下无需高亮的大纲解析、纯正文无标题合法性与目录约束。
"""

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from lib.config import load_project_config
from lib.engine import BuildError, UnifiedSynthesizer
from lib.source_strategies import build_outline, SourceStrategyError


class DocxDocumentStrategyTest(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_dir = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_document_extracts_headings_without_highlights(self):
        """测试 docx_document 策略提取标准大纲级标题，无需任何黄色高亮"""
        doc_path = self.test_dir / "standard_headings.docx"
        doc = Document()
        
        p1 = doc.add_paragraph("第一章 项目背景")
        p1._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
        
        doc.add_paragraph("这是背景描述文字。")
        
        p2 = doc.add_paragraph("1.1 需求说明")
        p2._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="1"/>'))
        
        doc.add_paragraph("这是需求内容。")
        doc.save(str(doc_path))

        manifest = {
            "project_name": "Standard Docx Test",
            "schema_version": 2,
            "source": {
                "strategy": "docx_document",
                "file": "standard_headings.docx"
            },
            "output": {
                "documents": [
                    {"id": "main", "filename": "main.docx", "parts": ["body"]}
                ]
            }
        }
        (self.test_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        config = load_project_config(self.test_dir)

        tree = build_outline(self.test_dir, config)
        self.assertEqual(len(tree), 1)
        root = tree[0]
        self.assertEqual(root["type"], "docx_document")
        self.assertEqual(root["content_block_count"], 4)
        self.assertEqual(root["heading_count"], 2)

        children = root["children"]
        self.assertEqual(len(children), 2)
        self.assertEqual(children[0]["title"], "第一章 项目背景")
        self.assertEqual(children[0]["level"], 1)
        self.assertEqual(children[1]["title"], "1.1 需求说明")
        self.assertEqual(children[1]["level"], 2)

    def test_body_only_document_is_valid_in_plan(self):
        """测试纯正文无标题 DOCX 是合法输入，plan 正常返回 content_block_count > 0 且 heading_count == 0"""
        doc_path = self.test_dir / "body_only.docx"
        doc = Document()
        doc.add_paragraph("没有任何大纲级别或标题样式的第一段正文。")
        doc.add_paragraph("第二段正文。")
        doc.save(str(doc_path))

        manifest = {
            "project_name": "Body Only Test",
            "schema_version": 2,
            "source": {
                "strategy": "docx_document",
                "file": "body_only.docx"
            },
            "output": {
                "documents": [
                    {"id": "main", "filename": "main.docx", "parts": ["body"]}
                ]
            }
        }
        (self.test_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

        plan = UnifiedSynthesizer.plan(self.test_dir)
        self.assertEqual(plan["content_block_count"], 2)
        self.assertEqual(plan["heading_count"], 0)
        # 不生成 toc 时无缺少标题的错误
        self.assertEqual(len([w for w in plan["warnings"] if "未包含各级标题" in w]), 0)

    def test_body_only_document_warns_when_toc_requested(self):
        """测试纯正文文档在请求生成目录 (toc) 时，plan 会给出警告"""
        doc_path = self.test_dir / "body_only_toc.docx"
        doc = Document()
        doc.add_paragraph("普通正文。")
        doc.save(str(doc_path))

        manifest = {
            "project_name": "Body Only Toc Test",
            "schema_version": 2,
            "source": {
                "strategy": "docx_document",
                "file": "body_only_toc.docx"
            },
            "output": {
                "documents": [
                    {"id": "main", "filename": "main.docx", "parts": ["toc", "body"]}
                ]
            }
        }
        (self.test_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

        plan = UnifiedSynthesizer.plan(self.test_dir)
        self.assertEqual(plan["heading_count"], 0)
        warnings = [w for w in plan["warnings"] if "未包含各级标题" in w]
        self.assertGreaterEqual(len(warnings), 1)

    def test_direct_single_file_source_path(self):
        """测试直接把 source_dir 指向单个 .docx 文件时 docx_document 策略的自适应"""
        doc_path = self.test_dir / "standalone.docx"
        doc = Document()
        p = doc.add_paragraph("独立文件一级标题")
        p._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
        doc.add_paragraph("正文内容")
        doc.save(str(doc_path))

        manifest = {
            "project_name": "Standalone Test",
            "schema_version": 2,
            "source": {
                "strategy": "docx_document"
            },
            "output": {
                "documents": [
                    {"id": "main", "filename": "main.docx", "parts": ["body"]}
                ]
            }
        }
        (self.test_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

        plan = UnifiedSynthesizer.plan(doc_path, manifest_path=self.test_dir / "manifest.json")
        self.assertEqual(plan["heading_count"], 1)
        self.assertEqual(plan["content_block_count"], 2)


if __name__ == "__main__":
    unittest.main()
