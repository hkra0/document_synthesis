# -*- coding: utf-8 -*-
"""P3 真实 Windows Word 页码测量与故障注入测试。

只在 Windows、``DOCUMENT_SYNTHESIS_WORD_TEST=1`` 且检测到 Word 时运行；否则整体跳过
（跳过不算通过）。覆盖计划第 8 节退出条件中的四种故障：

1. 强制结束 Word：操作进行中按 pid 文件结束本次启动的 WINWORD.EXE
2. 目标文件被占用：目标 PDF 被其他进程打开
3. 书签缺失
4. 超时：Word 启动阶段与文档处理阶段各一次

每种故障都要求：异常信息可读（中文、说明原因）；目标文件不被半成品替换；
测试结束时没有本次启动的 WINWORD.EXE 残留。
"""

import os
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from typing import Set
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from lib.office import OfficeBackendError, OfficeTimeoutError, set_backend
from lib.office import win_com
from lib.office.win_com import WinComBackend


def _bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _restart_page_numbers(section, start: int = 1) -> None:
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), str(start))
    section._sectPr.append(pg)


def build_sample(path: Path, filler_pages: int = 0) -> None:
    """封面 2 页（第 1 节）+ 正文（第 2 节，页码从 1 重新开始）。

    书签：cover → 物理 1 / 打印 1；body → 物理 3 / 打印 1；later → 物理 4 / 打印 2。
    ``filler_pages`` 在末尾追加大量文字，让 Word 处理变慢，用于故障注入。
    """
    doc = Document()
    _bookmark(doc.add_paragraph("封面"), "cover", 1)
    doc.add_page_break()
    doc.add_paragraph("前言")
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    _restart_page_numbers(body_section, 1)
    _bookmark(doc.add_paragraph("正文第一页"), "body", 2)
    doc.add_page_break()
    _bookmark(doc.add_paragraph("正文第二页"), "later", 3)
    for index in range(filler_pages * 40):
        doc.add_paragraph(f"填充段落 {index}：用于让 Word 处理时间变长的测试文字。" * 3)
    doc.save(path)


def _word_pids() -> Set[int]:
    return set(win_com.list_process_ids(win_com.WORD_IMAGE))


@unittest.skipUnless(sys.platform == "win32", "Windows Word COM 故障注入只在 Windows 上运行")
class WindowsWordFaultTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if os.environ.get("DOCUMENT_SYNTHESIS_WORD_TEST") != "1":
            raise unittest.SkipTest("BLOCKED: set DOCUMENT_SYNTHESIS_WORD_TEST=1 to run Windows Word fault tests")
        status = WinComBackend().static_status()
        if not status.available:
            raise unittest.SkipTest(f"BLOCKED: {status.reason}")

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.work = self.dir / "word-access"
        self.env = patch.dict(os.environ, {"DOCUMENT_SYNTHESIS_WORD_ACCESS_DIR": str(self.work)})
        self.env.start()
        self.before = _word_pids()
        self.backend = WinComBackend()

    def tearDown(self) -> None:
        self.env.stop()
        set_backend(None)
        self.tmp.cleanup()

    def assertNoWordLeft(self) -> None:  # noqa: N802
        deadline = time.monotonic() + 30
        leftover = _word_pids() - self.before
        while leftover and time.monotonic() < deadline:
            time.sleep(0.5)
            leftover = _word_pids() - self.before
        self.assertEqual(leftover, set(), "本次启动的 WINWORD.EXE 未被回收")
        self.assertEqual(list(self.work.glob("job-*")), [], "工作目录未清理")

    def test_measures_physical_and_printed_pages(self) -> None:
        docx, pdf = self.dir / "sample.docx", self.dir / "sample.pdf"
        build_sample(docx)
        report = self.backend.inspect_bookmarks(docx, pdf, ["cover", "body", "later"])
        self.assertEqual(report.fidelity, "exact")
        self.assertEqual(report.total_pages, 4)
        self.assertEqual(report.bookmarks, {
            "cover": {"physical_page": 1, "printed_page": 1},
            "body": {"physical_page": 3, "printed_page": 1},
            "later": {"physical_page": 4, "printed_page": 2},
        })
        self.assertTrue(pdf.is_file())
        self.assertNoWordLeft()

    def test_missing_bookmark(self) -> None:
        docx, pdf = self.dir / "sample.docx", self.dir / "sample.pdf"
        build_sample(docx)
        with self.assertRaises(OfficeBackendError) as ctx:
            self.backend.inspect_bookmarks(docx, pdf, ["cover", "absent_mark"])
        self.assertIn("缺少书签: absent_mark", str(ctx.exception))
        self.assertFalse(pdf.exists())
        self.assertNoWordLeft()

    def test_locked_target(self) -> None:
        docx, pdf = self.dir / "sample.docx", self.dir / "sample.pdf"
        build_sample(docx)
        pdf.write_bytes(b"previous")
        with open(pdf, "rb"):
            with self.assertRaises(OfficeBackendError) as ctx:
                self.backend.inspect_bookmarks(docx, pdf, ["cover"])
        self.assertIn("正在", str(ctx.exception))
        self.assertEqual(pdf.read_bytes(), b"previous")
        self.assertEqual(sorted(p.name for p in self.dir.iterdir() if p.suffix == ".tmp"), [])
        self.assertNoWordLeft()

    def test_timeout_during_launch_and_processing(self) -> None:
        docx, pdf = self.dir / "big.docx", self.dir / "big.pdf"
        build_sample(docx, filler_pages=150)
        for seconds in ("1", "8"):
            with self.subTest(timeout=seconds), patch.dict(os.environ, {"DOCUMENT_SYNTHESIS_OFFICE_TIMEOUT": seconds}):
                with self.assertRaises(OfficeTimeoutError) as ctx:
                    self.backend.inspect_bookmarks(docx, pdf, ["cover", "body", "later"])
                self.assertIn("超时", str(ctx.exception))
                self.assertFalse(pdf.exists())
                self.assertNoWordLeft()

    def test_word_killed_during_operation(self) -> None:
        docx, pdf = self.dir / "big.docx", self.dir / "big.pdf"
        build_sample(docx, filler_pages=150)
        outcome = {}

        def run() -> None:
            try:
                outcome["report"] = self.backend.inspect_bookmarks(docx, pdf, ["cover", "body", "later"])
            except Exception as exc:  # noqa: BLE001 - 断言在主线程进行
                outcome["error"] = exc

        worker = threading.Thread(target=run)
        worker.start()
        killed = None
        deadline = time.monotonic() + 120
        while worker.is_alive() and killed is None and time.monotonic() < deadline:
            for pid_file in self.work.glob("job-*/office.pid"):
                record = win_com.read_pid_record(pid_file)
                if record:
                    win_com.terminate_process(record["pid"])
                    killed = record["pid"]
                    break
            time.sleep(0.05)
        worker.join(timeout=300)
        self.assertIsNotNone(killed, "未能在操作期间结束 Word（文档处理过快）")
        self.assertIn("error", outcome, "Word 被结束后操作不应成功")
        self.assertIsInstance(outcome["error"], OfficeBackendError)
        self.assertRegex(str(outcome["error"]), "意外退出|异常退出|Microsoft Word")
        self.assertFalse(pdf.exists())
        self.assertNoWordLeft()

    def test_build_timeout_publishes_nothing(self) -> None:
        from lib.engine import BuildError, UnifiedSynthesizer
        from lib.qa import OfficeExportError

        output_dir = self.dir / "output"
        with patch.dict(os.environ, {"DOCUMENT_SYNTHESIS_OFFICE_TIMEOUT": "3"}):
            with self.assertRaises((BuildError, OfficeExportError)) as ctx:
                demo = ROOT / "examples" / "minimal-demo"
                UnifiedSynthesizer.synthesize(demo / "source", output_dir, demo / "manifest.json")
        self.assertIn("超时", str(ctx.exception))
        self.assertFalse((output_dir / "demo_complete.docx").exists())
        self.assertNoWordLeft()


def _find_open_document(path: Path):
    """在运行对象表中查找已打开的文档；找不到时返回 None。

    不用 ``GetObject(path)``：文档未打开时它会启动新的 Word 实例去打开文件。
    """
    import pythoncom
    import win32com.client

    target = os.path.normcase(str(path))
    rot = pythoncom.GetRunningObjectTable()
    context = pythoncom.CreateBindCtx(0)
    for moniker in rot.EnumRunning():
        try:
            name = moniker.GetDisplayName(context, None)
        except pythoncom.com_error:
            continue
        if os.path.normcase(name) == target:
            return win32com.client.Dispatch(rot.GetObject(moniker).QueryInterface(pythoncom.IID_IDispatch))
    return None


@unittest.skipUnless(sys.platform == "win32", "Windows Word COM 故障注入只在 Windows 上运行")
class UserWordIsolationTests(unittest.TestCase):
    """用户自己打开、带未保存修改的文档，在构建、超时和 Word 被结束后都不受影响。"""

    MARKER = "【未保存的用户修改】"

    @classmethod
    def setUpClass(cls) -> None:
        WindowsWordFaultTests.setUpClass()

    def setUp(self) -> None:
        import win32process

        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.work = self.dir / "word-access"
        self.env = patch.dict(os.environ, {"DOCUMENT_SYNTHESIS_WORD_ACCESS_DIR": str(self.work)})
        self.env.start()
        self.backend = WinComBackend()

        # 以与构建无关的方式打开“用户文档”：和双击文件一样经由 Shell 启动 Word。
        self.word_before = _word_pids()
        self.user_path = self.dir / "用户正在编辑.docx"
        user_doc = Document()
        user_doc.add_paragraph("用户正在编辑的文档")
        user_doc.save(self.user_path)
        self.user_bytes = self.user_path.read_bytes()
        os.startfile(str(self.user_path))
        self.user_document = None
        deadline = time.monotonic() + 90
        while self.user_document is None and time.monotonic() < deadline:
            self.user_document = _find_open_document(self.user_path)
            if self.user_document is None:
                time.sleep(0.5)
        if self.user_document is None:
            self.fail("无法打开模拟的用户文档")
        hwnd = int(self.user_document.Windows(1).Hwnd)
        self.user_pid = int(win32process.GetWindowThreadProcessId(hwnd)[1])
        self.user_identity = win_com.process_identity(self.user_pid)
        self.assertIsNotNone(self.user_identity)
        self.assertFalse(win_com.is_automation_instance(self.user_pid), "模拟的用户 Word 不应是 COM 启动的实例")
        self.user_document.Content.InsertAfter(self.MARKER)
        self.assertFalse(self.user_document.Saved)

    def tearDown(self) -> None:
        try:
            document = _find_open_document(self.user_path)
            if document is not None:
                application = document.Application
                document.Close(0)
                # 只退出本测试启动的 Word；用户原本就在运行的 Word 只关闭这份文档。
                if self.user_pid not in self.word_before and application.Documents.Count == 0:
                    application.Quit(0)
                del application, document
            self.user_document = None
        finally:
            self.env.stop()
            set_backend(None)
            deadline = time.monotonic() + 30
            while self.user_pid not in self.word_before and win_com.process_identity(self.user_pid) and time.monotonic() < deadline:
                time.sleep(0.5)
            self.tmp.cleanup()

    def assertUserWordUntouched(self, step: str) -> None:  # noqa: N802
        self.assertEqual(win_com.process_identity(self.user_pid), self.user_identity, f"{step}: 用户 Word 进程被结束或替换")
        document = _find_open_document(self.user_path)
        self.assertIsNotNone(document, f"{step}: 用户文档不再处于打开状态")
        self.assertFalse(document.Saved, f"{step}: 用户文档的未保存状态丢失")
        self.assertIn(self.MARKER, document.Content.Text, f"{step}: 用户的未保存修改丢失")
        self.assertEqual(self.user_path.read_bytes(), self.user_bytes, f"{step}: 用户文档文件被改写")
        leftover = _word_pids() - self.word_before - {self.user_pid}
        deadline = time.monotonic() + 30
        while leftover and time.monotonic() < deadline:
            time.sleep(0.5)
            leftover = _word_pids() - self.word_before - {self.user_pid}
        self.assertEqual(leftover, set(), f"{step}: 本次启动的 WINWORD.EXE 未被回收")

    def test_user_document_survives_build_timeout_and_killed_word(self) -> None:
        from lib.engine import UnifiedSynthesizer

        # 1. 正常的 exact 构建
        output_dir = self.dir / "output"
        demo = ROOT / "examples" / "minimal-demo"
        UnifiedSynthesizer.synthesize(demo / "source", output_dir, demo / "manifest.json")
        self.assertTrue((output_dir / "demo_complete.docx").is_file())
        self.assertUserWordUntouched("正常构建")

        # 2. 超时：启动阶段（孤儿判定路径）与处理阶段（按 PID 结束路径）
        big = self.dir / "big.docx"
        build_sample(big, filler_pages=150)
        for seconds in ("1", "8"):
            with patch.dict(os.environ, {"DOCUMENT_SYNTHESIS_OFFICE_TIMEOUT": seconds}):
                with self.assertRaises(OfficeTimeoutError):
                    self.backend.inspect_bookmarks(big, self.dir / "timeout.pdf", ["cover"])
            self.assertUserWordUntouched(f"{seconds} 秒超时")

        # 3. 操作进行中强制结束本次启动的 Word
        outcome = {}

        def run() -> None:
            try:
                self.backend.inspect_bookmarks(big, self.dir / "killed.pdf", ["cover"])
            except Exception as exc:  # noqa: BLE001 - 断言在主线程进行
                outcome["error"] = exc

        worker = threading.Thread(target=run)
        worker.start()
        killed = None
        deadline = time.monotonic() + 120
        while worker.is_alive() and killed is None and time.monotonic() < deadline:
            for pid_file in self.work.glob("job-*/office.pid"):
                record = win_com.read_pid_record(pid_file)
                if record:
                    self.assertNotEqual(record["pid"], self.user_pid, "pid 文件记录了用户的 Word")
                    win_com.terminate_process(record["pid"])
                    killed = record["pid"]
                    break
            time.sleep(0.05)
        worker.join(timeout=300)
        self.assertIsNotNone(killed)
        self.assertIsInstance(outcome.get("error"), OfficeBackendError)
        self.assertUserWordUntouched("Word 被强制结束")


if __name__ == "__main__":
    unittest.main()
