# -*- coding: utf-8 -*-
"""R7 样例校正、预览与 manifest 迁移闭环回归。"""

import json
import importlib.util
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt

from lib.config import ConfigError
from lib.engine import UnifiedSynthesizer
from lib.format_analysis import analyze_format_sample
from lib.format_review import build_decisions_example, compile_format_package, compile_role_mapping
from lib.manifest_migration import migrate_manifest_file, migrate_manifest_data
from lib.preview import build_preview


def _load_unstructured_evaluator():
    path = Path(__file__).resolve().parents[1] / "docs" / "acceptance" / "r7-evaluation" / "evaluate_unstructured_dataset.py"
    spec = importlib.util.spec_from_file_location("r7_unstructured_evaluator", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class R7WorkflowTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.sample = self.root / "sample.docx"
        doc = Document()
        title = doc.add_paragraph("样例秘密标题")
        title.alignment = 1
        title.runs[0].font.size = Pt(20)
        heading = doc.add_paragraph("一、样例标题")
        heading._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
        doc.add_paragraph("样例秘密正文，不应进入格式包。")
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "样例表格"
        doc.save(self.sample)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_analysis_exposes_field_evidence_usage_and_confidence(self):
        report = analyze_format_sample(self.sample)
        self.assertGreaterEqual(report["sample_counts"]["paragraphs"], 3)
        self.assertIn("style_inventory", report)
        self.assertIn("defined_only_style_ids", report["style_inventory"])
        self.assertTrue(report["field_evidence"])
        self.assertIn("confidence_summary", report)
        self.assertEqual(report["confidence_summary"]["automatic_accept_threshold"], 0.98)
        self.assertTrue(all("status" in item and "source" in item for item in report["field_evidence"]))

    def test_unstructured_metrics_record_score_calibration_diagnostics(self):
        sample = self.root / "unstructured-eval.docx"
        doc = Document()
        heading = doc.add_paragraph("一、无大纲候选标题")
        heading.runs[0].bold = True
        heading.runs[0].font.size = Pt(15)
        doc.add_paragraph("这是一段足够长的普通正文，用于验证评测结果会记录分数范围，而不是把启发式分数当作概率。")
        doc.save(sample)

        evaluator = _load_unstructured_evaluator()
        result = evaluator.evaluate(
            {
                "dataset_id": "r7-metric-contract",
                "document_count": 1,
                "categories": ["fixture"],
                "documents": [{
                    "filename": sample.name,
                    "category": "fixture",
                    "heading_paragraph_indices": [0],
                    "body_paragraph_indices": [1],
                }],
            },
            self.root,
        )
        overall = result["overall"]
        self.assertEqual(overall["automatic_acceptance_threshold"], 0.98)
        self.assertLessEqual(overall["selected_score_range"]["min"], overall["selected_score_range"]["max"])
        self.assertIs(overall["automatic_threshold_reachable"], False)

    def test_frozen_unstructured_baseline_is_independent_and_meets_gate(self):
        root = Path(__file__).resolve().parents[1] / "docs" / "acceptance" / "r7-evaluation"
        baseline = json.loads((root / "unstructured-human-baseline.json").read_text(encoding="utf-8"))
        metrics = json.loads((root / "unstructured-metrics.json").read_text(encoding="utf-8"))
        self.assertFalse(baseline["algorithm_output_used"])
        self.assertEqual(baseline["review_scope"]["documents"], 20)
        self.assertEqual(baseline["review_scope"]["heading_positives"], 100)
        self.assertEqual(baseline["review_scope"]["body_negatives"], 500)
        self.assertTrue(all(item["review_status"] == "manually_reviewed" for item in baseline["documents"]))
        self.assertGreaterEqual(metrics["overall"]["automatic_acceptance_coverage"], 0.50)
        self.assertGreaterEqual(metrics["overall"]["exact_role_auto_precision"], 0.98)
        self.assertGreaterEqual(metrics["overall"]["automatic_acceptance_precision"], 0.98)
        self.assertEqual(metrics["overall"]["body_as_heading_false_positives"], 0)

    def test_draft_decisions_are_pending_and_cannot_compile(self):
        report = analyze_format_sample(self.sample)
        decisions = build_decisions_example(report)
        self.assertEqual(decisions["decisions_schema_version"], 1)
        if report["missing_roles"]:
            self.assertIn("pending", decisions["missing_roles"].values())
            with self.assertRaisesRegex(ConfigError, "待确认"):
                compile_format_package(report, decisions)

    def test_preview_is_fictional_and_covers_required_roles(self):
        report = analyze_format_sample(self.sample)
        decisions = {
            "report_id": report["report_id"],
            "source_sha256": report["source_sha256"],
            "role_styles": {role: item["cluster_id"] for role, item in report["candidate_roles"].items()},
            "style_overrides": {},
            "missing_roles": {item["role"]: "inherit" for item in report["missing_roles"]},
        }
        package_path = self.root / "format.json"
        compile_format_package(report, decisions, output_path=package_path)
        preview_path = self.root / "preview.html"
        metadata_path = self.root / "preview.json"
        metadata = build_preview(package_path, preview_path, metadata_path=metadata_path)
        html_text = preview_path.read_text(encoding="utf-8")
        self.assertTrue(metadata["approximate"])
        self.assertTrue(metadata["fictional_content"])
        self.assertIn("heading.9", html_text)
        self.assertIn("data-role=\"table.body\"", html_text)
        self.assertIn("虚构图片占位", html_text)
        self.assertIn("近似预览", html_text)
        self.assertNotIn("样例秘密正文", html_text)
        self.assertEqual(json.loads(metadata_path.read_text(encoding="utf-8"))["source_content_included"], False)

    def test_migrate_v1_is_explicit_and_non_destructive(self):
        old = self.root / "old.json"
        new = self.root / "new.json"
        old_data = {
            "schema_version": 1,
            "project_name": "迁移测试",
            "output": {"compiled_document": "old-body.docx"},
        }
        old.write_text(json.dumps(old_data, ensure_ascii=False), encoding="utf-8")
        migrated = migrate_manifest_file(old, new)
        self.assertEqual(migrated["schema_version"], 2)
        self.assertEqual([item["id"] for item in migrated["output"]["documents"]], ["compiled_document", "cover_toc", "toc_body"])
        self.assertEqual(old.read_text(encoding="utf-8"), json.dumps(old_data, ensure_ascii=False))
        with self.assertRaises(ConfigError):
            migrate_manifest_file(old, new)

    def test_migrate_v2_default_documents(self):
        migrated = migrate_manifest_data({"schema_version": 2, "project_name": "v2"})
        self.assertEqual(migrated["output"]["documents"][0]["parts"], ["cover", "toc", "body"])

    def test_compiled_format_and_mapping_reuse_on_another_target(self):
        report = analyze_format_sample(self.sample)
        decisions = build_decisions_example(report)
        decisions["missing_roles"] = {
            role: "inherit" for role in decisions.get("missing_roles", {})
        }

        format_path = self.root / "reusable-format.json"
        compile_format_package(
            report,
            decisions,
            base_format_ref="preset:academic-basic@1.0.0",
            format_id="reusable-demo",
            format_version="1.0.0",
            output_path=format_path,
        )
        role_map_path = self.root / "reusable-role-map.json"
        decisions["node_roles"] = dict(report.get("node_roles", {}))
        compile_role_mapping(report, decisions, output_path=role_map_path)

        target = self.root / "target.docx"
        target.write_bytes(self.sample.read_bytes())
        manifest = self.root / "target-manifest.json"
        manifest.write_text(
            json.dumps(
                {
                    "schema_version": 3,
                    "project_name": "reusable-target",
                    "source": {"strategy": "docx_document", "file": target.name},
                    "format": {"ref": format_path.name},
                    "formatting": {"mode": "restyle", "role_map": role_map_path.name},
                    "output": {
                        "documents": [
                            {"id": "main", "filename": "target.docx", "parts": ["body"]}
                        ]
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        package_text = format_path.read_text(encoding="utf-8")
        self.assertNotIn("样例秘密正文", package_text)
        prepared = UnifiedSynthesizer.prepare_build(target, manifest)
        self.assertEqual(prepared.role_map_path, role_map_path.resolve())
        self.assertTrue(prepared.role_map_hash)
        self.assertEqual(len(prepared.format_hash), 64)
        self.assertTrue(any(item.confirmed for item in prepared.assignments))

    def test_hard_conflict_manual_decision_reuses_format_and_mapping_on_another_target(self):
        """A missing level stays pending until an explicit target decision."""
        source = self.root / "hard-source.docx"
        target = self.root / "hard-target.docx"

        def write_hard_case(path: Path, label: str) -> None:
            document = Document()
            heading_one = document.add_paragraph(f"1. {label} introduction")
            heading_one.runs[0].bold = True
            heading_one.runs[0].font.size = Pt(18)
            document.add_paragraph(f"{label} ordinary body material. " * 24)
            heading_three = document.add_paragraph(f"1.1.1. {label} detailed protocol")
            heading_three.runs[0].bold = True
            heading_three.runs[0].font.size = Pt(14)
            document.add_paragraph(f"{label} detailed body material. " * 24)
            document.save(path)

        write_hard_case(source, "source")
        write_hard_case(target, "target")

        source_report = analyze_format_sample(source)
        conflict_path = "/w:document/w:body/w:p[3]"
        conflict = source_report["unstructured_candidate_evidence"][conflict_path]
        self.assertEqual(conflict["selected_role"], "heading.3")
        self.assertEqual(conflict["review_status"], "pending_review")
        self.assertTrue(any(item["kind"] == "numbering_vs_size_level" for item in conflict["conflicts"]))

        source_decisions = build_decisions_example(source_report)
        source_decisions["missing_roles"] = {
            role: "inherit" for role in source_decisions.get("missing_roles", {})
        }
        source_decisions["review_audit"] = {
            "audit_schema_version": 1,
            "report_id": source_report["report_id"],
            "source_sha256": source_report["source_sha256"],
            "origin": "test_fixture",
            "automation": True,
            "reviewer_attested": False,
            "operation_count": 1,
            "operations": [{
                "sequence": 1,
                "operation": "select_cluster",
                "target": "heading.3",
                "from": "pending",
                "to": "cluster-hard-heading-3",
            }],
        }
        source_decisions_path = self.root / "hard-source.decisions.final.json"
        source_decisions_path.write_text(
            json.dumps(source_decisions, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        source_decisions = json.loads(source_decisions_path.read_text(encoding="utf-8"))
        format_path = self.root / "hard-case.format.json"
        compile_format_package(
            source_report,
            source_decisions,
            base_format_ref="preset:academic-basic@1.0.0",
            format_id="hard-case-reuse",
            format_version="1.0.0",
            output_path=format_path,
        )

        target_report = analyze_format_sample(target)
        target_decisions = build_decisions_example(target_report)
        target_decisions["missing_roles"] = {
            role: "inherit" for role in target_decisions.get("missing_roles", {})
        }
        target_decisions["node_roles"] = dict(target_report["node_roles"])
        target_decisions["node_roles"][conflict_path] = "heading.3"
        target_decisions["review_audit"] = {
            "audit_schema_version": 1,
            "report_id": target_report["report_id"],
            "source_sha256": target_report["source_sha256"],
            "origin": "test_fixture",
            "automation": True,
            "reviewer_attested": False,
            "operation_count": 1,
            "operations": [{
                "sequence": 1,
                "operation": "set_node_role",
                "target": conflict_path,
                "from": "pending_review",
                "to": "heading.3",
            }],
        }
        target_decisions_path = self.root / "hard-target.decisions.final.json"
        target_decisions_path.write_text(
            json.dumps(target_decisions, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        target_decisions = json.loads(target_decisions_path.read_text(encoding="utf-8"))
        mapping_path = self.root / "hard-target.roles.json"
        mapping = compile_role_mapping(target_report, target_decisions, output_path=mapping_path)
        self.assertEqual(
            next(item for item in mapping["assignments"] if item["node_ref"]["element_path"] == conflict_path)["role"],
            "heading.3",
        )

        manifest = self.root / "hard-target-manifest.json"
        manifest.write_text(
            json.dumps(
                {
                    "schema_version": 3,
                    "project_name": "hard-case-target",
                    "source": {"strategy": "docx_document", "file": target.name},
                    "format": {"ref": format_path.name},
                    "formatting": {
                        "mode": "restyle",
                        "role_map": mapping_path.name,
                        "on_unmapped": "error",
                    },
                    "output": {
                        "documents": [{"id": "main", "filename": "hard-target.docx", "parts": ["body"]}]
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        result = subprocess.run(
            [
                sys.executable,
                "synthesize.py",
                "--source",
                str(target),
                "--manifest",
                str(manifest),
                "--plan",
            ],
            cwd=Path(__file__).resolve().parents[1],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("hard-case-target", result.stdout)


if __name__ == "__main__":
    unittest.main()
