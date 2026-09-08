"""R2 contract, identity, registry, and provenance coverage."""

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from lib.config import ConfigError, ProjectConfig, merge_dict_with_provenance
from lib.contracts import (
    ContractError,
    SEMANTIC_ROLES,
    validate_analysis_data,
    validate_contract,
    validate_role_map_data,
)
from lib.format_analysis import analyze_format_sample
from lib.format_review import compile_format_package, compile_role_mapping
from lib.role_mapper import RoleMapper


class R2ContractTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.docx_path = self.root / "sample.docx"
        doc = Document()
        doc.add_heading("R2 heading", level=1)
        doc.add_paragraph("R2 body")
        doc.save(self.docx_path)
        self.analysis = analyze_format_sample(self.docx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_analysis_decisions_and_role_map_are_versioned_and_schema_valid(self):
        self.assertEqual(self.analysis["analysis_schema_version"], 1)
        validate_analysis_data(self.analysis)

        decisions = {
            "decisions_schema_version": 1,
            "report_id": self.analysis["report_id"],
            "source_sha256": self.analysis["source_sha256"],
            "role_styles": {
                role: candidate["cluster_id"]
                for role, candidate in self.analysis["candidate_roles"].items()
            },
            "style_overrides": {},
            "missing_roles": {item["role"]: "inherit" for item in self.analysis["missing_roles"]},
        }
        package = compile_format_package(self.analysis, decisions)
        self.assertEqual(package["format_schema_version"], 1)
        role_map = compile_role_mapping(self.analysis, decisions)
        validate_role_map_data(role_map)
        self.assertTrue(role_map["assignments"])
        assignment = role_map["assignments"][0]
        self.assertEqual(
            set(assignment["node_ref"]),
            {"source_sha256", "part_uri", "element_path", "text_hash"},
        )
        self.assertIn("provenance", assignment)
        self.assertIn("confirmed", assignment)

    def test_invalid_candidate_and_stale_node_are_rejected_with_pointer(self):
        bad_decisions = {
            "decisions_schema_version": 1,
            "report_id": self.analysis["report_id"],
            "source_sha256": self.analysis["source_sha256"],
            "role_styles": {"body": "missing-candidate"},
            "style_overrides": {},
            "missing_roles": {item["role"]: "inherit" for item in self.analysis["missing_roles"]},
        }
        with self.assertRaisesRegex(ConfigError, r"/role_styles/body"):
            compile_format_package(self.analysis, bad_decisions)

        stale_decisions = {
            "report_id": self.analysis["report_id"],
            "source_sha256": self.analysis["source_sha256"],
            "node_roles": {"/w:document/w:body/w:p[999]": "body"},
        }
        with self.assertRaises(ConfigError):
            compile_role_mapping(self.analysis, stale_decisions)

    def test_role_registry_distinguishes_table_structure_from_table_role(self):
        self.assertIn("table.body", SEMANTIC_ROLES)
        self.assertNotIn("table", SEMANTIC_ROLES)
        mapper = RoleMapper()
        assignments, _ = mapper.map_document(self.docx_path)
        self.assertNotIn("table", {assignment.role for assignment in assignments})

    def test_config_merge_preserves_false_zero_empty_list_and_leaf_provenance(self):
        merged, provenance = merge_dict_with_provenance(
            {"flags": {"enabled": True, "limit": 5}, "items": ["old"]},
            {"flags": {"enabled": False, "limit": 0}, "items": []},
            {"/flags/enabled": "manifest", "/flags/limit": "manifest", "/items": "manifest"},
            "/override.json",
        )
        self.assertFalse(merged["flags"]["enabled"])
        self.assertEqual(merged["flags"]["limit"], 0)
        self.assertEqual(merged["items"], [])
        self.assertEqual(provenance["/flags/enabled"], "/override.json")
        self.assertEqual(provenance["/flags/limit"], "/override.json")
        self.assertEqual(provenance["/items"], "/override.json")

    def test_relative_references_use_declaring_manifest_or_override_file(self):
        config_dir = self.root / "config"
        override_dir = config_dir / "overrides"
        formats_dir = self.root / "formats"
        maps_dir = config_dir / "maps"
        templates_dir = config_dir / "templates"
        for directory in (override_dir, formats_dir, maps_dir, templates_dir):
            directory.mkdir(parents=True)

        format_file = formats_dir / "custom.json"
        format_file.write_text(json.dumps({
            "format_schema_version": 1,
            "id": "r2-custom",
            "version": "1.0.0",
            "page": {"width_mm": 210, "height_mm": 297},
        }), encoding="utf-8")
        (maps_dir / "map.json").write_text("{}", encoding="utf-8")
        manifest = config_dir / "manifest.json"
        manifest.write_text(json.dumps({
            "schema_version": 3,
            "project_name": "R2 Project",
            "format": {"ref": "../formats/custom.json"},
            "formatting": {"mode": "restyle"},
            "source": {"strategy": "docx_document", "file": "../sample.docx"},
            "cover": {"template": "templates/cover.docx"},
        }), encoding="utf-8")
        override = override_dir / "paths.json"
        override.write_text(json.dumps({"formatting": {"role_map": "../maps/map.json"}}), encoding="utf-8")

        config = ProjectConfig(
            json.loads(manifest.read_text(encoding="utf-8")),
            project_dir=self.root,
            manifest_path=manifest,
        )
        # Apply the same provenance-aware merge used by load_project_config.
        from lib.config import load_project_config
        config = load_project_config(self.root, manifest_path=manifest, override_paths=[override])
        self.assertEqual(config.source["file"], str(self.docx_path.resolve()))
        self.assertEqual(config.formatting.role_map, str((maps_dir / "map.json").resolve()))
        self.assertEqual(config.cover["template"], str((templates_dir / "cover.docx").resolve()))
        self.assertEqual(config.provenance["/formatting/role_map"], str(override.resolve()))


if __name__ == "__main__":
    unittest.main()
