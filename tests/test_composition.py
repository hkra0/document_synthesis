"""Relationship-aware delivery assembly is testable without Office."""

from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image
from lxml import etree

from lib.composition import BOUNDARIES, add_bookmark, assemble_document
from lib.config import ProjectConfig
from lib.pagination import validate_document_structure


def bitmap(color):
    stream = BytesIO()
    Image.new("RGB", (20, 20), color).save(stream, format="PNG")
    stream.seek(0)
    return stream


def xml_text(element):
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


class CompositionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.body = self.root / "body.docx"
        self.nodes = [{"title": "Body heading", "level": 1, "bookmark_name": "HeadingOne"}]
        body = Document()
        add_bookmark(body.add_paragraph("Body heading")._p, "HeadingOne", 5)
        body.add_paragraph("Body text")
        body.save(self.body)

    def config(self, parts=None, cover=None, docs=None):
        raw = {"schema_version": 2, "project_name": "Fictional Project", "cover": {"template": False, "main_title": "Cover title", "author": "Fictional author", "date": "2026-09", **(cover or {})}}
        if docs is not None:
            raw["output"] = {"documents": docs}
        elif parts:
            raw["output"] = {"documents": [{"id": "main", "filename": "complete.docx", "parts": parts}]}
        return ProjectConfig(raw, self.root)

    def assemble(self, config, index=0):
        spec = config.documents[index]
        path = self.root / spec["filename"]
        assemble_document(config, spec, self.body, self.nodes, {"HeadingOne": 3}, path)
        return Document(path), path

    def test_complete_body_cover_and_body_toc_parts(self):
        original = self.body.read_bytes()
        for parts in (["cover", "toc", "body"], ["body"], ["cover"], ["toc", "body"], ["cover", "body"]):
            with self.subTest(parts=parts):
                doc, path = self.assemble(self.config(parts))
                text = xml_text(doc.element.body)
                self.assertEqual("Cover title" in text, "cover" in parts)
                self.assertEqual("目  录" in text, "toc" in parts)
                self.assertEqual("Body text" in text, "body" in parts)
                names = {b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))}
                for part, bookmark in BOUNDARIES.items():
                    self.assertEqual(bookmark in names, part in parts)
                validate_document_structure(path)
        self.assertEqual(self.body.read_bytes(), original)

    def test_standalone_toc_uses_reference_pages_without_local_links(self):
        config = self.config(docs=[
            {"id": "front", "filename": "front.docx", "parts": ["toc"], "toc": {"reference": "main"}},
            {"id": "main", "filename": "main.docx", "parts": ["body"]},
        ])
        doc, path = self.assemble(config)
        self.assertIn("Body heading", xml_text(doc.element.body))
        self.assertIn("3", xml_text(doc.element.body))
        self.assertNotIn("Body text", xml_text(doc.element.body))
        self.assertEqual(list(doc.element.body.iter(qn("w:hyperlink"))), [])
        validate_document_structure(path)

    def test_cover_author_date_and_xml_sensitive_toc_titles(self):
        self.nodes[0]["title"] = 'R&D <report> "quote"'
        doc, path = self.assemble(self.config())
        text = xml_text(doc.element.body)
        self.assertIn('R&D <report> "quote"', text)
        self.assertIn("Fictional author", text)
        self.assertIn("2026-09", text)
        validate_document_structure(path)

    def test_body_landscape_image_styles_numbering_and_header_retained(self):
        body = Document(self.body)
        custom = body.styles.add_style("Body custom", WD_STYLE_TYPE.PARAGRAPH)
        custom.font.size = Pt(13)
        body.add_paragraph("Original custom paragraph", style=custom)
        body.add_paragraph("Original numbered item", style="List Number")
        body.add_picture(bitmap("blue"), width=Inches(0.25))
        body.sections[0].header.paragraphs[0].text = "Original body header"
        landscape = body.add_section(WD_SECTION_START.NEW_PAGE)
        landscape.orientation = WD_ORIENT.LANDSCAPE
        landscape.page_width, landscape.page_height = Inches(11.7), Inches(8.3)
        body.add_paragraph("Landscape body")
        style_before = custom.element.xml
        numbering_before = [etree.tostring(child) for child in body.part.numbering_part.element]
        body.save(self.body)
        original = self.body.read_bytes()
        result, path = self.assemble(self.config())
        self.assertEqual(result.styles["Body custom"].element.xml, style_before)
        self.assertEqual([etree.tostring(child) for child in result.part.numbering_part.element][:len(numbering_before)], numbering_before)
        self.assertEqual(result.sections[-1].orientation, WD_ORIENT.LANDSCAPE)
        self.assertEqual(result.sections[-1].page_width, Inches(11.7))
        self.assertIn("Original body header", result.sections[-2].header.paragraphs[0].text)
        self.assertEqual(len(result.inline_shapes), 1)
        embed = next(result.element.body.iter(qn("a:blip"))).get(qn("r:embed"))
        self.assertTrue(result.part.rels[embed].target_part.blob.startswith(b"\x89PNG"))
        self.assertEqual(self.body.read_bytes(), original)
        validate_document_structure(path)

    def test_template_table_split_run_placeholders_image_header_and_hyperlink(self):
        template = Document()
        title = template.add_paragraph()
        title.add_run("{{MAIN_").bold = True
        title.add_run("TITLE}}")
        table = template.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "{{AUTHOR}} / {{DATE}}"
        template.add_picture(bitmap("red"), width=Inches(0.3))
        header = template.sections[0].header.paragraphs[0]
        header.text = "Static template header"
        header.add_run().add_picture(bitmap("green"), width=Inches(0.2))
        link = OxmlElement("w:hyperlink")
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        rid = template.part.relate_to("https://example.org/reference", RT.HYPERLINK, is_external=True)
        link.set(qn("r:id"), rid)
        template.add_paragraph()._p.append(link)
        template_path = self.root / "template.docx"
        template.save(template_path)
        doc, path = self.assemble(self.config(cover={"template": str(template_path)}))
        self.assertIn("Cover title", xml_text(doc.element.body))
        self.assertNotIn("{{", xml_text(doc.element.body))
        self.assertEqual(doc.tables[0].cell(0, 0).text, "Fictional author / 2026-09")
        self.assertIn("Static template header", doc.sections[0].header.paragraphs[0].text)
        header_part = doc.sections[0].header.part
        header_embed = next(header_part.element.iter(qn("a:blip"))).get(qn("r:embed"))
        self.assertTrue(header_part.rels[header_embed].target_part.blob.startswith(b"\x89PNG"))
        body_embed = next(doc.element.body.iter(qn("a:blip"))).get(qn("r:embed"))
        self.assertTrue(doc.part.rels[body_embed].target_part.blob.startswith(b"\x89PNG"))
        links = [item for item in doc.element.body.iter(qn("w:hyperlink")) if item.get(qn("r:id"))]
        self.assertEqual(doc.part.rels[links[0].get(qn("r:id"))].target_ref, "https://example.org/reference")
        validate_document_structure(path)

    def test_duplicate_bookmark_names_fail_closed(self):
        body = Document(self.body)
        add_bookmark(body.add_paragraph("Duplicate heading")._p, "HeadingOne", 7)
        body.save(self.body)
        config = self.config()
        path = self.root / config.documents[0]["filename"]
        with self.assertRaisesRegex(ValueError, "书签"):
            self.assemble(config)
        self.assertFalse(path.exists())

    def test_front_style_ids_do_not_collide_with_original_body_styles(self):
        body = Document(self.body)
        custom = body.styles.add_style("SynthFront_2", WD_STYLE_TYPE.PARAGRAPH)
        body.add_paragraph("Body with preserved style", style=custom)
        before = custom.element.xml
        body.save(self.body)
        result, _ = self.assemble(self.config())
        ids = [style.get(qn("w:styleId")) for style in result.part.styles.element if style.tag == qn("w:style")]
        self.assertEqual(len(ids), len(set(ids)), "Front style IDs must not shadow original body styles")
        self.assertEqual(result.styles["SynthFront_2"].element.xml, before)

    def test_template_implicit_normal_style_is_retained(self):
        template = Document()
        template.styles["Normal"].font.size = Pt(30)
        template.add_paragraph("{{MAIN_TITLE}}")
        template_path = self.root / "normal-template.docx"
        template.save(template_path)
        body = Document(self.body)
        body.styles["Normal"].font.size = Pt(11)
        body.save(self.body)
        result, _ = self.assemble(self.config(cover={"template": str(template_path)}))
        cover_paragraph = next(p for p in result.paragraphs if p.text == "Cover title")
        body_paragraph = next(p for p in result.paragraphs if p.text == "Body text")
        self.assertEqual(cover_paragraph.style.font.size, Pt(30))
        self.assertEqual(body_paragraph.style.font.size, Pt(11))


if __name__ == "__main__":
    unittest.main()
