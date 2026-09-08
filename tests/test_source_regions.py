# -*- coding: utf-8 -*-
"""
源选区与边界切片单元测试 (tests/test_source_regions.py)
"""

from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from lib.document_parts import (
    DocumentPart,
    InvalidBoundaryError,
    PartKind,
    SelectionOverlapError,
    SelectionValidator,
    ProtectedBoundaryError,
    SourceRegion,
    UnsupportedObjectError,
    UnassignedContentError,
    get_default_parts,
    get_part_boundary_bookmark,
)


class SourceRegionsTest(unittest.TestCase):
    def test_default_parts_and_boundary_bookmarks(self):
        parts = get_default_parts()
        self.assertIn("cover", parts)
        self.assertIn("toc", parts)
        self.assertIn("body", parts)
        self.assertEqual(parts["cover"].kind, PartKind.COVER.value)
        self.assertEqual(parts["toc"].kind, PartKind.GENERATED_TOC.value)
        self.assertEqual(parts["body"].kind, PartKind.CONTENT.value)

        # 历史兼容名称
        self.assertEqual(get_part_boundary_bookmark("cover"), "_Synth_cover")
        self.assertEqual(get_part_boundary_bookmark("toc"), "_Synth_toc")
        self.assertEqual(get_part_boundary_bookmark("body"), "_Synth_body")

        # 具名规范名称
        self.assertEqual(get_part_boundary_bookmark("abstract_cn", is_start=True), "_Synth_part_abstract_cn_start")
        self.assertEqual(get_part_boundary_bookmark("abstract_cn", is_start=False), "_Synth_part_abstract_cn_end")

    def test_selection_validator_happy_path(self):
        doc = Document()
        doc.add_paragraph("论文总标题")  # p1
        doc.add_paragraph("中文摘要第一段")  # p2
        doc.add_paragraph("中文摘要第二段")  # p3
        doc.add_paragraph("第一章 引言正文")  # p4
        doc.add_paragraph("1.1 研究背景")  # p5

        regions = {
            "title": SourceRegion(id="title", start="/w:document/w:body/w:p[1]", end="/w:document/w:body/w:p[2]"),
            "abstract": SourceRegion(id="abstract", start="/w:document/w:body/w:p[2]", end="/w:document/w:body/w:p[4]"),
            "main": SourceRegion(id="main", start="/w:document/w:body/w:p[4]", end="end_of_document"),
        }

        spans = SelectionValidator.validate_regions(doc, regions)
        self.assertEqual(spans["title"], (0, 1))
        self.assertEqual(spans["abstract"], (1, 3))
        self.assertEqual(spans["main"], (3, 5))

    def test_selection_validator_detects_overlap(self):
        doc = Document()
        doc.add_paragraph("段落 1")
        doc.add_paragraph("段落 2")
        doc.add_paragraph("段落 3")
        doc.add_paragraph("段落 4")

        regions = {
            "reg1": SourceRegion(id="reg1", start="/w:document/w:body/w:p[1]", end="/w:document/w:body/w:p[3]"),
            "reg2": SourceRegion(id="reg2", start="/w:document/w:body/w:p[2]", end="end_of_document"),
        }

        with self.assertRaises(SelectionOverlapError):
            SelectionValidator.validate_regions(doc, regions)

    def test_selection_validator_detects_unassigned_content(self):
        doc = Document()
        doc.add_paragraph("段落 1：摘要")
        doc.add_paragraph("段落 2：被遗漏的重要正文")
        doc.add_paragraph("段落 3：结论")

        regions = {
            "reg1": SourceRegion(id="reg1", start="/w:document/w:body/w:p[1]", end="/w:document/w:body/w:p[2]"),
            "reg3": SourceRegion(id="reg3", start="/w:document/w:body/w:p[3]", end="end_of_document"),
        }

        with self.assertRaises(UnassignedContentError) as ctx:
            SelectionValidator.validate_regions(doc, regions)
        self.assertIn("被遗漏的重要正文", str(ctx.exception))

    def test_selection_validator_accepts_excluded_unassigned_content(self):
        doc = Document()
        doc.add_paragraph("段落 1：正文")
        doc.add_paragraph("段落 2：草稿待删除内容")

        regions = {
            "main": SourceRegion(id="main", start="/w:document/w:body/w:p[1]", end="/w:document/w:body/w:p[2]"),
            "draft": SourceRegion(id="draft", start="/w:document/w:body/w:p[2]", end="end_of_document", exclude=True),
        }

        spans = SelectionValidator.validate_regions(doc, regions)
        self.assertIn("main", spans)
        self.assertIn("draft", spans)

    def test_selection_validator_rejects_table_interior_boundary(self):
        doc = Document()
        doc.add_paragraph("段落 1")
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "表格内部"

        regions = {
            "invalid": SourceRegion(
                id="invalid",
                start="/w:document/w:body/w:tbl[1]/w:tr[1]/w:tc[1]/w:p[1]",
                end="end_of_document",
            )
        }

        with self.assertRaises(InvalidBoundaryError) as ctx:
            SelectionValidator.validate_regions(doc, regions)
        self.assertIn("表格内部", str(ctx.exception))

    def test_slice_document_by_region(self):
        doc = Document()
        doc.add_paragraph("第一段")
        doc.add_paragraph("第二段：摘要")
        doc.add_paragraph("第三段：正文")

        sliced = SelectionValidator.slice_document_by_region(doc, (1, 2))
        self.assertEqual(len(sliced.paragraphs), 1)
        self.assertEqual(sliced.paragraphs[0].text, "第二段：摘要")

    def test_slice_returns_import_result_and_keeps_sectpr_last(self):
        doc = Document()
        doc.add_paragraph("第一段")
        doc.add_paragraph("第二段")

        sliced, result = SelectionValidator.slice_document_by_region_with_result(doc, (0, 1))

        self.assertEqual(sliced.paragraphs[0].text, "第一段")
        self.assertEqual(
            result.node_map["/w:document/w:body/w:p[1]"],
            "/w:document/w:body/w:p[1]",
        )
        self.assertEqual(sliced.element.body[-1].tag, qn("w:sectPr"))

    def test_cross_bookmark_boundary_is_rejected(self):
        doc = Document()
        first = doc.add_paragraph("书签开始")
        first._p.insert(0, parse_xml(
            f'<w:bookmarkStart {nsdecls("w")} w:id="17" w:name="cross"/>'
        ))
        second = doc.add_paragraph("书签结束")
        second._p.append(parse_xml(
            f'<w:bookmarkEnd {nsdecls("w")} w:id="17"/>'
        ))

        with self.assertRaises(ProtectedBoundaryError):
            SelectionValidator.validate_regions(doc, {
                "partial": SourceRegion(
                    id="partial",
                    start="/w:document/w:body/w:p[1]",
                    end="/w:document/w:body/w:p[2]",
                )
            })

    def test_cross_field_boundary_is_rejected(self):
        doc = Document()
        first = doc.add_paragraph("字段开始")
        first._p.append(parse_xml(
            f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'
        ))
        second = doc.add_paragraph("字段结束")
        second._p.append(parse_xml(
            f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'
        ))

        with self.assertRaises(ProtectedBoundaryError):
            SelectionValidator.validate_regions(doc, {
                "partial": SourceRegion(
                    id="partial",
                    start="/w:document/w:body/w:p[1]",
                    end="/w:document/w:body/w:p[2]",
                )
            })

    def test_unsupported_top_level_object_is_rejected(self):
        doc = Document()
        doc.add_paragraph("正文")
        doc.element.body.append(parse_xml(
            f'<w:altChunk {nsdecls("w", "r")} r:id="rIdUnsupported"/>'
        ))

        with self.assertRaises(UnsupportedObjectError):
            SelectionValidator.validate_regions(doc, {
                "body": SourceRegion(
                    id="body",
                    start="/w:document/w:body/w:p[1]",
                    end="end_of_document",
                )
            })


if __name__ == "__main__":
    unittest.main()
