# -*- coding: utf-8 -*-
"""P2 Office 后端抽象契约与隔离断言测试。

测试覆盖：
1. 假后端成功驱动完整构建全路径
2. 后端调用超时故障与回滚
3. 书签缺失检测与阻断
4. 导出 PDF 缺失检测与阻断
5. 非 exact 保真度 (draft) 一律拒绝，草稿交付在标注落地 (P5) 前不可发布
6. --doctor 退出码约定 (0 / 2 / 1)
7. 生产代码 (lib/) 中 osascript 与 darwin 关键字隔离断言
"""

import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from typing import Dict, List, Optional, Sequence
from unittest.mock import patch

from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.engine import BuildError, UnifiedSynthesizer
from lib.office import (
    BackendStatus,
    BookmarkReport,
    NullBackend,
    OfficeBackend,
    OfficeBackendError,
    OfficeTimeoutError,
    OfficeUnsupportedError,
    get_backend,
    set_backend,
)
from lib.pagination import OfficeExportError, inspect_document
from synthesize import run_doctor


class FakeOfficeBackend:
    """可编程模拟各种 Office 成功与故障行为的假后端。"""
    name: str = "fake-word"
    fidelity: str = "exact"

    def __init__(
        self,
        *,
        fidelity: str = "exact",
        should_timeout: bool = False,
        missing_bookmarks: bool = False,
        omit_pdf: bool = False,
    ) -> None:
        self.fidelity = fidelity
        self.should_timeout = should_timeout
        self.missing_bookmarks = missing_bookmarks
        self.omit_pdf = omit_pdf

    def static_status(self) -> BackendStatus:
        return BackendStatus(available=True, reason="Fake backend available", fidelity=self.fidelity)

    def probe(self) -> BackendStatus:
        return BackendStatus(available=True, reason="Fake backend probe ok", fidelity=self.fidelity)

    def convert_doc_to_docx(self, src: Path, dst: Path) -> None:
        from docx import Document
        doc = Document()
        doc.add_paragraph("Converted from doc")
        doc.save(dst)

    def convert_pptx_to_pdf(self, src: Path, dst: Path) -> None:
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 72), "PPTX 转换内容\n第二行内容\n第三行内容", fontsize=12)
        doc.save(str(dst))
        doc.close()

    def export_pdf(self, docx: Path, pdf: Path) -> None:
        if self.should_timeout:
            raise OfficeTimeoutError("Microsoft Word 导出 PDF 超时。")
        if not self.omit_pdf:
            import pymupdf
            doc = pymupdf.open()
            page = doc.new_page(width=595, height=842)
            page.insert_text((72, 72), "Word 导出内容\n第二行内容\n第三行内容", fontsize=12)
            doc.save(str(pdf))
            doc.close()

    def inspect_bookmarks(
        self, docx: Path, pdf: Path, names: Sequence[str]
    ) -> BookmarkReport:
        if self.should_timeout:
            raise OfficeTimeoutError("等待 Word 最终分页核验超时。")
        if not self.omit_pdf:
            import pymupdf
            doc = pymupdf.open()
            for p in range(4):
                page = doc.new_page(width=595, height=842)
                page.insert_text((72, 250), f"正文段落内容第 {p + 1} 页\n第二行内容\n第三行内容", fontsize=12)
            doc.save(str(pdf))
            doc.close()

        if self.missing_bookmarks:
            bookmarks: Dict[str, Dict[str, int]] = {}
        else:
            bookmarks = {}
            for name in names:
                if name == "_Synth_cover":
                    bookmarks[name] = {"physical_page": 1, "printed_page": 1}
                elif name == "_Synth_toc":
                    bookmarks[name] = {"physical_page": 2, "printed_page": 1}
                elif name == "_Synth_body":
                    bookmarks[name] = {"physical_page": 3, "printed_page": 1}
                else:
                    bookmarks[name] = {"physical_page": 3, "printed_page": 1}

        return BookmarkReport(
            bookmarks=bookmarks,
            total_pages=4,
            pdf_path=pdf,
            fidelity=self.fidelity,
        )


class OfficeBackendContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.original_backend_env = os.environ.get("DOCUMENT_SYNTHESIS_OFFICE_BACKEND")

    def tearDown(self) -> None:
        set_backend(None)
        if self.original_backend_env is not None:
            os.environ["DOCUMENT_SYNTHESIS_OFFICE_BACKEND"] = self.original_backend_env
        else:
            os.environ.pop("DOCUMENT_SYNTHESIS_OFFICE_BACKEND", None)

    def test_fake_backend_happy_path_pipeline(self) -> None:
        """测试假后端驱动完整合成管道成功发布文档与元数据。"""
        fake = FakeOfficeBackend(fidelity="exact")
        set_backend(fake)

        with tempfile.TemporaryDirectory() as out_dir:
            source_dir = ROOT / "examples" / "minimal-demo"
            output_dir = Path(out_dir) / "output"
            UnifiedSynthesizer.synthesize(source_dir, output_dir=output_dir)

            published_docx = output_dir / "demo_complete.docx"
            self.assertTrue(published_docx.is_file(), f"交付 DOCX 未生成: {list(output_dir.iterdir())}")

    def test_backend_timeout_aborts_cleanly(self) -> None:
        """测试后端调用超时抛出清晰异常，且不污染正式交付目录。"""
        fake = FakeOfficeBackend(should_timeout=True)
        set_backend(fake)

        with tempfile.TemporaryDirectory() as out_dir:
            source_dir = ROOT / "examples" / "minimal-demo"
            output_dir = Path(out_dir) / "output"
            with self.assertRaises((OfficeExportError, BuildError)) as ctx:
                UnifiedSynthesizer.synthesize(source_dir, output_dir=output_dir)
            self.assertIn("超时", str(ctx.exception))
            # 确认未生成假交付物
            self.assertFalse((output_dir / "demo_complete.docx").exists())

    def test_missing_bookmarks_rejected(self) -> None:
        """测试后端返回书签缺失时，门禁检查拒绝发布。"""
        fake = FakeOfficeBackend(missing_bookmarks=True)
        set_backend(fake)

        with tempfile.TemporaryDirectory() as out_dir:
            source_dir = ROOT / "examples" / "minimal-demo"
            output_dir = Path(out_dir) / "output"
            with self.assertRaises((OfficeExportError, BuildError)) as ctx:
                UnifiedSynthesizer.synthesize(source_dir, output_dir=output_dir)
            self.assertTrue("书签" in str(ctx.exception) or "缺失" in str(ctx.exception))

    def test_missing_exported_pdf_rejected(self) -> None:
        """测试后端未生成导出 PDF 时，构建流程及时停止。"""
        fake = FakeOfficeBackend(omit_pdf=True)
        set_backend(fake)

        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "input.docx"
            pdf_path = Path(tmp) / "output.pdf"
            from docx import Document
            doc = Document()
            doc.add_paragraph("Test")
            doc.save(docx_path)

            with self.assertRaises(OfficeExportError):
                inspect_document(docx_path, pdf_path, ["body"])

    def test_draft_fidelity_always_rejected(self) -> None:
        """草稿标注 (P5) 落地前，draft 页码不能进入发布，也不存在绕过开关。"""
        set_backend(FakeOfficeBackend(fidelity="draft"))

        with tempfile.TemporaryDirectory() as out_dir:
            source_dir = ROOT / "examples" / "minimal-demo"
            output_dir = Path(out_dir) / "output"
            with self.assertRaises((OfficeExportError, BuildError)) as ctx:
                UnifiedSynthesizer.synthesize(source_dir, output_dir=output_dir)
            self.assertIn("draft", str(ctx.exception))
            self.assertFalse((output_dir / "demo_complete.docx").exists())

    def test_cli_backend_option_precedence(self) -> None:
        """未传 --office-backend 时，环境变量必须生效；显式参数覆盖环境变量。"""
        import subprocess
        base = [sys.executable, "synthesize.py", "--doctor"]
        env = dict(os.environ, DOCUMENT_SYNTHESIS_OFFICE_BACKEND="none")
        kwargs = dict(cwd=ROOT, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)

        from_env = subprocess.run(base, env=env, **kwargs)
        self.assertEqual(from_env.returncode, 2, from_env.stdout + from_env.stderr)
        self.assertIn("Office Automation (none)", from_env.stdout)

        env["DOCUMENT_SYNTHESIS_OFFICE_BACKEND"] = "bogus"
        explicit = subprocess.run(base + ["--office-backend", "none"], env=env, **kwargs)
        self.assertEqual(explicit.returncode, 2, explicit.stdout + explicit.stderr)

        invalid = subprocess.run(base, env=env, **kwargs)
        self.assertEqual(invalid.returncode, 1)
        self.assertIn("Office 后端配置无效", invalid.stdout)

    def test_unimplemented_backends_are_reported_not_substituted(self) -> None:
        """尚未实现的 LibreOffice 后端必须明确报错，不能静默退回其他后端。"""
        with self.assertRaises(OfficeUnsupportedError):
            get_backend("libreoffice")
        self.assertNotIn("libreoffice", NullBackend().static_status().reason)
        self.assertNotIn("--allow-draft", NullBackend().static_status().reason)

    def test_doctor_exit_codes(self) -> None:
        """测试 --doctor 退出码规范: 0 (exact 就绪), 2 (草稿/分析就绪), 1 (依赖缺失或参数错误)。"""
        # Case 0: exact 就绪
        fake_exact = FakeOfficeBackend(fidelity="exact")
        set_backend(fake_exact)
        with patch("sys.stdout"):
            code_0 = run_doctor(source_dir=ROOT / "examples" / "minimal-demo")
        self.assertEqual(code_0, 0)

        # Case 2: 无 exact Office (NullBackend)
        set_backend(NullBackend())
        with patch("sys.stdout"):
            code_2 = run_doctor(source_dir=ROOT / "examples" / "minimal-demo")
        self.assertEqual(code_2, 2)

        # Case 2b: draft 后端 (只能生成草稿或分析)
        fake_draft = FakeOfficeBackend(fidelity="draft")
        set_backend(fake_draft)
        with patch("sys.stdout"):
            code_2b = run_doctor(source_dir=ROOT / "examples" / "minimal-demo")
        self.assertEqual(code_2b, 2)

        # Case 1: 缺少依赖或无效输入路径
        with patch("sys.stdout"):
            code_1 = run_doctor(source_dir=Path("/non_existent_path_xyz"))
        self.assertEqual(code_1, 1)

    def test_production_code_isolation_assertions(self) -> None:
        """严格断言生产代码 (lib/) 中 osascript 和 darwin 的封装隔离边界。"""
        lib_dir = ROOT / "lib"
        py_files = sorted(lib_dir.rglob("*.py"))

        allowed_osascript = {lib_dir / "office" / "mac_applescript.py"}
        allowed_darwin = {
            lib_dir / "office" / "mac_applescript.py",
            lib_dir / "office" / "__init__.py",
        }

        osascript_violations = []
        darwin_violations = []

        for py_path in py_files:
            content = py_path.read_text(encoding="utf-8")
            if "osascript" in content and py_path not in allowed_osascript:
                osascript_violations.append(str(py_path.relative_to(ROOT)))
            if "darwin" in content and py_path not in allowed_darwin:
                darwin_violations.append(str(py_path.relative_to(ROOT)))

        self.assertEqual(
            osascript_violations,
            [],
            f"以下生产代码文件中发现了未被隔离的 osascript 引用: {osascript_violations}",
        )
        self.assertEqual(
            darwin_violations,
            [],
            f"以下生产代码文件中发现了未被隔离的 darwin 平台检查: {darwin_violations}",
        )

        # Windows COM 调用同样只能出现在 Windows 后端内。
        allowed_com = {
            lib_dir / "office" / "win_com.py",
            lib_dir / "office" / "win_com_worker.py",
        }
        com_markers = ("win32com", "pythoncom", "DispatchEx", "WINWORD", "taskkill")
        com_violations = sorted(
            str(py_path.relative_to(ROOT))
            for py_path in py_files
            if py_path not in allowed_com
            and any(marker in py_path.read_text(encoding="utf-8") for marker in com_markers)
        )
        self.assertEqual(com_violations, [], f"以下生产代码文件中发现了未被隔离的 Windows COM 调用: {com_violations}")

        # 绝不能按映像名结束 Office：用户自己打开的 Word 会被一并关掉。
        for path in allowed_com:
            content = path.read_text(encoding="utf-8")
            self.assertNotIn("taskkill", content)
            self.assertNotIn("/IM", content)


if __name__ == "__main__":
    unittest.main()
