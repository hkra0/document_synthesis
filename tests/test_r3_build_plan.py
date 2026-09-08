# -*- coding: utf-8 -*-
"""R3：共享只读 PreparedBuild 与显式 NodeRef 映射验收。"""

from dataclasses import replace
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from lib.docx_inspector import NodeRef, inspect_docx
from lib.engine import BuildError, UnifiedSynthesizer
from lib.role_mapper import RoleAssignment, RoleMapper, serialize_role_map


class R3BuildPlanTest(unittest.TestCase):
    def _manifest(self, root: Path, source_name: str, role_map: str = None) -> Path:
        formatting = {"mode": "restyle", "on_unmapped": "error"}
        if role_map:
            formatting["role_map"] = role_map
        manifest = {
            "schema_version": 3,
            "project_name": "R3 Plan Fixture",
            "source": {"strategy": "docx_document", "file": source_name},
            "format": {"ref": "preset:academic-basic@1.0.0"},
            "formatting": formatting,
            "output": {"documents": [{"id": "main", "filename": "main.docx", "parts": ["body"]}]},
        }
        path = root / "manifest.json"
        path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        return path

    @staticmethod
    def _heading(document: Document, text: str, level: int = 1):
        paragraph = document.add_paragraph(text)
        paragraph._p.get_or_add_pPr().append(
            parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level - 1}"/>')
        )

    def test_prepare_is_shared_and_public_plan_omits_source_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            document = Document()
            self._heading(document, "Public heading")
            document.add_paragraph("CONFIDENTIAL BODY PROSE MUST NOT APPEAR IN PLAN")
            document.save(source)
            manifest = self._manifest(root, source.name)

            prepared = UnifiedSynthesizer.prepare_build(source, manifest)
            public = prepared.to_public_dict()
            plan = UnifiedSynthesizer.plan(source, manifest)

            self.assertEqual(prepared.source_hashes, plan["source_hashes"])
            self.assertEqual(prepared.format_hash, plan["format_hash"])
            self.assertTrue(plan["source_body_omitted"])
            self.assertNotIn("CONFIDENTIAL BODY PROSE MUST NOT APPEAR IN PLAN", json.dumps(public, default=str))
            self.assertEqual(len(prepared.inspections), 1)
            self.assertTrue(prepared.node_index)

    def test_explicit_role_map_controls_heading_style_and_unique_navigation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            document = Document()
            self._heading(document, "Repeated heading", 1)
            self._heading(document, "Repeated heading", 1)
            document.add_paragraph("Supporting body")
            document.save(source)

            inspection = inspect_docx(source)
            automatic = RoleMapper().map_inspection(inspection)
            assignments = [
                replace(
                    assignment,
                    role="heading.2" if index == 0 else assignment.role,
                    level=2 if index == 0 else assignment.level,
                    provenance="explicit" if index == 0 else assignment.provenance,
                    confirmed=index == 0,
                )
                for index, assignment in enumerate(automatic)
            ]
            role_map_path = root / "roles.json"
            role_map_path.write_text(
                json.dumps(
                    serialize_role_map(
                        assignments,
                        source_file=source.name,
                        on_unmapped="error",
                    ),
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            manifest = self._manifest(root, source.name, role_map_path.name)
            prepared = UnifiedSynthesizer.prepare_build(source, manifest)

            first = next(item for item in prepared.assignments if item.node_ref.element_path.endswith("p[1]"))
            second = next(item for item in prepared.assignments if item.node_ref.element_path.endswith("p[2]"))
            self.assertEqual(first.role, "heading.2")
            self.assertNotEqual(first.bookmark_name, second.bookmark_name)

            run_dir = root / "run"
            rendered = UnifiedSynthesizer.render_body(prepared, run_dir)
            output = Document(rendered.path)
            self.assertEqual(output.paragraphs[0].style.style_id, "SynthHeading2")
            names = {
                bookmark.get(qn("w:name"))
                for bookmark in output.paragraphs[0]._p.iter(qn("w:bookmarkStart"))
            }
            self.assertIn(first.bookmark_name, names)

    def test_stale_mapping_is_rejected_before_run_directory_and_render(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            document = Document()
            self._heading(document, "Stale heading")
            document.save(source)
            inspection = inspect_docx(source)
            block = inspection.blocks[0]
            stale = RoleAssignment(
                node_ref=NodeRef(
                    source_sha256="0" * 64,
                    part_uri=block.node.part_uri,
                    element_path=block.node.element_path,
                    text_hash=block.node.text_hash,
                ),
                role="heading.2",
                level=2,
                text_hash=block.node.text_hash,
            )
            role_map = root / "stale.json"
            role_map.write_text(
                json.dumps(serialize_role_map([stale], source_file=source.name, on_unmapped="preserve")),
                encoding="utf-8",
            )
            manifest = self._manifest(root, source.name, role_map.name)
            output = root / "output"

            with self.assertRaises(BuildError):
                UnifiedSynthesizer.synthesize(source, output, manifest)
            self.assertFalse(output.exists())

            prepared = UnifiedSynthesizer.prepare_build(source, self._manifest(root, source.name))
            source.write_bytes(source.read_bytes() + b"tampered")
            with self.assertRaises(BuildError):
                UnifiedSynthesizer.render_body(prepared, root / "tampered-run")
            self.assertFalse((root / "tampered-run").exists())


if __name__ == "__main__":
    unittest.main()
