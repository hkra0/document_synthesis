# -*- coding: utf-8 -*-
"""
DOCX 只读检查器与安全限额测试 (tests/test_docx_inspector.py)
"""

import unittest
import tempfile
import zipfile
from pathlib import Path

import docx
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from lib.config import ConfigError
from lib.docx_inspector import (
    inspect_docx,
    DocumentInspectionLimits,
    DocumentInspection,
    NodeRef,
)
from tests.fixtures.formatting.generate_samples import (
    create_complex_fields_document,
    create_mixed_fonts_document,
    create_sections_and_pagination_document,
)


class DocxInspectorTest(unittest.TestCase):

    def test_rejects_docm_and_doc_extensions(self):
        with tempfile.TemporaryDirectory() as td:
            p_docm = Path(td) / "malicious.docm"
            p_docm.touch()
            with self.assertRaises(ConfigError) as ctx:
                inspect_docx(p_docm)
            self.assertIn(".docm", str(ctx.exception))

            p_doc = Path(td) / "old.doc"
            p_doc.touch()
            with self.assertRaises(ConfigError) as ctx:
                inspect_docx(p_doc)
            self.assertIn(".doc", str(ctx.exception))

    def test_safety_limits_enforced(self):
        with tempfile.TemporaryDirectory() as td:
            tmpdir = Path(td)
            normal_doc = tmpdir / "sample.docx"
            create_mixed_fonts_document(normal_doc)

            # 1. 超过最大部件数量
            limits_parts = DocumentInspectionLimits(max_part_count=2)
            with self.assertRaises(ConfigError) as ctx:
                inspect_docx(normal_doc, limits=limits_parts)
            self.assertIn("文档部件数量", str(ctx.exception))

            # 2. 超过单部件上限
            limits_single = DocumentInspectionLimits(max_single_xml_bytes=100)
            with self.assertRaises(ConfigError) as ctx:
                inspect_docx(normal_doc, limits=limits_single)
            self.assertIn("单部件解压大小", str(ctx.exception))

    def test_identical_paragraphs_have_unique_noderefs(self):
        with tempfile.TemporaryDirectory() as td:
            doc_path = Path(td) / "duplicate_text.docx"
            doc = Document()
            doc.add_paragraph("重复正文内容。")
            doc.add_paragraph("重复正文内容。")
            doc.save(str(doc_path))

            inspection = inspect_docx(doc_path)
            self.assertEqual(len(inspection.blocks), 2)
            b1, b2 = inspection.blocks[0], inspection.blocks[1]

            # 相同可见文本
            self.assertEqual(b1.visible_text, b2.visible_text)
            # 但具有不同的 NodeRef 路径
            self.assertEqual(b1.node.element_path, "/w:document/w:body/w:p[1]")
            self.assertEqual(b2.node.element_path, "/w:document/w:body/w:p[2]")
            self.assertNotEqual(b1.node, b2.node)

    def test_table_nested_paragraphs_and_noderef_paths(self):
        with tempfile.TemporaryDirectory() as td:
            doc_path = Path(td) / "table_doc.docx"
            doc = Document()
            doc.add_paragraph("前置正文")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).paragraphs[0].text = "Cell 0-0"
            table.cell(1, 1).paragraphs[0].text = "Cell 1-1"
            doc.save(str(doc_path))

            inspection = inspect_docx(doc_path)
            # 前置段落 + 表格中 4 个单元格段落 = 5 个 blocks
            self.assertEqual(len(inspection.blocks), 5)
            # 检查表格单元格路径结构
            paths = [b.node.element_path for b in inspection.blocks]
            self.assertIn("/w:document/w:body/w:p[1]", paths)
            self.assertIn("/w:document/w:body/w:tbl[1]/w:tr[1]/w:tc[1]/w:p[1]", paths)
            self.assertIn("/w:document/w:body/w:tbl[1]/w:tr[2]/w:tc[2]/w:p[1]", paths)

    def test_protected_objects_detection(self):
        with tempfile.TemporaryDirectory() as td:
            # 1. 复杂字段
            p_fld = Path(td) / "fld.docx"
            create_complex_fields_document(p_fld)
            insp_fld = inspect_docx(p_fld)
            p2 = insp_fld.blocks[1]
            self.assertIn("field_complex", p2.protected_objects)

            # 2. 分节信息
            p_sec = Path(td) / "sec.docx"
            create_sections_and_pagination_document(p_sec)
            insp_sec = inspect_docx(p_sec)
            self.assertTrue(len(insp_sec.sections) >= 2)
            sec1 = insp_sec.sections[0]
            sec2 = insp_sec.sections[1]
            self.assertEqual(sec1.page_format, "upperRoman")
            self.assertEqual(sec2.section_type, "oddPage")
            self.assertEqual(sec2.page_start, 1)

    def test_field_cache_changes_do_not_change_noderef_identity(self):
        """A Word-updated PAGEREF cache is not author text identity."""
        with tempfile.TemporaryDirectory() as td:
            doc_path = Path(td) / "field-cache.docx"
            doc = Document()
            paragraph = doc.add_paragraph("前置")
            paragraph._p.append(parse_xml(
                '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'w:instr="PAGEREF ordinary_target"><w:r><w:t>1</w:t></w:r></w:fldSimple>'
            ))
            paragraph.add_run("后置")
            doc.save(str(doc_path))

            before = inspect_docx(doc_path).blocks[0]
            self.assertEqual(before.visible_text, "前置1后置")
            self.assertIn("field_complex", before.protected_objects)

            doc = Document(str(doc_path))
            field = next(doc.element.body.iter(qn("w:fldSimple")))
            next(field.iter(qn("w:t"))).text = "99"
            doc.save(str(doc_path))

            after = inspect_docx(doc_path).blocks[0]
            self.assertEqual(after.visible_text, "前置99后置")
            self.assertEqual(before.node.text_hash, after.node.text_hash)

    def test_story_bindings_include_linked_to_previous_sections(self):
        with tempfile.TemporaryDirectory() as td:
            doc_path = Path(td) / "linked-stories.docx"
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = "共享页眉"
            doc.add_paragraph("第一节正文")
            second = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
            self.assertTrue(second.header.is_linked_to_previous)
            doc.add_paragraph("第二节正文")
            doc.save(doc_path)

            inspection = inspect_docx(doc_path)
            header_bindings = [
                item for item in inspection.section_bindings
                if item.story_type == "header" and item.variant == "default"
            ]
            self.assertEqual(len(header_bindings), 2)
            self.assertFalse(header_bindings[0].linked_to_previous)
            self.assertTrue(header_bindings[1].linked_to_previous)
            self.assertEqual(header_bindings[0].part_uri, header_bindings[1].part_uri)
            self.assertTrue(inspection.sections[0].has_header)
            self.assertTrue(inspection.sections[1].has_header)
            self.assertEqual(len(inspection.stories), 1)
            self.assertEqual(inspection.stories[0].blocks[0].visible_text, "共享页眉")

    def test_story_table_is_kept_out_of_body_and_has_stable_identity(self):
        with tempfile.TemporaryDirectory() as td:
            doc_path = Path(td) / "story-table.docx"
            doc = Document()
            table = doc.sections[0].header.add_table(rows=1, cols=2, width=doc.sections[0].page_width)
            table.cell(0, 0).text = "页眉左单元格"
            table.cell(0, 1).text = "页眉右单元格"
            doc.add_paragraph("正文不应混入页眉表格")
            doc.save(doc_path)

            inspection = inspect_docx(doc_path)
            self.assertEqual([block.visible_text for block in inspection.blocks], ["正文不应混入页眉表格"])
            story = next(item for item in inspection.stories if item.story_type == "header")
            table_block = next(block for block in story.blocks if block.structure_type == "table")
            self.assertEqual(table_block.node.part_uri, story.part_uri)
            self.assertEqual(table_block.node.element_path, "/w:hdr/w:tbl[1]")
            self.assertIn("页眉左单元格", table_block.visible_text)
            self.assertIn("页眉右单元格", table_block.visible_text)


if __name__ == "__main__":
    unittest.main()
