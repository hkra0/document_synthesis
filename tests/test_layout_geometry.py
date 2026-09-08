# -*- coding: utf-8 -*-
"""
版心几何度量与渲染上下文单元测试 (tests/test_layout_geometry.py)
"""

import unittest
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.shared import Inches, Pt, Cm

from lib.layout import (
    ContentBox,
    compute_content_box,
    constrain_image_box,
    adapt_table_to_content_box,
    isolate_section_boundaries,
    RenderContext,
    DEFAULT_A4_WIDTH_TWIP,
    DEFAULT_A4_HEIGHT_TWIP,
)
from lib.format_schema import (
    PageSpec,
    LengthValue,
    ResolvedFormat,
)


class LayoutGeometryTest(unittest.TestCase):

    def test_content_box_properties_and_orientation(self):
        # 宽 10000 twip, 高 15000 twip, 边距各 1000 twip
        box = ContentBox(
            page_width_twip=10000,
            page_height_twip=15000,
            margin_top_twip=1000,
            margin_bottom_twip=1000,
            margin_left_twip=1000,
            margin_right_twip=1000,
        )
        self.assertEqual(box.content_width_twip, 8000)
        self.assertEqual(box.content_height_twip, 13000)
        self.assertAlmostEqual(box.content_width_pt, 400.0)
        self.assertAlmostEqual(box.content_height_pt, 650.0)
        self.assertEqual(box.orientation, "portrait")

        # 横版测试
        box_landscape = ContentBox(
            page_width_twip=15000,
            page_height_twip=10000,
            margin_top_twip=1000,
            margin_bottom_twip=1000,
            margin_left_twip=1000,
            margin_right_twip=1000,
        )
        self.assertEqual(box_landscape.orientation, "landscape")

    def test_compute_content_box_defaults_and_page_spec(self):
        # 默认 None 返回标准 A4
        default_box = compute_content_box(None)
        self.assertEqual(default_box.page_width_twip, DEFAULT_A4_WIDTH_TWIP)
        self.assertEqual(default_box.page_height_twip, DEFAULT_A4_HEIGHT_TWIP)
        self.assertEqual(default_box.orientation, "portrait")

        # 从 PageSpec 计算
        page_spec = PageSpec(
            width_mm=200.0,
            height_mm=300.0,
            margin_top_mm=20.0,
            margin_bottom_mm=20.0,
            margin_left_mm=30.0,
            margin_right_mm=30.0,
            orientation="portrait",
        )
        box = compute_content_box(page_spec)
        self.assertAlmostEqual(box.content_width_mm, 140.0, places=1)
        self.assertAlmostEqual(box.content_height_mm, 260.0, places=1)

    def test_compute_content_box_from_docx_section(self):
        doc = Document()
        sec = doc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)

        box = compute_content_box(sec)
        self.assertEqual(box.content_width_twip, int(6.5 * 1440))  # 6.5 inches * 1440 twips/inch
        self.assertEqual(box.content_height_twip, int(9.0 * 1440))
        self.assertEqual(box.orientation, "portrait")

    def test_constrain_image_box_aspect_ratio(self):
        box = ContentBox(
            page_width_twip=11906,
            page_height_twip=16838,
            margin_top_twip=2000,
            margin_bottom_twip=2000,
            margin_left_twip=2000,
            margin_right_twip=2000,
        )
        # content_width_cm ≈ 13.94 cm, content_height_cm ≈ 22.64 cm
        # 1. 超宽图片 (宽 2000px, 高 1000px -> 纵横比 0.5)
        w_cm, h_cm = constrain_image_box(2000, 1000, box)
        self.assertAlmostEqual(w_cm, box.content_width_cm, places=2)
        self.assertAlmostEqual(h_cm, w_cm * 0.5, places=2)

        # 2. 超长高图 (宽 500px, 高 2000px -> 纵横比 4.0)
        w_cm2, h_cm2 = constrain_image_box(500, 2000, box, max_height_ratio=0.8)
        max_h = box.content_height_cm * 0.8
        self.assertAlmostEqual(h_cm2, max_h, places=2)
        self.assertAlmostEqual(w_cm2, h_cm2 / 4.0, places=2)

    def test_adapt_table_to_content_box_scales_oversized_table(self):
        doc = Document()
        # 创建 2 列的表格
        tbl = doc.add_table(rows=2, cols=2)
        tbl.columns[0].width = Inches(5.0)  # 5 inches = 7200 twip
        tbl.columns[1].width = Inches(5.0)  # 5 inches = 7200 twip
        # 总宽 14400 twip

        # 版心可用宽仅 8000 twip
        box = ContentBox(
            page_width_twip=12000,
            page_height_twip=16000,
            margin_top_twip=2000,
            margin_bottom_twip=2000,
            margin_left_twip=2000,
            margin_right_twip=2000,
        )
        self.assertEqual(box.content_width_twip, 8000)

        # 缩放前
        modified = adapt_table_to_content_box(tbl, box)
        self.assertTrue(modified)

        # 验证缩放后表格网格列宽总和 == 8000 twip
        from docx.oxml.ns import qn
        gridCols = tbl._tbl.tblGrid.findall(qn("w:gridCol"))
        col_widths = [int(col.get(qn("w:w"))) for col in gridCols]
        self.assertEqual(sum(col_widths), 8000)

        # 已经适应版心的表格再次调用不重复修改
        modified_again = adapt_table_to_content_box(tbl, box)
        self.assertFalse(modified_again)

    def test_render_context_state_isolation_and_landscape_tracking(self):
        ctx1 = RenderContext()
        ctx2 = RenderContext()

        ctx1.track_landscape(True)
        self.assertTrue(ctx1.last_rendered_landscape)
        self.assertFalse(ctx2.last_rendered_landscape, "不同 RenderContext 实例必须完全独立")

        # 消费状态
        consumed = ctx1.consume_landscape()
        self.assertTrue(consumed)
        self.assertFalse(ctx1.last_rendered_landscape, "消费后横版标志必须重置")

    def test_isolate_section_boundaries_disconnects_headers(self):
        doc = Document()
        sec1 = doc.sections[0]
        sec1.header.paragraphs[0].text = "Header 1"
        sec2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
        
        # 初始情况下 sec2 的 header 可能是 linked
        isolate_section_boundaries(sec2)
        self.assertFalse(sec2.header.is_linked_to_previous)
        self.assertFalse(sec2.footer.is_linked_to_previous)


if __name__ == "__main__":
    unittest.main()
