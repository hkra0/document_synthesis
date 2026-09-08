# -*- coding: utf-8 -*-
"""NW01–NW12 production-CLI Word/PDF integration matrix.

This suite deliberately does not replace Word with a mock.  It is opt-in and
is marked BLOCKED when the host cannot control Word; the post-review status
file records that condition separately from the offline test counts.
"""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Mm, Pt

from lib.build_plan import prepare_project_build
from lib.config import load_project_config
from lib.content_integrity import verify_delivery_format
from lib.docx_inspector import inspect_docx
from lib.engine import UnifiedSynthesizer
from lib.field_updater import build_field_index, verify_final_field_caches
from lib.manifest_migration import migrate_manifest_file
from lib.pagination import extract_pdf_page_labels
from lib.qa import export_docx_to_pdf, word_automation_status
from lib.verification_contracts import build_verification_context


REPO = Path(__file__).resolve().parents[1]
FOLLOWUP = REPO / "docs" / "acceptance" / "word-followup"


def _write_manifest(
    root: Path,
    *,
    filename: str,
    project: str,
    parts: list[str] | None = None,
    source_strategy: str = "docx_document",
    formatting: dict | None = None,
    layout: dict | None = None,
    cover: dict | None = None,
    source_file: str | None = None,
) -> Path:
    data = {
        "schema_version": 3,
        "project_name": project,
        "source": {"strategy": source_strategy},
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": formatting or {"mode": "restyle"},
        "output": {"documents": [{"id": "main", "filename": filename, "parts": parts or ["body"]}]},
    }
    if source_file:
        data["source"]["file"] = source_file
    if layout:
        data["layout"] = layout
    if cover:
        data["cover"] = cover
    path = root / f"{Path(filename).stem}.manifest.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _write_basic_source(path: Path, *, headings: bool = True, paragraphs: int = 12) -> None:
    doc = Document()
    if headings:
        doc.add_heading("第一章 匿名研究", level=1)
    for index in range(paragraphs):
        doc.add_paragraph(
            f"匿名验收正文段落 {index + 1}。该段仅用于真实 Word 分页、格式和内容门禁验证，"
            "不含任何业务材料。" * 4
        )
    if headings:
        doc.add_heading("第二章 匿名结果", level=1)
        # Keep the synthetic heading attached to real body content so the
        # page-level orphan-tail gate tests pagination defects rather than an
        # intentionally empty fixture heading.
        doc.add_paragraph(
            "第二章正文仅用于真实 Word 页级验收，确保章节标题后存在可分页的正文内容。" * 4
        )
    doc.save(path)


def _append_simple_field(paragraph, instruction: str, cached: str = "1") -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = cached
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _part_text(path: Path, part_id: str) -> str:
    doc = Document(str(path))
    children = list(doc.element.body)
    start_name = {"cover": "_Synth_cover", "toc": "_Synth_toc", "body": "_Synth_body"}[part_id]
    end_name = {"cover": "_Synth_part_cover_end", "toc": "_Synth_part_toc_end", "body": "_Synth_part_body_end"}[part_id]
    start = next(i for i, child in enumerate(children) if any(
        item.get(qn("w:name")) == start_name for item in child.iter(qn("w:bookmarkStart"))
    ))
    end = next(i for i in range(start, len(children)) if any(
        item.get(qn("w:name")) == end_name for item in children[i].iter(qn("w:bookmarkStart"))
    ))
    return "".join(
        node.text or ""
        for child in children[start:end + 1]
        for node in child.iter(qn("w:t"))
    )


class PostReviewWordMatrixTest(unittest.TestCase):
    """Real Word cases are intentionally fail-closed when capability is absent."""

    @classmethod
    def setUpClass(cls):
        if os.environ.get("DOCUMENT_SYNTHESIS_WORD_TEST") != "1":
            raise unittest.SkipTest("BLOCKED: set DOCUMENT_SYNTHESIS_WORD_TEST=1 to run NW01–NW12")
        available, reason = word_automation_status()
        if not available:
            raise unittest.SkipTest(f"BLOCKED: Word Automation unavailable: {reason}")

    def _cli(self, source: Path, manifest: Path, output: Path) -> Path:
        common = [sys.executable, "synthesize.py", "--source", str(source), "--manifest", str(manifest)]
        plan = subprocess.run(common + ["--plan"], cwd=REPO, capture_output=True, text=True, timeout=120)
        self.assertEqual(plan.returncode, 0, plan.stdout + plan.stderr)
        doctor = subprocess.run(common + ["--doctor"], cwd=REPO, capture_output=True, text=True, timeout=120)
        self.assertEqual(doctor.returncode, 0, doctor.stdout + doctor.stderr)
        result = subprocess.run(common + ["--output-dir", str(output)], cwd=REPO, capture_output=True, text=True, timeout=900)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        delivered = output / json.loads(manifest.read_text(encoding="utf-8"))["output"]["documents"][0]["filename"]
        self.assertTrue(delivered.is_file(), result.stdout + result.stderr)
        return delivered

    def test_NW01_no_toc_pageref_is_measured_and_cached(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = root / "fixture"
            fixture.mkdir()
            subprocess.run([sys.executable, str(REPO / "docs" / "acceptance" / "create_word_followup_fixtures.py"), str(fixture)], check=True)
            raw = json.loads((fixture / "pageref-manifest.json").read_text(encoding="utf-8"))
            raw["layout"]["parts"].pop("toc", None)
            raw["output"]["documents"][0]["parts"] = ["body"]
            manifest = root / "no-toc.json"
            manifest.write_text(json.dumps(raw, indent=2), encoding="utf-8")
            delivered = self._cli(fixture / "pageref-source.docx", manifest, root / "out")
            index = build_field_index(delivered)
            self.assertTrue(index.pageref_targets)
            metadata = json.loads((root / "out" / "build-metadata.json").read_text(encoding="utf-8"))
            evidence = metadata["deliveries"]["main"]
            target = index.pageref_targets[0]
            self.assertIn(target, evidence["page_map"])
            self.assertNotEqual(verify_final_field_caches(delivered, evidence["page_map"])["fields"][0]["actual"], "1")

    def test_NW02_toc_and_reference_document_use_final_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = root / "fixture"
            fixture.mkdir()
            subprocess.run([sys.executable, str(REPO / "docs" / "acceptance" / "create_word_followup_fixtures.py"), str(fixture)], check=True)
            delivered = self._cli(fixture / "numbering-source.docx", fixture / "numbering-manifest.json", root / "out")
            pdf = root / "out.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            labels = extract_pdf_page_labels(pdf)
            self.assertIn("i", labels.values())
            self.assertIn("1", labels.values())

    def test_NW03_v3_default_cover_ignores_legacy_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            old = root / "封面+目录.docx"
            old_doc = Document()
            old_doc.add_paragraph("旧业务封面")
            old_doc.save(old)
            source = root / "source.docx"
            _write_basic_source(source, paragraphs=4)
            manifest = _write_manifest(root, filename="cover.docx", project="匿名封面", parts=["cover", "body"], cover={"main_title": "NW03 新标题", "author": "匿名", "date": "2026"})
            delivered = self._cli(source, manifest, root / "out")
            cover = _part_text(delivered, "cover")
            self.assertIn("NW03 新标题", cover)
            self.assertNotIn("旧业务封面", cover)

    def test_NW04_preserve_keeps_geometry_and_inline_emphasis(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            section = doc.sections[0]
            section.page_width = Mm(180)
            section.page_height = Mm(240)
            p = doc.add_paragraph()
            p.add_run("保留加粗").bold = True
            p.add_run("普通强调")
            doc.save(source)
            manifest = _write_manifest(root, filename="preserve.docx", project="NW04", formatting={"mode": "preserve", "page_policy": "source"})
            delivered = self._cli(source, manifest, root / "out")
            out = Document(delivered)
            self.assertAlmostEqual(out.sections[0].page_width.mm, 180, delta=0.1)
            self.assertTrue(out.paragraphs[0].runs[0].bold)
            pdf = root / "out.pdf"
            self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
            with __import__("pymupdf").open(pdf) as rendered:
                self.assertAlmostEqual(rendered[0].rect.width * 25.4 / 72, 180, delta=0.2)

    def test_NW05_mixed_managed_and_protected_scope_is_distinct(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.add_paragraph("托管正文")
            protected = doc.add_paragraph()
            protected.add_run("保护正文").bold = True
            doc.save(source)
            manifest = _write_manifest(
                root,
                filename="mixed.docx",
                project="NW05",
                formatting={"mode": "mixed", "page_policy": "target", "scopes": [{"file": str(source), "node_range": "/w:document/w:body/w:p[1]", "mode": "restyle"}]},
            )
            delivered = self._cli(source, manifest, root / "out")
            inspected = inspect_docx(delivered)
            self.assertTrue(inspected.blocks[0].p_style_id.startswith("Synth"))
            self.assertTrue(inspected.blocks[1].effective_run.bold)

    def test_NW06_complete_anonymous_thesis_has_all_declared_objects(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = root / "thesis"
            created = subprocess.run([sys.executable, str(REPO / "docs" / "acceptance" / "create_complete_thesis_fixture.py"), str(fixture)], capture_output=True, text=True)
            self.assertEqual(created.returncode, 0, created.stdout + created.stderr)
            delivered = self._cli(fixture, fixture / "manifest.json", root / "out")
            package = __import__("zipfile").ZipFile(delivered)
            names = set(package.namelist())
            document_xml = package.read("word/document.xml")
            self.assertIn("word/footnotes.xml", names)
            self.assertIn("word/endnotes.xml", names)
            for token in (b"PAGEREF ", b"SEQ Figure", b"SEQ Table", b"w:drawing", b"w:tbl", b"oMath"):
                self.assertIn(token, document_xml)
            self.assertIn("匿名学位论文样例", _part_text(delivered, "cover"))

    def test_NW07_roman_front_sequence_and_odd_page_boundary_are_explicit(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = root / "fixture"
            fixture.mkdir()
            subprocess.run([sys.executable, str(REPO / "docs" / "acceptance" / "create_word_followup_fixtures.py"), str(fixture)], check=True)
            delivered = self._cli(fixture / "numbering-source.docx", fixture / "numbering-manifest.json", root / "out")
            metadata = json.loads((root / "out" / "build-metadata.json").read_text(encoding="utf-8"))
            record = metadata["deliveries"]["main"]["page_records"]
            labels = extract_pdf_page_labels(root / "out.pdf") if (root / "out.pdf").is_file() else {}
            if not labels:
                pdf = root / "out.pdf"
                self.assertTrue(export_docx_to_pdf(str(delivered), str(pdf)))
                labels = extract_pdf_page_labels(pdf)
            self.assertIn("i", labels.values())
            body_record = record.get("_Synth_body")
            self.assertIsNotNone(body_record)
            self.assertEqual(body_record["physical_page"] % 2, 1)

    def test_NW08_preserve_landscape_section_and_header_variants(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = "首页页眉"
            doc.add_paragraph("纵向节正文" * 20)
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.orientation = 1
            section.page_width = Mm(297)
            section.page_height = Mm(210)
            section.header.paragraphs[0].text = "横向节页眉"
            doc.add_paragraph("横向表格节正文" * 20)
            doc.save(source)
            manifest = _write_manifest(root, filename="sections.docx", project="NW08", formatting={"mode": "preserve", "page_policy": "source"})
            delivered = self._cli(source, manifest, root / "out")
            out = Document(delivered)
            self.assertGreaterEqual(len(out.sections), 2)
            self.assertAlmostEqual(out.sections[1].page_width.mm, 297, delta=0.2)
            self.assertIn("横向节页眉", out.sections[1].header.paragraphs[0].text)

    def test_NW09_two_sources_with_same_note_id_keep_note_relationships(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            generated = root / "generated"
            generated.mkdir()
            subprocess.run([sys.executable, str(REPO / "docs" / "acceptance" / "create_complete_thesis_fixture.py"), str(generated)], check=True)
            sources = root / "sources"
            sources.mkdir()
            shutil.copy2(generated / "source" / "thesis-source.docx", sources / "one.docx")
            shutil.copy2(generated / "source" / "thesis-source.docx", sources / "two.docx")
            manifest = _write_manifest(root, filename="notes.docx", project="NW09", source_strategy="directory_tree")
            delivered = self._cli(sources, manifest, root / "out")
            package = __import__("zipfile").ZipFile(delivered)
            self.assertIn("word/footnotes.xml", package.namelist())
            self.assertGreaterEqual(package.read("word/footnotes.xml").count("匿名论文第一条虚构脚注".encode("utf-8")), 2)

    def test_NW10_objects_and_field_instructions_survive_production_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = root / "thesis"
            subprocess.run([sys.executable, str(REPO / "docs" / "acceptance" / "create_complete_thesis_fixture.py"), str(fixture)], check=True)
            delivered = self._cli(fixture, fixture / "manifest.json", root / "out")
            index = build_field_index(delivered)
            self.assertIn("PAGEREF", [item.command for item in index.observations])
            self.assertIn("SEQ", [item.command for item in index.observations])
            xml = __import__("zipfile").ZipFile(delivered).read("word/document.xml")
            for token in (b"w:drawing", b"w:tbl", b"oMath", b"w:numPr"):
                self.assertIn(token, xml)

    def test_NW11_sample_correction_package_mapping_and_target_build(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            work = root / "workflow"
            analysis = subprocess.run([sys.executable, "synthesize.py", "--analyze-format", str(REPO / "examples" / "custom-format-demo" / "sample.docx"), "--analysis-dir", str(work / "sample"), "--replace-output"], cwd=REPO, capture_output=True, text=True)
            self.assertEqual(analysis.returncode, 0, analysis.stdout + analysis.stderr)
            fmt = work / "academic-demo.json"
            result = subprocess.run([sys.executable, "synthesize.py", "--compile-format", str(work / "sample" / "analysis.json"), "--decisions", str(REPO / "examples" / "custom-format-demo" / "decisions.final.json"), "--format-base", "preset:academic-basic@1.0.0", "--format-id", "nw11", "--format-version", "1.0.0", "--format-out", str(fmt), "--replace-output"], cwd=REPO, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            target_analysis = subprocess.run([sys.executable, "synthesize.py", "--analyze-source", str(REPO / "examples" / "custom-format-demo" / "manuscript.docx"), "--analysis-dir", str(work / "target"), "--replace-output"], cwd=REPO, capture_output=True, text=True)
            self.assertEqual(target_analysis.returncode, 0, target_analysis.stdout + target_analysis.stderr)
            mapping = work / "roles.json"
            result = subprocess.run([sys.executable, "synthesize.py", "--compile-mapping", str(work / "target" / "analysis.json"), "--decisions", str(REPO / "examples" / "custom-format-demo" / "manuscript-decisions.final.json"), "--mapping-out", str(mapping), "--replace-output"], cwd=REPO, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            manifest = _write_manifest(root, filename="nw11.docx", project="NW11", formatting={"mode": "restyle", "role_map": str(mapping), "on_unmapped": "error"})
            manifest_data = json.loads(manifest.read_text(encoding="utf-8"))
            manifest_data["format"]["ref"] = str(fmt)
            manifest.write_text(json.dumps(manifest_data, indent=2), encoding="utf-8")
            delivered = self._cli(REPO / "examples" / "custom-format-demo" / "manuscript.docx", manifest, root / "out")
            self.assertTrue(delivered.is_file())

    def test_NW12_fault_injection_fails_closed_and_does_not_change_published_hash(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            # Leave enough body material for the real Word pagination gate to
            # establish a page boundary; this test targets fail-closed format
            # verification after corruption, not a short-document tail page.
            _write_basic_source(source, paragraphs=12)
            manifest = _write_manifest(root, filename="published.docx", project="NW12")
            delivered = self._cli(source, manifest, root / "out")
            before = hashlib.sha256(delivered.read_bytes()).hexdigest()
            corrupted = root / "corrupted.docx"
            doc = Document(delivered)
            doc.paragraphs[0].runs[0].font.size = Pt(70)
            doc.save(corrupted)
            config = load_project_config(source.parent, manifest, project_name="NW12")
            prepared = prepare_project_build(source, config)
            context = build_verification_context(prepared, None, config.documents[0], corrupted)
            report = verify_delivery_format(corrupted, config.resolved_format, verification_context=context)
            self.assertFalse(report.passed)
            self.assertEqual(before, hashlib.sha256(delivered.read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
