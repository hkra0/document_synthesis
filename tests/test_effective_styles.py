# -*- coding: utf-8 -*-
"""
OOXML 有效排版属性层叠求值测试 (tests/test_effective_styles.py)
"""

import unittest
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml

from lib.docx_inspector import (
    inspect_docx,
    EffectiveStyleEvaluator,
    StyleNode,
    ThemeFonts,
)
from lib.format_schema import (
    RunStyle,
    ParagraphStyle,
    LengthValue,
    LineSpacing,
)


class EffectiveStylesTest(unittest.TestCase):

    def test_multi_level_style_inheritance(self):
        """测试 Grandparent -> Parent -> Child 的多级属性层叠继承"""
        styles = {
            "Grandparent": StyleNode(
                style_id="Grandparent",
                style_type="paragraph",
                name="Grandparent",
                run_style=RunStyle(size_pt=14.0, bold=True, color="0000ff", east_asia="宋体"),
                paragraph_style=ParagraphStyle(alignment="center", space_before_pt=10.0),
            ),
            "Parent": StyleNode(
                style_id="Parent",
                style_type="paragraph",
                name="Parent",
                based_on="Grandparent",
                run_style=RunStyle(size_pt=16.0),  # 覆写字号
                paragraph_style=ParagraphStyle(space_before_pt=15.0),  # 覆写段前
            ),
            "Child": StyleNode(
                style_id="Child",
                style_type="paragraph",
                name="Child",
                based_on="Parent",
                run_style=RunStyle(east_asia="黑体"),  # 覆写字体
                paragraph_style=ParagraphStyle(alignment="left"),  # 覆写对齐
            ),
        }
        evaluator = EffectiveStyleEvaluator(styles, ThemeFonts())

        eff_para = evaluator.compute_effective_paragraph("Child", None)
        self.assertEqual(eff_para.alignment, "left")  # 来自 Child
        self.assertEqual(eff_para.space_before_pt, 15.0)  # 来自 Parent

        eff_run = evaluator.compute_effective_run("Child", None, None)
        self.assertEqual(eff_run.east_asia, "黑体")  # 来自 Child
        self.assertEqual(eff_run.size_pt, 16.0)  # 来自 Parent
        self.assertTrue(eff_run.bold)  # 来自 Grandparent
        self.assertEqual(eff_run.color, "0000ff")  # 来自 Grandparent

    def test_direct_formatting_toggle_override(self):
        """测试直接格式 <w:b w:val="0"/> 能够成功覆盖样式的 bold=True"""
        styles = {
            "BoldHeading": StyleNode(
                style_id="BoldHeading",
                style_type="paragraph",
                name="BoldHeading",
                run_style=RunStyle(bold=True, size_pt=16.0),
            )
        }
        evaluator = EffectiveStyleEvaluator(styles, ThemeFonts())

        # 1. 样式默认加粗
        run_default = evaluator.compute_effective_run("BoldHeading", None, None)
        self.assertTrue(run_default.bold)

        # 2. 直接格式显式取消加粗 <w:b w:val="0"/>
        rPr_unbold = parse_xml(r'<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:b w:val="0"/></w:rPr>')
        run_unbold = evaluator.compute_effective_run("BoldHeading", None, rPr_unbold)
        self.assertFalse(run_unbold.bold)

    def test_theme_font_resolution(self):
        """测试主题字体引用的解析"""
        theme = ThemeFonts(
            major_east_asia="华文中宋",
            major_ascii="Arial",
            minor_east_asia="仿宋",
            minor_ascii="Calibri",
        )
        evaluator = EffectiveStyleEvaluator({}, theme)

        rPr = parse_xml(r'''
            <w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                <w:rFonts w:eastAsiaTheme="minorEastAsia" w:asciiTheme="majorHAnsi"/>
            </w:rPr>
        ''')
        run = evaluator.compute_effective_run(None, None, rPr)
        self.assertEqual(run.east_asia, "仿宋")
        self.assertEqual(run.latin, "Arial")

    def test_full_document_effective_properties_inspection(self):
        """测试真实生成的 DOCX 文档中段距、行距与字符缩进的解析"""
        with tempfile.TemporaryDirectory() as td:
            doc_path = Path(td) / "props_test.docx"
            doc = Document()
            p = doc.add_paragraph()
            p.alignment = 3  # justify
            p_xml = p._p
            pPr = p_xml.get_or_add_pPr()

            # 首行缩进 2 字符 (w:firstLineChars="200")
            ind = parse_xml(r'<w:ind xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:firstLineChars="200"/>')
            pPr.append(ind)
            # 1.5 倍行距 (w:line="360" w:lineRule="auto")
            sp = parse_xml(r'<w:spacing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:before="240" w:after="120" w:line="360" w:lineRule="auto"/>')
            pPr.append(sp)

            r = p.add_run("测试正文文本")
            doc.save(str(doc_path))

            inspection = inspect_docx(doc_path)
            self.assertEqual(len(inspection.blocks), 1)
            b = inspection.blocks[0]

            self.assertEqual(b.visible_text, "测试正文文本")
            self.assertIsNotNone(b.effective_paragraph)
            self.assertEqual(b.effective_paragraph.alignment, "justify")
            self.assertEqual(b.effective_paragraph.first_line_indent.value, 2.0)
            self.assertEqual(b.effective_paragraph.first_line_indent.unit, "char")
            self.assertEqual(b.effective_paragraph.space_before_pt, 12.0)
            self.assertEqual(b.effective_paragraph.space_after_pt, 6.0)
            self.assertEqual(b.effective_paragraph.line_spacing.mode, "multiple")
            self.assertEqual(b.effective_paragraph.line_spacing.value, 1.5)


if __name__ == "__main__":
    unittest.main()
