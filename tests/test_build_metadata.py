# -*- coding: utf-8 -*-
"""
构建元数据与发布事务原子回滚单元测试 (tests/test_build_metadata.py)
"""

import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
import pymupdf

from lib.config import ProjectConfig, load_project_config
from lib.engine import BuildError, UnifiedSynthesizer, _publish_deliveries, generate_build_metadata
from lib.format_schema import (
    LengthValue,
    LineSpacing,
    PageSpec,
    ParagraphStyle,
    ResolvedFormat,
    RoleSpec,
    RunStyle,
    StyleDefinition,
    TocSpec,
    HeaderFooterSpec,
)


class BuildMetadataTest(unittest.TestCase):
    def test_generate_build_metadata_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            f1 = tmp_path / "main.docx"
            Document().save(str(f1))

            config = ProjectConfig(
                {
                    "schema_version": 3,
                    "project_name": "metadata-test",
                    "format": {"ref": "preset:academic-basic@1.0.0"},
                    "formatting": {"mode": "restyle"},
                    "output": {"documents": [{"id": "main", "filename": "main.docx", "parts": ["body"]}]},
                },
                tmp_path,
            )

            staged = {"main": f1}
            source_hashes = {"input.docx": "abcdef1234567890"}
            qa_summary = {
                "content_integrity": "passed",
                "format_verification": "passed",
                "unverified_attributes": [],
            }

            meta = generate_build_metadata(config, staged, source_hashes, qa_summary=qa_summary)

            self.assertEqual(meta["generator"], "document-synthesis")
            self.assertEqual(meta["engine_version"], "3.0.0")
            self.assertEqual(meta["schema_version"], 3)
            self.assertEqual(meta["project_name"], "metadata-test")
            self.assertIn("build_timestamp", meta)
            self.assertEqual(meta["source_hashes"], source_hashes)
            self.assertIn("main", meta["deliveries"])
            self.assertEqual(meta["deliveries"]["main"]["filename"], "main.docx")
            self.assertTrue(len(meta["deliveries"]["main"]["sha256"]) == 64)
            self.assertIn("python_version", meta["environment"])
            self.assertIn("platform", meta["environment"])
            self.assertEqual(meta["qa"]["content_integrity"], "passed")
            self.assertEqual(meta["qa"]["format_verification"], "passed")
            self.assertEqual(len(meta["configuration"]["config_contract_sha256"]), 64)

    def test_metadata_records_pdf_page_evidence_and_configuration_hashes(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            manifest_path = tmp_path / "manifest.json"
            manifest_path.write_text(
                json.dumps(
                    {
                        "schema_version": 3,
                        "project_name": "evidence-test",
                        "format": {"ref": "preset:academic-basic@1.0.0"},
                        "formatting": {"mode": "restyle"},
                        "output": {
                            "documents": [
                                {"id": "main", "filename": "main.docx", "parts": ["body"]}
                            ]
                        },
                    }
                ),
                encoding="utf-8",
            )
            config = load_project_config(tmp_path, manifest_path)
            doc_path = tmp_path / "main.docx"
            Document().save(doc_path)
            pdf_path = tmp_path / "qa-main.pdf"
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 750), "- 1 -", fontsize=10)
                pdf.save(pdf_path)

            page_map = {
                "_Synth_body": {
                    "physical_page": 1,
                    "printed_page": 1,
                    "expected_label": "1",
                    "observed_label": "1",
                    "verification_status": "verified",
                }
            }
            page_records = {
                "_Synth_body": {
                    "physical_page": 1,
                    "section_id": 1,
                    "sequence_id": "main",
                    "number_value": 1,
                    "expected_label": "1",
                    "observed_label": "1",
                    "label_verified": True,
                    "verification_status": "verified",
                }
            }
            meta = generate_build_metadata(
                config,
                {"main": doc_path},
                {str(doc_path): "source-hash"},
                qa_summary={"content_integrity": "passed", "format_verification": "passed"},
                page_evidence={
                    "main": {
                        "pdf_path": str(pdf_path),
                        "page_map": page_map,
                        "page_records": page_records,
                    }
                },
            )

            delivery = meta["deliveries"]["main"]
            self.assertEqual(len(delivery["sha256"]), 64)
            self.assertEqual(len(delivery["pdf"]["sha256"]), 64)
            self.assertEqual(delivery["page_records"]["_Synth_body"]["verification_status"], "verified")
            self.assertEqual(len(meta["configuration"]["manifest_sha256"]), 64)
            self.assertIn("word_automation", meta["environment"])

    def test_publish_deliveries_moves_metadata_atomically(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            run_dir = tmp_path / "work"
            run_dir.mkdir()
            out_dir = tmp_path / "output"
            out_dir.mkdir()

            doc_file = run_dir / "deliver.docx"
            Document().save(str(doc_file))
            meta_file = run_dir / "build-metadata.json"
            meta_file.write_text('{"test": true}', encoding="utf-8")

            staged = {
                "main": doc_file,
                "__metadata__": meta_file,
            }

            targets = _publish_deliveries(staged, out_dir, run_dir)
            self.assertTrue((out_dir / "deliver.docx").is_file())
            self.assertTrue((out_dir / "build-metadata.json").is_file())
            self.assertEqual(targets["main"], out_dir / "deliver.docx")
            self.assertEqual(targets["__metadata__"], out_dir / "build-metadata.json")

    def test_publish_deliveries_rollback_on_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            run_dir = tmp_path / "work"
            run_dir.mkdir()
            out_dir = tmp_path / "output"
            out_dir.mkdir()

            # 模拟交付目标已存在旧版文件
            old_doc = out_dir / "target.docx"
            old_doc.write_bytes(b"OLD_DOC_BYTES")
            old_meta = out_dir / "build-metadata.json"
            old_meta.write_text('{"version": "old"}', encoding="utf-8")

            # 模拟工作区准备就绪的新文件
            new_doc = run_dir / "target.docx"
            new_doc.write_bytes(b"NEW_DOC_BYTES")
            new_meta = run_dir / "build-metadata.json"
            new_meta.write_text('{"version": "new"}', encoding="utf-8")

            staged = {
                "main": new_doc,
                "__metadata__": new_meta,
            }

            original_replace = os.replace
            replace_count = [0]

            def faulty_replace(src, dst):
                if Path(src).name == "build-metadata.json":
                    raise PermissionError("模拟第 2 个文件替换时磁盘权限故障")
                original_replace(src, dst)

            with patch("os.replace", side_effect=faulty_replace):
                with self.assertRaises(PermissionError):
                    _publish_deliveries(staged, out_dir, run_dir)

            # 验证回滚：第 1 个被替换的文件必须已被还原回旧内容！
            self.assertEqual(old_doc.read_bytes(), b"OLD_DOC_BYTES")
            self.assertEqual(old_meta.read_text(encoding="utf-8"), '{"version": "old"}')


if __name__ == "__main__":
    unittest.main()
