# -*- coding: utf-8 -*-
"""
具名样式安装、角色应用与行内语义保护测试 (tests/test_style_applier.py)
"""

import unittest
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from lib.format_resolver import resolve_format_package
from lib.format_schema import FormattingPolicy
from lib.style_applier import (
    install_styles,
    apply_roles,
    add_styled_heading,
    add_styled_toc_entry,
    setup_styled_footer,
    ROLE_STYLE_ID_MAP,
)


class StyleApplierTest(unittest.TestCase):

    def setUp(self):
        self.report_format = resolve_format_package("preset:report-basic@1.0.0")
        self.academic_format = resolve_format_package("preset:academic-basic@1.0.0")

    def test_install_styles_creates_synth_styles_and_outline_levels(self):
        doc = Document()
        role_map = install_styles(doc, self.academic_format)

        self.assertIn("body", role_map)
        self.assertEqual(role_map["body"], "SynthBody")
        self.assertEqual(role_map["heading.1"], "SynthHeading1")

        styles_el = doc.part.styles._element
        # 验证 SynthBody
        s_body = styles_el.find(f".//{qn('w:style')}[@{qn('w:styleId')}='SynthBody']")
        self.assertIsNotNone(s_body)
        name_body = s_body.find(qn("w:name")).get(qn("w:val"))
        self.assertEqual(name_body, "Synth Body")

        # 验证 SynthHeading1 大纲级别 (outlineLvl = 0)
        s_h1 = styles_el.find(f".//{qn('w:style')}[@{qn('w:styleId')}='SynthHeading1']")
        self.assertIsNotNone(s_h1)
        out_lvl = s_h1.find(f".//{qn('w:outlineLvl')}")
        self.assertIsNotNone(out_lvl)
        self.assertEqual(out_lvl.get(qn("w:val")), "0")

    def test_apply_roles_clears_direct_formatting_and_sets_pstyle(self):
        doc = Document()
        p1 = doc.add_paragraph()
        # 直接添加带直接格式的标题
        p1_pr = p1._p.get_or_add_pPr()
        p1_pr.append(parse_xml(r'<w:jc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="center"/>'))
        p1_pr.append(parse_xml(r'<w:spacing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:before="500"/>'))
        r1 = p1.add_run("第一章 绪论")
        r1_pr = r1._r.get_or_add_rPr()
        r1_pr.append(parse_xml(r'<w:sz xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="48"/>'))
        r1_pr.append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:eastAsia="华文细黑"/>'))

        report = apply_roles(
            doc,
            role_assignments={"0": "heading.1"},
            resolved_format=self.academic_format,
        )

        self.assertEqual(report.styled_paragraphs_count, 1)
        self.assertTrue(report.cleared_direct_formats_count > 0)

        # 验证 pStyle 已应用
        pStyle = p1._p.find(f".//{qn('w:pStyle')}")
        self.assertIsNotNone(pStyle)
        self.assertEqual(pStyle.get(qn("w:val")), "SynthHeading1")

        # 验证冲突的直接格式已按属性所有权清理
        self.assertIsNone(p1._p.find(f".//{qn('w:jc')}"))
        self.assertIsNone(p1._p.find(f".//{qn('w:spacing')}"))
        self.assertIsNone(r1._r.find(f".//{qn('w:sz')}"))
        self.assertIsNone(r1._r.find(f".//{qn('w:rFonts')}"))

        # 验证文本内容毫发无损
        self.assertEqual(p1.text, "第一章 绪论")

    def test_apply_roles_preserves_style_inherited_numbering(self):
        doc = Document()
        paragraph = doc.add_paragraph("列表语义", style="List Bullet")

        apply_roles(doc, {"0": "body"}, self.academic_format)

        self.assertEqual(
            paragraph._p.find(f".//{qn('w:pStyle')}").get(qn("w:val")),
            "SynthBody",
        )
        self.assertIsNotNone(paragraph._p.find(f".//{qn('w:numPr')}/{qn('w:numId')}"))

    def test_preserve_inline_semantics_subscript_and_emphasis(self):
        """测试上下标、局部斜体强调、公式在样式应用后零丢失"""
        doc = Document()
        p = doc.add_paragraph()

        # Run 1: 普通文本
        r1 = p.add_run("水分子化学式为 H")
        # Run 2: 下标 2
        r2 = p.add_run("2")
        r2_pr = r2._r.get_or_add_rPr()
        r2_pr.append(parse_xml(r'<w:vertAlign xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="subscript"/>'))
        # Run 3: 普通文本
        r3 = p.add_run("O，其中 ")
        # Run 4: 局部斜体强调
        r4 = p.add_run("特别强调")
        r4.italic = True
        # Run 5: 结束
        r5 = p.add_run(" 该特性。")

        # 插入数学公式元素
        omml = parse_xml(r'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+y=z</m:t></m:r></m:oMath>')
        p._p.append(omml)

        report = apply_roles(
            doc,
            role_assignments={"0": "body"},
            resolved_format=self.academic_format,
            policy=FormattingPolicy(inline_emphasis="preserve"),
        )

        self.assertTrue(report.preserved_inline_objects_count >= 2)

        # 验证下标未被清理
        vert_align = r2._r.find(f".//{qn('w:vertAlign')}")
        self.assertIsNotNone(vert_align)
        self.assertEqual(vert_align.get(qn("w:val")), "subscript")

        # 验证斜体强调未被清理
        self.assertTrue(r4.italic)

        # 验证数学公式结构完好
        math_el = p._p.find(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        self.assertIsNotNone(math_el)

        # 验证文字完整
        self.assertIn("水分子化学式为 H2O", p.text)

    def test_apply_roles_styles_table_cell_paragraphs(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                cell.text = f"表格单元格 {row_index}-{cell_index}"

        role_map = {
            f"/w:document/w:body/w:tbl[1]/w:tr[{row_index}]"
            f"/w:tc[{cell_index}]/w:p[1]": "table.body"
            for row_index in range(1, 3)
            for cell_index in range(1, 3)
        }
        report = apply_roles(doc, role_map, self.academic_format)

        self.assertEqual(report.styled_paragraphs_count, 4)
        for row in table.rows:
            for cell in row.cells:
                p_style = cell.paragraphs[0]._p.find(f".//{qn('w:pStyle')}")
                self.assertIsNotNone(p_style)
                self.assertEqual(p_style.get(qn("w:val")), "SynthTableBody")

    def test_switching_between_formats_is_idempotent(self):
        """测试在同一文档上切换格式包，能够正确更新为新样式且幂等"""
        doc = Document()
        p = doc.add_paragraph("测试切换格式包的正文段落。")

        # 1. 首次应用 report-basic
        apply_roles(doc, {"0": "body"}, self.report_format)
        pStyle = p._p.find(f".//{qn('w:pStyle')}")
        self.assertEqual(pStyle.get(qn("w:val")), "SynthBody")
        s_body = doc.part.styles._element.find(f".//{qn('w:style')}[@{qn('w:styleId')}='SynthBody']")
        ea_font = s_body.find(f".//{qn('w:rFonts')}").get(qn("w:eastAsia"))
        self.assertEqual(ea_font, "微软雅黑")

        # 2. 切换应用 academic-basic
        apply_roles(doc, {"0": "body"}, self.academic_format)
        s_body2 = doc.part.styles._element.find(f".//{qn('w:style')}[@{qn('w:styleId')}='SynthBody']")
        ea_font2 = s_body2.find(f".//{qn('w:rFonts')}").get(qn("w:eastAsia"))
        self.assertEqual(ea_font2, "宋体")

        # 3. 再次应用 academic-basic 保持幂等
        apply_roles(doc, {"0": "body"}, self.academic_format)
        self.assertEqual(p.text, "测试切换格式包的正文段落。")

    def test_output_adapters(self):
        """测试标题、目录与页脚输出适配器"""
        doc = Document()

        # 标题适配器
        p_head = add_styled_heading(doc, "第一章 概述", level=1, resolved_format=self.academic_format, bookmark_name="_Toc_001")
        self.assertEqual(p_head.text, "第一章 概述")
        self.assertIsNotNone(p_head._p.find(f".//{qn('w:bookmarkStart')}"))

        # 目录适配器
        p_toc = add_styled_toc_entry(doc, "第一章 概述", page_num=1, level=1, resolved_format=self.academic_format, bookmark_name="_Toc_001")
        self.assertIn("第一章 概述", p_toc.text)
        self.assertIsNotNone(p_toc._p.find(f".//{qn('w:tabs')}"))

        # 页脚适配器
        sec = doc.sections[0]
        setup_styled_footer(sec, self.academic_format, start_page=1)
        self.assertIsNotNone(sec.footer.paragraphs[0]._p.find(f".//{qn('w:fldSimple')}"))


if __name__ == "__main__":
    unittest.main()
