"""Real assembly and PDF QA with only the Office measurement boundary mocked."""

import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib.composition import BOUNDARIES, add_bookmark
from lib.config import ProjectConfig
from lib.delivery import build_deliveries, validate_delivery_structure
from lib.engine import BuildError, UnifiedSynthesizer, _publish_deliveries
from lib.build_plan import PreparedBuild
from lib.qa import OfficeExportError


class DeliveryPipelineTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.body = self.root / "body.docx"
        self.run = self.root / "run"
        self.run.mkdir()
        self.nodes = [
            {"title": "Repeated title", "level": 1, "bookmark_name": "_Toc_auto_001", "bm_id": 1},
            {"title": "Repeated title", "level": 1, "bookmark_name": "_Toc_auto_002", "bm_id": 2},
        ]
        document = Document()
        for node in self.nodes:
            paragraph = document.add_paragraph(node["title"])
            add_bookmark(paragraph._p, node["bookmark_name"], node["bm_id"])
            document.add_paragraph("Fictional supporting text with enough content for testing.")
        document.save(self.body)

    def config(self, documents=None):
        data = {"schema_version": 2, "project_name": "Fictional Pipeline", "cover": {"template": False}}
        if documents is not None:
            data["output"] = {"documents": documents}
        return ProjectConfig(data)

    def split_config(self):
        return self.config([
            {"id": "front", "filename": "front.docx", "parts": ["cover", "toc"], "toc": {"reference": "main", "links": "none"}},
            {"id": "main", "filename": "main.docx", "parts": ["toc", "body"]},
        ])

    def measured_map(self, spec, last_printed=2):
        page_map = {}
        physical = 1
        for part in spec["parts"]:
            page_map[BOUNDARIES[part]] = {"physical_page": physical, "printed_page": 1}
            physical += 2 if part == "toc" else 1
        if "body" in spec["parts"]:
            start = page_map[BOUNDARIES["body"]]["physical_page"]
            page_map[self.nodes[0]["bookmark_name"]] = {"physical_page": start, "printed_page": 1}
            page_map[self.nodes[1]["bookmark_name"]] = {"physical_page": start + last_printed - 1, "printed_page": last_printed}
        return page_map

    def export_pdf(self, path, page_map, blank_page=None, orphan_page=None):
        count = max(item["physical_page"] for item in page_map.values())
        with pymupdf.open() as pdf:
            for number in range(1, count + 1):
                page = pdf.new_page()
                if number == blank_page:
                    continue
                text = "Orphan line" if number == orphan_page else "Fictional document\nSupporting content\nAdditional content\nFinal content"
                page.insert_text((72, 72), text, fontsize=12)
            pdf.save(str(path))

    def measurement(self, config, values, calls, blank=None, orphan=None, missing=None):
        counts = {}

        def inspect(path, pdf_path, names):
            spec = next(item for item in config.documents if item["filename"] == Path(path).name)
            identity = spec["id"]
            index = counts.get(identity, 0)
            counts[identity] = index + 1
            calls.append(identity)
            printed = values[min(index, len(values) - 1)]
            page_map = self.measured_map(spec, printed)
            self.assertEqual(set(names), set(page_map))
            self.export_pdf(pdf_path, page_map,
                            blank_page=(blank or {}).get(identity),
                            orphan_page=(orphan or {}).get(identity))
            if missing:
                page_map.pop(missing, None)
            return page_map

        return inspect

    def test_complete_document_converges_after_three_measured_passes(self):
        config, calls = self.config(), []
        original = self.body.read_bytes()
        with patch("lib.delivery.inspect_document", side_effect=self.measurement(config, [2, 3, 3], calls)):
            staged = build_deliveries(config, self.body, self.nodes, self.run)
        self.assertEqual(calls, ["main", "main", "main"])
        self.assertEqual(set(staged), {"main"})
        self.assertEqual(self.body.read_bytes(), original)
        validate_delivery_structure(staged["main"], config.documents[0], self.nodes,
                                    self.measured_map(config.documents[0], 3))

    def test_split_print_directory_uses_referenced_final_body_map(self):
        config, calls = self.split_config(), []
        with patch("lib.delivery.inspect_document", side_effect=self.measurement(config, [2, 2], calls)):
            staged = build_deliveries(config, self.body, self.nodes, self.run)
        self.assertEqual(calls, ["main", "main", "front"])
        self.assertEqual(set(staged), {"main", "front"})
        main = next(spec for spec in config.documents if spec["id"] == "main")
        front = next(spec for spec in config.documents if spec["id"] == "front")
        validate_delivery_structure(staged["front"], front, self.nodes, self.measured_map(main))
        from docx.oxml.ns import qn
        self.assertEqual(list(Document(staged["front"]).element.body.iter(qn("w:hyperlink"))), [])

    def test_nonconvergent_toc_stops_at_three_passes(self):
        config, calls = self.config(), []
        with patch("lib.delivery.inspect_document", side_effect=self.measurement(config, [2, 3, 4], calls)):
            with self.assertRaisesRegex(OfficeExportError, "3.*未收敛"):
                build_deliveries(config, self.body, self.nodes, self.run)
        self.assertEqual(calls, ["main"] * 3)

    def test_missing_measurement_never_becomes_guessed_page_one(self):
        config, calls = self.config(), []
        with patch("lib.delivery.inspect_document", side_effect=self.measurement(config, [2], calls, missing="_Toc_auto_002")):
            with self.assertRaises(OfficeExportError):
                build_deliveries(config, self.body, self.nodes, self.run)
        self.assertEqual(calls, ["main"])

    def test_missing_source_bookmark_rejected_before_word(self):
        config = self.config()
        Document().save(self.body)
        with patch("lib.delivery.inspect_document") as inspect:
            with self.assertRaisesRegex(OfficeExportError, "缺少书签"):
                build_deliveries(config, self.body, self.nodes, self.run)
        inspect.assert_not_called()

    def test_blank_cover_or_toc_in_split_artifact_fails_real_pdf_qa(self):
        config = self.split_config()
        for blank in (1, 2):
            with self.subTest(blank=blank):
                calls = []
                with patch("lib.delivery.inspect_document", side_effect=self.measurement(config, [2], calls, blank={"front": blank})):
                    with self.assertRaisesRegex(OfficeExportError, "页级 QA.*front.docx"):
                        build_deliveries(config, self.body, self.nodes, self.run)
                self.assertEqual(calls, ["main", "main", "front"])

    def test_body_orphan_in_complete_artifact_fails_real_pdf_qa(self):
        config, calls = self.config(), []
        with patch("lib.delivery.inspect_document", side_effect=self.measurement(config, [2], calls, orphan={"main": 5})):
            with self.assertRaisesRegex(OfficeExportError, "页级 QA"):
                build_deliveries(config, self.body, self.nodes, self.run)

    def test_pipeline_failure_does_not_publish_or_delete_prior_artifacts(self):
        config = self.config()
        source, output = self.root / "source", self.root / "output"
        source.mkdir()
        output.mkdir()
        existing = output / config.documents[0]["filename"]
        existing.write_bytes(b"previously verified delivery")
        unrelated = output / "unrelated.docx"
        unrelated.write_bytes(b"user-owned file")
        prepared = PreparedBuild(
            source_path=source.resolve(),
            config=config,
            deliveries=tuple(config.documents),
        )
        with patch.object(UnifiedSynthesizer, "prepare_build", return_value=prepared), \
                patch("lib.engine.prepare_conversions", return_value={}), \
                patch("lib.delivery.build_deliveries", side_effect=OfficeExportError("final artifact failed QA")), \
                patch("lib.engine._publish_deliveries") as publish:
            with self.assertRaisesRegex(BuildError, "failed QA"):
                UnifiedSynthesizer.synthesize(source, output)
        publish.assert_not_called()
        self.assertEqual(existing.read_bytes(), b"previously verified delivery")
        self.assertEqual(unrelated.read_bytes(), b"user-owned file")
        self.assertFalse((output / ".work").exists())

    def test_publish_second_replacement_failure_restores_first_original(self):
        output = self.root / "published"
        output.mkdir()
        staged = {}
        for name in ("first", "second"):
            (output / f"{name}.docx").write_bytes(f"old {name}".encode())
            stage = self.run / f"{name}.docx"
            stage.write_bytes(f"new {name}".encode())
            staged[name] = stage
        replace = os.replace
        count = 0

        def fail_second(source, target):
            nonlocal count
            count += 1
            if count == 2:
                raise OSError("simulated publication failure")
            return replace(source, target)

        with patch("lib.engine.os.replace", side_effect=fail_second):
            with self.assertRaisesRegex(OSError, "publication failure"):
                _publish_deliveries(staged, output, self.run)
        self.assertEqual(count, 3)
        for name in ("first", "second"):
            self.assertEqual((output / f"{name}.docx").read_bytes(), f"old {name}".encode())

    def test_publish_failure_restores_all_deliveries_and_metadata(self):
        """S7: publication rollback covers every artifact in the batch."""
        output = self.root / "published-with-metadata"
        output.mkdir()
        staged = {}
        old_contents = {
            "first.docx": b"old first",
            "second.docx": b"old second",
            "build-metadata.json": b'{"status":"previous"}',
        }
        for filename, content in old_contents.items():
            (output / filename).write_bytes(content)
            stage = self.run / filename
            stage.write_bytes((b"new " + filename.encode()))
            staged[filename] = stage
        replace = os.replace

        def fail_second(source, target):
            if Path(source).name == "second.docx":
                raise OSError("simulated publication failure")
            return replace(source, target)

        with patch("lib.engine.os.replace", side_effect=fail_second):
            with self.assertRaisesRegex(OSError, "publication failure"):
                _publish_deliveries(staged, output, self.run)
        for filename, content in old_contents.items():
            self.assertEqual((output / filename).read_bytes(), content)

    def test_publish_rollback_removes_new_first_artifact_without_old_version(self):
        output = self.root / "published"
        output.mkdir()
        staged = {}
        for name in ("first", "second"):
            stage = self.run / f"{name}.docx"
            stage.write_bytes(f"new {name}".encode())
            staged[name] = stage
        replace = os.replace

        def fail_second(source, target):
            if Path(source) == staged["second"]:
                raise OSError("simulated publication failure")
            return replace(source, target)

        with patch("lib.engine.os.replace", side_effect=fail_second):
            with self.assertRaises(OSError):
                _publish_deliveries(staged, output, self.run)
        self.assertFalse((output / "first.docx").exists())
        self.assertFalse((output / "second.docx").exists())

    def create_engine_fixture(self, strategy, schema=2):
        fixture = self.root / f"{strategy}-v{schema}"
        fixture.mkdir()
        source_dir = fixture / "source"
        source_dir.mkdir()
        source_doc = source_dir / "1_evidence.docx"
        document = Document()
        run = document.add_paragraph().add_run("Fictional Evidence Heading")
        if strategy == "highlighted_docx":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        document.add_paragraph("Original fictional evidence retained without modification.")
        document.add_paragraph("Additional fictional content for the source document.")
        document.save(source_doc)
        data = {"schema_version": schema, "project_name": "Fictional Engine",
                "cover": {"template": False}, "source": {"strategy": strategy}}
        if strategy == "explicit_tree":
            data["source"]["tree"] = [{
                "level": 1, "title": "Configured Evidence", "type": "docx_outline",
                "file": source_doc.name,
                "outline_rules": [{"level": 2, "pattern": "^Fictional Evidence"}],
            }]
        manifest = fixture / "manifest.json"
        manifest.write_text(json.dumps(data), encoding="utf-8")
        source = source_doc if strategy == "highlighted_docx" else source_dir
        return source, manifest, fixture / "output", source_doc

    def flat_engine_measurement(self, calls):
        """Measure each declared part as one page, keeping headings on body page 1."""
        def inspect(path, pdf_path, names):
            calls.append(Path(path).name)
            page_map, physical = {}, 1
            for part in ("cover", "toc", "body"):
                name = BOUNDARIES[part]
                if name in names:
                    page_map[name] = {"physical_page": physical, "printed_page": 1}
                    physical += 1
            for name in names:
                if name not in page_map:
                    page_map[name] = dict(page_map[BOUNDARIES["body"]])
            self.export_pdf(pdf_path, page_map)
            return page_map
        return inspect

    def test_all_source_strategies_render_validate_and_publish_one_complete_doc(self):
        for strategy in ("directory_tree", "explicit_tree", "highlighted_docx"):
            with self.subTest(strategy=strategy):
                source, manifest, output, original_doc = self.create_engine_fixture(strategy)
                original_bytes, manifest_bytes = original_doc.read_bytes(), manifest.read_bytes()
                calls = []
                with patch("lib.delivery.inspect_document", side_effect=self.flat_engine_measurement(calls)):
                    result = UnifiedSynthesizer.synthesize(source, output, manifest)
                self.assertEqual(set(result), {"main"})
                self.assertEqual(calls, [result["main"].name])
                self.assertEqual(sorted(output.iterdir()), [result["main"]])
                content = "\n".join(p.text for p in Document(result["main"]).paragraphs)
                self.assertIn("Original fictional evidence retained without modification.", content)
                from docx.oxml.ns import qn
                bookmarks = {b.get(qn("w:name")) for b in Document(result["main"]).element.body.iter(qn("w:bookmarkStart"))}
                self.assertTrue(set(BOUNDARIES.values()).issubset(bookmarks))
                self.assertEqual(original_doc.read_bytes(), original_bytes)
                self.assertEqual(manifest.read_bytes(), manifest_bytes)
                self.assertEqual(list(original_doc.parent.iterdir()), [original_doc])
                self.assertFalse((output / ".work").exists())

    def test_v1_full_engine_retains_all_three_legacy_deliveries(self):
        source, manifest, output, original_doc = self.create_engine_fixture("directory_tree", schema=1)
        original_bytes = original_doc.read_bytes()
        calls = []
        expected = ProjectConfig(json.loads(manifest.read_text()), source).documents
        with patch("lib.delivery.inspect_document", side_effect=self.flat_engine_measurement(calls)):
            result = UnifiedSynthesizer.synthesize(source, output, manifest)
        self.assertEqual(set(result), {spec["id"] for spec in expected})
        self.assertEqual({p.name for p in output.iterdir()}, {spec["filename"] for spec in expected})
        self.assertEqual(set(calls), {spec["filename"] for spec in expected})
        self.assertEqual(len(calls), 3)
        self.assertEqual(original_doc.read_bytes(), original_bytes)
        self.assertFalse((output / ".work").exists())

    def test_hidden_outline_parent_gets_anchor_without_extra_heading(self):
        source, manifest, output, _ = self.create_engine_fixture("explicit_tree")
        data = json.loads(manifest.read_text())
        data["source"]["tree"][0]["add_heading_before_content"] = False
        manifest.write_text(json.dumps(data), encoding="utf-8")
        with patch("lib.delivery.inspect_document", side_effect=self.flat_engine_measurement([])):
            result = UnifiedSynthesizer.synthesize(source, output, manifest)
        from docx.oxml.ns import qn
        document = Document(result["main"])
        parent = next(b for b in document.element.body.iter(qn("w:bookmarkStart"))
                      if b.get(qn("w:name")) == "_Toc_auto_001")
        texts = [t.text or "" for t in parent.getparent().iter(qn("w:t"))]
        self.assertEqual("".join(texts), "Fictional Evidence Heading")

    def test_rollback_failure_preserves_backups_through_engine_cleanup(self):
        source, manifest, output, original_doc = self.create_engine_fixture("directory_tree", schema=1)
        output.mkdir()
        expected = ProjectConfig(json.loads(manifest.read_text()), source).documents
        for spec in expected:
            (output / spec["filename"]).write_bytes(f"old {spec['id']}".encode())
        replace, calls = os.replace, []
        count = 0

        def fail_after_first(source, target):
            nonlocal count
            count += 1
            if count >= 2:
                raise OSError("simulated filesystem failure including rollback")
            return replace(source, target)

        with patch("lib.delivery.inspect_document", side_effect=self.flat_engine_measurement(calls)), \
                patch("lib.engine.os.replace", side_effect=fail_after_first):
            with self.assertRaisesRegex(BuildError, "回滚未完成") as caught:
                UnifiedSynthesizer.synthesize(source, output, manifest)
        self.assertTrue(caught.exception.preserve_work)
        run_dirs = list((output / ".work").glob("run-*"))
        self.assertEqual(len(run_dirs), 1)
        backups = run_dirs[0] / "publish-backups"
        for spec in expected:
            self.assertEqual((backups / spec["filename"]).read_bytes(), f"old {spec['id']}".encode())
        self.assertIn(str(backups), str(caught.exception))


if __name__ == "__main__":
    unittest.main()
