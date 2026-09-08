# -*- coding: utf-8 -*-
"""R0 验收契约回归。

这些测试把 F01-F10 的最小反例固定为实现必须满足的断言。标记为
``integration without Word`` 的测试会替换 Word 边界或只调用离线校验器，
绝不计入真实 Word 验收。真实 Word/PDF 边界由独立的 opt-in 测试组执行，
并在验收状态文件中单独记录，不以离线测试替代。
"""

import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt
from PIL import Image

from lib.composition import add_bookmark, assemble_document, mark_start
from lib.config import ProjectConfig
from lib.content_integrity import (
    ContentIntegrityError,
    extract_semantic_inventory,
    verify_content_integrity,
    verify_delivery_format,
)
from lib.document_parts import (
    InvalidBoundaryError,
    SelectionValidator,
    SourceRegion,
    UnassignedContentError,
)
from lib.engine import BuildError, UnifiedSynthesizer
from lib.docx_inspector import inspect_docx, NodeRef
from lib.role_mapper import RoleAssignment, serialize_role_map
from lib.field_updater import FieldUpdater
from lib.format_schema import (
    HeaderFooterSpec,
    PageSpec,
    ParagraphStyle,
    ResolvedFormat,
    RoleSpec,
    RunStyle,
    StyleDefinition,
    TocSpec,
)
from lib.pagination import validate_document_structure
from lib.package_importer import validate_relationship_closure
from lib.qa import OfficeExportError


REPO = Path(__file__).resolve().parents[1]
CONTRACTS_PATH = REPO / "tests" / "fixtures" / "formatting" / "r0_contracts.json"


def _load_contracts():
    return json.loads(CONTRACTS_PATH.read_text(encoding="utf-8"))


def _recipe_hash(case_id, fixture, recipe_version):
    payload = {
        "case_id": case_id,
        "fixture": fixture,
        "recipe_version": recipe_version,
    }
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _file_sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _base_manifest(source_name, *, strategy="docx_document", parts=None):
    return {
        "schema_version": 3,
        "project_name": "R0ContractFixture",
        "source": {"strategy": strategy, "file": source_name},
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": {"mode": "restyle"},
        "output": {
            "documents": [{
                "id": "main",
                "filename": "r0-output.docx",
                "parts": parts or ["body"],
            }]
        },
    }


class _CapturedBody(Exception):
    pass


def _capture_body(source, manifest, root, *, output_name="captured.docx"):
    """Capture the staging DOCX while replacing only the delivery boundary."""
    captured = {}

    def capture(config, body, nodes, run_dir, **kwargs):
        target = root / output_name
        shutil.copy2(body, target)
        captured["path"] = target
        captured["nodes"] = nodes
        raise _CapturedBody()

    with patch("lib.delivery.build_deliveries", side_effect=capture):
        try:
            UnifiedSynthesizer.synthesize(source, root / "output", manifest)
        except _CapturedBody:
            pass
    return captured


def _custom_format(*, body_style="style_body", page=None):
    return ResolvedFormat(
        id="r0-custom",
        version="1.0.0",
        page=page or PageSpec(),
        styles={
            body_style: StyleDefinition(
                name="R0 Body",
                run=RunStyle(east_asia="宋体", latin="Times New Roman", size_pt=12.0),
                paragraph=ParagraphStyle(alignment="justify"),
            )
        },
        roles={"body": RoleSpec(style=body_style)},
        toc=TocSpec(),
        header=HeaderFooterSpec(),
        footer=HeaderFooterSpec(),
        required_capabilities=["styles.fonts.basic"],
        provenance={},
        content_hash="r0-contract",
    )


class R0ContractInventoryTest(unittest.TestCase):
    def test_contract_inventory_has_all_b_cases_and_word_followups(self):
        data = _load_contracts()
        self.assertEqual(data["schema_version"], 1)
        self.assertEqual({case["id"] for case in data["cases"]}, {f"F{i:02d}" for i in range(1, 11)})
        self.assertEqual({case["id"] for case in data["word_followup"]}, {"W04", "W05", "W06"})

        for case in data["cases"]:
            with self.subTest(case=case["id"]):
                self.assertEqual(
                    case["recipe_sha256"],
                    _recipe_hash(case["id"], case["fixture"], case["recipe_version"]),
                )
                self.assertIn("contract", case)
                self.assertIn("test_kind", case)
                self.assertIn("expected_state", case)

        for case in data["word_followup"]:
            with self.subTest(case=case["id"]):
                source = REPO / case["path"]
                self.assertTrue(source.is_file(), source)
                self.assertEqual(_file_sha256(source), case["input_sha256"])
                self.assertEqual(case["test_kind"], "integration with Word")

    def test_plan_contract_is_read_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            manifest_path = root / "manifest.json"
            doc = Document()
            doc.add_paragraph("R0 fictional source paragraph")
            doc.save(source)
            manifest_path.write_text(
                json.dumps(_base_manifest(source.name), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            before = {
                path.relative_to(root).as_posix(): _file_sha256(path)
                for path in root.rglob("*")
                if path.is_file()
            }
            plan = UnifiedSynthesizer.plan(source, manifest_path)
            after = {
                path.relative_to(root).as_posix(): _file_sha256(path)
                for path in root.rglob("*")
                if path.is_file()
            }
            self.assertEqual(before, after)
            serializable_plan = json.dumps(plan, ensure_ascii=False, default=str)
            self.assertNotIn("R0 fictional source paragraph", serializable_plan)


class R0KnownDefectContractTest(unittest.TestCase):
    """Known negative contracts; each expected failure is tied to an R batch."""

    def test_F01_region_import_preserves_image_relationship_and_style(self):
        """F01 integration without Word: 选区必须导入关系闭包和自定义样式。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / "red.png"
            Image.new("RGB", (24, 24), "red").save(image_path)

            source = Document()
            custom = source.styles.add_style("R0CustomStyle", WD_STYLE_TYPE.PARAGRAPH)
            custom.font.size = Pt(13)
            paragraph = source.add_paragraph("Region with image", style=custom)
            paragraph.add_run().add_picture(str(image_path), width=Mm(10))

            sliced = SelectionValidator.slice_document_by_region(source, (0, 1))
            source_path = root / "source.docx"
            target_path = root / "slice.docx"
            source.save(source_path)
            sliced.save(target_path)

            source_inventory = extract_semantic_inventory(source_path)
            target_inventory = extract_semantic_inventory(target_path)
            self.assertEqual(source_inventory.media_hashes, target_inventory.media_hashes)
            self.assertIn("R0CustomStyle", {style.style_id for style in sliced.styles})
            validate_relationship_closure(sliced)
            validate_document_structure(target_path)

    def test_F02_corrupt_format_rejects_publication_and_preserves_previous_result(self):
        """F02 integration without Word: 12pt→70pt 不得发布。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source_doc = Document()
            source_doc.add_paragraph("F02 body")
            source_doc.save(source)
            manifest_path = root / "manifest.json"
            manifest_path.write_text(json.dumps(_base_manifest(source.name)), encoding="utf-8")
            output = root / "output"
            output.mkdir()
            old_delivery = output / "r0-output.docx"
            old_meta = output / "build-metadata.json"
            Document().save(old_delivery)
            old_meta.write_text('{"status":"previously-published"}', encoding="utf-8")
            old_hashes = (old_delivery.read_bytes(), old_meta.read_bytes())

            def corrupt_stage(config, body, nodes, run_dir, **kwargs):
                target = Path(run_dir) / config.documents[0]["filename"]
                corrupted = Document(body)
                corrupted.paragraphs[0].runs[0].font.size = Pt(70)
                corrupted.save(target)
                return {"main": target}

            with patch("lib.delivery.build_deliveries", side_effect=corrupt_stage):
                with self.assertRaises(BuildError):
                    UnifiedSynthesizer.synthesize(source, output, manifest_path)

            self.assertEqual((old_delivery.read_bytes(), old_meta.read_bytes()), old_hashes)

    def test_F03_stale_explicit_role_map_is_rejected_before_delivery(self):
        """F03 integration without Word: 过期映射必须在运行目录前拒绝。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.add_paragraph("Mapped heading")
            doc.save(source)
            role_map = root / "stale-map.json"
            inspection = inspect_docx(source)
            block = inspection.blocks[0]
            stale_ref = NodeRef(
                source_sha256="0" * 64,
                part_uri=block.node.part_uri,
                element_path=block.node.element_path,
                text_hash=block.node.text_hash,
            )
            stale_map = serialize_role_map(
                [RoleAssignment(
                    node_ref=stale_ref,
                    role="heading.2",
                    level=2,
                    text_hash=block.node.text_hash,
                )],
                source_file=source.name,
                on_unmapped="preserve",
            )
            role_map.write_text(json.dumps(stale_map), encoding="utf-8")
            manifest = _base_manifest(source.name)
            manifest["formatting"]["role_map"] = role_map.name
            manifest_path = root / "manifest.json"
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

            def stage(config, body, nodes, run_dir, **kwargs):
                target = Path(run_dir) / config.documents[0]["filename"]
                shutil.copy2(body, target)
                return {"main": target}

            with patch("lib.delivery.build_deliveries", side_effect=stage):
                with self.assertRaises(BuildError):
                    UnifiedSynthesizer.synthesize(source, root / "output", manifest_path)

    def test_F04_multi_source_override_reaches_rendered_heading(self):
        """F04 integration without Word: 多文件声明的 37pt 标题必须实际生效。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_dir = root / "sources"
            source_dir.mkdir()
            source = source_dir / "one.docx"
            doc = Document()
            doc.add_paragraph("Multi-source heading")
            doc.add_paragraph("Fictional supporting paragraph")
            doc.save(source)
            manifest = _base_manifest(source.name, strategy="directory_tree")
            manifest["source"] = {"strategy": "directory_tree"}
            manifest["format"]["overrides"] = {"styles": {"heading.1": {"run": {"size_pt": 37}}}}
            manifest_path = root / "manifest.json"
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

            captured = _capture_body(source_dir, manifest_path, root)
            self.assertAlmostEqual(Document(captured["path"]).paragraphs[0].runs[0].font.size.pt, 37.0)

    def test_F04_preserve_keeps_source_geometry_and_footer(self):
        """F04 integration without Word: preserve/source 不得覆盖页面和页脚。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            section = doc.sections[0]
            section.page_width = Mm(180)
            section.page_height = Mm(240)
            section.bottom_margin = Mm(30)
            section.footer.paragraphs[0].text = "R0 ORIGINAL FOOTER"
            doc.add_paragraph("Preserved source paragraph")
            doc.save(source)
            manifest = _base_manifest(source.name)
            manifest["formatting"] = {"mode": "preserve", "page_policy": "source"}
            manifest_path = root / "manifest.json"
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

            captured = _capture_body(source, manifest_path, root)
            result = Document(captured["path"])
            self.assertAlmostEqual(result.sections[0].page_width.mm, 180.0, delta=0.1)
            self.assertAlmostEqual(result.sections[0].page_height.mm, 240.0, delta=0.1)
            self.assertAlmostEqual(result.sections[0].bottom_margin.mm, 30.0, delta=0.1)
            self.assertEqual(result.sections[0].footer.paragraphs[0].text, "R0 ORIGINAL FOOTER")

    def test_F05_bad_observed_page_label_rejects_delivery(self):
        """F05 validator only; synthetic pages are not Word acceptance evidence."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            raw = _base_manifest("source.docx")
            raw["layout"] = {
                "page_sequences": {"front": {"format": "upperRoman", "start": 1}},
                "parts": {"body": {"kind": "content", "page_sequence": "front"}},
            }
            cfg = ProjectConfig(raw, root)
            doc = Document()
            doc.add_paragraph("Synthetic page label body")
            mark_start(doc, "body")
            docx_path = root / "delivery.docx"
            pdf_path = root / "delivery.pdf"
            doc.save(docx_path)
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((60, 70), "Synthetic page label body")
                pdf.save(pdf_path)
            page_map = {
                "_Synth_body": {
                    "physical_page": 1,
                    "printed_page": 7,
                    "observed_label": "999",
                }
            }
            with patch("lib.delivery.run_qa_assertions", return_value=True):
                with self.assertRaises(OfficeExportError):
                    from lib.delivery import validate_measured_delivery
                    validate_measured_delivery(
                        docx_path,
                        cfg.documents[0],
                        [],
                        pdf_path,
                        page_map,
                        parts_registry=cfg.parts,
                        page_sequences=cfg.page_sequences,
                    )

    def test_F06_seq_roman_and_alphabetic_parameters_are_supported(self):
        """F06 integration without Word: SEQ 格式参数不得触发类型错误。"""
        updater = FieldUpdater()
        self.assertEqual(updater._evaluate_instruction(r"SEQ Figure \\* ROMAN"), "I")
        self.assertEqual(updater._evaluate_instruction(r"SEQ Figure \\* ALPHABETIC"), "B")

    def test_F06_ref_uses_only_bookmarked_range(self):
        """F06 integration without Word: REF 不得吞入书签范围外的正文。"""
        doc = Document()
        paragraph = doc.add_paragraph("Figure ")
        paragraph._p.append(parse_xml(
            f'<w:bookmarkStart {nsdecls("w")} w:id="9" w:name="fig_number"/>'
        ))
        paragraph._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="SEQ Figure">'
            f'<w:r><w:t>0</w:t></w:r></w:fldSimple>'
        ))
        paragraph._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="9"/>'))
        paragraph.add_run(" Caption outside bookmark")
        ref = doc.add_paragraph()
        ref._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="REF fig_number">'
            f'<w:r><w:t>old</w:t></w:r></w:fldSimple>'
        ))
        FieldUpdater().update_document_fields(doc)
        self.assertEqual("".join(t.text or "" for t in ref._p.iter(qn("w:t"))), "1")

    def test_F07_unassigned_image_is_not_silently_accepted(self):
        """F07 integration without Word: 未分配图片必须进入选区完整性诊断。"""
        with tempfile.TemporaryDirectory() as tmp:
            image = Path(tmp) / "image.png"
            Image.new("RGB", (20, 20), "blue").save(image)
            doc = Document()
            doc.add_picture(str(image), width=Mm(10))
            doc.add_paragraph("Selected text")
            with self.assertRaises(UnassignedContentError):
                SelectionValidator.validate_regions(doc, {
                    "text": SourceRegion("text", "/w:document/w:body/w:p[2]", "end_of_document")
                })

    def test_F07_invalid_region_does_not_fall_back_to_full_document(self):
        """F07 integration without Word: 无效边界必须在装配前报错。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.add_paragraph("Only source content")
            doc.save(source)
            raw = _base_manifest(source.name, parts=["main_text"])
            raw["layout"] = {"parts": {"main_text": {"kind": "content", "source_region": "bad"}}}
            raw["source"]["regions"] = {
                "bad": {"start": "/w:document/w:body/w:p[99]", "end": "end_of_document"}
            }
            cfg = ProjectConfig(raw, root)
            with self.assertRaises(InvalidBoundaryError):
                assemble_document(cfg, cfg.documents[0], source, [], {}, root / "output.docx")

    def test_F08_removed_footnote_is_detected_by_content_integrity(self):
        """F08 integration without Word: 脚注引用和定义都属于内容清单。"""
        doc = Document()
        paragraph = doc.add_paragraph("Body with note")
        fn_ref = parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="1"/>')
        paragraph.add_run()._r.append(fn_ref)
        note_xml = (
            f'<w:footnotes {nsdecls("w")}><w:footnote w:id="1">'
            f'<w:p><w:r><w:footnoteRef/><w:t>Required note text</w:t></w:r></w:p>'
            f'</w:footnote></w:footnotes>'
        )
        note_part = Part(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            note_xml.encode("utf-8"),
            doc.part.package,
        )
        note_rid = doc.part.relate_to(note_part, RT.FOOTNOTES)
        inventory = extract_semantic_inventory(doc)
        fn_ref.getparent().remove(fn_ref)
        doc.part.drop_rel(note_rid)
        with self.assertRaises(ContentIntegrityError):
            verify_content_integrity(inventory, doc)

    def test_F09_role_style_alias_is_resolved_for_format_verification(self):
        """F09 integration without Word: roles.style 可引用不同名的样式 ID。"""
        resolved = _custom_format()
        doc = Document()
        from lib.style_applier import apply_roles
        section = doc.sections[0]
        section.page_width = Mm(resolved.page.width_mm)
        section.page_height = Mm(resolved.page.height_mm)
        section.top_margin = Mm(resolved.page.margin_top_mm)
        section.bottom_margin = Mm(resolved.page.margin_bottom_mm)
        section.left_margin = Mm(resolved.page.margin_left_mm)
        section.right_margin = Mm(resolved.page.margin_right_mm)
        doc.add_paragraph("Alias style body")
        apply_roles(doc, {"/w:document/w:body/w:p[1]": "body"}, resolved)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "alias.docx"
            doc.save(path)
            report = verify_delivery_format(path, resolved)
            self.assertTrue(report.passed)

    def test_F10_assembly_and_word_layers_are_explicitly_separated(self):
        """F10 test-layer contract: 装配测试和真实 Word 测试不能互相冒充。"""
        thesis_source = (REPO / "tests" / "test_thesis_acceptance.py").read_text(encoding="utf-8")
        word_source_path = REPO / "tests" / "test_word_formatting.py"
        word_source = word_source_path.read_text(encoding="utf-8")
        self.assertIn("装配级", thesis_source)
        self.assertIn("integration with Word", word_source)
        self.assertIn("DOCUMENT_SYNTHESIS_WORD_TEST", word_source)


if __name__ == "__main__":
    unittest.main()
