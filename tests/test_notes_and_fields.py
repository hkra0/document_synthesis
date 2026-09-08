# -*- coding: utf-8 -*-
"""
脚注、尾注与受控字段单元测试套件 (tests/test_notes_and_fields.py)
验证 M2 能力: 'notes.merge.v1' 与 'fields.managed_update.v1'
"""

import tempfile
from pathlib import Path
import unittest

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from lib.delivery import validate_delivery_structure
from lib.document_parts import SelectionValidator
from lib.field_updater import FieldUpdater
from lib.notes_merger import NotesMerger
from lib.package_importer import PackageImporter, validate_relationship_closure
from lib.qa import OfficeExportError


def _create_doc_with_footnote(text: str, fn_id: int, fn_text: str) -> Document:
    doc = Document()
    fn_xml = f"""<w:footnotes {nsdecls('w')}>
      <w:footnote w:type="separator" w:id="-1">
        <w:p><w:r><w:separator/></w:r></w:p>
      </w:footnote>
      <w:footnote w:type="continuationSeparator" w:id="0">
        <w:p><w:r><w:continuationSeparator/></w:r></w:p>
      </w:footnote>
      <w:footnote w:id="{fn_id}">
        <w:p><w:r><w:footnoteRef/><w:t>{fn_text}</w:t></w:r></w:p>
      </w:footnote>
    </w:footnotes>"""
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        fn_xml.encode("utf-8"),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.FOOTNOTES)
    p = doc.add_paragraph(text)
    r = p.add_run()
    r._r.append(parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="{fn_id}"/>'))
    return doc


class NotesMergerTest(unittest.TestCase):
    """测试跨文档脚注与尾注合并"""

    def test_merge_two_docs_with_footnotes(self):
        doc1 = _create_doc_with_footnote("文档一正文", 1, "文档一脚注")
        doc2 = _create_doc_with_footnote("文档二正文", 1, "文档二脚注")

        importer = PackageImporter(source=doc2, target=doc1)
        # 导入 doc2 的正文段落
        for p in doc2.element.body.iter(qn("w:p")):
            doc1.element.body.append(importer.element(p))

        # 检查 doc1 目标脚注部件
        fn_part = doc1.part.part_related_by(RT.FOOTNOTES)
        root = parse_xml(fn_part.blob)
        footnotes = root.findall(qn("w:footnote"))

        # 应该有 4 个条目: separator(-1), continuationSeparator(0), id=1, id=2
        self.assertEqual(len(footnotes), 4)
        ids = {fn.get(qn("w:id")) for fn in footnotes}
        self.assertIn("-1", ids)
        self.assertIn("0", ids)
        self.assertIn("1", ids)
        self.assertIn("2", ids)

        # 检查正文中的引用已被重映射
        refs = [r.get(qn("w:id")) for r in doc1.element.body.iter(qn("w:footnoteReference"))]
        self.assertEqual(refs, ["1", "2"])

    def test_footnote_internal_relationship_cloned(self):
        doc1 = Document()
        doc2 = Document()
        part2 = Part(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            b"",
            doc2.part.package,
        )
        new_rid = part2.relate_to(
            "https://deepmind.google",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        fn_xml = f"""<w:footnotes {nsdecls('w')}>
          <w:footnote w:id="1">
            <w:p><w:hyperlink r:id="{new_rid}" {nsdecls('r')}><w:r><w:t>LinkInFootnote</w:t></w:r></w:hyperlink></w:p>
          </w:footnote>
        </w:footnotes>"""
        part2._blob = fn_xml.encode("utf-8")
        doc2.part.relate_to(part2, RT.FOOTNOTES)

        mapping = NotesMerger.merge_notes(doc2, doc1, is_endnote=False)
        self.assertEqual(mapping, {"1": "1"})

        target_part = doc1.part.part_related_by(RT.FOOTNOTES)
        target_root = parse_xml(target_part.blob)
        hyperlinks = list(target_root.iter(qn("w:hyperlink")))
        self.assertEqual(len(hyperlinks), 1)
        cloned_rid = hyperlinks[0].get(qn("r:id"))
        self.assertIsNotNone(cloned_rid)
        self.assertIn(cloned_rid, target_part.rels)
        self.assertEqual(target_part.rels[cloned_rid].target_ref, "https://deepmind.google")

    def test_region_import_clones_only_referenced_notes(self):
        source = Document()
        first = source.add_paragraph("选区内正文")
        first.add_run()._r.append(parse_xml(
            f'<w:footnoteReference {nsdecls("w")} w:id="1"/>'
        ))
        second = source.add_paragraph("选区外正文")
        second.add_run()._r.append(parse_xml(
            f'<w:footnoteReference {nsdecls("w")} w:id="2"/>'
        ))
        fn_xml = f"""<w:footnotes {nsdecls('w')}>
          <w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>
          <w:footnote w:type="continuationSeparator" w:id="0"><w:p/></w:footnote>
          <w:footnote w:id="1"><w:p><w:r><w:t>选区内脚注</w:t></w:r></w:p></w:footnote>
          <w:footnote w:id="2"><w:p><w:r><w:t>选区外脚注</w:t></w:r></w:p></w:footnote>
        </w:footnotes>"""
        part = Part(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            fn_xml.encode("utf-8"),
            source.part.package,
        )
        source.part.relate_to(part, RT.FOOTNOTES)

        sliced = SelectionValidator.slice_document_by_region(source, (0, 1))
        target_part = sliced.part.part_related_by(RT.FOOTNOTES)
        target_root = parse_xml(target_part.blob)
        target_ids = {node.get(qn("w:id")) for node in target_root.findall(qn("w:footnote"))}
        self.assertIn("1", target_ids)
        self.assertNotIn("2", target_ids)

    def test_footnote_internal_media_is_cloned_into_target_package(self):
        source = Document()
        note_part = Part(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            b"",
            source.part.package,
        )
        image_part = Part(
            PackURI("/word/media/note-image.png"),
            "image/png",
            b"not-a-real-png-but-a-package-payload",
            source.part.package,
        )
        image_rid = note_part.relate_to(image_part, RT.IMAGE)
        note_xml = f"""<w:footnotes {nsdecls('w')} {nsdecls('a')} {nsdecls('r')}>
          <w:footnote w:id="1"><w:p><w:r><w:drawing>
            <a:blip r:embed="{image_rid}"/>
          </w:drawing></w:r></w:p></w:footnote>
        </w:footnotes>"""
        note_part._blob = note_xml.encode("utf-8")
        source.part.relate_to(note_part, RT.FOOTNOTES)

        target = Document()
        NotesMerger.merge_notes(source, target, is_endnote=False, strict_rels=True)
        target_note_part = target.part.part_related_by(RT.FOOTNOTES)
        target_root = parse_xml(target_note_part.blob)
        # Read the relationship by looking up the actual namespaced attribute.
        blip = next(target_root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
        target_rid = blip.get(qn("r:embed"))
        self.assertIsNotNone(target_rid)
        self.assertIn(target_rid, target_note_part.rels)
        self.assertEqual(target_note_part.rels[target_rid].target_part.content_type, "image/png")
        validate_relationship_closure(target)


class FieldUpdaterTest(unittest.TestCase):
    """测试受控字段状态机更新与数学公式保留"""

    def test_seq_figure_and_table_increment(self):
        doc = Document()
        # 添加图 1, 图 2
        for i in range(2):
            p = doc.add_paragraph("图 ")
            r_begin = p.add_run()
            r_begin._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
            r_instr = p.add_run()
            r_instr._r.append(parse_xml(rf'<w:instrText {nsdecls("w")} xml:space="preserve"> SEQ Figure \* ARABIC </w:instrText>'))
            r_sep = p.add_run()
            r_sep._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
            r_res = p.add_run("0")
            r_end = p.add_run()
            r_end._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
            p.add_run(" 架构图示")

        # 添加表 1 (独立计数)
        p_tbl = doc.add_paragraph("表 ")
        fld_simple = parse_xml(rf'<w:fldSimple {nsdecls("w")} w:instr="SEQ Table \* ARABIC"><w:r><w:t>0</w:t></w:r></w:fldSimple>')
        p_tbl._p.append(fld_simple)
        p_tbl.add_run(" 性能对比表")

        updater = FieldUpdater()
        updater.update_document_fields(doc)

        self.assertEqual(doc.paragraphs[0].text, "图 1 架构图示")
        self.assertEqual(doc.paragraphs[1].text, "图 2 架构图示")
        tbl_text = "".join(t.text or "" for t in doc.paragraphs[2]._p.iter(qn("w:t")))
        self.assertEqual(tbl_text, "表 1 性能对比表")

    def test_pageref_resolution_from_page_map(self):
        doc = Document()
        p = doc.add_paragraph("详见附录第 ")
        r1 = p.add_run()
        r1._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
        r2 = p.add_run()
        r2._r.append(parse_xml(f'<w:instrText {nsdecls("w")}> PAGEREF appendix_sec </w:instrText>'))
        r3 = p.add_run()
        r3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
        r4 = p.add_run("1")
        r5 = p.add_run()
        r5._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
        p.add_run(" 页")

        page_map = {"appendix_sec": {"physical_page": 10, "printed_page": 8, "expected_label": "8"}}
        updater = FieldUpdater(page_map=page_map)
        updater.update_document_fields(doc)

        self.assertEqual(p.text, "详见附录第 8 页")

    def test_omml_math_preserved_intact(self):
        doc = Document()
        p = doc.add_paragraph("公式测试: ")
        # 构造一个原生 OMML 公式
        omml_elem = parse_xml(f'''<m:oMath {nsdecls("m")}>
            <m:r><m:t>E=mc^2</m:t></m:r>
        </m:oMath>''')
        p._p.append(omml_elem)

        updater = FieldUpdater()
        updater.update_document_fields(doc)

        found_math = list(doc.element.body.iter(qn("m:oMath")))
        self.assertEqual(len(found_math), 1)
        self.assertEqual(found_math[0].find(qn("m:r")).find(qn("m:t")).text, "E=mc^2")


class DeliveryDanglingNotesTest(unittest.TestCase):
    """测试交付结构检查对悬空未定义脚注的拦截"""

    def test_dangling_footnote_reference_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            doc = Document()
            p = doc.add_paragraph("带有悬空未定义脚注的正文")
            r = p.add_run()
            r._r.append(parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="999"/>'))
            # 建立一个只有 id=1 的脚注部件
            fn_xml = f"""<w:footnotes {nsdecls('w')}>
              <w:footnote w:id="1"><w:p><w:t>Only 1</w:t></w:p></w:footnote>
            </w:footnotes>"""
            part = Part(
                PackURI("/word/footnotes.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                fn_xml.encode("utf-8"),
                doc.part.package,
            )
            doc.part.relate_to(part, RT.FOOTNOTES)
            
            p_first = doc.paragraphs[0]._p
            b_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="10" w:name="_Synth_body"/>')
            p_first.insert(0, b_start)

            doc_path = Path(td) / "dangling.docx"
            doc.save(str(doc_path))

            spec = {"id": "main", "parts": ["body"], "filename": "main.docx"}
            nodes = []

            with self.assertRaises(OfficeExportError) as ctx:
                validate_delivery_structure(doc_path, spec, nodes)
            self.assertIn("未定义的脚注引用", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
