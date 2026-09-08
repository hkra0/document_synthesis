# -*- coding: utf-8 -*-
"""
排版接管模式与边界隔离单元测试 (tests/test_formatting_modes.py)
"""

import tempfile
import unittest
from pathlib import Path
import pymupdf

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from lib.layout import isolate_section_boundaries, ContentBox, compute_content_box, apply_section_spec
from lib.renderers import get_cached_pdf_page_image, set_pdf_render_cache_dir
from lib.format_resolver import resolve_format_package
from lib.format_schema import FormattingPolicy, FormatDiagnosticCode
from lib.style_applier import apply_roles


class FormattingModesTest(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)

    def test_section_boundary_isolation_preserves_independent_headers(self):
        """测试保留区与重排区跨节时，页眉页脚完全隔离不互相污染"""
        doc = Document()
        doc.sections[0].header.is_linked_to_previous = False
        doc.sections[0].header.paragraphs[0].text = "保留封面/前置页眉"
        doc.sections[0].footer.is_linked_to_previous = False
        doc.sections[0].footer.paragraphs[0].text = "前置页脚 - 罗马数字"

        # 添加新节（主文重排区）
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        isolate_section_boundaries(doc.sections[1])

        # 在新节设置不同页眉页脚
        doc.sections[1].header.paragraphs[0].text = "主文统一页眉"
        doc.sections[1].footer.paragraphs[0].text = "正文页码 - 阿拉伯数字"

        # 断言两节独立互不影响
        self.assertFalse(doc.sections[1].header.is_linked_to_previous)
        self.assertFalse(doc.sections[1].footer.is_linked_to_previous)
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "保留封面/前置页眉")
        self.assertEqual(doc.sections[1].header.paragraphs[0].text, "主文统一页眉")
        self.assertEqual(doc.sections[0].footer.paragraphs[0].text, "前置页脚 - 罗马数字")
        self.assertEqual(doc.sections[1].footer.paragraphs[0].text, "正文页码 - 阿拉伯数字")

    def test_pdf_cache_key_varies_with_sanitization_policy(self):
        """测试 PDF 渲染缓存键随清洗策略（去页码、裁边距）变化，杜绝策略更改后复用错误图像"""
        # 生成带页码的单页测试 PDF
        pdf_path = self.root / "sample.pdf"
        pdf_doc = pymupdf.open()
        page = pdf_doc.new_page(width=300, height=400)
        page.insert_text((50, 100), "Hello Test Page", fontsize=16)
        page.insert_text((140, 380), "1", fontsize=10)  # 底部页码
        pdf_doc.save(str(pdf_path))
        pdf_doc.close()

        cache_dir = self.root / "cache"
        set_pdf_render_cache_dir(cache_dir)

        # 1. 默认清洗页脚与裁切
        img_sanitized = get_cached_pdf_page_image(
            pdf_path, 0, dpi=150, sanitize_footer=True, crop_whitespace=True
        )

        # 2. 保留页脚（不抹除）
        img_unsanitized = get_cached_pdf_page_image(
            pdf_path, 0, dpi=150, sanitize_footer=False, crop_whitespace=True
        )

        # 3. 不裁切边缘
        img_uncropped = get_cached_pdf_page_image(
            pdf_path, 0, dpi=150, sanitize_footer=True, crop_whitespace=False
        )

        cached_files = list(cache_dir.glob("*.png"))
        # 不同的策略参数必须生成独立的缓存文件
        self.assertEqual(len(cached_files), 3, "三种不同清洗策略必须生成 3 个独立缓存条目")
        self.assertNotEqual(img_sanitized, img_unsanitized)
        self.assertNotEqual(img_sanitized, img_uncropped)

    def test_unsupported_capability_in_format_package_diagnosed(self):
        """测试格式包声明当前不支持的高阶能力时，在解析阶段给出明确诊断而非静默吞下"""
        import json
        from lib.config import ConfigError
        package_dict = {
            "format_schema_version": 1,
            "id": "future-academic",
            "version": "1.0.0",
            "required_capabilities": ["pagination.roman", "notes.footnotes.merge"],
            "page": {
                "width_mm": 210.0,
                "height_mm": 297.0,
                "orientation": "portrait",
                "margin_top_mm": 25.0,
                "margin_bottom_mm": 25.0,
                "margin_left_mm": 25.0,
                "margin_right_mm": 25.0
            },
            "styles": {
                "body": {"name": "正文", "run": {"size_pt": 12.0}}
            }
        }
        pkg_file = self.root / "future-academic.json"
        pkg_file.write_text(json.dumps(package_dict, ensure_ascii=False), encoding="utf-8")
        with self.assertRaises(ConfigError) as cm:
            resolve_format_package(str(pkg_file), declared_in=self.root)
        self.assertIn(FormatDiagnosticCode.UNSUPPORTED_CAPABILITY, str(cm.exception))

    def test_apply_section_spec_lands_page_geometry_and_grid(self):
        """页面规格必须真正写入节，而不是只留在配置对象里。"""
        from lib.format_schema import PageSpec

        doc = Document()
        spec = PageSpec(
            width_mm=180.0,
            height_mm=240.0,
            orientation="landscape",
            margin_top_mm=11.0,
            margin_bottom_mm=12.0,
            margin_left_mm=13.0,
            margin_right_mm=14.0,
            header_distance_mm=7.0,
            footer_distance_mm=8.0,
            snap_to_grid=True,
        )
        apply_section_spec(doc.sections[0], spec)
        section = doc.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 240.0, delta=0.1)
        self.assertAlmostEqual(section.page_height.mm, 180.0, delta=0.1)
        self.assertAlmostEqual(section.left_margin.mm, 13.0, delta=0.1)
        self.assertAlmostEqual(section.header_distance.mm, 7.0, delta=0.1)
        self.assertIsNotNone(section._sectPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}docGrid"))

    def test_preserve_mode_does_not_install_or_mutate_source_formatting(self):
        """preserve 不得写入 Synth pStyle 或清理源直接格式。"""
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("保留原格式")
        run.font.size = Pt(19)
        before = paragraph._p.xml
        resolved = resolve_format_package("preset:academic-basic@1.0.0")
        report = apply_roles(
            doc,
            {"0": "heading.1"},
            resolved,
            policy=FormattingPolicy(mode="preserve"),
        )
        self.assertEqual(report.styled_paragraphs_count, 0)
        self.assertEqual(paragraph._p.xml, before)
        self.assertIsNone(doc.part.styles._element.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}style[@{http://schemas.openxmlformats.org/wordprocessingml/2006/main}styleId='SynthBody']"))

    def test_mixed_unmapped_paragraph_is_left_untouched(self):
        """mixed 的未托管段落由 on_unmapped=preserve 保护。"""
        doc = Document()
        managed = doc.add_paragraph("托管标题")
        preserved = doc.add_paragraph("源格式段落")
        preserved.alignment = WD_ALIGN_PARAGRAPH.CENTER
        before = preserved._p.xml
        resolved = resolve_format_package("preset:academic-basic@1.0.0")
        apply_roles(
            doc,
            {"0": "heading.1"},
            resolved,
            policy=FormattingPolicy(mode="mixed", on_unmapped="preserve", page_policy="target"),
        )
        self.assertIsNotNone(managed._p.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle"))
        self.assertEqual(preserved._p.xml, before)


if __name__ == "__main__":
    unittest.main()
