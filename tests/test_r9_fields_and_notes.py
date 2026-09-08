# -*- coding: utf-8 -*-
"""R9 离线字段依赖、书签范围与脚注/尾注映射契约。"""

import unittest

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from lib.field_updater import FieldUpdateError, FieldUpdater
from lib.notes_merger import NotesMerger
from lib.pagination_types import PageRecord
from lib.package_importer import PackageImporter, validate_relationship_closure


class R9FieldDependencyTest(unittest.TestCase):
    def test_seq_switches_case_reset_and_continue(self):
        updater = FieldUpdater()
        self.assertEqual(updater._evaluate_instruction(r"SEQ Figure \\* ROMAN"), "I")
        self.assertEqual(updater._evaluate_instruction(r"SEQ Figure \\* roman"), "ii")
        self.assertEqual(updater._evaluate_instruction(r"SEQ Figure \\r 4 \\* ALPHABETIC"), "D")
        self.assertEqual(updater._evaluate_instruction(r"SEQ Figure \\c \\* alphabetic"), "d")
        self.assertEqual(updater._evaluate_instruction(r"SEQ Figure \\* ARABIC"), "5")
        with self.assertRaises(FieldUpdateError) as ctx:
            updater._evaluate_instruction(r"SEQ Figure \\r 0")
        self.assertEqual(ctx.exception.code, "INVALID_SEQ_RESET")

    def test_forward_ref_and_cross_paragraph_bookmark_use_exact_range(self):
        doc = Document()
        ref = doc.add_paragraph("前置引用 ")
        ref._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="REF future_caption">'
            '<w:r><w:t>old</w:t></w:r></w:fldSimple>'
        ))

        first = doc.add_paragraph("段落外前缀")
        first._p.append(parse_xml(
            f'<w:bookmarkStart {nsdecls("w")} w:id="17" w:name="future_caption"/>'
        ))
        first._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="SEQ Figure">'
            '<w:r><w:t>old</w:t></w:r></w:fldSimple>'
        ))
        second = doc.add_paragraph("书签内后半段")
        second._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="17"/>'))
        second.add_run("段落外后缀")

        FieldUpdater().update_document_fields(doc)
        ref_text = "".join(node.text or "" for node in ref._p.iter(qn("w:t")))
        first_text = "".join(node.text or "" for node in first._p.iter(qn("w:t")))
        second_text = "".join(node.text or "" for node in second._p.iter(qn("w:t")))
        self.assertEqual(ref_text, "前置引用 1书签内后半段")
        self.assertEqual(first_text, "段落外前缀1")
        self.assertEqual(second_text, "书签内后半段段落外后缀")

    def test_complex_forward_ref_chain_converges_within_bound(self):
        doc = Document()

        def add_complex_ref(paragraph, target, cached="old"):
            begin = paragraph.add_run()
            begin._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
            instr = paragraph.add_run()
            instr._r.append(parse_xml(
                f'<w:instrText {nsdecls("w")} xml:space="preserve"> REF {target} </w:instrText>'
            ))
            separate = paragraph.add_run()
            separate._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
            paragraph.add_run(cached)
            end = paragraph.add_run()
            end._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))

        first = doc.add_paragraph()
        first._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="31" w:name="a"/>'))
        add_complex_ref(first, "b")
        first._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="31"/>'))

        second = doc.add_paragraph()
        second._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="32" w:name="b"/>'))
        add_complex_ref(second, "c")
        second._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="32"/>'))

        third = doc.add_paragraph()
        third._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="33" w:name="c"/>'))
        third._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="SEQ Figure">'
            '<w:r><w:t>old</w:t></w:r></w:fldSimple>'
        ))
        third._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="33"/>'))

        FieldUpdater().update_document_fields(doc)
        self.assertEqual("".join(node.text or "" for node in first._p.iter(qn("w:t"))), "1")
        self.assertEqual("".join(node.text or "" for node in second._p.iter(qn("w:t"))), "1")

    def test_pageref_requires_typed_page_index_and_accepts_page_record(self):
        record = PageRecord(
            physical_page=3,
            section_id=2,
            sequence_id="front",
            number_value=4,
            expected_label="iv",
        )
        self.assertEqual(
            FieldUpdater({"target": record})._evaluate_instruction("PAGEREF target"),
            "iv",
        )
        with self.assertRaises(FieldUpdateError) as ctx:
            FieldUpdater({"target": "4"})._evaluate_instruction("PAGEREF target")
        self.assertEqual(ctx.exception.code, "INVALID_PAGE_INDEX")

    def test_unknown_field_cache_is_preserved(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="DOCPROPERTY Custom">'
            '<w:r><w:t>cached value</w:t></w:r></w:fldSimple>'
        ))
        FieldUpdater().update_document_fields(doc)
        self.assertEqual(
            "".join(node.text or "" for node in paragraph._p.iter(qn("w:t"))),
            "cached value",
        )

    def test_ref_cycle_is_rejected_with_structured_diagnostic(self):
        doc = Document()
        first = doc.add_paragraph()
        first._p.append(parse_xml(
            f'<w:bookmarkStart {nsdecls("w")} w:id="21" w:name="a"/>'
        ))
        first._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="REF b"><w:r><w:t>a</w:t></w:r></w:fldSimple>'
        ))
        first._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="21"/>'))
        second = doc.add_paragraph()
        second._p.append(parse_xml(
            f'<w:bookmarkStart {nsdecls("w")} w:id="22" w:name="b"/>'
        ))
        second._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="REF a"><w:r><w:t>b</w:t></w:r></w:fldSimple>'
        ))
        second._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="22"/>'))

        with self.assertRaises(FieldUpdateError) as ctx:
            FieldUpdater().update_document_fields(doc)
        self.assertEqual(ctx.exception.code, "FIELD_DEPENDENCY_CYCLE")


def _doc_with_endnote(text, note_id, note_text):
    doc = Document()
    xml = f'''<w:endnotes {nsdecls("w")}>
      <w:endnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:endnote>
      <w:endnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:endnote>
      <w:endnote w:id="{note_id}"><w:p><w:r><w:endnoteRef/><w:t>{note_text}</w:t></w:r></w:p></w:endnote>
    </w:endnotes>'''
    part = Part(
        PackURI("/word/endnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
        xml.encode("utf-8"),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.ENDNOTES)
    p = doc.add_paragraph(text)
    p.add_run()._r.append(parse_xml(
        f'<w:endnoteReference {nsdecls("w")} w:id="{note_id}"/>'
    ))
    return doc


class R9NotesTest(unittest.TestCase):
    def test_endnotes_are_remapped_and_relationship_closure_is_preserved(self):
        target = _doc_with_endnote("目标正文", 1, "目标尾注")
        source = _doc_with_endnote("源正文", 1, "源尾注")
        importer = PackageImporter(source=source, target=target, strict_rels=True)
        imported = importer.import_body_contents_with_result()
        for element in imported.elements:
            target.element.body.append(element)

        self.assertEqual(imported.endnote_map, {"1": "2"})
        refs = [
            node.get(qn("w:id"))
            for node in target.element.body.iter(qn("w:endnoteReference"))
        ]
        self.assertEqual(refs, ["1", "2"])
        part = target.part.part_related_by(RT.ENDNOTES)
        ids = {node.get(qn("w:id")) for node in parse_xml(part.blob).findall(qn("w:endnote"))}
        self.assertEqual(ids, {"-1", "0", "1", "2"})
        validate_relationship_closure(target)


if __name__ == "__main__":
    unittest.main()
