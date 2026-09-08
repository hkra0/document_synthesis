# -*- coding: utf-8 -*-
"""
排版与公文样式模块
提供标准版面配置、页脚页码、样式合并、格式内联展开、结构化目录控件与标题段落构建
"""
import io
import re
import copy
from typing import Dict, Any, Optional, List
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import lxml.etree
from xml.sax.saxutils import escape

# 公文标准版面参数
DEFAULT_PAGE_SETUP = {
    "width_cm": 21.0,
    "height_cm": 29.7,
    "margin_top_cm": 3.7,
    "margin_bottom_cm": 3.5,
    "margin_left_cm": 2.8,
    "margin_right_cm": 2.6
}

# 默认字体体系
DEFAULT_FONTS = {
    "cover_title": "方正小标宋简体",
    "cover_sub": "仿宋_GB2312",
    "toc_title": "黑体",
    "h1": "黑体",
    "h2": "楷体_GB2312",
    "h3": "仿宋_GB2312",
    "body": "仿宋_GB2312",
    "en": "Times New Roman"
}

def add_toc_entry_with_hyperlink(doc: Document, text: str, page_num: int, level: int = 1, bookmark_name: str = None, fonts: Dict[str, str] = None):
    """向文档插入带有点号制表符与书签超链接跳转的目录条目"""
    fonts = fonts or DEFAULT_FONTS
    p = doc.add_paragraph()
    pf = p.paragraph_format
    
    font_cn = fonts.get("h1", "黑体") if level == 1 else (fonts.get("h2", "楷体_GB2312") if level == 2 else fonts.get("h3", "仿宋_GB2312"))
    font_en = fonts.get("en", "Times New Roman")
    is_bold = level in (1, 2)
    
    if level == 1:
        pf.left_indent = Cm(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(2)
        pf.line_spacing = Pt(28)
    elif level == 2:
        pf.left_indent = Cm(0.74)
        pf.space_before = Pt(3)
        pf.space_after = Pt(1)
        pf.line_spacing = Pt(26)
    else:
        pf.left_indent = Cm(1.48)
        pf.space_before = Pt(2)
        pf.space_after = Pt(1)
        pf.line_spacing = Pt(24)
        
    pPr = p._element.get_or_add_pPr()
    tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="8900"/></w:tabs>')
    sp_el = pPr.find(qn('w:spacing'))
    if sp_el is not None:
        sp_el.addprevious(tabs)
    else:
        pPr.append(tabs)
    
    if bookmark_name:
        text = escape(text)
        hlink = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="{bookmark_name}" w:history="1"/>')
        r_t = parse_xml(f'<w:r {nsdecls("w")}><w:t>{text}</w:t></w:r>')
        r_tab = parse_xml(f'<w:r {nsdecls("w")}><w:tab/></w:r>')
        r_p = parse_xml(f'<w:r {nsdecls("w")}><w:t>{page_num}</w:t></w:r>')
        
        for r_el in [r_t, r_tab, r_p]:
            rPr = r_el.get_or_add_rPr()
            rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_en}" w:eastAsia="{font_cn}" w:hAnsi="{font_en}" w:cs="{font_en}"/>')
            rPr.append(rFonts)
            if is_bold:
                rPr.append(parse_xml(f'<w:b {nsdecls("w")}/>'))
            rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="000000"/>'))
            rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="32"/>'))
            rPr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="32"/>'))
            hlink.append(r_el)
        p._element.append(hlink)
    else:
        r1 = p.add_run(text)
        set_run_fonts(r1, font_cn, 16, bold=is_bold, font_en=font_en)
        r2 = p.add_run("\t")
        set_run_fonts(r2, font_cn, 16, bold=is_bold, font_en=font_en)
        r3 = p.add_run(str(page_num))
        set_run_fonts(r3, font_cn, 16, bold=is_bold, font_en=font_en)


def apply_standard_page_setup(section, cfg: Dict[str, float] = None):
    """应用公文标准页边距与纸张尺寸"""
    cfg = cfg or DEFAULT_PAGE_SETUP
    section.page_width = Cm(cfg["width_cm"])
    section.page_height = Cm(cfg["height_cm"])
    section.top_margin = Cm(cfg["margin_top_cm"])
    section.bottom_margin = Cm(cfg["margin_bottom_cm"])
    section.left_margin = Cm(cfg["margin_left_cm"])
    section.right_margin = Cm(cfg["margin_right_cm"])


def set_run_fonts(run, font_name_cn: str, font_size_pt: float, bold: bool = False, color_rgb: RGBColor = None, font_en: str = "Times New Roman"):
    """设置文本块的中西文字体、字号、加粗与颜色"""
    run.font.name = font_en
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    if color_rgb is not None:
        run.font.color.rgb = color_rgb
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)
        
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_en)
    rFonts.set(qn('w:hAnsi'), font_en)
    rFonts.set(qn('w:eastAsia'), font_name_cn)
    rFonts.set(qn('w:cs'), font_en)


def setup_footer(section, start_page: Optional[int] = None, font_en: str = "Times New Roman"):
    """设置公文页脚，居中显示 - 页码 -"""
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    
    r1 = p.add_run("- ")
    set_run_fonts(r1, "宋体", 10.5, font_en=font_en)
    
    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="{font_en}" w:hAnsi="{font_en}" w:eastAsia="宋体" w:cs="{font_en}"/><w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    p._element.append(fld1)
    
    r3 = p.add_run(" -")
    set_run_fonts(r3, "宋体", 10.5, font_en=font_en)
    
    if start_page is not None:
        secPr = section._sectPr
        pgNumType = secPr.find(qn('w:pgNumType'))
        if pgNumType is None:
            pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")}/>')
            secPr.append(pgNumType)
        pgNumType.set(qn('w:start'), str(start_page))
    else:
        secPr = section._sectPr
        pgNumType = secPr.find(qn('w:pgNumType'))
        if pgNumType is not None and qn('w:start') in pgNumType.attrib:
            del pgNumType.attrib[qn('w:start')]


BLUE_THEME_COLORS = {
    '4F81BD', '365F91', '1F497D', '376092', '17365D', '244061',
    '2E74B5', '0070C0', '002060', '4472C4', '5B9BD5', '70AD47', 'ED7D31'
}


def neutralize_document_styles(doc: Document):
    """彻底清除目标文档内置样式库中的 Office 默认蓝色与主题色，确保所有默认标题和正文均为纯黑色"""
    styles_el = doc.part.styles._element
    for s in styles_el.findall(qn('w:style')):
        rPr = s.find(qn('w:rPr'))
        if rPr is not None:
            for col in rPr.findall(qn('w:color')):
                val = col.get(qn('w:val'), '').upper()
                if val in BLUE_THEME_COLORS:
                    col.set(qn('w:val'), '000000')
                col.attrib.pop(qn('w:themeColor'), None)
                col.attrib.pop(qn('w:themeShade'), None)
                col.attrib.pop(qn('w:themeTint'), None)


def merge_styles_into_doc(doc_src: Document, doc_dst: Document):
    """合并源文档样式库至目标文档，清除 Office 默认的蓝色主题色"""
    dst_styles_el = doc_dst.part.styles._element
    src_styles_el = doc_src.part.styles._element
    
    dst_styles_map = {}
    for s in dst_styles_el.findall(qn('w:style')):
        sid = s.get(qn('w:styleId'))
        if sid:
            dst_styles_map[sid] = s
            
    for src_s in src_styles_el.findall(qn('w:style')):
        sid = src_s.get(qn('w:styleId'))
        if not sid:
            continue
        s_type = src_s.get(qn('w:type'), 'paragraph')
        if s_type in ('paragraph', 'character', 'table'):
            if sid in dst_styles_map:
                dst_s = dst_styles_map[sid]
                dst_styles_el.remove(dst_s)
            cloned_s = copy.deepcopy(src_s)
            for rPr in cloned_s.findall('.//' + qn('w:rPr')):
                col = rPr.find(qn('w:color'))
                if col is not None:
                    c_val = col.get(qn('w:val'), '').upper()
                    if c_val in BLUE_THEME_COLORS:
                        col.set(qn('w:val'), '000000')
                    col.attrib.pop(qn('w:themeColor'), None)
                    col.attrib.pop(qn('w:themeShade'), None)
                    col.attrib.pop(qn('w:themeTint'), None)
            dst_styles_el.append(cloned_s)
            dst_styles_map[sid] = cloned_s


PPR_ORDER = [
    'pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr',
    'widowControl', 'numPr', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
    'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE',
    'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing',
    'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap',
    'jc', 'textDirection', 'textAlignment', 'textboxTightWrap',
    'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'
]

RPR_ORDER = [
    'rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps',
    'strike', 'dstrike', 'outline', 'shadow', 'emboss', 'imprint',
    'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing',
    'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect',
    'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang',
    'eastAsianLayout', 'specVanish', 'oMath', 'rPrChange'
]


def normalize_openxml_element_order(parent, order_list):
    """根据 OpenXML ECMA-376 官方标准严格重排子元素顺序，防止 Word 打开时报文档损坏"""
    from lxml import etree
    children = list(parent)
    if len(children) <= 1:
        return
    def key_func(c):
        name = etree.QName(c).localname
        return order_list.index(name) if name in order_list else 999
    sorted_children = sorted(children, key=key_func)
    if children != sorted_children:
        for c in children:
            parent.remove(c)
        for c in sorted_children:
            parent.append(c)


def inline_style_properties(element, doc_src: Document):
    """将样式中的对齐、缩进、行距与字体直接内联写入节点，防止跨文档合并时格式失效或继承目标默认蓝色"""
    src_styles_el = doc_src.part.styles._element
    styles_by_id = {}
    for s in src_styles_el.findall(qn('w:style')):
        sid = s.get(qn('w:styleId'))
        if sid:
            styles_by_id[sid] = s
            
    p_elements = [element] if element.tag.endswith('p') else element.findall('.//' + qn('w:p'))
    for p_elem in p_elements:
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p_elem.insert(0, pPr)
            
        # 清除前导多余空格并补全首行缩进
        runs = p_elem.findall(qn('w:r'))
        if len(runs) > 0:
            first_r = runs[0]
            t_elem = first_r.find(qn('w:t'))
            if t_elem is not None and t_elem.text:
                orig_t = t_elem.text
                stripped = re.sub(r'^(?:[ \t]{2,}|\u3000+)', '', orig_t)
                if stripped != orig_t:
                    t_elem.text = stripped
                    ind = pPr.find(qn('w:ind'))
                    if ind is None:
                        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="560" w:firstLineChars="200"/>'))
                    elif ind.get(qn('w:firstLine')) is None and ind.get(qn('w:firstLineChars')) is None:
                        ind.set(qn('w:firstLine'), '560')
                        ind.set(qn('w:firstLineChars'), '200')
                        
        pStyle = pPr.find(qn('w:pStyle'))
        sid = pStyle.get(qn('w:val')) if pStyle is not None else 'Normal'
        has_src_color = False
        style_el = styles_by_id.get(sid)
        if style_el is None:
            style_el = styles_by_id.get('Normal')
        if style_el is not None:
            style_pPr = style_el.find(qn('w:pPr'))
            if style_pPr is not None:
                for prop in ['spacing', 'ind', 'jc']:
                    val_el = style_pPr.find(qn(f'w:{prop}'))
                    if val_el is not None and pPr.find(qn(f'w:{prop}')) is None:
                        pPr.append(copy.deepcopy(val_el))
            style_rPr = style_el.find(qn('w:rPr'))
            if style_rPr is not None:
                col_el = style_rPr.find(qn('w:color'))
                if col_el is not None:
                    has_src_color = True
                for r_elem in p_elem.findall('.//' + qn('w:r')):
                    rPr = r_elem.get_or_add_rPr()
                    for prop in ['rFonts', 'b', 'bCs', 'color', 'sz', 'szCs']:
                        val_el = style_rPr.find(qn(f'w:{prop}'))
                        if val_el is not None and rPr.find(qn(f'w:{prop}')) is None:
                            rPr.append(copy.deepcopy(val_el))
                            
        # 针对长篇正文自然段（非标题、非图表、非目录），统一规范为 1.5 倍行距
        p_txt = ''.join(p_elem.itertext()).strip()
        is_heading_or_special = bool(
            re.match(r'^[一二三四五六七八九十]+[、\.]', p_txt) or 
            re.match(r'^（[一二三四五六七八九十]+）', p_txt) or 
            re.match(r'^\d+\.\s*', p_txt) or 
            p_txt.startswith('图') or 
            p_txt.startswith('表') or 
            p_txt in ['目  录', '目录', '目 录'] or 
            '\t' in p_txt or 
            bool(re.search(r'\.{4,}\s*\d+\s*$', p_txt))
        )
        if not is_heading_or_special and len(p_txt) > 20:
            sp = pPr.find(qn('w:spacing'))
            if sp is None:
                sp = parse_xml(f'<w:spacing {nsdecls("w")}/>')
                pPr.append(sp)
            if sp.get(qn('w:line')) is None or sp.get(qn('w:line')) in ['240', '260', '280']:
                sp.set(qn('w:line'), '360')
                sp.set(qn('w:lineRule'), 'auto')
                
        for r_elem in p_elem.findall('.//' + qn('w:r')):
            rPr = r_elem.get_or_add_rPr()
            col = rPr.find(qn('w:color'))
            if col is not None:
                c_val = col.get(qn('w:val'), '').upper()
                theme_c = col.get(qn('w:themeColor'))
                if c_val in BLUE_THEME_COLORS or (theme_c and theme_c.startswith('accent')):
                    col.set(qn('w:val'), '000000')
                    col.attrib.pop(qn('w:themeColor'), None)
                    col.attrib.pop(qn('w:themeShade'), None)
                    col.attrib.pop(qn('w:themeTint'), None)
            elif not has_src_color:
                # 若源样式与 run 均无显式着色，注入纯黑避免被目标文档标题默认蓝色覆盖
                rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="000000"/>'))

        # 严格执行 OpenXML ECMA-376 规范元素顺序重排，防止 Word 提示文档损坏
        normalize_openxml_element_order(pPr, PPR_ORDER)
        for r_elem in p_elem.findall('.//' + qn('w:r')):
            rPr = r_elem.find(qn('w:rPr'))
            if rPr is not None:
                normalize_openxml_element_order(rPr, RPR_ORDER)


def add_heading_paragraph(doc: Document, node: Dict[str, Any], fonts: Dict[str, str], need_page_break: bool = False):
    """
    添加大纲标题段落
    使用段落属性 page_break_before 控制换节分页，同时设置 keep_with_next 与书签锚点。
    """
    lvl = node["level"]
    title = node["title"]
    bm_name = node["bookmark_name"]
    bm_id = node.get("bm_id", 1)
    
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.left_indent = Cm(0)
    if need_page_break:
        pf.page_break_before = True
    
    bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_id}" w:name="{bm_name}"/>')
    p._element.append(bm_start)
    
    font_en = fonts.get("en", "Times New Roman")
    if lvl == 1:
        pf.space_before = Pt(14)
        pf.space_after = Pt(6)
        pf.line_spacing = Pt(30)
        r = p.add_run(title)
        set_run_fonts(r, fonts.get("h1", "黑体"), 16, bold=True, font_en=font_en)
    elif lvl == 2:
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.line_spacing = Pt(30)
        r = p.add_run(title)
        set_run_fonts(r, fonts.get("h2", "楷体_GB2312"), 16, bold=True, font_en=font_en)
    else:
        pf.space_before = Pt(6)
        pf.space_after = Pt(3)
        pf.line_spacing = Pt(28)
        r = p.add_run(title)
        set_run_fonts(r, fonts.get("h3", "仿宋_GB2312"), 16, bold=True, font_en=font_en)
        
    bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_id}"/>')
    p._element.append(bm_end)


def build_toc_sdt(doc: Document, items: List[Dict[str, Any]], heading_pages: Dict[str, int], fonts: Dict[str, str]):
    """构建 Word 结构化文档标签目录控件"""
    sdt_xml = f'''
    <w:sdt {nsdecls("w")}>
      <w:sdtPr>
        <w:docPartObj>
          <w:docPartGallery w:val="Table of Contents"/>
          <w:docPartUnique/>
        </w:docPartObj>
      </w:sdtPr>
      <w:sdtContent>
      </w:sdtContent>
    </w:sdt>
    '''
    sdt_el = parse_xml(sdt_xml)
    sdt_content = sdt_el.find(qn('w:sdtContent'))
    
    # 目录标题
    p_title = parse_xml(f'''
    <w:p {nsdecls("w")}>
      <w:pPr>
        <w:jc w:val="center"/>
        <w:spacing w:before="280" w:after="120" w:line="400" w:lineRule="exact"/>
        <w:rPr>
          <w:rFonts w:ascii="黑体" w:eastAsia="黑体" w:hAnsi="黑体"/>
          <w:b/>
          <w:sz w:val="32"/>
          <w:szCs w:val="32"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:pPr>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="黑体" w:eastAsia="黑体" w:hAnsi="黑体"/>
          <w:b/>
          <w:sz w:val="32"/>
          <w:szCs w:val="32"/>
          <w:color w:val="000000"/>
        </w:rPr>
        <w:t>目  录</w:t>
      </w:r>
    </w:p>
    ''')
    sdt_content.append(p_title)
    
    for item in items:
        lvl = item["level"]
        title_text = item["title"]
        bm_name = item["bookmark_name"]
        page_num = heading_pages.get(title_text, 1)
        
        font_cn = fonts.get("h1", "黑体") if lvl == 1 else (fonts.get("h2", "楷体_GB2312") if lvl == 2 else fonts.get("h3", "仿宋_GB2312"))
        font_sz_half = "32"
        is_bold_xml = "<w:b/><w:bCs/>" if lvl in (1, 2) else ""
        
        if lvl == 1:
            ind_w = "0"
            sp_before = "120"
            sp_after = "40"
            line_sp = "380"
        elif lvl == 2:
            ind_w = "480"
            sp_before = "60"
            sp_after = "20"
            line_sp = "360"
        else:
            ind_w = "960"
            sp_before = "30"
            sp_after = "10"
            line_sp = "340"
            
        p_entry = parse_xml(f'''
        <w:p {nsdecls("w")}>
          <w:pPr>
            <w:pStyle w:val="TOC{lvl}"/>
            <w:tabs>
              <w:tab w:val="right" w:leader="dot" w:pos="8900"/>
            </w:tabs>
            <w:ind w:left="{ind_w}"/>
            <w:spacing w:before="{sp_before}" w:after="{sp_after}" w:line="{line_sp}" w:lineRule="exact"/>
          </w:pPr>
          <w:hyperlink w:anchor="{bm_name}" w:history="1">
            <w:r>
              <w:rPr>
                <w:rFonts w:ascii="{fonts.get("en", "Times New Roman")}" w:eastAsia="{font_cn}" w:hAnsi="{fonts.get("en", "Times New Roman")}"/>
                {is_bold_xml}
                <w:sz w:val="{font_sz_half}"/>
                <w:szCs w:val="{font_sz_half}"/>
                <w:color w:val="000000"/>
              </w:rPr>
              <w:t>{title_text}</w:t>
            </w:r>
            <w:r>
              <w:rPr>
                <w:rFonts w:ascii="{fonts.get("en", "Times New Roman")}" w:eastAsia="{font_cn}" w:hAnsi="{fonts.get("en", "Times New Roman")}"/>
                {is_bold_xml}
                <w:sz w:val="{font_sz_half}"/>
                <w:szCs w:val="{font_sz_half}"/>
                <w:color w:val="000000"/>
              </w:rPr>
              <w:tab/>
            </w:r>
            <w:r>
              <w:rPr>
                <w:rFonts w:ascii="{fonts.get("en", "Times New Roman")}" w:eastAsia="{font_cn}" w:hAnsi="{fonts.get("en", "Times New Roman")}"/>
                {is_bold_xml}
                <w:sz w:val="{font_sz_half}"/>
                <w:szCs w:val="{font_sz_half}"/>
                <w:color w:val="000000"/>
              </w:rPr>
              <w:t>{page_num}</w:t>
            </w:r>
          </w:hyperlink>
        </w:p>
        ''')
        sdt_content.append(p_entry)
        
    doc.element.body.append(sdt_el)
