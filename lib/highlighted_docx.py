# -*- coding: utf-8 -*-
"""高亮标题 DOCX 的通用预处理器。

此模块只处理原文档本身：提取被配置标记的标题、统一标题格式、修复容易导致
空白页的分节空段，并连续化页脚。封面、目录、精确页码与发布仍由统一引擎负责。
"""

import re
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
import pymupdf

from .styles import set_run_fonts, setup_footer


DEFAULT_HIGHLIGHT_SETTINGS = {
    "colors": ["yellow", "ffff00", "ff0"],
    "heading_rules": [
        {"level": 1, "pattern": r"^([一二三四五六七八九十]+)[、.\s]*(.*)", "separator": "、"},
        {"level": 2, "pattern": r"^(\d+)[、.\s]*(.*)", "separator": "."},
    ],
    "title_replacements": {},
}


def _is_marked(paragraph, colors: set[str]) -> bool:
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is not None:
        shd = p_pr.find(qn("w:shd"))
        if shd is not None and shd.get(qn("w:fill"), "").lower() in colors:
            return True
    for run in paragraph.runs:
        r_pr = run._r.find(qn("w:rPr"))
        if r_pr is None:
            continue
        highlight = r_pr.find(qn("w:highlight"))
        if highlight is not None and highlight.get(qn("w:val"), "").lower() in colors:
            return True
        shd = r_pr.find(qn("w:shd"))
        if shd is not None and shd.get(qn("w:fill"), "").lower() in colors:
            return True
    return False


def _strip_marks(paragraph) -> None:
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is not None:
        shd = p_pr.find(qn("w:shd"))
        if shd is not None:
            p_pr.remove(shd)
    for run in paragraph.runs:
        r_pr = run._r.find(qn("w:rPr"))
        if r_pr is None:
            continue
        for tag in (qn("w:highlight"), qn("w:shd")):
            element = r_pr.find(tag)
            if element is not None:
                r_pr.remove(element)


def extract_and_format_highlighted_headings(
    doc: Document,
    settings: Dict[str, Any],
    fonts: Dict[str, str],
    normalize: bool = True,
) -> List[Dict[str, Any]]:
    """读取配置化标记标题；``normalize=False`` 可用于只读构建计划。"""
    merged = {**DEFAULT_HIGHLIGHT_SETTINGS, **(settings or {})}
    colors = {str(color).lower() for color in merged["colors"]}
    rules = merged["heading_rules"]
    replacements = merged.get("title_replacements", {})
    headings = []

    for index, paragraph in enumerate(doc.paragraphs):
        if not _is_marked(paragraph, colors):
            continue
        raw_text = paragraph.text.strip()
        if not raw_text:
            continue
        raw_text = replacements.get(raw_text, raw_text)
        rule_match = None
        for rule in rules:
            match = re.match(rule["pattern"], raw_text)
            if match:
                rule_match = (rule, match)
                break
        if rule_match is None:
            level, title = 1, raw_text
        else:
            rule, match = rule_match
            level = rule["level"]
            title = f"{match.group(1)}{rule.get('separator', '')}{match.group(2).strip()}"

        node = {
            "idx": index,
            "level": level,
            "title": title,
            "toc_title": title,
            "type": "embedded_heading",
            "bookmark_name": f"_Toc_auto_{len(headings) + 1:03d}",
            "bm_id": len(headings) + 1,
            "raw_text": paragraph.text.strip(),
        }
        headings.append(node)
        if not normalize:
            continue

        _strip_marks(paragraph)
        paragraph_format = paragraph.paragraph_format
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph_format.left_indent = Cm(0)
        paragraph_format.keep_with_next = True
        if level == 1:
            paragraph_format.space_before, paragraph_format.space_after = Pt(14), Pt(6)
            paragraph_format.line_spacing = Pt(30)
            font_name = fonts.get("h1", "黑体")
        else:
            paragraph_format.space_before, paragraph_format.space_after = Pt(10), Pt(4)
            paragraph_format.line_spacing = Pt(30)
            font_name = fonts.get("h2", "楷体_GB2312")
        for run in list(paragraph.runs):
            paragraph._p.remove(run._r)
        run = paragraph.add_run(title)
        set_run_fonts(run, font_name, 16, bold=True, color_rgb=RGBColor(0, 0, 0), font_en=fonts.get("en", "Times New Roman"))
        p_pr = paragraph._p.get_or_add_pPr()
        start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{node["bm_id"]}" w:name="{node["bookmark_name"]}"/>')
        end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{node["bm_id"]}"/>')
        paragraph._p.insert(list(paragraph._p).index(p_pr) + 1, start)
        paragraph._p.append(end)
    return headings


def normalize_document_pagination(doc: Document, start_page: int = 1) -> None:
    """清除分节周围冗余空段，并统一页脚与连续页码。"""
    for section in doc.sections:
        pg_mar = section._sectPr.find(qn("w:pgMar"))
        if pg_mar is not None:
            if pg_mar.get(qn("w:bottom")) == "0" or int(pg_mar.get(qn("w:bottom"), "1395")) < 500:
                pg_mar.set(qn("w:bottom"), "1395")
                pg_mar.set(qn("w:footer"), "1166")
    body = doc._body._element
    children = list(body)
    for index, element in enumerate(children):
        if element.tag != qn("w:p") or element.find(".//" + qn("w:sectPr")) is None:
            continue
        for neighbor in (children[index - 1] if index else None, children[index + 1] if index + 1 < len(children) else None):
            if neighbor is None or neighbor.tag != qn("w:p"):
                continue
            if not "".join(neighbor.itertext()).strip() and not neighbor.findall(".//" + qn("w:drawing")):
                if neighbor.getparent() is body:
                    body.remove(neighbor)
        p_pr = element.find(qn("w:pPr"))
        if p_pr is not None:
            spacing = p_pr.find(qn("w:spacing"))
            if spacing is None:
                spacing = OxmlElement("w:spacing")
                p_pr.append(spacing)
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "20")
            spacing.set(qn("w:lineRule"), "exact")
    # 删除文末最后一个实质内容之后的空段，避免分节符在满页后额外生成空白页。
    children = list(body)
    last_content_index = 0
    for index, element in enumerate(children):
        if element.tag == qn("w:tbl"):
            last_content_index = index
        elif element.tag == qn("w:p") and (
            "".join(element.itertext()).strip()
            or element.findall(".//" + qn("w:drawing"))
            or element.findall(".//{urn:schemas-microsoft-com:vml}shape")
        ):
            last_content_index = index
    for element in children[last_content_index + 1:]:
        if element.tag == qn("w:p") and element.getparent() is body:
            body.remove(element)

    final_section = body.find(qn("w:sectPr"))
    if final_section is not None:
        pg_mar = final_section.find(qn("w:pgMar"))
        if pg_mar is not None:
            pg_mar.set(qn("w:bottom"), "1395")
            pg_mar.set(qn("w:footer"), "1166")

    for index, section in enumerate(doc.sections):
        pg_num = section._sectPr.find(qn("w:pgNumType"))
        if pg_num is not None and qn("w:start") in pg_num.attrib:
            del pg_num.attrib[qn("w:start")]
        if index == 0:
            for paragraph in list(section.footer.paragraphs)[1:]:
                paragraph._element.getparent().remove(paragraph._element)
            setup_footer(section, start_page=start_page, font_en="Times New Roman")
        else:
            section.footer.is_linked_to_previous = True


def find_heading_pages_in_pdf(pdf_path: str, headings: List[Dict[str, Any]]) -> Dict[str, int]:
    """使用 PyMuPDF 反查保留原版式单文档的标题页码，避免损坏 PDF 目录对象影响解析。"""
    pdf = pymupdf.open(pdf_path)
    try:
        page_texts = [re.sub(r"[\s\.、，。（）()|｜·/\-_:\—\t\n\r‘’“”+]", "", page.get_text()) for page in pdf]
        result: Dict[str, int] = {}
        current_page = 0
        for heading in headings:
            key = re.sub(r"[\s\.、，。（）()|｜·/\-_:\—\t\n\r‘’“”+]", "", heading["title"])
            probe = key[:12]
            for index in range(current_page, len(page_texts)):
                if probe and probe in page_texts[index]:
                    current_page = index
                    break
            result[heading["title"]] = current_page + 1
        return result
    finally:
        pdf.close()
