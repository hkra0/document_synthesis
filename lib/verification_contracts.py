# -*- coding: utf-8 -*-
"""构建预期与最终 DOCX 之间的可审计验证契约。

这些值对象只保存身份、顺序、策略和属性摘要，不复制源正文。它们把
``PreparedBuild`` 产生的预期传给 engine、内容门禁和格式门禁，避免从输出
样式或字符串反推“应该检查什么”。
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Tuple

from .docx_inspector import BlockInspection, DocumentInspection, NodeRef
from .document_parts import get_part_boundary_bookmark, generated_content_bookmark_name
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _compute_sha256(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(frozen=True)
class SourceOccurrence:
    source_sha256: str
    part_uri: str
    element_path: str
    delivery_id: str
    part_id: str
    occurrence_index: int
    source_order_index: int

    @property
    def occurrence_id(self) -> str:
        return "|".join(
            (
                self.source_sha256,
                self.part_uri,
                self.element_path,
                self.delivery_id,
                self.part_id,
                str(self.occurrence_index),
            )
        )

    def to_dict(self) -> Dict[str, Any]:
        data = asdict(self)
        data["occurrence_id"] = self.occurrence_id
        return data


def occurrence_bookmark_name(occurrence_id: str) -> str:
    """Return a stable, Word-compatible, non-visible provenance bookmark name."""
    import hashlib

    # Word bookmark names are limited in length and must start with a letter
    # or underscore.  The full occurrence identity remains in the sidecar;
    # the digest is only the on-document lookup key.
    digest = hashlib.sha256(str(occurrence_id).encode("utf-8")).hexdigest()[:28]
    return f"_SynthSrc_{digest}"


@dataclass(frozen=True)
class ExpectedNode:
    occurrence: SourceOccurrence
    role: str
    story_type: str
    structure_type: str
    semantic_summary: Dict[str, Any] = field(default_factory=dict)
    managed_properties: Tuple[str, ...] = ()
    protected_properties: Tuple[str, ...] = ()
    provenance: str = "prepared"

    def to_dict(self) -> Dict[str, Any]:
        return {
            "occurrence": self.occurrence.to_dict(),
            "role": self.role,
            "story_type": self.story_type,
            "structure_type": self.structure_type,
            "semantic_summary": dict(self.semantic_summary),
            "managed_properties": list(self.managed_properties),
            "protected_properties": list(self.protected_properties),
            "provenance": self.provenance,
        }


@dataclass(frozen=True)
class ExpectedDelivery:
    delivery_id: str
    parts: Tuple[str, ...]
    ordered_nodes: Tuple[ExpectedNode, ...]
    source_order: Tuple[str, ...] = ()
    page_expectations: Tuple[Dict[str, Any], ...] = ()
    generated_content: Tuple[Dict[str, Any], ...] = ()
    field_dependencies: Tuple[str, ...] = ()
    policy_mode: str = "restyle"
    page_policy: str = "target"
    inline_emphasis: str = "preserve"
    story_expectations: Tuple[Dict[str, Any], ...] = ()
    required_checks: Tuple[str, ...] = (
        "content_integrity",
        "format_verification",
        "page_qa",
    )

    def to_dict(self) -> Dict[str, Any]:
        return {
            "delivery_id": self.delivery_id,
            "parts": list(self.parts),
            "ordered_nodes": [node.to_dict() for node in self.ordered_nodes],
            "source_order": list(self.source_order),
            "page_expectations": [dict(item) for item in self.page_expectations],
            "generated_content": [dict(item) for item in self.generated_content],
            "field_dependencies": list(self.field_dependencies),
            "policy_mode": self.policy_mode,
            "page_policy": self.page_policy,
            "inline_emphasis": self.inline_emphasis,
            "story_expectations": [dict(item) for item in self.story_expectations],
            "required_checks": list(self.required_checks),
        }


@dataclass(frozen=True)
class OutputNodeMap:
    """来源 occurrence 到最终 DOCX 节点的位置映射。"""

    locations: Dict[str, Tuple[Dict[str, Any], ...]] = field(default_factory=dict)
    transformations: Tuple[Dict[str, Any], ...] = ()
    missing_occurrences: Tuple[str, ...] = ()
    duplicate_occurrences: Tuple[str, ...] = ()

    def to_dict(self) -> Dict[str, Any]:
        return {
            "locations": {key: list(value) for key, value in self.locations.items()},
            "transformations": [dict(item) for item in self.transformations],
            "missing_occurrences": list(self.missing_occurrences),
            "duplicate_occurrences": list(self.duplicate_occurrences),
        }


@dataclass(frozen=True)
class VerificationContext:
    expected_delivery: ExpectedDelivery
    output_node_map: OutputNodeMap
    source_hashes: Dict[str, str]
    format_hash: str
    delivery_id: str
    config_hash: str = ""
    delivery_sha256: Optional[str] = None
    format_source_hashes: Dict[str, str] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "contract_version": 1,
            "delivery_id": self.delivery_id,
            "expected_delivery": self.expected_delivery.to_dict(),
            "output_node_map": self.output_node_map.to_dict(),
            "source_hashes": dict(self.source_hashes),
            "format_hash": self.format_hash,
            "config_hash": self.config_hash,
            "format_source_hashes": dict(self.format_source_hashes),
            "delivery_sha256": self.delivery_sha256,
            "coverage": {
                "expected_count": len(self.expected_delivery.ordered_nodes),
                "mapped_count": len(self.output_node_map.locations),
                "missing_count": len(self.output_node_map.missing_occurrences),
                "duplicate_count": len(self.output_node_map.duplicate_occurrences),
            },
        }


def _assignment_by_key(prepared: Any) -> Dict[Tuple[str, str, str], Any]:
    result = {}
    for assignment in getattr(prepared, "assignments", ()) or ():
        ref = assignment.node_ref
        result[(ref.source_sha256, ref.part_uri, ref.element_path)] = assignment
    return result


def _blocks_in_source_order(prepared: Any) -> List[BlockInspection]:
    blocks: List[BlockInspection] = []
    for inspection in getattr(prepared, "inspections", ()) or ():
        blocks.extend(inspection.blocks)
    return blocks


def _part_value(prepared: Any, part_id: str, name: str, default: Any = None) -> Any:
    part = (getattr(prepared, "parts", {}) or {}).get(part_id)
    if isinstance(part, Mapping):
        return part.get(name, default)
    return getattr(part, name, default) if part is not None else default


def _blocks_for_content_part(prepared: Any, part_id: str) -> List[BlockInspection]:
    """Return source blocks selected by one content part, in source order."""
    all_blocks = _blocks_in_source_order(prepared)
    region_id = _part_value(prepared, part_id, "source_region")
    regions = getattr(getattr(prepared, "config", None), "regions", {}) or {}
    if not region_id or region_id == "entire_document":
        return list(all_blocks)
    if region_id not in regions:
        return []
    region = regions[region_id]
    selected: List[BlockInspection] = []
    for inspection in getattr(prepared, "inspections", ()) or ():
        block_list = list(inspection.blocks)
        paths = [block.node.element_path for block in block_list]
        if region.start not in paths:
            continue
        start = paths.index(region.start)
        end = len(block_list) if region.end == "end_of_document" else (
            paths.index(region.end) if region.end in paths else None
        )
        if end is not None:
            selected.extend(block_list[start:end])
    return selected


def _node_is_restyled(prepared: Any, assignment: Any) -> bool:
    policy = getattr(getattr(prepared, "config", None), "formatting", None)
    mode = getattr(policy, "mode", "restyle")
    if mode == "restyle":
        return True
    if mode == "preserve":
        return False
    scopes = list(getattr(policy, "scopes", ()) or ())
    if not scopes:
        return True
    if assignment is None:
        return False
    source_path = Path(next(
        (inspection.file_path for inspection in getattr(prepared, "inspections", ())
         if inspection.source_sha256 == assignment.node_ref.source_sha256),
        "",
    )).resolve()
    applicable = []
    for scope in scopes:
        declared = scope.get("file") if isinstance(scope, Mapping) else None
        if not declared:
            continue
        try:
            if Path(declared).resolve() != source_path:
                continue
        except OSError:
            if Path(declared).name != source_path.name:
                continue
        node_range = scope.get("node_range")
        if not node_range or node_range in {"*", "all"}:
            in_range = True
        else:
            import re
            numbers = [int(item) for item in re.findall(r"\[(\d+)\]", node_range)]
            path_numbers = [int(item) for item in re.findall(r"\[(\d+)\]", assignment.node_ref.element_path)]
            in_range = bool(numbers and path_numbers and numbers[0] <= path_numbers[0] <= numbers[-1])
        if in_range:
            applicable.append(scope)
    return bool(applicable and applicable[-1].get("mode") == "restyle")


def _serialize_run_spans(block: BlockInspection) -> List[Dict[str, Any]]:
    """Serialize logical text intervals without retaining inspector objects."""
    spans = []
    for span in getattr(block, "effective_run_spans", ()) or ():
        style = span.get("style") if isinstance(span, Mapping) else None
        if hasattr(style, "__dataclass_fields__"):
            style = asdict(style)
        spans.append({
            "start": int(span.get("start", 0)),
            "end": int(span.get("end", 0)),
            "style": dict(style or {}),
        })
    return spans


def _serialize_inline_emphasis_spans(block: BlockInspection) -> List[Dict[str, Any]]:
    """Serialize only direct run-level emphasis evidence."""
    return [
        {
            "start": int(span.get("start", 0)),
            "end": int(span.get("end", 0)),
            "bold": span.get("bold"),
            "italic": span.get("italic"),
        }
        for span in (getattr(block, "inline_emphasis_spans", ()) or ())
    ]


def build_expected_delivery(prepared: Any, spec: Mapping[str, Any]) -> ExpectedDelivery:
    """从准备结果冻结一个交付物的有序来源实例。"""

    delivery_id = str(spec.get("id", "main"))
    parts = tuple(str(item) for item in spec.get("parts", ()))
    content_parts = tuple(
        part_id
        for part_id in parts
        if part_id == "body"
        or getattr((getattr(prepared, "parts", {}) or {}).get(part_id), "kind", None) == "content"
    )
    assignments = _assignment_by_key(prepared)
    source_config = getattr(getattr(prepared, "config", None), "source", {}) or {}
    source_strategy = (
        source_config.get("strategy", "directory_tree")
        if isinstance(source_config, Mapping)
        else getattr(source_config, "strategy", "directory_tree")
    )
    # Directory-tree imports add an outer file heading before each source.
    # The first imported source paragraph therefore intentionally overrides a
    # Heading 1 style's inherited page break so the wrapper heading is not left
    # alone on a title-only page.  Record that narrow layout override in the
    # frozen contract instead of making the verifier infer it from output.
    boundary_first_keys = set()
    if source_strategy == "directory_tree":
        for inspection in getattr(prepared, "inspections", ()) or ():
            for block in inspection.blocks:
                if block.structure_type == "paragraph" and block.visible_text:
                    boundary_first_keys.add((
                        block.node.source_sha256,
                        block.node.part_uri,
                        block.node.element_path,
                    ))
                    break
                if block.structure_type in {"paragraph", "table"}:
                    break
    expected_nodes: List[ExpectedNode] = []
    occurrence_counts: Dict[Tuple[str, str, str], int] = {}
    source_order_index = 0
    for part_id in content_parts:
        for block in _blocks_for_content_part(prepared, part_id):
            ref = block.node
            key = (ref.source_sha256, ref.part_uri, ref.element_path)
            occurrence_index = occurrence_counts.get(key, 0)
            occurrence_counts[key] = occurrence_index + 1
            assignment = assignments.get(key)
            role = getattr(assignment, "role", None) or (
                "table.body" if block.structure_type == "table" else "body"
            )
            occurrence = SourceOccurrence(
                source_sha256=ref.source_sha256,
                part_uri=ref.part_uri,
                element_path=ref.element_path,
                delivery_id=delivery_id,
                part_id=part_id,
                occurrence_index=occurrence_index,
                source_order_index=source_order_index,
            )
            source_order_index += 1
            managed = (
                ("paragraph", "run")
                if assignment is not None and _node_is_restyled(prepared, assignment)
                else ()
            )
            protected = ("semantic_objects", "inline_emphasis", "source_effective_style")
            semantic_summary = {
                "text_hash": ref.text_hash,
                "runs_count": block.runs_count,
                "protected_objects": list(block.protected_objects),
                "source_effective_run": asdict(block.effective_run) if block.effective_run else None,
                "source_effective_runs": [asdict(item) for item in block.effective_runs],
                "source_effective_run_spans": _serialize_run_spans(block),
                "source_inline_emphasis_spans": _serialize_inline_emphasis_spans(block),
                "source_effective_paragraph": asdict(block.effective_paragraph) if block.effective_paragraph else None,
            }
            if (
                key in boundary_first_keys
                and role.startswith("heading.")
                and managed
            ):
                semantic_summary["format_overrides"] = {"page_break_before": False}
            expected_nodes.append(ExpectedNode(
                occurrence=occurrence,
                role=role,
                story_type=block.story_type,
                structure_type=block.structure_type,
                semantic_summary=semantic_summary,
                managed_properties=managed,
                protected_properties=protected,
                provenance=getattr(assignment, "provenance", "prepared"),
            ))

    format_policy = getattr(getattr(prepared, "config", None), "formatting", None)
    page_expectations: List[Dict[str, Any]] = []
    if getattr(format_policy, "page_policy", "target") == "source":
        for inspection in getattr(prepared, "inspections", ()) or ():
            page_expectations.extend(asdict(section.page_spec) for section in inspection.sections)
    field_dependencies: List[str] = []
    for inspection in getattr(prepared, "inspections", ()) or ():
        for block in inspection.blocks:
            field_dependencies.extend(
                str(item) for item in getattr(block, "protected_objects", ())
                if str(item).startswith("field")
            )
    generated_content: List[Dict[str, Any]] = []
    if "cover" in parts:
        import hashlib
        cover_spec = getattr(getattr(prepared, "config", None), "cover_spec", None)
        fallback_cover = getattr(getattr(prepared, "config", None), "cover", {}) or {}
        template_counts: Dict[str, int] = {}
        template_path = getattr(cover_spec, "template_path", None)
        if (
            getattr(cover_spec, "mode", "generated") == "template"
            and template_path
            and Path(template_path).is_file()
        ):
            # Count placeholders from the frozen template so an intentional
            # repeated field remains legal while an extra copy in the final
            # cover is still detectable.
            try:
                template_doc = Document(str(template_path))
                template_text = "".join(
                    node.text or "" for node in template_doc.element.body.iter(qn("w:t"))
                )
                template_counts = {
                    "header_title": template_text.count("{{HEADER_TITLE}}"),
                    "sub_title": template_text.count("{{SUB_TITLE}}"),
                    "main_title": template_text.count("{{MAIN_TITLE}}"),
                    "author": template_text.count("{{AUTHOR}}"),
                    "date": template_text.count("{{DATE}}"),
                }
            except (OSError, ValueError, TypeError):
                # build_cover performs the authoritative template validation;
                # this audit sidecar must not make a valid build fail merely
                # because an optional count could not be re-read.
                template_counts = {}
        dynamic_fields = dict(getattr(cover_spec, "dynamic_fields", {}) or {})
        if not dynamic_fields and isinstance(fallback_cover, Mapping):
            dynamic_fields = {
                "header_title": str(fallback_cover.get("header_title") or ""),
                "sub_title": str(fallback_cover.get("sub_title") or ""),
                "main_title": str(fallback_cover.get("main_title") or getattr(getattr(prepared, "config", None), "project_name", "")),
                "author": str(fallback_cover.get("author") or ""),
                "date": str(fallback_cover.get("date") or ""),
            }
        for field_name, value in dynamic_fields.items():
            if value:
                generated_content.append({
                    "part_id": "cover",
                    "kind": "cover_field",
                    "field": field_name,
                    "expected_text": str(value),
                    "expected_text_hash": hashlib.sha256(str(value).encode("utf-8")).hexdigest(),
                    "expected_count": max(1, template_counts.get(field_name, 1)),
                    "position": {"part_id": "cover", "field": field_name},
                    "marker_names": [
                        generated_content_bookmark_name(
                            "cover", field_name, index
                        )
                        for index in range(max(1, template_counts.get(field_name, 1)))
                    ],
                    "mode": getattr(cover_spec, "mode", "generated"),
                })
    source_order = tuple(str(item) for item in getattr(prepared, "source_order", ()) or ())
    story_expectations: List[Dict[str, Any]] = []
    for inspection in getattr(prepared, "inspections", ()) or ():
        story_by_key = {
            (story.story_type, story.variant, story.part_uri): story
            for story in getattr(inspection, "stories", ()) or ()
        }
        for binding in getattr(inspection, "section_bindings", ()) or ():
            story = story_by_key.get((binding.story_type, binding.variant, binding.part_uri))
            if story is None:
                continue
            section = next(
                (item for item in inspection.sections if item.section_index == binding.section_index),
                None,
            )
            story_expectations.append({
                "source_sha256": inspection.source_sha256,
                "section_index": binding.section_index,
                "story_type": binding.story_type,
                "variant": binding.variant,
                "part_uri": binding.part_uri,
                "texts": [block.visible_text for block in story.blocks],
                "blocks": [
                    {
                        "text": block.visible_text,
                        "spans": _serialize_run_spans(block),
                    }
                    for block in story.blocks
                ],
                "header_distance_mm": section.page_spec.header_distance_mm if section else None,
                "footer_distance_mm": section.page_spec.footer_distance_mm if section else None,
                "different_first_page": section.different_first_page if section else False,
                "linked_to_previous": binding.linked_to_previous,
            })
    return ExpectedDelivery(
        delivery_id=delivery_id,
        parts=parts,
        ordered_nodes=tuple(expected_nodes),
        source_order=source_order,
        generated_content=tuple(generated_content),
        page_expectations=tuple(page_expectations),
        field_dependencies=tuple(field_dependencies),
        policy_mode=getattr(format_policy, "mode", "restyle"),
        page_policy=getattr(format_policy, "page_policy", "target"),
        inline_emphasis=getattr(format_policy, "inline_emphasis", "preserve"),
        story_expectations=tuple(story_expectations),
    )


def _content_output_paths(delivery_path: Optional[Path], expected: ExpectedDelivery) -> Optional[set[str]]:
    """Limit mapping to declared content-part boundaries, excluding cover/TOC."""
    if not delivery_path or not Path(delivery_path).is_file() or not expected.ordered_nodes:
        return None if expected.ordered_nodes else set()
    content_parts = [
        part for part in expected.parts
        if part == "body" or part not in {"cover", "toc"}
    ]
    if not content_parts:
        return set()
    doc = Document(str(delivery_path))
    children = list(doc.element.body)
    starts: Dict[str, int] = {}
    for index, child in enumerate(children):
        for marker in child.iter(qn("w:bookmarkStart")):
            name = marker.get(qn("w:name"))
            for part_id in content_parts:
                if name == get_part_boundary_bookmark(part_id, is_start=True):
                    starts[part_id] = index
    intervals = []
    for part_id in content_parts:
        start = starts.get(part_id)
        if start is None:
            continue
        end = None
        end_name = get_part_boundary_bookmark(part_id, is_start=False)
        for index in range(start, len(children)):
            if any(marker.get(qn("w:name")) == end_name for marker in children[index].iter(qn("w:bookmarkStart"))):
                end = index
                break
        if end is None:
            following = [value for value in starts.values() if value > start]
            end = min(following) if following else len(children) - 1
        intervals.append((start, end))

    selected = set()
    tag_counts = {qn("w:p"): 0, qn("w:tbl"): 0}
    for index, child in enumerate(children):
        if child.tag not in tag_counts:
            continue
        tag_counts[child.tag] += 1
        if any(start <= index <= end for start, end in intervals):
            prefix = "p" if child.tag == qn("w:p") else "tbl"
            selected.add(f"/w:document/w:body/w:{prefix}[{tag_counts[child.tag]}]")
    return selected


def _element_for_path(document: Document, element_path: str):
    """Resolve the inspector's stable body/table paragraph path in a DOCX."""
    import re

    body = document.element.body
    top_match = re.fullmatch(r"/w:document/w:body/w:p\[(\d+)\]", element_path)
    if top_match:
        paragraphs = [child for child in body if child.tag == qn("w:p")]
        index = int(top_match.group(1)) - 1
        return paragraphs[index] if 0 <= index < len(paragraphs) else None

    table_match = re.fullmatch(
        r"/w:document/w:body/w:tbl\[(\d+)\]/w:tr\[(\d+)\]/w:tc\[(\d+)\]/w:p\[(\d+)\]",
        element_path,
    )
    if not table_match:
        return None
    table_index, row_index, cell_index, paragraph_index = (
        int(value) - 1 for value in table_match.groups()
    )
    tables = [child for child in body if child.tag == qn("w:tbl")]
    if not 0 <= table_index < len(tables):
        return None
    rows = tables[table_index].findall(qn("w:tr"))
    if not 0 <= row_index < len(rows):
        return None
    cells = rows[row_index].findall(qn("w:tc"))
    if not 0 <= cell_index < len(cells):
        return None
    paragraphs = cells[cell_index].findall(qn("w:p"))
    return paragraphs[paragraph_index] if 0 <= paragraph_index < len(paragraphs) else None


def _bookmark_names(element: Any) -> set[str]:
    return {
        marker.get(qn("w:name"))
        for marker in element.iter(qn("w:bookmarkStart"))
        if marker.get(qn("w:name"))
    }


def _logical_text_range(block: BlockInspection) -> Optional[Dict[str, int]]:
    spans = list(getattr(block, "effective_run_spans", ()) or ())
    if not spans:
        return None
    return {
        "start": min(int(span.get("start", 0)) for span in spans),
        "end": max(int(span.get("end", 0)) for span in spans),
    }


def stamp_output_occurrences(
    delivery_path: Path,
    expected: ExpectedDelivery,
) -> OutputNodeMap:
    """Stamp each v3 source occurrence with an invisible provenance bookmark.

    The input is already assembled and is rewritten in place before page
    measurement.  This gives the post-save verifier an identity stronger than
    a repeated paragraph string while keeping the marker out of visible text.
    """
    delivery_path = Path(delivery_path)
    if not expected.ordered_nodes:
        return OutputNodeMap()
    from .docx_inspector import inspect_docx

    output_map = _find_output_locations(expected, inspect_docx(delivery_path), delivery_path)
    if output_map.missing_occurrences or output_map.duplicate_occurrences:
        raise ValueError(
            "无法为每个预期来源实例建立 provenance 标记: "
            f"missing={len(output_map.missing_occurrences)}, "
            f"duplicate={len(output_map.duplicate_occurrences)}"
        )

    document = Document(str(delivery_path))
    existing_names = _bookmark_names(document.element.body)
    bookmark_ids = [
        int(marker.get(qn("w:id")))
        for marker in document.element.body.iter(qn("w:bookmarkStart"))
        if (marker.get(qn("w:id")) or "").lstrip("-").isdigit()
    ]
    next_id = max(bookmark_ids, default=-1) + 1
    transformations = []
    for expected_node in expected.ordered_nodes:
        occurrence_id = expected_node.occurrence.occurrence_id
        location = output_map.locations.get(occurrence_id, ())
        if len(location) != 1:
            raise ValueError(f"来源实例 provenance 定位不唯一: {occurrence_id}")
        element_path = location[0]["element_path"]
        element = _element_for_path(document, element_path)
        if element is None:
            raise ValueError(f"来源实例 provenance 节点无法定位: {element_path}")
        name = occurrence_bookmark_name(occurrence_id)
        names = _bookmark_names(element)
        if name not in names:
            if name in existing_names:
                raise ValueError(f"来源实例 provenance 书签重复: {name}")
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(next_id))
            start.set(qn("w:name"), name)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(next_id))
            ppr = element.find(qn("w:pPr"))
            element.insert(1 if ppr is not None else 0, start)
            element.append(end)
            existing_names.add(name)
            next_id += 1
        transformations.append({
            "kind": "provenance_bookmark",
            "occurrence_id": occurrence_id,
            "bookmark_name": name,
            "element_path": element_path,
        })
    document.save(str(delivery_path))
    return OutputNodeMap(
        locations=output_map.locations,
        transformations=tuple(transformations),
        missing_occurrences=output_map.missing_occurrences,
        duplicate_occurrences=output_map.duplicate_occurrences,
    )


def _find_output_locations(
    expected: ExpectedDelivery,
    output_inspection: DocumentInspection,
    delivery_path: Optional[Path] = None,
) -> OutputNodeMap:
    blocks = list(output_inspection.blocks)
    allowed_paths = _content_output_paths(delivery_path, expected)
    if allowed_paths is not None:
        blocks = [
            block for block in blocks
            if any(
                block.node.element_path == path
                or block.node.element_path.startswith(path + "/")
                for path in allowed_paths
            )
        ]
    unused = set(range(len(blocks)))
    marker_paths: Dict[str, List[int]] = {}
    if delivery_path is not None and Path(delivery_path).is_file():
        document = Document(str(delivery_path))
        for index, block in enumerate(blocks):
            element = _element_for_path(document, block.node.element_path)
            if element is None:
                continue
            for name in _bookmark_names(element):
                marker_paths.setdefault(name, []).append(index)
    locations: Dict[str, Tuple[Dict[str, Any], ...]] = {}
    missing: List[str] = []
    duplicates: List[str] = []
    transformations: List[Dict[str, Any]] = []
    for expected_node in expected.ordered_nodes:
        occurrence_id = expected_node.occurrence.occurrence_id
        expected_text_hash = expected_node.semantic_summary.get("text_hash")
        marker_name = occurrence_bookmark_name(occurrence_id)
        marker_candidates = [index for index in marker_paths.get(marker_name, []) if index in unused]
        # If a v3 output carries at least one provenance marker, use markers
        # for every expected instance.  A missing marker must not silently
        # fall back to a same-text block and shrink the denominator.
        has_provenance = any(name.startswith("_SynthSrc_") for name in marker_paths)
        candidates = marker_candidates if has_provenance else [
            index for index in sorted(unused)
            if blocks[index].structure_type == expected_node.structure_type
            and blocks[index].node.text_hash == expected_text_hash
        ]
        # A prepared source occurrence carries a semantic text fingerprint.
        # Never bind it to an arbitrary same-shaped output block when that
        # fingerprint is absent: doing so turns a missing/replaced source node
        # into a false mapping and silently shrinks the verification denominator.
        # The structure-only fallback remains only for legacy contexts that do
        # not provide a fingerprint at all.
        if not candidates and expected_text_hash is None:
            candidates = [
                index for index in sorted(unused)
                if blocks[index].structure_type == expected_node.structure_type
            ]
        if not candidates:
            missing.append(occurrence_id)
            continue
        selected = candidates if has_provenance else candidates[:1]
        for index in selected:
            unused.remove(index)
        locations[occurrence_id] = tuple({
            "part_uri": blocks[index].node.part_uri,
            "element_path": blocks[index].node.element_path,
            "text_hash": blocks[index].node.text_hash,
            "source_text_hash": expected_text_hash,
            "text_hash_matches": expected_text_hash is None or blocks[index].node.text_hash == expected_text_hash,
            "text_range": _logical_text_range(blocks[index]),
            "structure_type": blocks[index].structure_type,
            "story_type": blocks[index].story_type,
            "output_index": index,
            "provenance_marker": marker_name if has_provenance else None,
            "provenance_matches": index in marker_candidates if has_provenance else False,
        } for index in selected)
        if len(selected) > 1:
            # Multiple identical nodes are legal only when the source expected
            # list contains the corresponding repeated occurrences.  Keep the
            # information for an audit report instead of silently collapsing it.
            duplicates.append(occurrence_id)
        if has_provenance and selected:
            transformations.append({
                "kind": "provenance_bookmark",
                "occurrence_id": occurrence_id,
                "bookmark_name": marker_name,
                "element_path": blocks[selected[0]].node.element_path,
            })
    expected_signatures = {
        (node.structure_type, node.semantic_summary.get("text_hash"))
        for node in expected.ordered_nodes
    }
    for index in sorted(unused):
        block = blocks[index]
        if (block.structure_type, block.node.text_hash) in expected_signatures:
            duplicates.append(f"output:{block.node.part_uri}:{block.node.element_path}")
    return OutputNodeMap(
        locations=locations,
        transformations=tuple(transformations),
        missing_occurrences=tuple(missing),
        duplicate_occurrences=tuple(duplicates),
    )


def build_verification_context(
    prepared: Any,
    rendered: Any,
    delivery_spec: Mapping[str, Any],
    delivery_path: Optional[Path] = None,
) -> VerificationContext:
    """构造 engine/smoke 共用的预期与输出映射上下文。"""

    expected = build_expected_delivery(prepared, delivery_spec)
    output_map = OutputNodeMap()
    if delivery_path is not None and Path(delivery_path).is_file():
        from .docx_inspector import inspect_docx

        output_map = _find_output_locations(expected, inspect_docx(delivery_path), Path(delivery_path))
    return VerificationContext(
        expected_delivery=expected,
        output_node_map=output_map,
        source_hashes=dict(getattr(prepared, "source_hashes", {}) or {}),
        format_hash=str(getattr(prepared, "format_hash", "") or ""),
        format_source_hashes=dict(getattr(prepared, "format_source_hashes", {}) or {}),
        delivery_id=str(delivery_spec.get("id", "main")),
        config_hash=str(getattr(prepared, "config_hash", "") or ""),
        delivery_sha256=_compute_sha256(Path(delivery_path)) if delivery_path and Path(delivery_path).is_file() else None,
    )
