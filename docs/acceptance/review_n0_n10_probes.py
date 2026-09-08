"""Independent final-assembly probes. No mocks or Word required.

Run: python3 docs/acceptance/review_n0_n10_probes.py SCRATCH OUTPUT_JSON
"""
import json
import importlib.util
import copy
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from docx import Document
from docx.shared import Pt, Mm
from lib.engine import UnifiedSynthesizer
from lib.composition import assemble_document
from lib.content_integrity import verify_delivery_format, verify_content_integrity, extract_semantic_inventory, ExpectedInventory
from docx.oxml.ns import qn
from lib.verification_contracts import build_verification_context
from lib.format_analysis import analyze_format_sample


def main():
    root = Path(sys.argv[1]).resolve()
    root.mkdir(parents=True, exist_ok=True)
    results = {}
    for name, sizes in [('mixed_sizes', [12, 24]), ('uniform', [12, 12])]:
        case = root / name
        case.mkdir(exist_ok=True)
        source = case / 'source.docx'
        doc = Document()
        p = doc.add_paragraph()
        for i, size in enumerate(sizes):
            r = p.add_run(f'Anonymous source span {i} for preservation verification. ')
            r.font.name = 'Arial'
            r.font.size = Pt(size)
        doc.sections[0].header.paragraphs[0].text = 'EXPECTED ANONYMOUS HEADER'
        doc.sections[0].footer.paragraphs[0].text = 'EXPECTED ANONYMOUS FOOTER'
        doc.save(source)
        raw = {'schema_version': 3, 'project_name': 'Independent N review',
               'source': {'strategy': 'docx_document'},
               'format': {'ref': 'preset:academic-basic@1.0.0'},
               'formatting': {'mode': 'preserve', 'page_policy': 'source'},
               'output': {'documents': [{'id': 'main', 'filename': 'result.docx', 'parts': ['body']}]}}
        manifest = case / 'manifest.json'
        manifest.write_text(json.dumps(raw))
        plan = UnifiedSynthesizer.prepare_build(source, manifest)
        rendered = UnifiedSynthesizer.render_body(plan, case / 'rendered')
        out = case / 'assembled.docx'
        assemble_document(plan.config, plan.config.documents[0], rendered.path, rendered.nodes, {}, out)

        def check(path):
            ctx = build_verification_context(plan, rendered, plan.config.documents[0], path)
            rep = verify_delivery_format(path, plan.config.resolved_format, verification_context=ctx)
            return {'passed': rep.passed, 'violations': rep.violations,
                    'verified_count': rep.verified_count, 'coverage': ctx.to_dict()['coverage']}

        results[name] = {'unchanged': check(out)}
        changed = Document(out)
        if name == 'mixed_sizes':
            changed.paragraphs[0].runs[1].font.size = Pt(12)
        else:
            changed.sections[0].header.paragraphs[0].text = 'WRONG HEADER'
            changed.sections[0].footer.paragraphs[0].text = 'WRONG FOOTER'
            changed.sections[0].header_distance = Mm(40)
            changed.sections[0].footer_distance = Mm(40)
        bad = case / 'corrupted.docx'
        changed.save(bad)
        results[name]['corrupted'] = check(bad)
    case = root / 'duplicate_image'
    case.mkdir(exist_ok=True)
    source = case / 'source.docx'
    doc = Document()
    p = doc.add_paragraph()
    text_run = p.add_run('Anonymous image preservation contract.')
    text_run.font.name = 'Arial'
    text_run.font.size = Pt(12)
    pic_run = p.add_run()
    pic_run.font.name = 'Arial'
    pic_run.font.size = Pt(12)
    pic_run.add_picture(str(Path(__file__).parent/'r0-r11-review/thesis-fixture/figure.png'), width=Mm(15))
    doc.save(source)
    expected = ExpectedInventory(extract_semantic_inventory(source))
    clone = copy.deepcopy(next(doc.element.body.iter(qn('w:drawing'))))
    for node in clone.iter(qn('wp:docPr')):
        node.set('id', '999')
        node.set('name', 'Anonymous duplicate')
    pic_run._r.append(clone)
    bad = case / 'extra-image.docx'
    doc.save(bad)
    actual = extract_semantic_inventory(bad)
    results['duplicate_image'] = {'source_instances': len(expected.semantic.media_instances),
                                  'actual_instances': len(actual.media_instances),
                                  'content_gate_accepted': verify_content_integrity(expected, bad)}
    # An explicit missing intermediate heading level must not be compressed.
    case = root / 'missing_level'
    case.mkdir(exist_ok=True)
    source = case / 'source.docx'
    doc = Document()
    for heading, size in [('1. Introduction', 18), ('1.1.1. Detailed protocol', 14)]:
        run = doc.add_paragraph().add_run(heading)
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.bold = True
        for _ in range(3):
            run = doc.add_paragraph().add_run('This anonymous body paragraph describes the procedure and contains sufficient ordinary text to establish the body typography baseline. ' * 2)
            run.font.name = 'Arial'
            run.font.size = Pt(12)
    doc.save(source)
    analysis = analyze_format_sample(source)
    results['missing_level'] = {path: {'selected_role': item['selected_role'], 'selected_score': item['selected_score']}
                                for path, item in analysis['unstructured_candidate_evidence'].items()
                                if path in ['/w:document/w:body/w:p[1]', '/w:document/w:body/w:p[5]']}
    dataset = {'dataset_id': 'independent-missing-level', 'document_count': 1, 'categories': ['anonymous'],
               'documents': [{'filename': source.name, 'category': 'anonymous', 'heading_paragraph_indices': [0, 4]}]}
    baseline = {'documents': [{'filename': source.name, 'labels': {str(i): ('heading.1' if i == 0 else 'heading.3' if i == 4 else 'body') for i in range(8)}}]}
    spec = importlib.util.spec_from_file_location('audit_evaluation', Path(__file__).parent/'r7-evaluation/evaluate_unstructured_dataset.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    results['missing_level_metrics'] = module.evaluate(dataset, case, baseline)['overall']
    Path(sys.argv[2]).write_text(json.dumps(results, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
