"""Independent acceptance probes for R0-R11. No production changes or Word mocks.

Run from repo root with an empty anonymous output directory as the argument.
render_body probes stop before delivery; native CLI builds are run separately.
"""
import copy
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from docx import Document
from docx.shared import Mm, Pt
from lib.engine import UnifiedSynthesizer
from lib.content_integrity import extract_semantic_inventory, verify_content_integrity, verify_delivery_format
from lib.manifest_migration import migrate_manifest_file


def base():
    return {'schema_version': 3, 'project_name': 'AnonymousReview',
            'source': {'strategy': 'docx_document'},
            'format': {'ref': 'preset:academic-basic@1.0.0'},
            'formatting': {'mode': 'restyle'},
            'output': {'documents': [{'id': 'main', 'filename': 'result.docx', 'parts': ['body']}]}}


def main():
    root = Path(sys.argv[1]).resolve()
    root.mkdir(parents=True, exist_ok=True)
    obs = {}
    for mode in ['preserve', 'emphasis', 'normal']:
        case = root / mode
        case.mkdir(exist_ok=True)
        doc = Document()
        if mode == 'preserve':
            doc.sections[0].page_width = Mm(180)
            doc.sections[0].page_height = Mm(240)
            doc.sections[0].footer.paragraphs[0].text = 'ORIGINAL FOOTER'
        for i in range(5):
            p = doc.add_paragraph(f'Anonymous paragraph {i + 1} documents a reproducible formatting contract. ')
            if mode == 'emphasis' and i == 0:
                p.add_run('Required emphasis').bold = True
            p.add_run(' The remaining words are ordinary body text for this independent acceptance review.')
        source = case / 'source.docx'
        doc.save(source)
        raw = base()
        if mode == 'preserve':
            raw['formatting'] = {'mode': 'preserve', 'page_policy': 'source'}
        manifest = case / 'manifest.json'
        manifest.write_text(json.dumps(raw))
        plan = UnifiedSynthesizer.prepare_build(source, manifest)
        result = UnifiedSynthesizer.render_body(plan, case / 'rendered')
        rep = verify_delivery_format(result.path, plan.config.resolved_format)
        obs[mode] = {'stage': 'real render_body + real verifier, no Word',
                     'format_passed': rep.passed, 'verified_count': rep.verified_count,
                     'violations': rep.violations}
        if mode == 'normal':
            altered = Document(result.path)
            altered.paragraphs[0].style = altered.styles['Normal']
            for run in altered.paragraphs[0].runs:
                run.font.size = Pt(70)
            bad = case / 'style-removed.docx'
            altered.save(bad)
            check = verify_delivery_format(bad, plan.config.resolved_format)
            obs['removed_managed_style'] = {'actual_first_run_pt': 70, 'accepted': check.passed,
                                           'verified_count': check.verified_count,
                                           'total_blocks': check.details['total_blocks']}

    source = Document()
    source.add_paragraph('First required paragraph')
    source.add_paragraph('Second required paragraph')
    inv = extract_semantic_inventory(source)
    output = Document()
    output.add_paragraph('Second required paragraph')
    output.add_paragraph('First required paragraph')
    output.add_paragraph('First required paragraph')
    obs['reordered_duplicated_content_accepted'] = verify_content_integrity(inv, output)

    migration = root / 'anonymous-migration.json'
    raw = {'schema_version': 1, 'project_name': 'AnonymousMigration'}
    migration.write_text(json.dumps(raw))
    before = migration.read_bytes()
    migrate_manifest_file(migration, migration, replace=True)
    obs['migration_overwrites_own_input'] = migration.read_bytes() != before

    raw = base()
    (root / 'body-pageref-manifest.json').write_text(json.dumps(raw))
    (root / 'observations.json').write_text(json.dumps(obs, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(obs, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
