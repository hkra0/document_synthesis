#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""P4 真实 Word 验收矩阵（WW01–WW12）与跨平台页码对照。

在 macOS 与 Windows 上运行同一套匿名夹具（与 tests/test_post_review_word.py 的
NW01–NW12 相同），每项按 plan → doctor → build → smoke 串行执行生产 CLI，
复刻 NW 的断言，并记录书签物理页/打印页码、导出 PDF 的页数与页码标签。

子命令：

    python docs/acceptance/public-release/word_matrix.py run --label windows-word
    python docs/acceptance/public-release/word_matrix.py environment --label windows-word
    python docs/acceptance/public-release/word_matrix.py compare macos-word windows-word

``run`` 与 ``environment`` 的结果写入 ``docs/acceptance/public-release/<label>/``；
``compare`` 写入 ``docs/acceptance/public-release/cross-platform-pages.json``。
证据中的仓库路径、临时目录与用户目录都替换为占位符，不含个人信息。

状态口径：``verified`` 仅表示全部步骤退出码为 0 且全部检查通过；
doctor 报告精确 Office 不可用时为 ``blocked``；其余为 ``failed``。skip 不算通过。
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
from typing import Any, Callable, Dict, List, Optional, Tuple

REPO = Path(__file__).resolve().parents[3]
EVIDENCE_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(REPO))

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION_START  # noqa: E402
from docx.shared import Mm, Pt  # noqa: E402

SCHEMA_VERSION = 1
STEP_TIMEOUT = 1800


# ---------------------------------------------------------------------------
# 脱敏
# ---------------------------------------------------------------------------

_EMAIL = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")


class Scrubber:
    """把机器相关路径替换为占位符：<repo>、<work>、<tmp>、<user-home>。"""

    def __init__(self, work: Optional[Path] = None) -> None:
        pairs = [(REPO, "<repo>")]
        if work is not None:
            pairs.append((work, "<work>"))
        pairs.append((Path(tempfile.gettempdir()), "<tmp>"))
        pairs.append((Path.home(), "<user-home>"))
        self.pairs: List[Tuple[str, str]] = []
        for path, token in pairs:
            for variant in {str(path), str(path.resolve()), path.resolve().as_posix()}:
                self.pairs.append((variant, token))
        self.pairs.sort(key=lambda item: len(item[0]), reverse=True)

    def text(self, value: str) -> str:
        for raw, token in self.pairs:
            value = re.sub(re.escape(raw), token, value, flags=re.IGNORECASE)
        value = re.sub(r"[A-Za-z]:[\\/]+Users[\\/]+[^\\/\s\"']+", "<user-home>", value)
        value = re.sub(r"/(?:Users|home)/[^/\s\"']+", "<user-home>", value)
        return _EMAIL.sub("<email>", value)

    def data(self, value: Any) -> Any:
        if isinstance(value, str):
            return self.text(value)
        if isinstance(value, list):
            return [self.data(item) for item in value]
        if isinstance(value, dict):
            return {self.text(str(key)): self.data(item) for key, item in value.items()}
        return value


# ---------------------------------------------------------------------------
# 夹具（与 tests/test_post_review_word.py 保持一致）
# ---------------------------------------------------------------------------

def _write_manifest(root: Path, *, filename: str, project: str, parts: Optional[List[str]] = None,
                    source_strategy: str = "docx_document", formatting: Optional[dict] = None,
                    cover: Optional[dict] = None) -> Path:
    data: Dict[str, Any] = {
        "schema_version": 3,
        "project_name": project,
        "source": {"strategy": source_strategy},
        "format": {"ref": "preset:academic-basic@1.0.0"},
        "formatting": formatting or {"mode": "restyle"},
        "output": {"documents": [{"id": "main", "filename": filename, "parts": parts or ["body"]}]},
    }
    if cover:
        data["cover"] = cover
    path = root / f"{Path(filename).stem}.manifest.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _write_basic_source(path: Path, *, headings: bool = True, paragraphs: int = 12) -> None:
    doc = Document()
    if headings:
        doc.add_heading("第一章 匿名研究", level=1)
    for index in range(paragraphs):
        doc.add_paragraph(
            f"匿名验收正文段落 {index + 1}。该段仅用于真实 Word 分页、格式和内容门禁验证，"
            "不含任何业务材料。" * 4
        )
    if headings:
        doc.add_heading("第二章 匿名结果", level=1)
        doc.add_paragraph("第二章正文仅用于真实 Word 页级验收，确保章节标题后存在可分页的正文内容。" * 4)
    doc.save(path)


def _run_script(script: str, target: Path) -> None:
    subprocess.run(
        [sys.executable, str(REPO / "docs" / "acceptance" / script), str(target)],
        check=True, capture_output=True, text=True, encoding="utf-8", errors="replace",
    )


def _part_text(path: Path, part_id: str) -> str:
    from docx.oxml.ns import qn

    doc = Document(str(path))
    children = list(doc.element.body)
    start_name = {"cover": "_Synth_cover", "toc": "_Synth_toc", "body": "_Synth_body"}[part_id]
    end_name = {"cover": "_Synth_part_cover_end", "toc": "_Synth_part_toc_end", "body": "_Synth_part_body_end"}[part_id]
    start = next(i for i, child in enumerate(children) if any(
        item.get(qn("w:name")) == start_name for item in child.iter(qn("w:bookmarkStart"))))
    end = next(i for i in range(start, len(children)) if any(
        item.get(qn("w:name")) == end_name for item in children[i].iter(qn("w:bookmarkStart"))))
    return "".join(node.text or "" for child in children[start:end + 1] for node in child.iter(qn("w:t")))


# ---------------------------------------------------------------------------
# 用例
# ---------------------------------------------------------------------------

class Context:
    def __init__(self, root: Path, source: Path, manifest: Path) -> None:
        self.root = root
        self.source = source
        self.manifest = manifest
        self.out = root / "out"
        self.delivered: Optional[Path] = None
        self.metadata: Dict[str, Any] = {}
        self.exported_pdf: Optional[Path] = None
        self.labels: Optional[Dict[int, str]] = None
        self.checks: List[Dict[str, Any]] = []

    def check(self, name: str, passed: bool, detail: str = "") -> None:
        self.checks.append({"name": name, "passed": bool(passed), "detail": detail})

    def delivery(self) -> Dict[str, Any]:
        return self.metadata.get("deliveries", {}).get("main", {})


def _zip_members(path: Path) -> Tuple[set, bytes, bytes]:
    import zipfile

    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        document = package.read("word/document.xml")
        footnotes = package.read("word/footnotes.xml") if "word/footnotes.xml" in names else b""
    return names, document, footnotes


def prepare_ww01(root: Path) -> Tuple[Path, Path]:
    fixture = root / "fixture"
    fixture.mkdir()
    _run_script("create_word_followup_fixtures.py", fixture)
    raw = json.loads((fixture / "pageref-manifest.json").read_text(encoding="utf-8"))
    raw["layout"]["parts"].pop("toc", None)
    raw["output"]["documents"][0]["parts"] = ["body"]
    manifest = root / "no-toc.json"
    manifest.write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")
    return fixture / "pageref-source.docx", manifest


def check_ww01(ctx: Context) -> None:
    from lib.field_updater import build_field_index, verify_final_field_caches

    index = build_field_index(ctx.delivered)
    ctx.check("存在 PAGEREF 目标", bool(index.pageref_targets))
    page_map = ctx.delivery().get("page_map", {})
    target = index.pageref_targets[0] if index.pageref_targets else None
    ctx.check("PAGEREF 目标有最终页码", target in page_map, str(target))
    fields = verify_final_field_caches(ctx.delivered, page_map)["fields"] if target in page_map else []
    ctx.check("PAGEREF 缓存已回填（不是占位值 1）", bool(fields) and fields[0]["actual"] != "1",
              fields[0]["actual"] if fields else "")


def prepare_numbering(root: Path) -> Tuple[Path, Path]:
    fixture = root / "fixture"
    fixture.mkdir()
    _run_script("create_word_followup_fixtures.py", fixture)
    return fixture / "numbering-source.docx", fixture / "numbering-manifest.json"


def check_ww02(ctx: Context) -> None:
    labels = set((ctx.labels or {}).values())
    ctx.check("PDF 页码标签含罗马数字 i", "i" in labels, ",".join(sorted(labels)))
    ctx.check("PDF 页码标签含阿拉伯数字 1", "1" in labels)


def prepare_ww03(root: Path) -> Tuple[Path, Path]:
    old = root / "封面+目录.docx"
    old_doc = Document()
    old_doc.add_paragraph("旧业务封面")
    old_doc.save(old)
    source = root / "source.docx"
    _write_basic_source(source, paragraphs=4)
    manifest = _write_manifest(root, filename="cover.docx", project="匿名封面", parts=["cover", "body"],
                               cover={"main_title": "WW03 新标题", "author": "匿名", "date": "2026"})
    return source, manifest


def check_ww03(ctx: Context) -> None:
    cover = _part_text(ctx.delivered, "cover")
    ctx.check("封面使用 v3 新标题", "WW03 新标题" in cover)
    ctx.check("不引用旧封面模板", "旧业务封面" not in cover)


def prepare_ww04(root: Path) -> Tuple[Path, Path]:
    source = root / "source.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(180)
    section.page_height = Mm(240)
    paragraph = doc.add_paragraph()
    paragraph.add_run("保留加粗").bold = True
    paragraph.add_run("普通强调")
    doc.save(source)
    return source, _write_manifest(root, filename="preserve.docx", project="WW04",
                                   formatting={"mode": "preserve", "page_policy": "source"})


def check_ww04(ctx: Context) -> None:
    import pymupdf

    out = Document(str(ctx.delivered))
    ctx.check("DOCX 页宽保持 180 mm", abs(out.sections[0].page_width.mm - 180) <= 0.1,
              f"{out.sections[0].page_width.mm:.2f}")
    ctx.check("加粗保留", bool(out.paragraphs[0].runs[0].bold))
    with pymupdf.open(ctx.exported_pdf) as rendered:
        width = rendered[0].rect.width * 25.4 / 72
    ctx.check("Word 导出 PDF 页宽 180 mm", abs(width - 180) <= 0.2, f"{width:.2f}")


def prepare_ww05(root: Path) -> Tuple[Path, Path]:
    source = root / "source.docx"
    doc = Document()
    doc.add_paragraph("托管正文")
    doc.add_paragraph().add_run("保护正文").bold = True
    doc.save(source)
    return source, _write_manifest(root, filename="mixed.docx", project="WW05", formatting={
        "mode": "mixed", "page_policy": "target",
        "scopes": [{"file": str(source), "node_range": "/w:document/w:body/w:p[1]", "mode": "restyle"}],
    })


def check_ww05(ctx: Context) -> None:
    from lib.docx_inspector import inspect_docx

    inspected = inspect_docx(ctx.delivered)
    ctx.check("托管段落使用 Synth 样式", inspected.blocks[0].p_style_id.startswith("Synth"),
              inspected.blocks[0].p_style_id)
    ctx.check("保护段落保留加粗", bool(inspected.blocks[1].effective_run.bold))


def prepare_thesis(root: Path) -> Tuple[Path, Path]:
    fixture = root / "thesis"
    _run_script("create_complete_thesis_fixture.py", fixture)
    return fixture, fixture / "manifest.json"


def check_ww06(ctx: Context) -> None:
    names, document, _ = _zip_members(ctx.delivered)
    ctx.check("含脚注部件", "word/footnotes.xml" in names)
    ctx.check("含尾注部件", "word/endnotes.xml" in names)
    for token in (b"PAGEREF ", b"SEQ Figure", b"SEQ Table", b"w:drawing", b"w:tbl", b"oMath"):
        ctx.check(f"正文含 {token.decode().strip()}", token in document)
    ctx.check("封面含论文标题", "匿名学位论文样例" in _part_text(ctx.delivered, "cover"))


def check_ww07(ctx: Context) -> None:
    labels = set((ctx.labels or {}).values())
    ctx.check("前置部分使用罗马页码 i", "i" in labels, ",".join(sorted(labels)))
    body = ctx.delivery().get("page_records", {}).get("_Synth_body")
    ctx.check("正文起始书签有页码记录", body is not None)
    ctx.check("正文从物理奇数页开始", bool(body) and body.get("physical_page", 0) % 2 == 1,
              str(body.get("physical_page") if body else ""))


def prepare_ww08(root: Path) -> Tuple[Path, Path]:
    source = root / "source.docx"
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "首页页眉"
    doc.add_paragraph("纵向节正文" * 20)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = 1
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.header.paragraphs[0].text = "横向节页眉"
    doc.add_paragraph("横向表格节正文" * 20)
    doc.save(source)
    return source, _write_manifest(root, filename="sections.docx", project="WW08",
                                   formatting={"mode": "preserve", "page_policy": "source"})


def check_ww08(ctx: Context) -> None:
    out = Document(str(ctx.delivered))
    ctx.check("至少两个节", len(out.sections) >= 2, str(len(out.sections)))
    ctx.check("第二节为横向 297 mm", len(out.sections) >= 2 and abs(out.sections[1].page_width.mm - 297) <= 0.2)
    ctx.check("第二节页眉保留", len(out.sections) >= 2 and "横向节页眉" in out.sections[1].header.paragraphs[0].text)


def prepare_ww09(root: Path) -> Tuple[Path, Path]:
    generated = root / "generated"
    generated.mkdir()
    _run_script("create_complete_thesis_fixture.py", generated)
    sources = root / "sources"
    sources.mkdir()
    shutil.copy2(generated / "source" / "thesis-source.docx", sources / "one.docx")
    shutil.copy2(generated / "source" / "thesis-source.docx", sources / "two.docx")
    return sources, _write_manifest(root, filename="notes.docx", project="WW09", source_strategy="directory_tree")


def check_ww09(ctx: Context) -> None:
    names, _, footnotes = _zip_members(ctx.delivered)
    ctx.check("含脚注部件", "word/footnotes.xml" in names)
    count = footnotes.count("匿名论文第一条虚构脚注".encode("utf-8"))
    ctx.check("两个来源的同号脚注都保留", count >= 2, str(count))


def check_ww10(ctx: Context) -> None:
    from lib.field_updater import build_field_index

    commands = [item.command for item in build_field_index(ctx.delivered).observations]
    ctx.check("保留 PAGEREF 域", "PAGEREF" in commands)
    ctx.check("保留 SEQ 域", "SEQ" in commands)
    _, document, _ = _zip_members(ctx.delivered)
    for token in (b"w:drawing", b"w:tbl", b"oMath", b"w:numPr"):
        ctx.check(f"保留 {token.decode()}", token in document)


def prepare_ww11(root: Path) -> Tuple[Path, Path]:
    work = root / "workflow"
    demo = REPO / "examples" / "custom-format-demo"
    fmt = work / "academic-demo.json"
    mapping = work / "roles.json"
    steps = [
        ["--analyze-format", str(demo / "sample.docx"), "--analysis-dir", str(work / "sample"), "--replace-output"],
        ["--compile-format", str(work / "sample" / "analysis.json"), "--decisions", str(demo / "decisions.final.json"),
         "--format-base", "preset:academic-basic@1.0.0", "--format-id", "ww11", "--format-version", "1.0.0",
         "--format-out", str(fmt), "--replace-output"],
        ["--analyze-source", str(demo / "manuscript.docx"), "--analysis-dir", str(work / "target"), "--replace-output"],
        ["--compile-mapping", str(work / "target" / "analysis.json"), "--decisions",
         str(demo / "manuscript-decisions.final.json"), "--mapping-out", str(mapping), "--replace-output"],
    ]
    for args in steps:
        result = _cli(args)
        if result["returncode"] != 0:
            raise RuntimeError(f"格式工作流步骤失败: {args[0]}: {result['tail']}")
    manifest = _write_manifest(root, filename="ww11.docx", project="WW11",
                               formatting={"mode": "restyle", "role_map": str(mapping), "on_unmapped": "error"})
    data = json.loads(manifest.read_text(encoding="utf-8"))
    data["format"]["ref"] = str(fmt)
    manifest.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    source = root / "manuscript.docx"
    shutil.copy2(demo / "manuscript.docx", source)
    return source, manifest


def check_ww11(ctx: Context) -> None:
    ctx.check("按纠正后的格式包与角色映射完成构建", ctx.delivered is not None and ctx.delivered.is_file())


def prepare_ww12(root: Path) -> Tuple[Path, Path]:
    source = root / "source.docx"
    _write_basic_source(source, paragraphs=12)
    return source, _write_manifest(root, filename="published.docx", project="WW12")


def check_ww12(ctx: Context) -> None:
    from lib.build_plan import prepare_project_build
    from lib.config import load_project_config
    from lib.content_integrity import verify_delivery_format
    from lib.verification_contracts import build_verification_context

    before = hashlib.sha256(ctx.delivered.read_bytes()).hexdigest()
    corrupted = ctx.root / "corrupted.docx"
    doc = Document(str(ctx.delivered))
    doc.paragraphs[0].runs[0].font.size = Pt(70)
    doc.save(corrupted)
    config = load_project_config(ctx.source.parent, ctx.manifest, project_name="WW12")
    prepared = prepare_project_build(ctx.source, config)
    context = build_verification_context(prepared, None, config.documents[0], corrupted)
    report = verify_delivery_format(corrupted, config.resolved_format, verification_context=context)
    ctx.check("篡改格式后核验失败（fail closed）", not report.passed)
    ctx.check("已发布交付物哈希不变", before == hashlib.sha256(ctx.delivered.read_bytes()).hexdigest())


CASES: List[Tuple[str, str, str, Callable[[Path], Tuple[Path, Path]], Callable[[Context], None]]] = [
    ("WW01", "NW01", "无目录 PAGEREF 测量与缓存回填", prepare_ww01, check_ww01),
    ("WW02", "NW02", "目录与引用使用最终页码标签", prepare_numbering, check_ww02),
    ("WW03", "NW03", "v3 默认封面不引用旧模板", prepare_ww03, check_ww03),
    ("WW04", "NW04", "preserve 保留页面几何与行内强调", prepare_ww04, check_ww04),
    ("WW05", "NW05", "mixed 模式托管与保护范围区分", prepare_ww05, check_ww05),
    ("WW06", "NW06", "完整匿名论文含全部声明对象", prepare_thesis, check_ww06),
    ("WW07", "NW07", "罗马数字前置页码与奇数页正文起始", prepare_numbering, check_ww07),
    ("WW08", "NW08", "preserve 横向节与页眉变体", prepare_ww08, check_ww08),
    ("WW09", "NW09", "两个来源同号脚注保持关联", prepare_ww09, check_ww09),
    ("WW10", "NW10", "对象与域指令经生产 CLI 保留", prepare_thesis, check_ww10),
    ("WW11", "NW11", "样例纠正、格式包、角色映射与目标构建", prepare_ww11, check_ww11),
    ("WW12", "NW12", "篡改注入 fail closed，已发布哈希不变", prepare_ww12, check_ww12),
]


# ---------------------------------------------------------------------------
# 执行
# ---------------------------------------------------------------------------

def _cli(args: List[str], script: str = "synthesize.py") -> Dict[str, Any]:
    started = time.monotonic()
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    result = subprocess.run(
        [sys.executable, script] + args, cwd=REPO, capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=STEP_TIMEOUT, env=env,
    )
    output = (result.stdout or "") + (result.stderr or "")
    return {
        "returncode": result.returncode,
        "seconds": round(time.monotonic() - started, 1),
        "tail": "\n".join(output.strip().splitlines()[-8:]),
    }


def _export_and_measure(ctx: Context) -> Dict[str, Any]:
    """用当前平台的 Word 导出交付物 PDF，记录页数与页码标签（跨平台对照用）。"""
    from pypdf import PdfReader
    from lib.pagination import extract_pdf_page_labels
    from lib.qa import export_docx_to_pdf

    pdf = ctx.root / "export.pdf"
    if not export_docx_to_pdf(str(ctx.delivered), str(pdf)):
        raise RuntimeError("Word 导出交付物 PDF 失败")
    ctx.exported_pdf = pdf
    ctx.labels = extract_pdf_page_labels(pdf)
    return {
        "page_count": len(PdfReader(str(pdf)).pages),
        "page_labels": {str(page): label for page, label in sorted(ctx.labels.items())},
    }


def run_case(case_id: str, nw_id: str, title: str, prepare, check, keep: Optional[Path]) -> Dict[str, Any]:
    root = Path(tempfile.mkdtemp(prefix=f"{case_id.lower()}-"))
    scrub = Scrubber(root)
    record: Dict[str, Any] = {"id": case_id, "mirrors": nw_id, "title": title, "steps": {}, "checks": []}
    started = time.monotonic()
    try:
        source, manifest = prepare(root)
        ctx = Context(root, source, manifest)
        common = ["--source", str(source), "--manifest", str(manifest)]
        for step, args, script in (
            ("plan", common + ["--plan"], "synthesize.py"),
            ("doctor", common + ["--doctor"], "synthesize.py"),
            ("build", common + ["--output-dir", str(ctx.out)], "synthesize.py"),
            ("smoke", common + ["--output-dir", str(ctx.out)], "smoke_test.py"),
        ):
            result = _cli(args, script)
            record["steps"][step] = result
            if step == "doctor" and result["returncode"] == 2:
                record["status"] = "blocked"
                record["reason"] = "doctor 报告精确 Office 自动化不可用"
                break
            if result["returncode"] != 0:
                record["status"] = "failed"
                record["reason"] = f"{step} 退出码 {result['returncode']}"
                break
            if step == "build":
                filename = json.loads(manifest.read_text(encoding="utf-8"))["output"]["documents"][0]["filename"]
                ctx.delivered = ctx.out / filename
                metadata_path = ctx.out / "build-metadata.json"
                if metadata_path.is_file():
                    ctx.metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        else:
            # 四个步骤全部成功后才导出 PDF 并执行用例检查。
            delivery = ctx.delivery()
            record["delivered_sha256"] = delivery.get("sha256")
            record["page_map"] = delivery.get("page_map", {})
            record["export"] = _export_and_measure(ctx)
            check(ctx)
            record["checks"] = ctx.checks
            failed = [item["name"] for item in ctx.checks if not item["passed"]]
            record["status"] = "failed" if failed or not ctx.checks else "verified"
            if failed:
                record["reason"] = "检查未通过: " + "；".join(failed)
    except Exception as exc:  # noqa: BLE001 - 记录为失败，不中断矩阵
        record["status"] = "failed"
        record["reason"] = f"{type(exc).__name__}: {exc}"
    finally:
        record["seconds"] = round(time.monotonic() - started, 1)
        if keep is not None:
            target = keep / case_id
            shutil.rmtree(target, ignore_errors=True)
            shutil.copytree(root, target, ignore=shutil.ignore_patterns("~$*"))
        shutil.rmtree(root, ignore_errors=True)
    return scrub.data(record)


def _git_commit() -> str:
    try:
        return subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=REPO, text=True).strip()
    except Exception:
        return ""


def command_run(args: argparse.Namespace) -> int:
    selected = {item.strip().upper() for item in args.cases.split(",")} if args.cases else None
    keep = Path(args.keep_work).resolve() if args.keep_work else None
    results = []
    for case in CASES:
        if selected and case[0] not in selected:
            continue
        print(f"[{case[0]}] {case[2]} ...", flush=True)
        record = run_case(*case, keep=keep)
        print(f"[{case[0]}] {record['status']} ({record['seconds']} 秒) {record.get('reason', '')}", flush=True)
        results.append(record)
    summary = {status: sum(1 for item in results if item["status"] == status)
               for status in ("verified", "failed", "blocked")}
    document = {
        "schema_version": SCHEMA_VERSION,
        "label": args.label,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "git_commit": _git_commit(),
        "office_backend": os.environ.get("DOCUMENT_SYNTHESIS_OFFICE_BACKEND", "auto"),
        "note": args.note or "",
        "summary": summary,
        "cases": results,
    }
    out_dir = EVIDENCE_ROOT / args.label
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / (args.output or "ww-results.json")
    out_file.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(f"结果: {summary}，已写入 {out_file.relative_to(REPO).as_posix()}")
    return 0 if summary["failed"] == 0 and summary["blocked"] == 0 else 1


# ---------------------------------------------------------------------------
# 环境
# ---------------------------------------------------------------------------

#: 预设与示例使用的字体，以及它们在各平台注册表或字体文件中常见的名称。
FONT_ALIASES: Dict[str, List[str]] = {
    "宋体": ["SimSun", "宋体", "Songti"],
    "黑体": ["SimHei", "黑体", "Heiti"],
    "仿宋": ["FangSong", "仿宋", "STFangsong"],
    "楷体": ["KaiTi", "楷体", "STKaiti"],
    "仿宋_GB2312": ["FangSong_GB2312", "仿宋_GB2312"],
    "楷体_GB2312": ["KaiTi_GB2312", "楷体_GB2312"],
    "方正小标宋简体": ["FZXiaoBiaoSong-B05S", "FZXBSJW", "方正小标宋"],
    "Times New Roman": ["Times New Roman", "TimesNewRoman"],
}


def _windows_font_names() -> List[str]:
    import winreg

    names = []
    for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            with winreg.OpenKey(root, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts") as key:
                index = 0
                while True:
                    try:
                        names.append(winreg.EnumValue(key, index)[0])
                    except OSError:
                        break
                    index += 1
        except OSError:
            pass
    return names


def _mac_font_names() -> List[str]:
    names = []
    for directory in (Path("/Library/Fonts"), Path("/System/Library/Fonts"), Path.home() / "Library" / "Fonts",
                      Path("/Applications/Microsoft Word.app/Contents/Resources/DFonts")):
        if directory.is_dir():
            names.extend(path.name for path in directory.rglob("*") if path.suffix.lower() in (".ttf", ".ttc", ".otf"))
    return names


def _font_report() -> Dict[str, Any]:
    installed = _windows_font_names() if sys.platform == "win32" else _mac_font_names() if sys.platform == "darwin" else []
    lowered = [name.lower() for name in installed]
    report = {}
    for family, aliases in FONT_ALIASES.items():
        matches = sorted({installed[i] for i, name in enumerate(lowered)
                          for alias in aliases if alias.lower() in name})
        report[family] = {"installed": bool(matches), "matched_names": matches[:4]}
    return {"method": "Windows 字体注册表" if sys.platform == "win32" else "字体文件名", "families": report}


_C2R_CHANNELS = {
    "492350f6-3a01-4f97-b9c0-c7c6ddf67d60": "Current Channel",
    "64256afe-f5d9-4f86-8936-8840a6a4f5be": "Current Channel (Preview)",
    "55336b82-a18d-4dd6-b5f6-9e5095c314a6": "Monthly Enterprise Channel",
    "7ffbc6bf-bc32-4f92-8982-f9dd17fd3114": "Semi-Annual Enterprise Channel",
    "b8f9b850-328d-4355-9145-c59439a0c4cf": "Semi-Annual Enterprise Channel (Preview)",
    "5440fd1f-7ecb-4221-8110-145efaa6372f": "Beta Channel",
}


def _windows_environment() -> Dict[str, Any]:
    import winreg

    def read(root, path, name):
        try:
            with winreg.OpenKey(root, path) as key:
                return winreg.QueryValueEx(key, name)[0]
        except OSError:
            return None

    version = sys.getwindowsversion()
    c2r = r"SOFTWARE\Microsoft\Office\ClickToRun\Configuration"
    cdn = read(winreg.HKEY_LOCAL_MACHINE, c2r, "CDNBaseUrl") or ""
    channel_id = cdn.rstrip("/").rsplit("/", 1)[-1].lower() if cdn else ""
    from lib.office.win_com import WinComBackend

    probe = WinComBackend().probe()
    try:
        import win32print

        printer = win32print.GetDefaultPrinter()
    except Exception:
        printer = None
    addins = {}
    for root_name, root in (("HKCU", winreg.HKEY_CURRENT_USER), ("HKLM", winreg.HKEY_LOCAL_MACHINE)):
        path = r"Software\Microsoft\Office\Word\Addins"
        try:
            with winreg.OpenKey(root, path) as key:
                index = 0
                while True:
                    try:
                        progid = winreg.EnumKey(key, index)
                    except OSError:
                        break
                    index += 1
                    addins.setdefault(progid, {})[root_name] = read(root, path + "\\" + progid, "LoadBehavior")
        except OSError:
            pass
    return {
        "os": {
            "system": "Windows",
            "release": platform.release(),
            "version": platform.version(),
            "build": version.build,
            "edition": read(winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion", "EditionID"),
            "display_version": read(winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion", "DisplayVersion"),
            "long_paths_enabled": bool(read(winreg.HKEY_LOCAL_MACHINE, r"SYSTEM\CurrentControlSet\Control\FileSystem", "LongPathsEnabled")),
        },
        "word": {
            "probe_available": probe.available,
            "version": probe.details.get("word_version"),
            "build": probe.details.get("word_build"),
            "bitness": probe.details.get("word_bitness"),
            "click_to_run": probe.details.get("click_to_run"),
            "version_to_report": read(winreg.HKEY_LOCAL_MACHINE, c2r, "VersionToReport"),
            "channel": _C2R_CHANNELS.get(channel_id, channel_id or None),
            "platform": read(winreg.HKEY_LOCAL_MACHINE, c2r, "Platform"),
            "powerpoint_registered": probe.details.get("powerpoint_registered"),
        },
        "word_com_addins_load_behavior": addins,
        "default_printer": printer,
    }


def _mac_environment() -> Dict[str, Any]:
    import plistlib

    info = {}
    plist = Path("/Applications/Microsoft Word.app/Contents/Info.plist")
    if plist.is_file():
        with plist.open("rb") as handle:
            data = plistlib.load(handle)
        info = {key: data.get(key) for key in ("CFBundleShortVersionString", "CFBundleVersion")}
    try:
        printer = subprocess.run(["lpstat", "-d"], capture_output=True, text=True, timeout=10).stdout.strip() or None
    except Exception:
        printer = None
    return {
        "os": {"system": "macOS", "release": platform.mac_ver()[0], "machine": platform.machine()},
        "word": info,
        "default_printer": printer,
    }


def command_environment(args: argparse.Namespace) -> int:
    from importlib import metadata

    packages = {}
    for name in ("python-docx", "pymupdf", "pypdf", "Pillow", "lxml", "jsonschema", "numpy", "pywin32"):
        try:
            packages[name] = metadata.version(name)
        except metadata.PackageNotFoundError:
            pass
    environment: Dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "label": args.label,
        "recorded_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "python": platform.python_version(),
        "platform": platform.platform(),
        "packages": packages,
    }
    if sys.platform == "win32":
        environment.update(_windows_environment())
    elif sys.platform == "darwin":
        environment.update(_mac_environment())
    environment["fonts"] = _font_report()
    out_dir = EVIDENCE_ROOT / args.label
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / "environment.json"
    out_file.write_text(json.dumps(Scrubber().data(environment), ensure_ascii=False, indent=2) + "\n",
                        encoding="utf-8", newline="\n")
    print(f"已写入 {out_file.relative_to(REPO).as_posix()}")
    return 0


# ---------------------------------------------------------------------------
# 跨平台对照
# ---------------------------------------------------------------------------

_COMPARED_FIELDS = ("physical_page", "printed_page", "observed_label")


_SOURCE_BOOKMARK = re.compile(r"^_SynthBm_([0-9a-f]+)_(\d+)$")


def _normalize_bookmark_names(page_map: Dict[str, Any]) -> Dict[str, Any]:
    """把 ``_SynthBm_<源路径哈希>_<序号>`` 中随机器路径变化的哈希换成稳定的组号。

    重名书签改名时使用源文件绝对路径的哈希（lib/renderers.py），不同运行、不同平台必然不同；
    组号按该组最早出现的物理页、再按序号排序，与路径无关。
    """
    groups: Dict[str, Tuple[int, int]] = {}
    for name, record in page_map.items():
        match = _SOURCE_BOOKMARK.match(name)
        if match:
            key = (int(record.get("physical_page") or 0), int(match.group(2)))
            token = match.group(1)
            groups[token] = min(groups.get(token, key), key)
    order = {token: index for index, token in enumerate(sorted(groups, key=lambda t: groups[t]), 1)}
    renamed = {}
    for name, record in page_map.items():
        match = _SOURCE_BOOKMARK.match(name)
        renamed[f"_SynthBm_s{order[match.group(1)]}_{match.group(2)}" if match else name] = record
    return renamed


def _page_view(page_map: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    """只取页码本身；核验状态等字段与平台无关，不参与对照。"""
    return {name: {key: record.get(key) for key in _COMPARED_FIELDS}
            for name, record in _normalize_bookmark_names(page_map).items()}


def compare_results(left: Dict[str, Any], right: Dict[str, Any]) -> Dict[str, Any]:
    """逐用例比较书签物理页/打印页码与导出 PDF 的页数、页码标签。"""
    left_cases = {item["id"]: item for item in left.get("cases", [])}
    right_cases = {item["id"]: item for item in right.get("cases", [])}
    cases = []
    for case_id in sorted(set(left_cases) | set(right_cases)):
        a, b = left_cases.get(case_id), right_cases.get(case_id)
        entry: Dict[str, Any] = {"id": case_id, "differences": []}
        if a is None or b is None:
            entry["comparable"] = False
            entry["reason"] = "只有一个平台有结果"
            cases.append(entry)
            continue
        if a.get("status") != "verified" or b.get("status") != "verified":
            entry["comparable"] = False
            entry["reason"] = f"状态 {a.get('status')} / {b.get('status')}，非 verified 不对照"
            cases.append(entry)
            continue
        entry["comparable"] = True
        pages_a, pages_b = _page_view(a.get("page_map", {})), _page_view(b.get("page_map", {}))
        for name in sorted(set(pages_a) | set(pages_b)):
            if pages_a.get(name) != pages_b.get(name):
                entry["differences"].append({"bookmark": name, left["label"]: pages_a.get(name), right["label"]: pages_b.get(name)})
        export_a, export_b = a.get("export", {}), b.get("export", {})
        if export_a.get("page_count") != export_b.get("page_count"):
            entry["differences"].append({"export_page_count": {left["label"]: export_a.get("page_count"), right["label"]: export_b.get("page_count")}})
        if export_a.get("page_labels") != export_b.get("page_labels"):
            entry["differences"].append({"export_page_labels": {left["label"]: export_a.get("page_labels"), right["label"]: export_b.get("page_labels")}})
        entry["bookmarks_compared"] = len(set(pages_a) | set(pages_b))
        cases.append(entry)
    comparable = [item for item in cases if item["comparable"]]
    return {
        "schema_version": SCHEMA_VERSION,
        "left": {"label": left.get("label"), "git_commit": left.get("git_commit"), "generated_at": left.get("generated_at")},
        "right": {"label": right.get("label"), "git_commit": right.get("git_commit"), "generated_at": right.get("generated_at")},
        "summary": {
            "cases": len(cases),
            "comparable": len(comparable),
            "identical": sum(1 for item in comparable if not item["differences"]),
            "with_differences": sum(1 for item in comparable if item["differences"]),
        },
        "cases": cases,
    }


def command_compare(args: argparse.Namespace) -> int:
    left = json.loads((EVIDENCE_ROOT / args.left / "ww-results.json").read_text(encoding="utf-8"))
    right = json.loads((EVIDENCE_ROOT / args.right / "ww-results.json").read_text(encoding="utf-8"))
    report = compare_results(left, right)
    out_file = EVIDENCE_ROOT / "cross-platform-pages.json"
    out_file.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(f"对照结果: {report['summary']}，已写入 {out_file.relative_to(REPO).as_posix()}")
    return 0 if report["summary"]["with_differences"] == 0 and report["summary"]["comparable"] == report["summary"]["cases"] else 1


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)
    run = sub.add_parser("run", help="运行 WW01–WW12")
    run.add_argument("--label", required=True, help="证据目录名，例如 windows-word、macos-word")
    run.add_argument("--cases", help="只运行指定用例，逗号分隔，例如 WW01,WW07")
    run.add_argument("--output", help="结果文件名，默认 ww-results.json")
    run.add_argument("--note", help="写入结果的环境说明，例如 COM 插件已禁用")
    run.add_argument("--keep-work", help="把每个用例的工作目录复制到此处（不入库）")
    env = sub.add_parser("environment", help="记录环境信息")
    env.add_argument("--label", required=True)
    compare = sub.add_parser("compare", help="对照两个平台的 ww-results.json")
    compare.add_argument("left")
    compare.add_argument("right")
    args = parser.parse_args(argv)
    return {"run": command_run, "environment": command_environment, "compare": command_compare}[args.command](args)


if __name__ == "__main__":
    sys.exit(main())
