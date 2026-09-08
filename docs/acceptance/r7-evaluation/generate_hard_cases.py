#!/usr/bin/env python3
"""Generate the frozen no-outline numbering/conflict stress fixtures."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


CATEGORIES = {
    "公文": {"font": "SimSun", "align": WD_ALIGN_PARAGRAPH.LEFT},
    "报告": {"font": "Aptos", "align": WD_ALIGN_PARAGRAPH.LEFT},
    "论文": {"font": "宋体", "align": WD_ALIGN_PARAGRAPH.CENTER},
    "手工格式": {"font": "仿宋", "align": WD_ALIGN_PARAGRAPH.RIGHT},
}


def _font(run, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for script in ("eastAsia", "ascii", "hAnsi"):
        fonts.set(qn(f"w:{script}"), name)


def _paragraph(document: Document, text: str, category: str, size: float, *, bold: bool = False, alignment=None):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = CATEGORIES[category]["align"] if alignment is None else alignment
    run = paragraph.add_run(text)
    _font(run, CATEGORIES[category]["font"], size, bold=bold)
    return len(document.paragraphs) - 1


def _write_case(category: str, output: Path) -> dict[str, object]:
    document = Document()
    checks: list[dict[str, object]] = []

    index = _paragraph(document, f"1. {category}缺失中间层级", category, 18, bold=True)
    checks.append({"paragraph_index": index, "case": "missing_intermediate_level", "expected_role": "heading.1"})
    _paragraph(document, f"{category}缺失层级正文。" * 18, category, 12)
    # 16pt is intentionally the second local heading size (18pt is level 1),
    # while the numbering says depth 3.  This creates a real missing-level
    # conflict instead of letting the size ordering silently compress it.
    index = _paragraph(document, f"1.1.1. {category}深层冲突标题", category, 16, bold=True)
    checks.append({
        "paragraph_index": index,
        "case": "missing_intermediate_level",
        "expected_role": "heading.3",
        "expected_review_status": "pending_review",
        "required_conflict": "numbering_vs_size_level",
    })
    _paragraph(document, f"{category}深层标题后的正文。" * 18, category, 12)

    index = _paragraph(document, "2026.09.07 版本说明与发布日期", category, 18, bold=True)
    checks.append({
        "paragraph_index": index,
        "case": "date_or_version",
        "expected_role": "body",
        "expected_numbering_kind": "date_or_version",
        "expected_review_status": "pending_review",
    })
    index = _paragraph(document, "1. 普通编号正文项目，内容足够长，不应仅因为开头有编号就被识别成标题。" * 3, category, 12)
    checks.append({
        "paragraph_index": index,
        "case": "numbered_body",
        "expected_role": "body",
        "expected_review_status": "selected",
    })

    index = _paragraph(document, "注：短正文粗体说明", category, 12, bold=True)
    checks.append({
        "paragraph_index": index,
        "case": "short_bold_body",
        "expected_role": "body",
        "expected_review_status": "selected",
    })
    index = _paragraph(document, "图 1 研究框架题注", category, 12, bold=True)
    checks.append({
        "paragraph_index": index,
        "case": "caption",
        "expected_role": "body",
        "expected_review_status": "selected",
    })

    index = _paragraph(document, "2.1. 同层标题 A", category, 15, bold=True)
    checks.append({
        "paragraph_index": index,
        "case": "same_level_different_size",
        "expected_role": "heading.2",
        "expected_review_status": "pending_review",
    })
    _paragraph(document, f"同层标题 A 正文。" * 18, category, 12)
    index = _paragraph(document, "2.2. 同层标题 B", category, 14, bold=True)
    checks.append({
        "paragraph_index": index,
        "case": "same_level_different_size",
        "expected_role": "heading.2",
        "expected_review_status": "pending_review",
    })
    _paragraph(document, f"同层标题 B 正文。" * 18, category, 12)

    document.save(output)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    return {
        "filename": output.name,
        "category": category,
        "source_sha256": digest,
        "paragraph_count": len(document.paragraphs),
        "checks": checks,
        "all_documents_have_no_outlineLvl": True,
    }


def generate(output_dir: Path, manifest_path: Path) -> dict[str, object]:
    output_dir.mkdir(parents=True, exist_ok=True)
    documents = [
        _write_case(category, output_dir / f"{category}-hard-cases.docx")
        for category in CATEGORIES
    ]
    payload = {
        "dataset_schema_version": 1,
        "dataset_id": "r7-independent-anonymous-unstructured-hard-v1",
        "fictional_only": True,
        "frozen": True,
        "label_source": "fixture_author_manual_review",
        "document_count": len(documents),
        "categories": list(CATEGORIES),
        "case_types": sorted({check["case"] for item in documents for check in item["checks"]}),
        "documents": documents,
        "limitations": [
            "这是主冻结集之外的困难样例压力集，结论不外推为真实模板准确率。",
            "决策文件/人工操作日志仍需由实际审阅者在产品工作流中产生；本清单只冻结夹具事实与预期反例。",
        ],
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"dataset_id": payload["dataset_id"], "document_count": len(documents), "case_types": payload["case_types"]}, ensure_ascii=False))
    return payload


def main() -> int:
    parser = argparse.ArgumentParser()
    root = Path(__file__).parent
    parser.add_argument("--output-dir", type=Path, default=root / "unstructured-hard-samples")
    parser.add_argument("--manifest", type=Path, default=root / "unstructured-hard-cases.json")
    args = parser.parse_args()
    generate(args.output_dir.resolve(), args.manifest.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
