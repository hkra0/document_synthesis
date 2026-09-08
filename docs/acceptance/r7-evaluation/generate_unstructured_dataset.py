#!/usr/bin/env python3
"""Generate the independent no-outline evaluation set used by N7."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


CATEGORIES = {
    "公文": {"font": "SimSun", "body_size": 12, "align": WD_ALIGN_PARAGRAPH.LEFT},
    "报告": {"font": "Aptos", "body_size": 10.5, "align": WD_ALIGN_PARAGRAPH.LEFT},
    "论文": {"font": "宋体", "body_size": 12, "align": WD_ALIGN_PARAGRAPH.CENTER},
    "手工格式": {"font": "仿宋", "body_size": 11, "align": WD_ALIGN_PARAGRAPH.RIGHT},
}


def _font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for script in ("eastAsia", "ascii", "hAnsi"):
        fonts.set(qn(f"w:{script}"), name)


def _make(category: str, ordinal: int, output: Path) -> dict:
    spec = CATEGORIES[category]
    doc = Document()
    heading_indices = []
    body_indices = []
    for level in range(1, 6):
        heading = doc.add_paragraph(style="Normal")
        heading.alignment = spec["align"]
        prefix = f"{level}." if category in {"报告", "论文"} else f"第{level}部分"
        run = heading.add_run(f"{prefix} {category}无大纲评估标题 {ordinal}")
        _font(run, spec["font"], 16 - level * 0.5, bold=True)
        if category == "手工格式" and level % 2 == 0:
            run.italic = True
        heading_indices.append(len(doc.paragraphs) - 1)
        for body_no in range(1, 6):
            body = doc.add_paragraph(style="Normal")
            body_run = body.add_run(
                f"{category}无大纲正文负例 {ordinal}-{level}-{body_no}。"
                "这是一段较长的匿名虚构正文，用于检验候选角色的反证、长度、相邻段落和正文误报。"
                "它不应仅因为位于文档开头、包含数字或存在局部强调就被自动接受为标题。"
            )
            _font(body_run, spec["font"], spec["body_size"])
            if category == "手工格式" and body_no == 3:
                body_run.bold = True
            body_indices.append(len(doc.paragraphs) - 1)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {
        "category": category,
        "filename": output.name,
        "structured": False,
        "heading_paragraph_indices": heading_indices,
        "body_paragraph_indices": body_indices,
        "heading_count": len(heading_indices),
        "body_count": len(body_indices),
    }


def generate(output_dir: Path, dataset_path: Path) -> dict:
    entries = []
    for category in CATEGORIES:
        for ordinal in range(1, 6):
            entries.append(_make(category, ordinal, output_dir / f"{category}-unstructured-{ordinal:02d}.docx"))
    dataset = {
        "dataset_schema_version": 1,
        "dataset_id": "r7-independent-anonymous-unstructured-v1",
        "fictional_only": True,
        "structured": False,
        "categories": list(CATEGORIES),
        "document_count": len(entries),
        "heading_positive_count": sum(item["heading_count"] for item in entries),
        "body_negative_count": sum(item["body_count"] for item in entries),
        "documents": entries,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    dataset_path.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: dataset[key] for key in ("document_count", "heading_positive_count", "body_negative_count")}, ensure_ascii=False))
    return dataset


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).parent / "unstructured-samples")
    parser.add_argument("--dataset", type=Path, default=Path(__file__).parent / "unstructured-dataset.json")
    args = parser.parse_args()
    generate(args.output_dir.resolve(), args.dataset.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
