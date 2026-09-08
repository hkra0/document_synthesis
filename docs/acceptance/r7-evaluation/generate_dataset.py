# -*- coding: utf-8 -*-
"""Generate the independent, fictional R7 format-recognition evaluation set."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


CATEGORIES = {
    "公文": {"font": "SimSun", "body_size": 12, "heading_align": WD_ALIGN_PARAGRAPH.LEFT},
    "报告": {"font": "Aptos", "body_size": 10.5, "heading_align": WD_ALIGN_PARAGRAPH.LEFT},
    "论文": {"font": "宋体", "body_size": 12, "heading_align": WD_ALIGN_PARAGRAPH.CENTER},
    "手工格式": {"font": "仿宋", "body_size": 11, "heading_align": WD_ALIGN_PARAGRAPH.RIGHT},
}
DOCS_PER_CATEGORY = 5
HEADINGS_PER_DOC = 5
BODY_PER_HEADING = 5


def _set_run_font(run, font_name: str, size: float, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    east_asia = rpr.find(qn("w:rFonts"))
    if east_asia is None:
        east_asia = OxmlElement("w:rFonts")
        rpr.insert(0, east_asia)
    east_asia.set(qn("w:eastAsia"), font_name)
    east_asia.set(qn("w:ascii"), font_name)
    east_asia.set(qn("w:hAnsi"), font_name)


def _set_direct_outline(paragraph, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:outlineLvl"))
    if old is not None:
        ppr.remove(old)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level - 1))
    ppr.append(outline)


def _make_document(category: str, ordinal: int, output: Path) -> Dict[str, Any]:
    spec = CATEGORIES[category]
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(f"{category}匿名格式样例 {ordinal}")
    _set_run_font(title_run, spec["font"], 18, bold=True)

    heading_indices = []
    body_indices = []
    for level in range(1, HEADINGS_PER_DOC + 1):
        heading = doc.add_paragraph()
        heading.alignment = spec["heading_align"]
        # Keep direct formatting intentionally varied across categories.
        heading_run = heading.add_run(f"{category}第{level}级评估标题 {ordinal}")
        _set_run_font(heading_run, spec["font"], 16 - level * 0.5, bold=True)
        if category == "手工格式":
            heading_run.italic = level % 2 == 0
        if category == "报告":
            heading.style = "Heading 1" if level == 1 else "Normal"
        _set_direct_outline(heading, level)
        heading_indices.append(len(doc.paragraphs) - 1)

        for body_no in range(1, BODY_PER_HEADING + 1):
            body = doc.add_paragraph()
            body_run = body.add_run(
                f"{category}匿名正文负例 {ordinal}-{level}-{body_no}，用于识别评估的虚构语料。"
            )
            _set_run_font(body_run, spec["font"], spec["body_size"])
            if category == "手工格式" and body_no == 3:
                body_run.bold = True
            body_indices.append(len(doc.paragraphs) - 1)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return {
        "category": category,
        "filename": output.name,
        "heading_paragraph_indices": heading_indices,
        "body_paragraph_indices": body_indices,
        "heading_count": len(heading_indices),
        "body_count": len(body_indices),
    }


def generate_dataset(output_dir: Path) -> Dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    entries = []
    for category in CATEGORIES:
        for ordinal in range(1, DOCS_PER_CATEGORY + 1):
            filename = f"{category}-{ordinal:02d}.docx"
            entries.append(_make_document(category, ordinal, output_dir / filename))
    return {
        "dataset_schema_version": 1,
        "dataset_id": "r7-independent-anonymous-evaluation-v1",
        "fictional_only": True,
        "categories": list(CATEGORIES),
        "document_count": len(entries),
        "heading_positive_count": sum(item["heading_count"] for item in entries),
        "body_negative_count": sum(item["body_count"] for item in entries),
        "documents": entries,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).parent / "samples")
    parser.add_argument("--dataset", type=Path, default=Path(__file__).parent / "dataset.json")
    args = parser.parse_args()
    dataset = generate_dataset(args.output_dir.resolve())
    args.dataset.resolve().write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"generated {dataset['document_count']} documents, "
        f"{dataset['heading_positive_count']} heading positives, "
        f"{dataset['body_negative_count']} body negatives"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
