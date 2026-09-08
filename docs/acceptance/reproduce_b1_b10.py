"""Independent B1-B10 acceptance probes; never edits production code or inputs.

Run from repository root:
  python3 docs/acceptance/reproduce_b1_b10.py --output-dir output/acceptance-b1-b10-20260905

Prints observed results, not a claim of acceptance. Some probes replace the Word
delivery stage to isolate integration defects; they explicitly record this.
"""
import argparse
import contextlib
import io
import json
import shutil
import sys
from pathlib import Path
from unittest.mock import patch

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt
from PIL import Image
import pymupdf
from lib.composition import add_bookmark, assemble_document, mark_start
from lib.config import ProjectConfig
from lib.content_integrity import extract_semantic_inventory, verify_content_integrity
from lib.delivery import validate_delivery_structure, validate_measured_delivery
from lib.document_parts import SelectionValidator, SourceRegion
from lib.engine import UnifiedSynthesizer
from lib.field_updater import FieldUpdater


def base():
    return {
        'schema_version': 3, 'project_name': 'AcceptanceProbe',
        'source': {'strategy': 'docx_document'},
        'format': {'ref': 'preset:academic-basic@1.0.0'},
        'formatting': {'mode': 'restyle'}, 'cover': {'template': False},
        'output': {'documents': [{'id': 'main', 'filename': 'probe.docx', 'parts': ['body']}]},
    }


class Captured(Exception):
    pass


def capture_body(root, source, raw, case_name):
    case = root / case_name
    case.mkdir(exist_ok=True)
    manifest = case / 'manifest.json'
    manifest.write_text(json.dumps(raw), encoding='utf-8')
    result = {}

    def capture(config, body, nodes, run_dir):
        path = case / 'captured.docx'
        shutil.copy2(body, path)
        result.update(doc=Document(path), nodes=nodes)
        raise Captured()

    with patch('lib.delivery.build_deliveries', side_effect=capture), contextlib.redirect_stdout(io.StringIO()):
        try:
            UnifiedSynthesizer.synthesize(source, case / 'out', manifest)
        except Captured:
            pass
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir', type=Path, required=True)
    root = parser.parse_args().output_dir.resolve()
    root.mkdir(parents=True, exist_ok=True)
    results = {}
    source = root / 'fixture.docx'
    doc = Document()
    doc.sections[0].page_width = Mm(180)
    doc.sections[0].page_height = Mm(240)
    doc.sections[0].bottom_margin = Mm(30)
    doc.sections[0].footer.paragraphs[0].text = 'ORIGINAL FOOTER'
    p = doc.add_paragraph('Manually mapped heading')
    p.add_run(' source emphasis').font.size = Pt(37)
    for i in range(4):
        doc.add_paragraph(f'Body paragraph {i} with enough unique substantive text.')
    doc.save(source)

    mapping = root / 'stale-map.json'
    mapping.write_text(json.dumps({
        'mapping_schema_version': 1, 'source_sha256': '0' * 64,
        'assignments': [{'element_path': '/w:document/w:body/w:p[1]', 'role': 'heading.2'}],
    }), encoding='utf-8')
    raw = base()
    raw['formatting']['role_map'] = str(mapping)
    captured = capture_body(root, source, raw, 'mapping')
    results['mapping'] = {
        'stale_map_rejected': False,
        'expected_style_if_mapping_used': 'SynthHeading2',
        'actual_style': captured['doc'].paragraphs[0].style.style_id,
        'heading_nodes': len(captured['nodes']),
    }

    raw = base()
    raw['formatting'] = {'mode': 'preserve', 'page_policy': 'source'}
    d = capture_body(root, source, raw, 'preserve')['doc']
    results['preserve'] = {
        'source_style': 'Normal', 'actual_style': d.paragraphs[0].style.style_id,
        'source_bottom_mm': 30, 'actual_bottom_mm': round(d.sections[0].bottom_margin.mm, 3),
        'source_footer': 'ORIGINAL FOOTER', 'actual_footer_text': d.sections[0].footer.paragraphs[0].text,
    }

    directory = root / 'multi'
    directory.mkdir(exist_ok=True)
    shutil.copy2(source, directory / 'one.docx')
    raw = base()
    raw['source'] = {'strategy': 'directory_tree'}
    raw['format']['overrides'] = {'styles': {'heading.1': {'run': {'size_pt': 37}}}}
    d = capture_body(root, directory, raw, 'directory')['doc']
    results['directory'] = {'expected_heading_pt': 37, 'actual_heading_pt': d.paragraphs[0].runs[0].font.size.pt}

    pic = root / 'fixture.png'
    Image.new('RGB', (40, 40), 'red').save(pic)
    doc = Document()
    style = doc.styles.add_style('AuditCustomStyle', WD_STYLE_TYPE.PARAGRAPH)
    style.font.size = Pt(31)
    p = doc.add_paragraph('Region content', style=style)
    p.add_run().add_picture(str(pic))
    path = root / 'region.docx'
    doc.save(path)
    raw = base()
    raw['source']['regions'] = {'selected': {'start': '/w:document/w:body/w:p[1]', 'end': 'end_of_document'}}
    raw['layout'] = {'parts': {'main_text': {'kind': 'content', 'source_region': 'selected'}}}
    raw['output']['documents'][0]['parts'] = ['main_text']
    cfg = ProjectConfig(raw, root)
    out = root / 'region-result.docx'
    assemble_document(cfg, cfg.documents[0], path, [], {}, out)
    d = Document(out)
    rid = next(d.element.body.iter(qn('a:blip'))).get(qn('r:embed'))
    relation = d.part.rels.get(rid)
    validate_delivery_structure(out, cfg.documents[0], [], parts_registry=cfg.parts)
    results['region'] = {
        'image_rid': rid, 'actual_relation_type': relation.reltype if relation else None,
        'actual_target': relation.target_ref if relation else None,
        'media_parts': [str(p.partname) for p in d.part.package.parts if '/media/' in str(p.partname)],
        'source_style_exists': 'AuditCustomStyle' in {s.style_id for s in d.styles},
        'structure_validator_accepted': True,
    }
    doc = Document()
    doc.add_picture(str(pic))
    doc.add_paragraph('Selected text')
    SelectionValidator.validate_regions(doc, {'text': SourceRegion('text', '/w:document/w:body/w:p[2]', 'end_of_document')})
    results['unassigned_image_accepted'] = True

    raw = base()
    raw['source']['regions'] = {'invalid': {'start': '/w:document/w:body/w:p[99]', 'end': 'end_of_document'}}
    raw['layout'] = {'parts': {'main_text': {'kind': 'content', 'source_region': 'invalid'}}}
    raw['output']['documents'][0]['parts'] = ['main_text']
    cfg = ProjectConfig(raw, root)
    bad_region_path = root / 'bad-region.docx'
    assemble_document(cfg, cfg.documents[0], source, [], {}, bad_region_path)
    results['invalid_region_at_assembly'] = {
        'assembly_called_directly': True,
        'invalid_boundary_accepted': True,
        'source_paragraphs': len(Document(source).paragraphs),
        'actual_paragraphs': len(Document(bad_region_path).paragraphs),
    }

    raw = base()
    raw['layout'] = {'page_sequences': {'main': {'format': 'upperRoman', 'start': 1}},
                     'parts': {'body': {'kind': 'content', 'page_sequence': 'main'}}}
    cfg = ProjectConfig(raw, root)
    doc = Document()
    p = doc.add_paragraph('Heading target')
    add_bookmark(p._p, '_Audit_heading', 2)
    mark_start(doc, 'body')
    path = root / 'page-invalid.docx'
    doc.save(path)
    pdfpath = root / 'page-invalid.pdf'
    pdf = pymupdf.open()
    for i in range(2):
        page = pdf.new_page()
        for j in range(5):
            page.insert_text((50, 70 + j * 20), 'Visible substantive content for acceptance probe')
    pdf.save(pdfpath)
    pdf.close()
    nodes = [{'title': 'Heading target', 'level': 1, 'bookmark_name': '_Audit_heading'}]
    pages = {'_Synth_body': {'physical_page': 1, 'printed_page': 7, 'observed_label': '999'},
             '_Audit_heading': {'physical_page': 2, 'printed_page': 9, 'observed_label': '999'}}
    with contextlib.redirect_stdout(io.StringIO()):
        validate_measured_delivery(path, cfg.documents[0], nodes, pdfpath, pages,
                                   parts_registry=cfg.parts, page_sequences=cfg.page_sequences)
    results['page_validation'] = {'synthetic_measurement': True, 'expected_start': 1, 'invalid_map_accepted': True, 'map': pages}

    results['seq'] = {}
    for format_name in ('ROMAN', 'ALPHABETIC'):
        try:
            value = FieldUpdater()._evaluate_instruction('SEQ Figure \\* ' + format_name)
        except Exception as exc:
            value = f'{type(exc).__name__}: {exc}'
        results['seq'][format_name] = value
    doc = Document()
    p = doc.add_paragraph('Figure ')
    p._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="9" w:name="fig_number"/>'))
    p._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="SEQ Figure"><w:r><w:t>0</w:t></w:r></w:fldSimple>'))
    p._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="9"/>'))
    p.add_run(' Caption outside bookmark')
    ref = doc.add_paragraph()
    ref._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="REF fig_number"><w:r><w:t>old</w:t></w:r></w:fldSimple>'))
    FieldUpdater().update_document_fields(doc)
    results['ref'] = {'expected': '1', 'actual': ''.join(t.text or '' for t in ref._p.iter(qn('w:t')))}
    try:
        value = FieldUpdater({'target': 'III'})._evaluate_instruction('PAGEREF target')
    except Exception as exc:
        value = f'{type(exc).__name__}: {exc}'
    results['pageref_with_delivery_guess'] = value

    doc = Document()
    p = doc.add_paragraph('Unchanged body text')
    fn = parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="1"/>')
    p.add_run()._r.append(fn)
    note_xml = f'<w:footnotes {nsdecls("w")}><w:footnote w:id="1"><w:p><w:r><w:footnoteRef/><w:t>Required note text</w:t></w:r></w:p></w:footnote></w:footnotes>'
    note_part = Part(PackURI('/word/footnotes.xml'),
                     'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',
                     note_xml.encode('utf-8'), doc.part.package)
    note_rid = doc.part.relate_to(note_part, RT.FOOTNOTES)
    inventory = extract_semantic_inventory(doc)
    fn.getparent().remove(fn)
    doc.part.drop_rel(note_rid)
    results['removed_note_accepted_by_integrity'] = verify_content_integrity(inventory, doc)

    case = root / 'qa-gate'
    case.mkdir(exist_ok=True)
    manifest = case / 'manifest.json'
    manifest.write_text(json.dumps(base()), encoding='utf-8')

    def corrupt_stage(config, body, nodes, run_dir):
        target = Path(run_dir) / 'probe.docx'
        d = Document(body)
        d.paragraphs[0].runs[0].font.size = Pt(70)
        d.save(target)
        return {'main': target}

    with patch('lib.delivery.build_deliveries', side_effect=corrupt_stage), \
         patch('lib.qa.word_automation_status', return_value=(False, 'audit mock')), \
         contextlib.redirect_stdout(io.StringIO()):
        published = UnifiedSynthesizer.synthesize(source, case / 'out', manifest)
    meta = json.loads((case / 'out/build-metadata.json').read_text())
    results['format_gate'] = {'delivery_stage_mocked': True, 'published': published['main'].is_file(), 'qa': meta['qa']}
    (root / 'repro-results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
