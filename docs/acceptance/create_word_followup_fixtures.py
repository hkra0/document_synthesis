"""Generate anonymous M2 probes; build them with synthesize.py, never a mock.

Usage: python3 docs/acceptance/create_word_followup_fixtures.py OUTPUT_DIRECTORY
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from lib.role_mapper import RoleMapper


def main():
    root = Path(sys.argv[1]).resolve()
    root.mkdir(parents=True, exist_ok=True)
    source = root / 'numbering-source.docx'
    doc = Document()
    for chapter in (1, 2):
        doc.add_heading(f'Chapter {chapter} Anonymous Study', level=1)
        for i in range(6):
            doc.add_paragraph(f'Chapter {chapter}, observation {i + 1}. ' +
                              'This fictional study compares two layout configurations. '
                              'All statements are anonymous test material. '
                              'The document verifies page labels and stable cross references. ' * 5)
    doc.save(source)
    manifest = {
        'schema_version': 3, 'project_name': 'AnonymousM2Probe',
        'source': {'strategy': 'docx_document'},
        'format': {'ref': 'preset:academic-basic@1.0.0'},
        'formatting': {'mode': 'restyle'},
        'layout': {
            'page_sequences': {'front': {'format': 'lowerRoman', 'start': 1},
                               'main': {'format': 'decimal', 'start': 1}},
            'parts': {'toc': {'kind': 'generated_toc', 'page_sequence': 'front'},
                      'body': {'kind': 'content', 'page_sequence': 'main', 'section_type': 'oddPage'}}
        },
        'output': {'documents': [{'id': 'main', 'filename': 'm2.docx', 'parts': ['toc', 'body']}]}
    }
    (root / 'numbering-manifest.json').write_text(json.dumps(manifest, indent=2))
    headings, _ = RoleMapper().map_document(source)
    target = [a for a in headings if a.role == 'heading.1'][-1].bookmark_name
    p = doc.add_paragraph('The second chapter starts on page ')
    p._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGEREF {target}"><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    doc.save(root / 'pageref-source.docx')
    (root / 'pageref-manifest.json').write_text(json.dumps(manifest, indent=2))
    (root / 'fixture-expectations.json').write_text(json.dumps({
        'scope': 'M2 numbering and PAGEREF integration probes, not a complete thesis',
        'target_bookmark': target, 'toc_label': 'i', 'body_first_label': '1',
        'paper_mm': [210, 297], 'second_chapter': 'PAGEREF must match the measured label',
    }, indent=2))


if __name__ == '__main__':
    main()
