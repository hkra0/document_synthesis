#!/usr/bin/env python3
"""Create the deliberately anonymous DOCX inputs for the minimal demo."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source"
FONT = "Arial Unicode MS"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)


def set_font(run, size, color=None, bold=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    subtitle = doc.styles.add_style("Demo Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = FONT
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor(89, 89, 89)
    subtitle.paragraph_format.space_after = Pt(18)

    props = doc.core_properties
    props.author = ""
    props.title = "Anonymous Demo Material"
    props.subject = "Document Synthesis Test"
    props.keywords = "demo, anonymous, test"
    props.comments = ""
    props.category = ""


def add_title(doc, title, subtitle):
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(42)
    title_p.paragraph_format.space_after = Pt(4)
    run = title_p.add_run(title)
    set_font(run, 26, RGBColor(11, 37, 69), True)

    subtitle_p = doc.add_paragraph(subtitle, style="Demo Subtitle")
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    set_font(paragraph.add_run(text), 11)


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    set_font(paragraph.add_run(text), 11)


def add_overview():
    doc = Document()
    style_document(doc)
    add_title(doc, "Document Synthesis Demo", "For verifying outline discovery, body merging, and style handling")

    doc.add_paragraph("1. Demo objective", style="Heading 1")
    doc.add_paragraph("This material demonstrates the basic flow for reading several Word files, producing an outline plan, and keeping inputs read-only.")

    doc.add_paragraph("2. Included items", style="Heading 1")
    add_bullet(doc, "A short process overview")
    add_bullet(doc, "A practical checklist")
    add_bullet(doc, "A reviewable JSON configuration")

    doc.add_paragraph("3. Verification principle", style="Heading 1")
    doc.add_paragraph("Run the read-only plan first, then confirm the environment. Run a full build only after the environment check passes.")
    doc.save(SOURCE / "1_demo_overview.docx")


def add_checklist():
    doc = Document()
    style_document(doc)
    add_title(doc, "Operational Checklist", "Showing how ordered inputs appear in a synthesis plan")

    doc.add_paragraph("1. Prepare", style="Heading 1")
    add_numbered(doc, "Confirm that the input files are in the example source directory.")
    add_numbered(doc, "Confirm that manifest.json selects the directory_tree strategy.")

    doc.add_paragraph("2. Preview", style="Heading 1")
    add_numbered(doc, "Run synthesize.py with the --plan flag.")
    add_numbered(doc, "Review the titles, order, and node count.")

    doc.add_paragraph("3. Check before building", style="Heading 1")
    doc.add_paragraph("A full build requires local Word automation. If the environment check fails, stop after plan verification.")
    doc.save(SOURCE / "2_demo_checklist.docx")


def main():
    SOURCE.mkdir(parents=True, exist_ok=True)
    add_overview()
    add_checklist()


if __name__ == "__main__":
    main()
