#!/usr/bin/env python3
"""Generate Chinese demo inputs and an editable portrait cover, not the final build."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
COPY = json.loads((ROOT / "copy.json").read_text(encoding="utf-8"))


def font(run, name, size, color="000000"):
    name = {"仿宋_GB2312": "FangSong_GB2312", "楷体_GB2312": "KaiTi_GB2312",
            "黑体": "SimHei", "方正小标宋简体": "FZXiaoBiaoSong-B05S"}.get(name, name)
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().rFonts
    for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{kind}"), name)


def metadata(doc, title):
    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.core_properties.subject = "中文材料合成示例"
    doc.core_properties.comments = ""


def create_cover():
    doc = Document()
    title_style = doc.styles["Title"]
    for border in list(title_style._element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(28)
    section.top_margin, section.bottom_margin = Cm(2.8), Cm(2.5)
    section.left_margin = section.right_margin = Cm(1.8)
    cover = COPY["cover"]
    # Paragraph spacing keeps every line editable; no rasterized title or text box.
    for key, face, size, color, before, after in (
        ("eyebrow", "仿宋_GB2312", 16, "333333", 0, 0),
        ("title", "方正小标宋简体", 58, "B51E24", 100, 14),
        ("subtitle", "黑体", 25, "151515", 0, 0),
        ("inputs", "仿宋_GB2312", 16, "555555", 54, 0),
        ("feature_one", "楷体_GB2312", 20, "252525", 45, 10),
        ("feature_two", "楷体_GB2312", 20, "252525", 0, 0),
        ("footer", "仿宋_GB2312", 15, "555555", 85, 0),
    ):
        paragraph = doc.add_paragraph(style="Title" if key == "title" else "Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = paragraph.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = Pt(size * 1.4)
        pf.keep_with_next = False
        pf.widow_control = True
        font(paragraph.add_run(cover[key]), face, size, color)
    metadata(doc, cover["title"])
    path = ROOT / "templates" / "cover.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def create_source(section_copy):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(3.7), Cm(3.5)
    section.left_margin, section.right_margin = Cm(2.8), Cm(2.6)
    for block in section_copy["blocks"]:
        heading = "heading" in block
        paragraph = doc.add_paragraph(style="Heading 2" if heading else "Normal")
        pf = paragraph.paragraph_format
        pf.line_spacing = Pt(28)
        pf.space_before, pf.space_after = Pt(8 if heading else 0), Pt(0)
        pf.first_line_indent = Pt(32)
        pf.keep_with_next = heading
        pf.keep_together = heading
        pf.widow_control = True
        font(paragraph.add_run(block.get("heading", block.get("text"))),
             "楷体_GB2312" if heading else "仿宋_GB2312", 16)
    metadata(doc, section_copy["title"])
    path = ROOT / "source" / section_copy["file"]
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main():
    create_cover()
    for section_copy in COPY["sections"]:
        create_source(section_copy)
    print("已生成中文源材料与竖版封面模板。请使用 synthesize.py 合成。")


if __name__ == "__main__":
    main()
